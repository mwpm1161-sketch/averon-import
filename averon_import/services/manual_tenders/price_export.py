from __future__ import annotations

import hashlib
import json
import os
import posixpath
import re
import tempfile
import threading
import uuid
import zipfile
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

from averon_import.core.unit_normalization import normalize_unit_family
from averon_import.services.sourcing.history_identity import history_model_characteristic_conflicts
from averon_import.services.one_c_history.xlsx_import import OneCImportError, preflight_xlsx

from .parser import (
    MAX_UPLOAD_BYTES,
    TenderWorkbookParser,
    _package_preservation_inventory,
    _json_value,
    _style_fingerprint,
    is_semantically_empty_cell,
    parse_unit_basis,
    preflight_tender_xlsx,
)
from .repository import MAX_TENDER_STORAGE_BYTES, TenderWorkspaceError, TenderWorkspaceRepository
from .history_fuzzy_eligibility import (
    MAX_FUZZY_RETRIEVAL_RANK,
    fuzzy_confirmation_eligibility_reason,
    manual_search_confirmation_eligibility_reason,
)

TENDER_EXPORT_POLICY_REVISION = "xlsx-price-export-v1"
MAX_TENDER_EXPORT_BYTES = 10 * 1024 * 1024
MAX_COMPLETED_EXPORTS_PER_WORKSPACE = 3
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_EXPORT_ID_RE = re.compile(r"^[a-f0-9]{32}$")
_RUN_ID_RE = re.compile(r"^[a-f0-9]{32}$")
_LIVE_DECISIONS = {"MATCH", "LIKELY_MATCH"}
_SAFE_HISTORY_BASES = {"EXACT_ARTICLE", "EXACT_SOURCE_NAME_UNIT"}
_NON_CONVERTIBLE_UNIT_DIMENSIONS = {"package", "set"}
_HISTORICAL_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")


def _history_identity_text(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).casefold().replace("ё", "е").replace("\u00a0", " ").split())


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _bounded_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (str, Decimal, int, float)) and len(str(value)) > 100:
        return None
    try:
        result = value if isinstance(value, Decimal) else Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return result if result.is_finite() else None


@dataclass(frozen=True)
class TenderPriceDecision:
    source_row_id: str
    excel_row: int
    eligible: bool
    reason_code: str | None
    source_unit_price: Decimal | None
    total_price: Decimal | None
    source_kind: str | None
    historical: bool
    audit_summary: dict[str, Any]
    human_confirmed: bool = False

    def safe_summary(self) -> dict[str, Any]:
        return {
            "source_row_id": self.source_row_id,
            "excel_row": self.excel_row,
            "eligible": self.eligible,
            "reason_code": self.reason_code,
            "source_unit_price": str(self.source_unit_price) if self.source_unit_price is not None else None,
            "total_price": str(self.total_price) if self.total_price is not None else None,
            "source_kind": self.source_kind,
            "historical": self.historical,
            "authority": "HUMAN_CONFIRMED_HISTORY" if self.human_confirmed else ("AUTO_SAFE_HISTORY" if self.historical else None),
            "audit_summary": self.audit_summary,
        }


def _reason(
    code: str,
    row: dict[str, Any],
    *,
    historical: bool = False,
    human_confirmed: bool = False,
    audit: dict[str, Any] | None = None,
) -> TenderPriceDecision:
    return TenderPriceDecision(
        source_row_id=str(row.get("source_row_id") or ""),
        excel_row=int(row.get("excel_row") or 0),
        eligible=False,
        reason_code=code,
        source_unit_price=None,
        total_price=None,
        source_kind="one_c_history" if historical else None,
        historical=historical,
        audit_summary=audit or {},
        human_confirmed=human_confirmed,
    )


class TenderPriceResolver:
    """Pure deterministic policy over immutable workspace and durable run rows."""

    @staticmethod
    def validate_run(
        workspace: dict[str, Any],
        run: dict[str, Any],
        *,
        tender_id: str,
        run_id: str,
    ) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
        if not _RUN_ID_RE.fullmatch(str(run_id or "")) or run.get("run_id") != run_id:
            raise TenderWorkspaceError("Запуск подбора не найден.", 404, "TENDER_RUN_NOT_FOUND")
        if run.get("tender_id") != tender_id or workspace.get("tender_id") != tender_id:
            raise TenderWorkspaceError("Запуск подбора не найден.", 404, "TENDER_RUN_NOT_FOUND")
        if run.get("status") != "completed":
            raise TenderWorkspaceError("Экспорт доступен только для завершённого подбора.", 409, "TENDER_EXPORT_RUN_NOT_COMPLETED")
        try:
            revisions_match = int(run.get("workspace_revision", -1)) == int(workspace.get("revision", -2))
        except (TypeError, ValueError):
            revisions_match = False
        if run.get("source_sha256") != workspace.get("source_sha256") or not revisions_match:
            raise TenderWorkspaceError("Запуск не соответствует текущей версии тендера.", 409, "TENDER_RUN_WORKSPACE_MISMATCH")

        workspace_rows = workspace.get("rows")
        selected_ids = run.get("selected_source_row_ids")
        canonical_rows = run.get("rows")
        if not isinstance(workspace_rows, list) or not isinstance(selected_ids, list) or not isinstance(canonical_rows, list):
            raise TenderWorkspaceError("Данные подбора повреждены.", 409, "TENDER_RUN_CORRUPT")
        source_by_id: dict[str, dict[str, Any]] = {}
        for source in workspace_rows:
            if not isinstance(source, dict):
                raise TenderWorkspaceError("Строки тендера повреждены.", 409, "TENDER_WORKSPACE_CORRUPT")
            row_id = str(source.get("source_row_id") or "")
            if not row_id or row_id in source_by_id:
                raise TenderWorkspaceError("Строки тендера повреждены.", 409, "TENDER_WORKSPACE_CORRUPT")
            source_by_id[row_id] = source

        if any(not isinstance(item, str) or not item for item in selected_ids):
            raise TenderWorkspaceError("Данные подбора повреждены.", 409, "TENDER_RUN_CORRUPT")
        if len(selected_ids) != len(set(selected_ids)) or len(canonical_rows) != len(selected_ids):
            raise TenderWorkspaceError("Строки подбора не прошли проверку соответствия.", 409, "TENDER_RESULT_CORRELATION_FAILED")
        durable_by_id: dict[str, dict[str, Any]] = {}
        for canonical in canonical_rows:
            if not isinstance(canonical, dict):
                raise TenderWorkspaceError("Данные подбора повреждены.", 409, "TENDER_RUN_CORRUPT")
            row_id = str(canonical.get("source_row_id") or "")
            if not row_id or row_id in durable_by_id:
                raise TenderWorkspaceError("Строки подбора не прошли проверку соответствия.", 409, "TENDER_RESULT_CORRELATION_FAILED")
            durable_by_id[row_id] = canonical
        if set(selected_ids) != set(durable_by_id):
            raise TenderWorkspaceError("Строки подбора не прошли проверку соответствия.", 409, "TENDER_RESULT_CORRELATION_FAILED")
        for row_id in selected_ids:
            source = source_by_id.get(row_id)
            canonical = durable_by_id[row_id]
            if source is None:
                raise TenderWorkspaceError("Строка подбора отсутствует в тендере.", 409, "TENDER_RESULT_CORRELATION_FAILED")
            if source.get("row_type") != "item":
                raise TenderWorkspaceError("В запуск попала не строка позиции.", 409, "TENDER_RESULT_CORRELATION_FAILED")
            try:
                same_physical_row = int(canonical.get("physical_excel_row", -1)) == int(source.get("excel_row", -2))
            except (TypeError, ValueError):
                same_physical_row = False
            if not same_physical_row or canonical.get("status") != "completed":
                raise TenderWorkspaceError("Физическая строка подбора не совпадает с источником.", 409, "TENDER_RESULT_CORRELATION_FAILED")
        return source_by_id, durable_by_id

    def resolve(
        self,
        source: dict[str, Any],
        canonical: dict[str, Any],
        run: dict[str, Any],
        *,
        include_historical_prices: bool,
    ) -> TenderPriceDecision:
        row_id = str(source.get("source_row_id") or "")
        offer = canonical.get("recommended_offer")
        if not isinstance(offer, dict):
            return _reason("NO_RECOMMENDED_OFFER", source)

        match = canonical.get("recommended_match")
        route = canonical.get("route") if isinstance(canonical.get("route"), dict) else {}
        provenance = canonical.get("price_provenance") if isinstance(canonical.get("price_provenance"), dict) else {}
        provider = str(offer.get("provider") or "")
        final_kind = str(route.get("final_source_kind") or "")
        historical = final_kind == "historical_purchase"
        price = _bounded_decimal(offer.get("price"))
        audit: dict[str, Any] = {
            "provider": provider,
            "provenance_source": provenance.get("source"),
            "price_field": provenance.get("price_field"),
            "match_decision": match.get("decision") if isinstance(match, dict) else None,
            "currency": offer.get("currency"),
            "price_basis": provenance.get("price_basis"),
            "purchase_date": provenance.get("purchase_date"),
        }

        if price is None or price <= 0:
            return _reason("PRICE_INVALID", source, historical=historical, audit=audit)
        if offer.get("currency") != "RUB":
            return _reason("CURRENCY_UNSUPPORTED", source, historical=historical, audit=audit)
        exact_offer_match = (
            isinstance(match, dict)
            and isinstance(offer.get("offer_id"), str)
            and bool(offer.get("offer_id"))
            and match.get("offer_id") == offer.get("offer_id")
        )

        if historical:
            if provider != "one_c_history" or not exact_offer_match:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            if route.get("history_outcome") != "SAFE_MATCH" or route.get("history_safe_basis") not in _SAFE_HISTORY_BASES:
                return _reason("HISTORY_NOT_SAFE", source, historical=True, audit=audit)
            if run.get("source_mode") not in {"one_c_only", "one_c_then_provider"}:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            if provenance.get("source") != "one_c_history" or provenance.get("source_kind") != "historical_purchase":
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            if provenance.get("price_basis") != "gross_including_vat":
                return _reason("PRICE_BASIS_UNPROVEN", source, historical=True, audit=audit)
            history_item_id = str(provenance.get("history_item_id") or "")
            if not history_item_id or str(offer.get("source_item_id") or "") != history_item_id:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            event_id = str(provenance.get("selected_event_id") or "")
            routed_event_id = str(route.get("history_selected_event_id") or "")
            if not event_id or not routed_event_id or routed_event_id != event_id:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            route_version = str(route.get("history_catalog_version") or "")
            run_version = str(run.get("history_catalog_version") or "")
            provenance_version = str(provenance.get("snapshot_version") or "")
            if (
                not route_version or not run_version or not provenance_version
                or route_version != run_version
                or provenance_version != route_version
            ):
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            purchase_date = str(provenance.get("purchase_date") or "")
            try:
                parsed_date = date.fromisoformat(purchase_date)
            except ValueError:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            if parsed_date > datetime.now(timezone.utc).date() or str(route.get("history_purchase_date") or "") != purchase_date:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            effective_gross = _bounded_decimal(provenance.get("effective_unit_price_gross"))
            if effective_gross is None or effective_gross <= 0 or effective_gross != price:
                return _reason("PROVENANCE_MISMATCH", source, historical=True, audit=audit)
            if not include_historical_prices:
                return _reason("HISTORICAL_PRICE_NOT_INCLUDED", source, historical=True, audit=audit)
        else:
            source_mode = str(run.get("source_mode") or "")
            routed_provider = source_mode == "one_c_then_provider"
            # The direct provider_only path intentionally returns no OneC route;
            # one_c_then_provider still needs its explicit live-provider route.
            if (
                source_mode not in {"provider_only", "one_c_then_provider"}
                or (routed_provider and final_kind != "provider")
                or not exact_offer_match
            ):
                return _reason("MATCH_NOT_EXPORTABLE", source, audit=audit)
            if not isinstance(match, dict) or match.get("decision") not in _LIVE_DECISIONS:
                return _reason("MATCH_NOT_EXPORTABLE", source, audit=audit)
            if provenance.get("source") != provider or str((canonical.get("provider_source") or {}).get("provider") or "") != provider:
                return _reason("PROVENANCE_MISMATCH", source, audit=audit)
            source_item_id = str(offer.get("source_item_id") or "")
            if not source_item_id or str(provenance.get("source_item_id") or "") != source_item_id:
                return _reason("PROVENANCE_MISMATCH", source, audit=audit)
            if provider == "etm_ipro":
                # Client-supplied ETM API Product documentation (24.01.2025)
                # defines price as net and pricewnds as VAT-inclusive.
                if provenance.get("price_field") != "pricewnds":
                    return _reason("PRICE_BASIS_UNPROVEN", source, audit=audit)
            elif provider == "lemana_b2b":
                return _reason("PRICE_BASIS_UNPROVEN", source, audit=audit)
            else:
                return _reason("PRICE_BASIS_UNPROVEN", source, audit=audit)

        if source.get("quantity_trusted") is not True:
            return _reason("QUANTITY_UNTRUSTED", source, historical=historical, audit=audit)
        quantity = _bounded_decimal(source.get("quantity"))
        if quantity is None or quantity <= 0:
            return _reason("QUANTITY_UNTRUSTED", source, historical=historical, audit=audit)
        source_basis = source.get("unit_basis")
        reparsed_source = parse_unit_basis(source.get("raw_unit"))
        if not isinstance(source_basis, dict) or source_basis.get("trusted") is not True or reparsed_source.get("trusted") is not True:
            return _reason("SOURCE_UNIT_UNTRUSTED", source, historical=historical, audit=audit)
        basis_fields = ("raw_unit", "base_unit", "dimension", "scale", "conversion_basis", "trusted")
        if any(source_basis.get(key) != reparsed_source.get(key) for key in basis_fields):
            return _reason("UNIT_BASIS_UNPROVEN", source, historical=historical, audit=audit)
        if reparsed_source.get("dimension") in _NON_CONVERTIBLE_UNIT_DIMENSIONS:
            return _reason("SOURCE_UNIT_UNTRUSTED", source, historical=historical, audit=audit)
        offer_basis = parse_unit_basis(offer.get("price_unit"))
        if offer_basis.get("trusted") is not True:
            return _reason("OFFER_UNIT_UNTRUSTED", source, historical=historical, audit=audit)
        if offer_basis.get("dimension") in _NON_CONVERTIBLE_UNIT_DIMENSIONS:
            return _reason("OFFER_UNIT_UNTRUSTED", source, historical=historical, audit=audit)
        if reparsed_source.get("dimension") != offer_basis.get("dimension") or reparsed_source.get("base_unit") != offer_basis.get("base_unit"):
            return _reason("UNIT_INCOMPATIBLE", source, historical=historical, audit=audit)
        try:
            source_scale = Decimal(str(reparsed_source.get("scale")))
            offer_scale = Decimal(str(offer_basis.get("scale")))
            if source_scale <= 0 or offer_scale <= 0:
                raise InvalidOperation
            with localcontext() as context:
                context.prec = 50
                exported_price = (price * source_scale / offer_scale).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
                total = (exported_price * quantity).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, TypeError, ValueError):
            return _reason("PRICE_INVALID", source, historical=historical, audit=audit)
        audit.update({
            "source_unit": str(source.get("raw_unit") or ""),
            "offer_unit": str(offer.get("price_unit") or ""),
            "source_scale": str(source_scale),
            "offer_scale": str(offer_scale),
        })
        return TenderPriceDecision(
            source_row_id=row_id,
            excel_row=int(source.get("excel_row") or 0),
            eligible=True,
            reason_code=None,
            source_unit_price=exported_price,
            total_price=total,
            source_kind="one_c_history" if historical else provider,
            historical=historical,
            audit_summary=audit,
        )

    def resolve_human_history(
        self,
        source: dict[str, Any],
        canonical: dict[str, Any],
        run: dict[str, Any],
        confirmation: dict[str, Any],
        *,
        include_historical_prices: bool,
    ) -> TenderPriceDecision:
        """Resolve an exact durable review candidate under separate human authority."""
        candidate = confirmation.get("candidate") if isinstance(confirmation, dict) else None
        if not isinstance(candidate, dict):
            return _reason("HISTORY_CONFIRMATION_INVALID", source)
        offer = candidate.get("offer") if isinstance(candidate.get("offer"), dict) else {}
        provenance = candidate.get("price_provenance") if isinstance(candidate.get("price_provenance"), dict) else {}
        match = candidate.get("match") if isinstance(candidate.get("match"), dict) else {}
        route = canonical.get("route") if isinstance(canonical.get("route"), dict) else {}
        audit: dict[str, Any] = {
            "provider": offer.get("provider"),
            "provenance_source": provenance.get("source"),
            "match_decision": match.get("decision"),
            "currency": offer.get("currency"),
            "price_basis": provenance.get("price_basis"),
            "purchase_date": provenance.get("purchase_date"),
            "human_confirmation_id": confirmation.get("decision_id"),
            "human_confirmation_fingerprint": confirmation.get("evidence_fingerprint"),
            "human_confirmation_basis": confirmation.get("confirmation_basis"),
        }
        confirmation_basis = confirmation.get("confirmation_basis")
        manual_confirmation = confirmation_basis == "MANUAL_SEARCH_CONFIRMATION"
        if (
            run.get("source_mode") != "one_c_only"
            or route.get("source_mode") != "one_c_only"
            or (
                route.get("final_source_kind") not in {"history_review", "none"}
                or route.get("history_outcome") not in {"REVIEW", "NO_MATCH"}
                if manual_confirmation
                else route.get("final_source_kind") != "history_review" or route.get("history_outcome") != "REVIEW"
            )
            or route.get("history_safe_basis") not in (None, "")
            or not confirmation.get("decision_id")
            or not confirmation.get("evidence_fingerprint")
        ):
            return _reason("HISTORY_CONFIRMATION_INVALID", source, historical=True, human_confirmed=True, audit=audit)
        retrieval_classification = offer.get("retrieval_classification")
        fuzzy_confirmation = confirmation_basis == "FUZZY_MANUAL_CONFIRMATION"
        if (
            offer.get("provider") != "one_c_history"
            or retrieval_classification not in {"EXACT_ARTICLE", "EXACT_NAME_UNIT", "NORMALIZED_NAME_UNIT", "FUZZY", "MANUAL_HISTORY_SEARCH"}
            or (retrieval_classification == "NORMALIZED_NAME_UNIT" and confirmation_basis != "NORMALIZED_CONFIRMATION")
            or (retrieval_classification in {"EXACT_ARTICLE", "EXACT_NAME_UNIT"} and confirmation_basis != "EXACT_CONFIRMATION")
            or (retrieval_classification == "FUZZY" and (
                not fuzzy_confirmation
                or confirmation.get("identity_assertion") != "SAME_PRODUCT_V1"
                or isinstance(confirmation.get("candidate", {}).get("retrieval_rank"), bool)
                or not isinstance(confirmation.get("candidate", {}).get("retrieval_rank"), int)
                or not 1 <= confirmation["candidate"]["retrieval_rank"] <= MAX_FUZZY_RETRIEVAL_RANK
            ))
            or (retrieval_classification == "MANUAL_HISTORY_SEARCH" and (
                not manual_confirmation
                or confirmation.get("identity_assertion") != "SAME_PRODUCT_V1"
                or not isinstance(confirmation.get("candidate", {}).get("manual_variant_id"), str)
                or not confirmation["candidate"]["manual_variant_id"]
                or confirmation["candidate"].get("price_provenance", {}).get("manual_integrity_valid") is not True
            ))
            or (retrieval_classification != "FUZZY" and confirmation_basis == "FUZZY_MANUAL_CONFIRMATION")
            or (retrieval_classification != "MANUAL_HISTORY_SEARCH" and manual_confirmation)
            or match.get("offer_id") != offer.get("offer_id")
            or offer.get("offer_id") != confirmation.get("candidate_offer_id")
        ):
            return _reason("HISTORY_CONFIRMATION_INVALID", source, historical=True, human_confirmed=True, audit=audit)
        if retrieval_classification == "FUZZY":
            fuzzy_reason = fuzzy_confirmation_eligibility_reason(
                source, offer, provenance, match, route,
                expected_snapshot_version=run.get("history_catalog_version"),
                physical_excel_row=canonical.get("physical_excel_row"),
                retrieval_rank=candidate.get("retrieval_rank"),
            )
            if fuzzy_reason is not None:
                return _reason(fuzzy_reason, source, historical=True, human_confirmed=True, audit=audit)
        elif retrieval_classification == "MANUAL_HISTORY_SEARCH":
            manual_reason = manual_search_confirmation_eligibility_reason(
                source, offer, provenance, match, route,
                expected_snapshot_version=run.get("history_catalog_version"),
                physical_excel_row=canonical.get("physical_excel_row"),
            )
            if manual_reason is not None:
                return _reason(manual_reason, source, historical=True, human_confirmed=True, audit=audit)
        history_characteristic = offer.get("history_characteristic", "")
        if not isinstance(history_characteristic, str) or len(history_characteristic) > 300:
            return _reason("HISTORY_CONFIRMATION_INVALID", source, historical=True, human_confirmed=True, audit=audit)
        if history_model_characteristic_conflicts(source.get("model"), history_characteristic):
            return _reason("HISTORY_CANDIDATE_SOURCE_CONFLICT", source, historical=True, human_confirmed=True, audit=audit)
        source_article = _history_identity_text(source.get("article"))
        candidate_article = _history_identity_text(offer.get("article"))
        if source_article and (not candidate_article or source_article != candidate_article):
            return _reason("HISTORY_CANDIDATE_SOURCE_CONFLICT", source, historical=True, human_confirmed=True, audit=audit)
        source_manufacturer = _history_identity_text(source.get("manufacturer"))
        candidate_manufacturer = _history_identity_text(offer.get("manufacturer"))
        if source_manufacturer and candidate_manufacturer and source_manufacturer != candidate_manufacturer:
            return _reason("HISTORY_CANDIDATE_SOURCE_CONFLICT", source, historical=True, human_confirmed=True, audit=audit)
        if retrieval_classification == "NORMALIZED_NAME_UNIT":
            from averon_import.services.sourcing.history_identity import (
                HISTORY_IDENTITY_NORMALIZER_REVISION,
                history_name_signature,
                history_name_signature_digest,
            )

            source_signature = history_name_signature(source.get("name"))
            candidate_signature = history_name_signature(offer.get("title"))
            if (
                not source_signature
                or source_signature != candidate_signature
                or provenance.get("normalizer_revision") != HISTORY_IDENTITY_NORMALIZER_REVISION
                or provenance.get("normalized_name_signature") != history_name_signature_digest(source.get("name"))
            ):
                return _reason("HISTORY_CONFIRMATION_INVALID", source, historical=True, human_confirmed=True, audit=audit)
        snapshot = str(run.get("history_catalog_version") or "")
        if (
            not snapshot
            or confirmation.get("history_snapshot_version") != snapshot
            or provenance.get("snapshot_version") != snapshot
            or provenance.get("source") != "one_c_history"
            or provenance.get("source_kind") != "historical_purchase"
            or provenance.get("history_item_id") != offer.get("source_item_id")
            or provenance.get("history_item_id") != confirmation.get("history_item_id")
            or not provenance.get("selected_event_id")
            or provenance.get("selected_event_id") != confirmation.get("selected_event_id")
        ):
            return _reason("PROVENANCE_MISMATCH", source, historical=True, human_confirmed=True, audit=audit)
        if (
            route.get("history_selected_event_id") not in (None, "", provenance.get("selected_event_id"))
            or route.get("history_purchase_date") not in (None, "", provenance.get("purchase_date"))
        ):
            return _reason("PROVENANCE_MISMATCH", source, historical=True, human_confirmed=True, audit=audit)
        purchase_date = str(provenance.get("purchase_date") or "")
        try:
            parsed_date = date.fromisoformat(purchase_date)
        except ValueError:
            return _reason("PROVENANCE_MISMATCH", source, historical=True, human_confirmed=True, audit=audit)
        if parsed_date > datetime.now(timezone.utc).date():
            return _reason("PROVENANCE_MISMATCH", source, historical=True, human_confirmed=True, audit=audit)
        price = _bounded_decimal(offer.get("price"))
        effective = _bounded_decimal(provenance.get("effective_unit_price_gross"))
        if (
            price is None or effective is None or price <= 0 or price != effective
            or provenance.get("price_basis") != "gross_including_vat"
        ):
            return _reason("PRICE_BASIS_UNPROVEN", source, historical=True, human_confirmed=True, audit=audit)
        if offer.get("currency") != "RUB" or provenance.get("currency_basis") not in {"source", "company_default"}:
            return _reason("CURRENCY_UNSUPPORTED", source, historical=True, human_confirmed=True, audit=audit)
        if not include_historical_prices:
            return _reason("HISTORICAL_PRICE_NOT_INCLUDED", source, historical=True, human_confirmed=True, audit=audit)
        if source.get("quantity_trusted") is not True:
            return _reason("QUANTITY_UNTRUSTED", source, historical=True, human_confirmed=True, audit=audit)
        quantity = _bounded_decimal(source.get("quantity"))
        if quantity is None or quantity <= 0:
            return _reason("QUANTITY_UNTRUSTED", source, historical=True, human_confirmed=True, audit=audit)
        source_basis = source.get("unit_basis")
        reparsed_source = parse_unit_basis(source.get("raw_unit"))
        if not isinstance(source_basis, dict) or source_basis.get("trusted") is not True or reparsed_source.get("trusted") is not True:
            return _reason("SOURCE_UNIT_UNTRUSTED", source, historical=True, human_confirmed=True, audit=audit)
        basis_fields = ("raw_unit", "base_unit", "dimension", "scale", "conversion_basis", "trusted")
        if any(source_basis.get(key) != reparsed_source.get(key) for key in basis_fields):
            return _reason("UNIT_BASIS_UNPROVEN", source, historical=True, human_confirmed=True, audit=audit)
        if reparsed_source.get("dimension") in _NON_CONVERTIBLE_UNIT_DIMENSIONS:
            return _reason("SOURCE_UNIT_UNTRUSTED", source, historical=True, human_confirmed=True, audit=audit)
        offer_basis = parse_unit_basis(offer.get("price_unit"))
        if (
            offer_basis.get("trusted") is not True
            or offer_basis.get("dimension") in _NON_CONVERTIBLE_UNIT_DIMENSIONS
            or not normalize_unit_family(offer.get("price_unit"))
            or provenance.get("unit_family") != normalize_unit_family(offer.get("price_unit"))
        ):
            return _reason("OFFER_UNIT_UNTRUSTED", source, historical=True, human_confirmed=True, audit=audit)
        if reparsed_source.get("dimension") != offer_basis.get("dimension") or reparsed_source.get("base_unit") != offer_basis.get("base_unit"):
            return _reason("UNIT_INCOMPATIBLE", source, historical=True, human_confirmed=True, audit=audit)
        try:
            source_scale = Decimal(str(reparsed_source.get("scale")))
            offer_scale = Decimal(str(offer_basis.get("scale")))
            if source_scale <= 0 or offer_scale <= 0:
                raise InvalidOperation
            with localcontext() as context:
                context.prec = 50
                exported_price = (price * source_scale / offer_scale).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
                total = (exported_price * quantity).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, TypeError, ValueError):
            return _reason("PRICE_INVALID", source, historical=True, human_confirmed=True, audit=audit)
        audit.update({"source_unit": str(source.get("raw_unit") or ""), "offer_unit": str(offer.get("price_unit") or ""), "source_scale": str(source_scale), "offer_scale": str(offer_scale)})
        return TenderPriceDecision(
            source_row_id=str(source.get("source_row_id") or ""), excel_row=int(source.get("excel_row") or 0),
            eligible=True, reason_code=None, source_unit_price=exported_price, total_price=total,
            source_kind="one_c_history", historical=True, audit_summary=audit, human_confirmed=True,
        )

    def resolve_run(
        self,
        workspace: dict[str, Any],
        run: dict[str, Any],
        *,
        tender_id: str,
        run_id: str,
        include_historical_prices: bool,
        human_history_decisions: dict[str, dict[str, Any]] | None = None,
    ) -> list[TenderPriceDecision]:
        source_by_id, durable_by_id = self.validate_run(workspace, run, tender_id=tender_id, run_id=run_id)
        return [
            (
                self.resolve_human_history(
                    source_by_id[row_id], durable_by_id[row_id], run,
                    human_history_decisions[row_id], include_historical_prices=include_historical_prices,
                )
                if human_history_decisions and row_id in human_history_decisions
                else self.resolve(
                    source_by_id[row_id], durable_by_id[row_id], run,
                    include_historical_prices=include_historical_prices,
                )
            )
            for row_id in run["selected_source_row_ids"]
        ]


def summarize_decisions(decisions: list[TenderPriceDecision]) -> dict[str, Any]:
    reason_counts: dict[str, int] = {}
    for decision in decisions:
        if not decision.eligible and decision.reason_code:
            reason_counts[decision.reason_code] = reason_counts.get(decision.reason_code, 0) + 1
    eligible = [item for item in decisions if item.eligible]
    return {
        "selected_count": len(decisions),
        "priced_count": len(eligible),
        "historical_count": sum(1 for item in eligible if item.historical),
        "automatic_historical_count": sum(1 for item in eligible if item.historical and not item.human_confirmed),
        "human_confirmed_historical_count": sum(1 for item in eligible if item.historical and item.human_confirmed),
        "blank_count": len(decisions) - len(eligible),
        "reason_counts": dict(sorted(reason_counts.items())),
    }


def _canonical_xml(element: Any) -> Any:
    if element is None:
        return None
    if isinstance(element, list):
        return tuple(_canonical_xml(item) for item in element)
    tag = getattr(element, "tag", None)
    if tag is None:
        return str(element)
    attributes = tuple(sorted((str(key), str(value)) for key, value in element.attrib.items()))
    text = (element.text or "").strip()
    children = tuple(_canonical_xml(child) for child in list(element))
    return tag, attributes, text, children


def _dimension_style_fingerprint(
    sheet: Any,
    dimension: Any,
    row: int,
    column: int,
    *,
    default_style_is_unset: bool = False,
) -> str | None:
    style = getattr(dimension, "_style", None)
    # openpyxl can serialize a row or column with an explicit style index that
    # points to an all-zero/default XF when the source omitted that dimension
    # style. This is semantically identical to no dimension style; non-default
    # formatting remains part of the preservation comparison.
    if style is None or (default_style_is_unset and not any(style)):
        return None
    cell = sheet.cell(row, column)
    cell._style = copy(style)
    return _style_fingerprint(cell)


def _workbook_snapshot(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    try:
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            cells: dict[str, dict[str, Any]] = {}
            for cell in sheet._cells.values():
                comment = None if cell.comment is None else {"text": cell.comment.text, "author": cell.comment.author}
                hyperlink = None if cell.hyperlink is None else {
                    "target": cell.hyperlink.target,
                    "location": cell.hyperlink.location,
                    "tooltip": cell.hyperlink.tooltip,
                }
                if cell.value is None and comment is None and hyperlink is None and not cell.has_style:
                    continue
                cells[cell.coordinate] = {
                    "value": _json_value(cell.value),
                    "data_type": cell.data_type,
                    "comment": comment,
                    "hyperlink": hyperlink,
                    "style": _style_fingerprint(cell),
                }
            rows = {
                str(index): {
                    "height": dimension.height,
                    "hidden": bool(dimension.hidden),
                    "outlineLevel": int(dimension.outlineLevel or 0),
                    "collapsed": bool(dimension.collapsed),
                    "style": _dimension_style_fingerprint(
                        sheet, dimension, index, 1, default_style_is_unset=True,
                    ),
                    "thickTop": bool(dimension.thickTop),
                    "thickBottom": bool(dimension.thickBot),
                }
                for index, dimension in sheet.row_dimensions.items()
            }
            columns: dict[int, dict[str, Any]] = {}
            for dimension in sheet.column_dimensions.values():
                minimum = int(dimension.min or 0)
                maximum = int(dimension.max or minimum)
                if minimum <= 0:
                    continue
                state = {
                    "hidden": bool(dimension.hidden),
                    "width": dimension.width,
                    "bestFit": bool(dimension.bestFit),
                    "outlineLevel": int(dimension.outlineLevel or 0),
                    "collapsed": bool(dimension.collapsed),
                    "style": _dimension_style_fingerprint(
                        sheet, dimension, 1, minimum, default_style_is_unset=True,
                    ),
                    "customWidth": bool(dimension.customWidth),
                }
                for column in range(minimum, maximum + 1):
                    columns[column] = dict(state)
            tables = []
            for table in sheet.tables.values():
                tables.append({
                    "name": table.name,
                    "displayName": table.displayName,
                    "ref": table.ref,
                    "columns": [
                        {"id": column.id, "name": column.name, "totalsRowLabel": column.totalsRowLabel, "totalsRowFunction": column.totalsRowFunction}
                        for column in (table.tableColumns or [])
                    ],
                    "style": _canonical_xml(table.tableStyleInfo.to_tree() if table.tableStyleInfo else None),
                    "autoFilter": _canonical_xml(table.autoFilter.to_tree() if table.autoFilter else None),
                })
            sheets.append({
                "name": sheet.title,
                "state": sheet.sheet_state,
                "cells": cells,
                "defined_names": sorted(
                    [
                        {
                            "name": name,
                            "attr_text": str(defined.attr_text or ""),
                            "type": str(defined.type or ""),
                        }
                        for name, defined in sheet.defined_names.items()
                    ],
                    key=lambda item: item["name"],
                ),
                "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
                "tables": sorted(tables, key=lambda item: (item["name"], item["ref"])),
                "auto_filter": _canonical_xml(sheet.auto_filter.to_tree()),
                "data_validations": _canonical_xml(sheet.data_validations.to_tree()),
                "conditional_formatting": [
                    {
                        "sqref": str(formatting.sqref),
                        "rules": [_canonical_xml(rule.to_tree()) for rule in rules],
                    }
                    for formatting, rules in sheet.conditional_formatting._cf_rules.items()
                ],
                "protection": _canonical_xml(sheet.protection.to_tree()),
                "row_dimensions": rows,
                "column_dimensions": columns,
                "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                "print_area": str(sheet.print_area) if sheet.print_area else None,
                "print_title_rows": sheet.print_title_rows,
                "print_title_cols": sheet.print_title_cols,
                "page_settings": {
                    "page_setup": _canonical_xml(sheet.page_setup.to_tree()),
                    "page_margins": _canonical_xml(sheet.page_margins.to_tree()),
                    "print_options": _canonical_xml(sheet.print_options.to_tree()),
                    "sheet_properties": _canonical_xml(sheet.sheet_properties.to_tree()),
                    "sheet_format": _canonical_xml(sheet.sheet_format.to_tree()),
                    "views": _canonical_xml(sheet.views.to_tree()),
                },
            })
        names = [
            {
                "scope": "workbook",
                "name": name,
                "attr_text": str(defined.attr_text or ""),
                "type": str(defined.type or ""),
            }
            for name, defined in workbook.defined_names.items()
        ]
        names.extend(
            {
                "scope": sheet.title,
                "name": name,
                "attr_text": str(defined.attr_text or ""),
                "type": str(defined.type or ""),
            }
            for sheet in workbook.worksheets
            for name, defined in sheet.defined_names.items()
        )
        return {
            "sheet_order": list(workbook.sheetnames),
            "sheets": sheets,
            "defined_names": sorted(names, key=lambda item: (item["scope"], item["name"])),
            "calculation": _canonical_xml(workbook.calculation.to_tree()) if workbook.calculation else None,
        }
    finally:
        workbook.close()


def _package_snapshot(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = {item.filename.casefold(): item.filename for item in archive.infolist()}
        content_root = ElementTree.fromstring(archive.read(names["[content_types].xml"]))
        defaults = []
        overrides = []
        for child in list(content_root):
            if child.tag.endswith("Default"):
                defaults.append((child.attrib.get("Extension", "").casefold(), child.attrib.get("ContentType", "")))
            elif child.tag.endswith("Override"):
                overrides.append((child.attrib.get("PartName", "").casefold(), child.attrib.get("ContentType", "")))
        relationships: dict[str, list[tuple[str, str, str, str]]] = {}
        for lower_name, actual_name in names.items():
            if not lower_name.endswith(".rels"):
                continue
            root = ElementTree.fromstring(archive.read(actual_name))
            values = []
            identifiers: set[str] = set()
            rel_folder = posixpath.dirname(lower_name)
            base_folder = rel_folder[:-len("/_rels")] if rel_folder.endswith("/_rels") else ("" if rel_folder == "_rels" else rel_folder)
            for node in root.iter():
                if node.tag.endswith("Relationship"):
                    identifier = str(node.attrib.get("Id", ""))
                    target = str(node.attrib.get("Target", ""))
                    relationship_type = str(node.attrib.get("Type", ""))
                    target_mode = str(node.attrib.get("TargetMode", ""))
                    if (
                        not identifier
                        or identifier in identifiers
                        or not target
                        or not relationship_type
                        or target_mode.casefold() not in {"", "external"}
                    ):
                        raise TenderWorkspaceError(
                            "Связь компонента XLSX имеет неверный формат.",
                            409,
                            "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
                        )
                    identifiers.add(identifier)
                    if target_mode.casefold() != "external":
                        target = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base_folder, target)).casefold()
                        if target.startswith("../") or target not in names:
                            raise TenderWorkspaceError(
                                "Связь XLSX указывает на отсутствующий компонент.",
                                409,
                                "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
                            )
                    values.append((
                        identifier,
                        relationship_type,
                        target,
                        target_mode.casefold(),
                    ))
            relationships[lower_name] = sorted(values)
        relationship_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        for lower_name, actual_name in names.items():
            if not lower_name.endswith(".xml") or lower_name.endswith(".rels"):
                continue
            root = ElementTree.fromstring(archive.read(actual_name))
            if lower_name == "[content_types].xml":
                continue
            if lower_name == "_rels/.rels":
                continue
            folder, filename = posixpath.split(lower_name)
            rel_part = f"{folder}/_rels/{filename}.rels" if folder else f"_rels/{filename}.rels"
            rel_ids = {
                relationship[0]
                for relationship in relationships.get(rel_part, [])
            }
            for element in root.iter():
                for attribute, value in element.attrib.items():
                    if attribute.startswith(relationship_namespace) and value not in rel_ids:
                        raise TenderWorkspaceError(
                            "Ссылка XML на relationship XLSX не имеет соответствия.",
                            409,
                            "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
                        )
        effective_types: dict[str, str | None] = {}
        override_types: dict[str, str] = {}
        default_types: dict[str, str] = {}
        for extension, content_type in defaults:
            previous = default_types.get(extension)
            if previous is not None:
                raise TenderWorkspaceError(
                    "Типы компонентов XLSX заданы неоднозначно.",
                    409,
                    "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
                )
            default_types[extension] = content_type
        for part_name, content_type in overrides:
            previous = override_types.get(part_name)
            if previous is not None:
                raise TenderWorkspaceError(
                    "Типы компонентов XLSX заданы неоднозначно.",
                    409,
                    "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
                )
            override_types[part_name] = content_type
        for lower_name in names:
            if lower_name == "[content_types].xml":
                continue
            extension = lower_name.rsplit("/", 1)[-1].rsplit(".", 1)
            extension_key = extension[-1].casefold() if len(extension) == 2 else ""
            effective_types[lower_name] = override_types.get(
                f"/{lower_name}", default_types.get(extension_key),
            )
        table_semantics = {
            name: _canonical_xml(ElementTree.fromstring(archive.read(actual_name)))
            for name, actual_name in names.items()
            if re.fullmatch(r"xl/tables/table\d+\.xml", name)
        }
        return {
            "names": names,
            "defaults": sorted(defaults),
            "overrides": sorted(overrides),
            "effective_types": effective_types,
            "relationships": relationships,
            "table_semantics": table_semantics,
            "opaque_hashes": {key: hashlib.sha256(archive.read(actual)).hexdigest() for key, actual in names.items()},
        }


def _restore_source_owned_package_metadata(source: Path, output: Path) -> None:
    """Restore only validated source extension subtrees and original docProps parts."""
    source_extensions, unsupported = _package_preservation_inventory(source)
    if unsupported:
        raise TenderWorkspaceError(
            "Книга содержит неподдерживаемые объекты сохранения.",
            409,
            "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
        )
    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output) as output_zip:
        source_names = {item.filename.casefold(): item.filename for item in source_zip.infolist()}
        output_names = {item.filename.casefold(): item.filename for item in output_zip.infolist()}
        for part in ("docprops/app.xml", "docprops/core.xml"):
            source_part = source_names.get(part)
            if source_part is not None:
                if part not in output_names:
                    raise TenderWorkspaceError(
                        "Исходные метаданные XLSX не сохранены.", 409, "TENDER_EXPORT_PRESERVATION_FAILED",
                    )
                replacements[part] = source_zip.read(source_part)
        for part, extension_bytes in source_extensions.items():
            output_part = output_names.get(part)
            if output_part is None:
                raise TenderWorkspaceError(
                    "Компонент расширений XLSX не сохранён.", 409, "TENDER_EXPORT_PRESERVATION_FAILED",
                )
            output_root = ElementTree.fromstring(output_zip.read(output_part))
            for child in list(output_root):
                if child.tag.rsplit("}", 1)[-1] == "extLst":
                    output_root.remove(child)
            source_ext_list = ElementTree.fromstring(extension_bytes)
            output_root.append(source_ext_list)
            replacements[part] = ElementTree.tostring(
                output_root, encoding="utf-8", xml_declaration=True,
            )
    if not replacements:
        return
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".preserve-source-metadata.", suffix=".xlsx", dir=output.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(output, "r") as original, zipfile.ZipFile(temporary, "w") as rewritten:
            rewritten.comment = original.comment
            for info in original.infolist():
                payload = replacements.get(info.filename.casefold())
                if payload is None:
                    payload = original.read(info.filename)
                rewritten.writestr(info, payload)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def _worksheet_relationship_part_map(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            sheet_name: _worksheet_relationship_part(sheet_part)
            for sheet_name, sheet_part in _worksheet_part_map(archive).items()
        }


def _worksheet_relationship_part(sheet_part: str) -> str:
    folder, filename = posixpath.split(sheet_part)
    return f"{folder}/_rels/{filename}.rels".casefold()


def _sheet_related_parts(package: dict[str, Any], relationship_part: str, relationship_suffix: str) -> set[str]:
    return {
        target
        for _identifier, relationship_type, target, target_mode in package["relationships"].get(relationship_part, [])
        if relationship_type.rstrip("/").casefold().endswith(f"/{relationship_suffix.casefold()}")
        and not target_mode
    }


def _relationship_edges(values: list[tuple[str, str, str, str]]) -> list[tuple[str, str, str]]:
    """Normalize package relationships, whose IDs may be reassigned on save."""
    return sorted((relationship_type, target, target_mode) for _identifier, relationship_type, target, target_mode in values)


def _xml_without_extension_lists(path: Path, part_name: str) -> Any:
    with zipfile.ZipFile(path) as archive:
        names = {item.filename.casefold(): item.filename for item in archive.infolist()}
        actual_name = names.get(part_name.casefold())
        if actual_name is None:
            return None
        root = ElementTree.fromstring(archive.read(actual_name))
        for child in list(root):
            if child.tag.rsplit("}", 1)[-1] == "extLst":
                root.remove(child)
        return _canonical_xml(root)


def _cell_values_by_sheet(snapshot: dict[str, Any]) -> dict[str, dict[str, tuple[Any, str]]]:
    return {
        sheet["name"]: {
            coordinate: (cell["value"], cell["data_type"])
            for coordinate, cell in sheet["cells"].items()
        }
        for sheet in snapshot["sheets"]
    }


def _validate_shared_string_conversion(
    before: dict[str, Any],
    after: dict[str, Any],
    *,
    allowed_cells: set[str],
    target_sheet: str | None,
) -> None:
    old_sheets = _cell_values_by_sheet(before)
    new_sheets = _cell_values_by_sheet(after)
    if set(old_sheets) != set(new_sheets):
        raise TenderWorkspaceError(
            "Общая таблица строк XLSX изменила исходные значения.",
            409,
            "TENDER_EXPORT_PRESERVATION_FAILED",
        )
    for sheet_name, old_cells in old_sheets.items():
        for coordinate, old_value in old_cells.items():
            if sheet_name == target_sheet and coordinate in allowed_cells:
                continue
            if new_sheets[sheet_name].get(coordinate) != old_value:
                raise TenderWorkspaceError(
                    "Общая таблица строк XLSX изменила исходные значения.",
                    409,
                    "TENDER_EXPORT_PRESERVATION_FAILED",
                )


def _generated_docprops_are_safe(path: Path, added_parts: set[str], source_snapshot: dict[str, Any]) -> bool:
    if not added_parts:
        return True
    structural_names = {str(sheet["name"]) for sheet in source_snapshot["sheets"]}
    for sheet in source_snapshot["sheets"]:
        structural_names.update(str(table["name"]) for table in sheet["tables"])
    source_values = {
        str(cell["value"])
        for sheet in source_snapshot["sheets"]
        for cell in sheet["cells"].values()
        if isinstance(cell["value"], str) and cell["value"]
    } - structural_names
    app_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    core_namespace = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
    dc_namespace = "http://purl.org/dc/elements/1.1/"
    dcterms_namespace = "http://purl.org/dc/terms/"
    vtypes_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    xsi_type = "{http://www.w3.org/2001/XMLSchema-instance}type"
    allowed_app_fields = {
        "Application", "DocSecurity", "ScaleCrop", "HeadingPairs", "TitlesOfParts",
        "Manager", "Company", "LinksUpToDate", "SharedDoc", "HyperlinksChanged", "AppVersion",
    }
    allowed_variant_fields = {
        "variant", "vector", "lpstr", "lpwstr", "i1", "i2", "i4", "i8",
        "ui1", "ui2", "ui4", "ui8", "r4", "r8", "decimal", "bool",
        "filetime", "blob", "oblob", "empty", "null", "cy", "error",
    }
    with zipfile.ZipFile(path) as archive:
        names = {item.filename.casefold(): item.filename for item in archive.infolist()}
        for part in added_parts:
            actual_name = names.get(part)
            if actual_name is None:
                return False
            root = ElementTree.fromstring(archive.read(actual_name))
            if part == "docprops/app.xml":
                if root.tag != f"{{{app_namespace}}}Properties" or root.attrib:
                    return False
                parents = {child: parent for parent in root.iter() for child in list(parent)}
                for element in root.iter():
                    if element is root:
                        continue
                    namespace, _, local = element.tag[1:].partition("}") if element.tag.startswith("{") else ("", "", element.tag)
                    parent = parents.get(element)
                    parent_local = parent.tag.rsplit("}", 1)[-1] if parent is not None else ""
                    if namespace == app_namespace:
                        if local not in allowed_app_fields or parent is not root or element.attrib:
                            return False
                    elif namespace == vtypes_namespace:
                        if local not in allowed_variant_fields:
                            return False
                        if local == "vector":
                            if parent_local not in {"HeadingPairs", "TitlesOfParts"} or set(element.attrib) != {"size", "baseType"}:
                                return False
                            try:
                                size = int(element.attrib["size"])
                            except (TypeError, ValueError):
                                return False
                            if size < 0 or size > 1000 or element.attrib["baseType"] not in allowed_variant_fields:
                                return False
                        elif element.attrib:
                            return False
                        if local == "variant" and parent_local != "vector":
                            return False
                        if local in {"vector", "variant"} and (element.text or "").strip():
                            return False
                        if parent_local == "HeadingPairs" and local != "vector":
                            return False
                        if parent_local == "TitlesOfParts" and local != "vector":
                            return False
                        if parent_local == "vector" and local not in {"variant", "lpstr", "lpwstr", "i1", "i2", "i4", "i8", "ui1", "ui2", "ui4", "ui8", "r4", "r8", "decimal", "bool", "filetime", "blob", "oblob", "empty", "null", "cy", "error"}:
                            return False
                        if parent_local == "variant" and local not in {"lpstr", "lpwstr", "i1", "i2", "i4", "i8", "ui1", "ui2", "ui4", "ui8", "r4", "r8", "decimal", "bool", "filetime", "blob", "oblob", "empty", "null", "cy", "error"}:
                            return False
                    else:
                        return False
            elif part == "docprops/core.xml":
                if root.tag != f"{{{core_namespace}}}coreProperties" or root.attrib:
                    return False
                core_fields = {
                    (core_namespace, name)
                    for name in {"lastModifiedBy", "lastPrinted", "revision", "version", "category", "contentStatus", "identifier", "language", "keywords"}
                } | {
                    (dc_namespace, name)
                    for name in {"creator", "title", "subject", "description"}
                } | {
                    (dcterms_namespace, name)
                    for name in {"created", "modified"}
                }
                for element in list(root):
                    namespace, _, local = element.tag[1:].partition("}") if element.tag.startswith("{") else ("", "", element.tag)
                    if (namespace, local) not in core_fields or list(element):
                        return False
                    if (namespace, local) in {(dcterms_namespace, "created"), (dcterms_namespace, "modified")}:
                        if element.attrib and (set(element.attrib) != {xsi_type} or element.attrib[xsi_type] != "dcterms:W3CDTF"):
                            return False
                    elif element.attrib:
                        return False
            else:
                return False
            for element in root.iter():
                value = (element.text or "").strip()
                if value in source_values:
                    return False
    return True


def _worksheet_part_map(archive: zipfile.ZipFile) -> dict[str, str]:
    names = {item.filename.casefold(): item.filename for item in archive.infolist()}
    workbook_root = ElementTree.fromstring(archive.read(names["xl/workbook.xml"]))
    rels_root = ElementTree.fromstring(archive.read(names["xl/_rels/workbook.xml.rels"]))
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rels = {item.attrib.get("Id"): item.attrib.get("Target", "") for item in rels_root.findall(f"{{{package_ns}}}Relationship")}
    result: dict[str, str] = {}
    for sheet in workbook_root.iter():
        if not str(sheet.tag).endswith("}sheet"):
            continue
        title = str(sheet.attrib.get("name") or "")
        relationship_id = sheet.attrib.get(f"{{{rel_ns}}}id")
        target = rels.get(relationship_id, "")
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        actual = names.get(part.casefold())
        if not title or not actual or title in result:
            raise TenderWorkspaceError("Не удалось безопасно сопоставить листы XLSX.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
        result[title] = actual
    return result


def _cell_xml_pattern(coordinate: str) -> re.Pattern[bytes]:
    escaped = re.escape(coordinate.encode("ascii"))
    return re.compile(rb"<c\b(?=[^>]*\br=[\"']" + escaped + rb"[\"'])[^>]*(?:/>|>.*?</c>)", re.DOTALL)


def _row_xml_pattern(row_number: int) -> re.Pattern[bytes]:
    return re.compile(rb"<row\b(?=[^>]*\br=[\"']" + str(row_number).encode("ascii") + rb"[\"'])[^>]*(?:/>|>.*?</row>)", re.DOTALL)


def _restore_row_cells(output_xml: bytes, row_number: int, snippets: dict[str, bytes], source_xml: bytes) -> bytes:
    row_pattern = _row_xml_pattern(row_number)
    found = row_pattern.search(output_xml)
    if found:
        row_xml = found.group(0)
        if row_xml.endswith(b"/>"):
            opening, body, closing = row_xml[:-2] + b">", b"", b"</row>"
        else:
            open_end = row_xml.find(b">")
            close_start = row_xml.rfind(b"</row>")
            if close_start < 0:
                raise TenderWorkspaceError("Строка XLSX не прошла проверку сохранения.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
            opening, body, closing = row_xml[:open_end + 1], row_xml[open_end + 1:close_start], row_xml[close_start:]
        cells = {}
        for match in re.finditer(rb"<c\b[^>]*(?:/>|>.*?</c>)", body, re.DOTALL):
            cell_xml = match.group(0)
            coordinate_match = re.search(rb"\br=[\"']([A-Z]+[0-9]+)[\"']", cell_xml[:cell_xml.find(b">") + 1])
            if not coordinate_match:
                raise TenderWorkspaceError("Ячейка XLSX не прошла проверку сохранения.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
            cells[coordinate_match.group(1).decode("ascii")] = cell_xml
        if re.sub(rb"<c\b[^>]*(?:/>|>.*?</c>)", b"", body, flags=re.DOTALL).strip():
            raise TenderWorkspaceError("Строка XLSX содержит неподдерживаемые узлы.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
        cells.update(snippets)
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        ordered = sorted(
            cells.items(),
            key=lambda item: (coordinate_from_string(item[0])[1], column_index_from_string(coordinate_from_string(item[0])[0])),
        )
        return output_xml[:found.start()] + opening + b"".join(cell for _coordinate, cell in ordered) + closing + output_xml[found.end():]

    source_row = _row_xml_pattern(row_number).search(source_xml)
    if source_row is None:
        raise TenderWorkspaceError("Исходная строка XLSX не найдена.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
    source_row_xml = source_row.group(0)
    open_end = source_row_xml.find(b">")
    opening = source_row_xml[:open_end + 1]
    if opening.endswith(b"/>"):
        opening = opening[:-2] + b">"
    sheet_data = re.search(rb"<sheetData\b[^>]*>(.*?)</sheetData>", output_xml, re.DOTALL)
    if sheet_data is None:
        raise TenderWorkspaceError("Данные листа XLSX не найдены.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
    row_matches = list(re.finditer(rb"<row\b(?=[^>]*\br=[\"']([0-9]+)[\"'])[^>]*(?:/>|>.*?</row>)", sheet_data.group(1), re.DOTALL))
    insert_at = len(sheet_data.group(1))
    for match in row_matches:
        row_match = re.search(rb"\br=[\"']([0-9]+)[\"']", match.group(0)[:match.group(0).find(b">") + 1])
        if row_match and int(row_match.group(1)) > row_number:
            insert_at = match.start()
            break
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    ordered = sorted(
        snippets.items(),
        key=lambda item: (coordinate_from_string(item[0])[1], column_index_from_string(coordinate_from_string(item[0])[0])),
    )
    restored_row = opening + b"".join(cell for _coordinate, cell in ordered) + b"</row>"
    body = sheet_data.group(1)
    new_body = body[:insert_at] + restored_row + body[insert_at:]
    return output_xml[:sheet_data.start(1)] + new_body + output_xml[sheet_data.end(1):]


def _restore_empty_shared_string_cells(source_path: Path, output_path: Path, allowed_cells: set[str], target_sheet: str) -> None:
    source_book = load_workbook(source_path, data_only=False, read_only=False, keep_links=False)
    try:
        empty_by_sheet = {
            sheet.title: {
                cell.coordinate for cell in sheet._cells.values()
                if cell.value == "" and cell.data_type == "s" and cell.comment is None
            }
            for sheet in source_book.worksheets
        }
    finally:
        source_book.close()

    fd, temporary_name = tempfile.mkstemp(prefix=".restore-empty-cells.", suffix=".xlsx", dir=output_path.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(source_path, "r") as source_zip, zipfile.ZipFile(output_path, "r") as output_zip:
            source_names = {item.filename.casefold(): item.filename for item in source_zip.infolist()}
            output_infos = {item.filename: item for item in output_zip.infolist()}
            output_data = {item.filename: output_zip.read(item.filename) for item in output_zip.infolist()}
            source_sheet_parts = _worksheet_part_map(source_zip)
            output_sheet_parts = _worksheet_part_map(output_zip)
            for title, coordinates in empty_by_sheet.items():
                if not coordinates:
                    continue
                source_xml = source_zip.read(source_sheet_parts[title])
                output_part = output_sheet_parts[title]
                output_xml = output_data[output_part]
                restore: dict[int, dict[str, bytes]] = {}
                for coordinate in coordinates:
                    if title == target_sheet and coordinate in allowed_cells:
                        continue
                    source_cell = _cell_xml_pattern(coordinate).search(source_xml)
                    if source_cell is None:
                        raise TenderWorkspaceError("Пустая исходная ячейка XLSX не найдена.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
                    restore.setdefault(int(re.search(r"[0-9]+$", coordinate).group()), {})[coordinate] = source_cell.group(0)
                for row_number, snippets in restore.items():
                    output_xml = _restore_row_cells(output_xml, row_number, snippets, source_xml)
                output_data[output_part] = output_xml

            shared_names = [name for name in source_names if "sharedstrings.xml" == name.rsplit("/", 1)[-1]]
            if any(empty_by_sheet.values()) and shared_names:
                shared_source_name = source_names[shared_names[0]]
                if shared_source_name not in output_data:
                    shared_info = next(item for item in source_zip.infolist() if item.filename == shared_source_name)
                    output_infos[shared_source_name] = shared_info
                    output_data[shared_source_name] = source_zip.read(shared_source_name)

                rel_name = next(name for name in output_data if name.casefold() == "xl/_rels/workbook.xml.rels")
                rel_root = ElementTree.fromstring(output_data[rel_name])
                existing_shared = [item for item in rel_root if item.attrib.get("Type", "").endswith("/sharedStrings")]
                if not existing_shared:
                    source_rels_name = source_names["xl/_rels/workbook.xml.rels"]
                    source_rels = ElementTree.fromstring(source_zip.read(source_rels_name))
                    source_shared = next((item for item in source_rels if item.attrib.get("Type", "").endswith("/sharedStrings")), None)
                    if source_shared is not None:
                        relation = copy(source_shared)
                        used_ids = {item.attrib.get("Id") for item in rel_root}
                        if relation.attrib.get("Id") in used_ids:
                            number = 1
                            while f"rId{number}" in used_ids:
                                number += 1
                            relation.attrib["Id"] = f"rId{number}"
                        rel_root.append(relation)
                        output_data[rel_name] = ElementTree.tostring(rel_root, encoding="utf-8", xml_declaration=True)

                types_name = next(name for name in output_data if name.casefold() == "[content_types].xml")
                types_root = ElementTree.fromstring(output_data[types_name])
                content_override = f"/{shared_source_name}"
                if not any(item.attrib.get("PartName", "").casefold() == content_override.casefold() for item in types_root):
                    source_types_name = source_names["[content_types].xml"]
                    source_types = ElementTree.fromstring(source_zip.read(source_types_name))
                    original_override = next((item for item in source_types if item.attrib.get("PartName", "").casefold() == content_override.casefold()), None)
                    if original_override is None:
                        raise TenderWorkspaceError("Тип общей таблицы строк XLSX не найден.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
                    types_root.append(copy(original_override))
                    output_data[types_name] = ElementTree.tostring(types_root, encoding="utf-8", xml_declaration=True)

        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as rewritten:
            for name, data in output_data.items():
                info = output_infos.get(name)
                if info is None:
                    rewritten.writestr(name, data)
                else:
                    rewritten.writestr(info, data)
        os.replace(temporary, output_path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _verify_package_roundtrip(
    source: Path,
    output: Path,
    *,
    target_sheet: str | None = None,
    allowed_cells: set[str] | None = None,
    target_comment_cells: set[str] | None = None,
    semantic_before: dict[str, Any] | None = None,
    semantic_after: dict[str, Any] | None = None,
) -> None:
    before = _package_snapshot(source)
    after = _package_snapshot(output)
    source_names = set(before["names"])
    output_names = set(after["names"])
    allowed_cells = allowed_cells or set()
    target_comment_cells = target_comment_cells or set()

    source_extensions, source_unsupported = _package_preservation_inventory(source)
    output_extensions, output_unsupported = _package_preservation_inventory(output)
    if source_unsupported or output_unsupported:
        raise TenderWorkspaceError(
            "Книга содержит неподдерживаемые объекты сохранения.",
            409,
            "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
        )
    if {
        name: _canonical_xml(ElementTree.fromstring(payload))
        for name, payload in source_extensions.items()
    } != {
        name: _canonical_xml(ElementTree.fromstring(payload))
        for name, payload in output_extensions.items()
    }:
        raise TenderWorkspaceError(
            "Расширения XLSX не были сохранены.", 409, "TENDER_EXPORT_PRESERVATION_FAILED",
        )
    if "xl/theme/theme1.xml" in source_extensions and (
        _xml_without_extension_lists(source, "xl/theme/theme1.xml")
        != _xml_without_extension_lists(output, "xl/theme/theme1.xml")
    ):
        raise TenderWorkspaceError(
            "Исходная тема XLSX изменилась за пределами безопасного расширения.",
            409,
            "TENDER_EXPORT_PRESERVATION_FAILED",
        )

    # Shared-string indices and inline strings are alternate encodings of the
    # same cells.  Verify the cell values first, including formula text and
    # empty shared-string placeholders restored by the exporter.
    removed_shared_strings = "xl/sharedstrings.xml" in source_names - output_names
    if removed_shared_strings:
        semantic_before = semantic_before or _workbook_snapshot(source)
        semantic_after = semantic_after or _workbook_snapshot(output)
        _validate_shared_string_conversion(
            semantic_before,
            semantic_after,
            allowed_cells=allowed_cells,
            target_sheet=target_sheet,
        )

    target_sheet_rel: str | None = None
    source_sheet_rels: dict[str, str] = {}
    output_sheet_rels: dict[str, str] = {}
    output_target_comment_parts: set[str] = set()
    output_target_vml_parts: set[str] = set()
    mutable_target_comment_parts: set[str] = set()
    if target_sheet:
        try:
            with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output) as output_zip:
                source_sheet_parts = _worksheet_part_map(source_zip)
                output_sheet_parts = _worksheet_part_map(output_zip)
            source_sheet_rels = {
                name: _worksheet_relationship_part(part)
                for name, part in source_sheet_parts.items()
            }
            output_sheet_rels = {
                name: _worksheet_relationship_part(part)
                for name, part in output_sheet_parts.items()
            }
            if source_sheet_parts != output_sheet_parts or target_sheet not in source_sheet_rels:
                raise TenderWorkspaceError(
                    "Связи листов XLSX изменились.",
                    409,
                    "TENDER_EXPORT_PRESERVATION_FAILED",
                )
            target_sheet_rel = source_sheet_rels[target_sheet]
            output_target_comment_parts = _sheet_related_parts(after, output_sheet_rels[target_sheet], "comments")
            output_target_vml_parts = _sheet_related_parts(after, output_sheet_rels[target_sheet], "vmlDrawing")
            non_target_comment_parts = set().union(*(
                _sheet_related_parts(after, relationship_part, "comments")
                | _sheet_related_parts(after, relationship_part, "vmlDrawing")
                for sheet_name, relationship_part in output_sheet_rels.items()
                if sheet_name != target_sheet
            )) if len(output_sheet_rels) > 1 else set()
            mutable_target_comment_parts = (
                (output_target_comment_parts | output_target_vml_parts) - non_target_comment_parts
            )
        except (KeyError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
            raise TenderWorkspaceError(
                "Не удалось проверить связи листа XLSX.",
                409,
                "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
            ) from exc

    permitted_added = {
        name
        for name in output_names - source_names
        if (
            target_comment_cells and (
                name in mutable_target_comment_parts
                or name == target_sheet_rel
            )
        ) or name in {"docprops/app.xml", "docprops/core.xml"}
    }
    permitted_removed = {name for name in source_names - output_names if name == "xl/sharedstrings.xml"}
    if output_names - source_names != permitted_added or source_names - output_names != permitted_removed:
        raise TenderWorkspaceError("Структура XLSX изменилась за пределами разрешённых выходных данных.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")

    generated_docprops = {
        name for name in permitted_added
        if name in {"docprops/app.xml", "docprops/core.xml"}
    }
    if not _generated_docprops_are_safe(output, generated_docprops, semantic_before or _workbook_snapshot(source)):
        raise TenderWorkspaceError(
            "В метаданных XLSX появились неожиданные данные.",
            409,
            "TENDER_EXPORT_PRESERVATION_FAILED",
        )

    mutable_parts = {
        "[content_types].xml", "xl/workbook.xml", "xl/styles.xml",
        "xl/sharedstrings.xml", "xl/_rels/workbook.xml.rels", "_rels/.rels",
    }
    if "xl/theme/theme1.xml" in source_extensions:
        mutable_parts.add("xl/theme/theme1.xml")
    mutable_parts.update(name for name in source_names if re.fullmatch(r"xl/tables/table\d+\.xml", name))
    for name in source_names & output_names:
        if (
            name in mutable_parts
            or re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)
            or re.fullmatch(r"xl/worksheets/_rels/sheet\d+\.xml\.rels", name)
            or name in mutable_target_comment_parts
        ):
            continue
        if before["opaque_hashes"][name] != after["opaque_hashes"][name]:
            raise TenderWorkspaceError("Компонент XLSX изменился за пределами разрешённых выходных данных.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")

    if before["table_semantics"] != after["table_semantics"]:
        raise TenderWorkspaceError(
            "Определение таблицы XLSX изменилось.",
            409,
            "TENDER_EXPORT_PRESERVATION_FAILED",
        )

    for name in source_names - {"[content_types].xml"}:
        if name not in output_names:
            if name == "xl/sharedstrings.xml" and removed_shared_strings:
                continue
            raise TenderWorkspaceError("Из XLSX исчез исходный компонент.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        if before["effective_types"].get(name) != after["effective_types"].get(name):
            raise TenderWorkspaceError("Изменился тип исходного компонента XLSX.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
    for name in permitted_added:
        expected_type = None
        if name == "docprops/app.xml":
            expected_type = "application/vnd.openxmlformats-officedocument.extended-properties+xml"
        elif name == "docprops/core.xml":
            expected_type = "application/vnd.openxmlformats-package.core-properties+xml"
        elif re.fullmatch(r"xl/comments/comment\d+\.xml", name):
            expected_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"
        elif re.fullmatch(r"xl/drawings/commentsdrawing\d+\.vml", name):
            expected_type = "application/vnd.openxmlformats-officedocument.vmlDrawing"
        elif name == target_sheet_rel:
            expected_type = "application/vnd.openxmlformats-package.relationships+xml"
        if expected_type is None or after["effective_types"].get(name) != expected_type:
            raise TenderWorkspaceError("Появился компонент XLSX с неожиданным типом.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
    for package in (before, after):
        for part, content_type in package["effective_types"].items():
            if content_type is None:
                raise TenderWorkspaceError(
                    "Для компонента XLSX не определён тип содержимого.",
                    409,
                    "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
                )
        package_names = set(package["names"])
        if any(part_name.lstrip("/") not in package_names for part_name, _kind in package["overrides"]):
            raise TenderWorkspaceError(
                "Тип содержимого ссылается на отсутствующий компонент XLSX.",
                409,
                "TENDER_EXPORT_PRESERVATION_UNSUPPORTED",
            )

    # Existing document metadata belongs to the source workbook and is restored
    # byte-for-byte before this check. Packages without docProps may retain the
    # validated generated-props policy above.
    for name in {"docprops/app.xml", "docprops/core.xml"} & source_names:
        if before["opaque_hashes"].get(name) != after["opaque_hashes"].get(name):
            raise TenderWorkspaceError("Изменились исходные метаданные XLSX.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")

    # New historical comments are allowed only on the exact target cells, and
    # their relationship parts must remain attached to the selected sheet.
    if target_comment_cells:
        if (
            not target_sheet
            or len(output_target_comment_parts) != 1
            or len(output_target_vml_parts) != 1
            or any(not re.fullmatch(r"xl/comments/comment\d+\.xml", name) for name in output_target_comment_parts)
            or any(not re.fullmatch(r"xl/drawings/commentsdrawing\d+\.vml", name) for name in output_target_vml_parts)
        ):
            raise TenderWorkspaceError("Комментарии цены не связаны с целевым листом.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        if semantic_before is None:
            semantic_before = _workbook_snapshot(source)
        if semantic_after is None:
            semantic_after = _workbook_snapshot(output)
        original_sheets = {sheet["name"]: sheet for sheet in semantic_before["sheets"]}
        output_sheets = {sheet["name"]: sheet for sheet in semantic_after["sheets"]}
        if target_sheet not in output_sheets:
            raise TenderWorkspaceError("Целевой лист комментариев не найден.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        original_cells = original_sheets[target_sheet]["cells"]
        output_cells = output_sheets[target_sheet]["cells"]
        if any(
            coordinate not in output_cells
            or output_cells[coordinate].get("comment") is None
            for coordinate in target_comment_cells
        ):
            raise TenderWorkspaceError("Исторические комментарии XLSX не сохранены.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        if any(
            cell.get("comment") is not None
            and cell.get("comment") != original_cells.get(coordinate, {}).get("comment")
            and coordinate not in target_comment_cells
            for coordinate, cell in output_cells.items()
        ):
            raise TenderWorkspaceError("В XLSX появился посторонний комментарий.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        if any(
            cell.get("comment") is not None
            and (
                coordinate not in output_cells
                or output_cells[coordinate].get("comment") != cell.get("comment")
            )
            and coordinate not in target_comment_cells
            for coordinate, cell in original_cells.items()
        ):
            raise TenderWorkspaceError("Исходный комментарий XLSX не сохранён.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        for sheet_name, original_sheet in original_sheets.items():
            if sheet_name == target_sheet:
                continue
            current_sheet = output_sheets.get(sheet_name)
            if current_sheet is None:
                raise TenderWorkspaceError("Исходный лист комментариев XLSX не найден.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
            original_comments = {
                coordinate: cell["comment"]
                for coordinate, cell in original_sheet["cells"].items()
                if cell.get("comment") is not None
            }
            current_comments = {
                coordinate: cell["comment"]
                for coordinate, cell in current_sheet["cells"].items()
                if cell.get("comment") is not None
            }
            if original_comments != current_comments:
                raise TenderWorkspaceError("Комментарий на постороннем листе XLSX изменился.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")

    for part, original in before["relationships"].items():
        current = after["relationships"].get(part)
        if current is None:
            raise TenderWorkspaceError("Связи XLSX изменились за пределами разрешённых выходных данных.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        original_edges = _relationship_edges(original)
        current_edges = _relationship_edges(current)
        if part == "xl/_rels/workbook.xml.rels" and removed_shared_strings:
            shared_edge = next((
                edge for edge in original_edges
                if edge[0].endswith("/sharedStrings") and edge[1].casefold() == "xl/sharedstrings.xml" and edge[2] == ""
            ), None)
            if shared_edge is None:
                raise TenderWorkspaceError("Связь общей таблицы строк XLSX повреждена.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
            original_edges.remove(shared_edge)
        expected_edges = list(original_edges)
        if part == "_rels/.rels":
            if "docprops/core.xml" in generated_docprops:
                expected_edges.append((
                    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
                    after["names"]["docprops/core.xml"].casefold(), "",
                ))
            if "docprops/app.xml" in generated_docprops:
                expected_edges.append((
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
                    after["names"]["docprops/app.xml"].casefold(), "",
                ))
        if part == target_sheet_rel and target_comment_cells:
            expected_edges = [
                edge for edge in original_edges
                if not edge[0].casefold().endswith(("/comments", "/vmldrawing"))
            ]
            output_comment_edges = [
                edge for edge in current_edges
                if edge[0].casefold().endswith(("/comments", "/vmldrawing"))
            ]
            for relationship_type, target, target_mode in output_comment_edges:
                if target_mode or (
                    relationship_type.casefold().endswith("/comments")
                    and (target not in output_target_comment_parts or not re.fullmatch(r"xl/comments/comment\d+\.xml", target))
                ) or (
                    relationship_type.casefold().endswith("/vmldrawing")
                    and (target not in output_target_vml_parts or not re.fullmatch(r"xl/drawings/commentsdrawing\d+\.vml", target))
                ):
                    raise TenderWorkspaceError("Связь комментариев XLSX не соответствует целевому листу.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
            expected_edges.extend(output_comment_edges)
        if sorted(current_edges) != sorted(expected_edges):
            raise TenderWorkspaceError("Связи XLSX изменились за пределами разрешённых выходных данных.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
    for part, current in after["relationships"].items():
        if part in before["relationships"]:
            continue
        if part != target_sheet_rel or not target_comment_cells:
            raise TenderWorkspaceError("В XLSX появились неожиданные связи.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        edges = _relationship_edges(current)
        if (
            not edges
            or len(edges) != len(output_target_comment_parts) + len(output_target_vml_parts)
            or any(
                mode
                or not (
                    relationship_type.casefold().endswith("/comments")
                    and target in output_target_comment_parts
                    and re.fullmatch(r"xl/comments/comment\d+\.xml", target)
                    or relationship_type.casefold().endswith("/vmldrawing")
                    and target in output_target_vml_parts
                    and re.fullmatch(r"xl/drawings/commentsdrawing\d+\.vml", target)
                )
                for relationship_type, target, mode in edges
            )
        ):
            raise TenderWorkspaceError("В XLSX появились неожиданные связи.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")


def _verify_semantic_roundtrip(
    before: dict[str, Any],
    after: dict[str, Any],
    *,
    target_sheet: str,
    allowed_cells: set[str],
    target_columns: set[int],
) -> None:
    if before["sheet_order"] != after["sheet_order"] or before["defined_names"] != after["defined_names"] or before["calculation"] != after["calculation"]:
        raise TenderWorkspaceError("Настройки книги XLSX изменились за пределами разрешённых выходных данных.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
    before_by_name = {sheet["name"]: sheet for sheet in before["sheets"]}
    after_by_name = {sheet["name"]: sheet for sheet in after["sheets"]}
    if set(before_by_name) != set(after_by_name):
        raise TenderWorkspaceError("Состав листов XLSX изменился.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
    for name, original in before_by_name.items():
        current = after_by_name[name]
        if name != target_sheet:
            if _json_text(original) != _json_text(current):
                raise TenderWorkspaceError("Изменился лист за пределами целевого листа.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
            continue
        for field in (
            "state", "merged_ranges", "tables", "auto_filter", "data_validations",
            "conditional_formatting", "protection", "row_dimensions", "freeze_panes",
            "print_area", "print_title_rows", "print_title_cols", "page_settings",
        ):
            if original[field] != current[field]:
                raise TenderWorkspaceError("Структура листа XLSX изменилась за пределами целевых ячеек.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        old_cells = original["cells"]
        new_cells = current["cells"]
        if any(cell not in allowed_cells and new_cells.get(cell) != value for cell, value in old_cells.items()):
            raise TenderWorkspaceError("Исходные ячейки XLSX изменились.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        if any(cell not in old_cells and cell not in allowed_cells for cell in new_cells):
            raise TenderWorkspaceError("В XLSX появились посторонние ячейки.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        old_columns = original["column_dimensions"]
        new_columns = current["column_dimensions"]
        for column in set(old_columns) | set(new_columns):
            if column not in target_columns and old_columns.get(column) != new_columns.get(column):
                raise TenderWorkspaceError("Изменились настройки постороннего столбца.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")


def _copy_style(source_cell: Any, target_cell: Any) -> None:
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def _split_target_column_dimensions(sheet: Any, target_columns: set[int]) -> None:
    """Detach target columns from grouped dimensions while preserving neighbors."""
    for key, dimension in list(sheet.column_dimensions.items()):
        minimum = int(dimension.min or 0)
        maximum = int(dimension.max or minimum)
        selected = sorted(column for column in target_columns if minimum <= column <= maximum)
        if not selected:
            continue
        segments: list[tuple[int, int]] = []
        cursor = minimum
        for target in selected:
            if cursor < target:
                segments.append((cursor, target - 1))
            segments.append((target, target))
            cursor = target + 1
        if cursor <= maximum:
            segments.append((cursor, maximum))
        del sheet.column_dimensions[key]
        for start, end in segments:
            split = copy(dimension)
            split.min = start
            split.max = end
            split.index = get_column_letter(start)
            sheet.column_dimensions[split.index] = split


def _safe_output_filename(original: str) -> str:
    base = Path(str(original or "tender.xlsx").replace("\\", "/")).name
    stem = Path(base).stem
    stem = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", stem).strip(" ._")
    stem = re.sub(r"\s+", " ", stem)[:110] or "Тендер"
    return f"{stem}_Averon_цены.xlsx"


def _validate_source_and_manifest(source_path: Path, workspace: dict[str, Any], tender_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    if not source_path.is_file() or source_path.stat().st_size > MAX_UPLOAD_BYTES:
        raise TenderWorkspaceError("Исходная книга тендера недоступна.", 409, "TENDER_SOURCE_MISSING")
    source_sha = _sha256_file(source_path)
    if source_sha != workspace.get("source_sha256"):
        raise TenderWorkspaceError("Исходная книга тендера изменилась.", 409, "TENDER_SOURCE_CHANGED")
    mapping_override = None if workspace.get("table_name") else workspace.get("mapping")
    parsed = TenderWorkbookParser().parse(
        source_path,
        tender_id=tender_id,
        mapping_override=mapping_override,
        selected_sheet=str(workspace.get("sheet_name") or ""),
        selected_header_row=int(workspace.get("header_row") or 0),
    )
    def without_generated_ids(rows: Any) -> Any:
        if not isinstance(rows, list):
            return rows
        return [
            {key: value for key, value in row.items() if key != "source_row_id"}
            if isinstance(row, dict) else row
            for row in rows
        ]
    if (
        parsed.get("source_sha256") != source_sha
        or parsed.get("sheet_name") != workspace.get("sheet_name")
        or int(parsed.get("header_row") or 0) != int(workspace.get("header_row") or -1)
        or parsed.get("mapping") != workspace.get("mapping")
        or int(parsed.get("logical_right_edge") or 0) != int(workspace.get("logical_right_edge") or -1)
        or _json_text(parsed.get("manifest")) != _json_text(workspace.get("source_manifest"))
        or _json_text(without_generated_ids(parsed.get("rows"))) != _json_text(without_generated_ids(workspace.get("rows")))
    ):
        raise TenderWorkspaceError("Структура исходной книги не совпадает с подтверждённым тендером.", 409, "TENDER_SOURCE_MANIFEST_MISMATCH")
    return parsed, _workbook_snapshot(source_path)


def _check_target_collision(source_path: Path, workspace: dict[str, Any], target_sheet: str) -> tuple[int, int]:
    workbook = load_workbook(source_path, data_only=False, read_only=False, keep_links=False)
    try:
        if target_sheet not in workbook.sheetnames:
            raise TenderWorkspaceError("Целевой лист не найден.", 409, "TENDER_EXPORT_TARGET_INVALID")
        sheet = workbook[target_sheet]
        edge = int(workspace.get("logical_right_edge") or 0)
        first, second = edge + 1, edge + 2
        target_columns = {first, second}
        if first < 1 or second > 256:
            raise TenderWorkspaceError("Для цены и суммы нет допустимых выходных столбцов.", 409, "TENDER_EXPORT_TARGET_INVALID")

        def intersects_target(reference: Any) -> bool:
            ranges = getattr(reference, "ranges", None)
            values = list(ranges) if ranges is not None else [reference]
            for value in values:
                try:
                    min_col, _min_row, max_col, _max_row = range_boundaries(str(value))
                except (TypeError, ValueError):
                    return True
                if min_col is None or max_col is None or any(min_col <= column <= max_col for column in target_columns):
                    return True
            return False

        for column in (first, second):
            for cell in sheet._cells.values():
                if cell.column == column and not is_semantically_empty_cell(cell):
                    raise TenderWorkspaceError("В целевом столбце уже есть данные.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
            for merged in sheet.merged_cells.ranges:
                if merged.min_col <= column <= merged.max_col:
                    raise TenderWorkspaceError("Целевой столбец пересекает объединённый диапазон.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
        for table in sheet.tables.values():
            min_col, _min_row, max_col, _max_row = range_boundaries(table.ref)
            if any(min_col <= column <= max_col for column in (first, second)):
                raise TenderWorkspaceError("Целевой столбец пересекает таблицу XLSX.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
        if sheet.auto_filter.ref and intersects_target(sheet.auto_filter.ref):
            raise TenderWorkspaceError("Целевой столбец включён в существующий фильтр XLSX.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
        if any(intersects_target(validation.sqref) for validation in sheet.data_validations.dataValidation):
            raise TenderWorkspaceError("Целевой столбец участвует в проверке данных XLSX.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
        if any(intersects_target(formatting.sqref) for formatting in sheet.conditional_formatting):
            raise TenderWorkspaceError("Целевой столбец участвует в условном форматировании XLSX.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
        manifest = workspace.get("source_manifest") if isinstance(workspace.get("source_manifest"), dict) else {}
        if manifest.get("unsupported_preservation_sensitive_objects"):
            raise TenderWorkspaceError("Книга содержит неподдерживаемые объекты сохранения.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
        if sheet._charts or sheet._images:
            raise TenderWorkspaceError("Лист содержит объекты, для которых сохранение не подтверждено.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")
        return first, second
    finally:
        workbook.close()


class TenderXlsxPriceExporter:
    """Copy-only XLSX writer with semantic and package preservation checks."""

    def write(
        self,
        source_path: Path,
        output_path: Path,
        workspace: dict[str, Any],
        decisions: list[TenderPriceDecision],
    ) -> dict[str, Any]:
        _parsed, before = _validate_source_and_manifest(source_path, workspace, str(workspace["tender_id"]))
        target_sheet = str(workspace["sheet_name"])
        price_column, total_column = _check_target_collision(source_path, workspace, target_sheet)
        manifest = workspace.get("source_manifest") or {}
        if manifest.get("unsupported_preservation_sensitive_objects"):
            raise TenderWorkspaceError("Книга содержит неподдерживаемые объекты сохранения.", 409, "TENDER_EXPORT_PRESERVATION_UNSUPPORTED")

        normalized_path: Path | None = None
        workbook = None
        try:
            preflight_tender_xlsx(source_path)
            try:
                normalized_path = preflight_xlsx(source_path)
                workbook = load_workbook(normalized_path or source_path, data_only=False, read_only=False, keep_links=False)
            except OneCImportError as exc:
                raise TenderWorkspaceError("Исходная книга не прошла проверку безопасности.", 409, "TENDER_SOURCE_INVALID") from exc
            sheet = workbook[target_sheet]
            edge = int(workspace["logical_right_edge"])
            _split_target_column_dimensions(sheet, {edge + 1, edge + 2})
            row_by_id = {str(row.get("source_row_id") or ""): row for row in workspace["rows"] if isinstance(row, dict)}
            header_row = int(workspace["header_row"])
            style_source_header = sheet.cell(header_row, edge)
            price_header = sheet.cell(header_row, price_column)
            total_header = sheet.cell(header_row, total_column)
            _copy_style(style_source_header, price_header)
            _copy_style(style_source_header, total_header)
            price_header.value = "Цена за единицу"
            total_header.value = "Общая стоимость"
            price_header.number_format = "0.000000"
            total_header.number_format = "0.00"
            allowed_cells = {price_header.coordinate, total_header.coordinate}
            target_comment_cells: set[str] = set()
            historical_count = 0
            human_confirmed_history_count = 0
            for decision in decisions:
                if not decision.eligible:
                    continue
                source = row_by_id.get(decision.source_row_id)
                if source is None or source.get("row_type") != "item" or int(source.get("excel_row") or 0) != decision.excel_row:
                    raise TenderWorkspaceError("Строка цены не соответствует подтверждённому источнику.", 409, "TENDER_RESULT_CORRELATION_FAILED")
                row_index = decision.excel_row
                source_style = sheet.cell(row_index, edge)
                price_cell = sheet.cell(row_index, price_column)
                total_cell = sheet.cell(row_index, total_column)
                if not is_semantically_empty_cell(price_cell) or not is_semantically_empty_cell(total_cell):
                    raise TenderWorkspaceError("В целевой строке уже есть данные.", 409, "TENDER_EXPORT_TARGET_OCCUPIED")
                _copy_style(source_style, price_cell)
                _copy_style(source_style, total_cell)
                price_cell.value = decision.source_unit_price
                total_cell.value = decision.total_price
                price_cell.number_format = "0.000000"
                total_cell.number_format = "0.00"
                allowed_cells.update((price_cell.coordinate, total_cell.coordinate))
                if decision.historical:
                    historical_count += 1
                    human_confirmed_history_count += int(decision.human_confirmed)
                    price_cell.fill = _HISTORICAL_FILL
                    total_cell.fill = _HISTORICAL_FILL
                    purchase_date = str((decision.audit_summary or {}).get("purchase_date") or "")
                    if decision.human_confirmed and (decision.audit_summary or {}).get("human_confirmation_basis") == "MANUAL_SEARCH_CONFIRMATION":
                        human_note = " Совпадение позиции вручную найдено и подтверждено пользователем."
                    elif decision.human_confirmed and (decision.audit_summary or {}).get("human_confirmation_basis") == "FUZZY_MANUAL_CONFIRMATION":
                        human_note = " Совпадение позиции явно подтверждено пользователем."
                    else:
                        human_note = " Совпадение позиции подтверждено пользователем." if decision.human_confirmed else ""
                    price_cell.comment = Comment(
                        f"Историческая цена по предыдущей покупке 1С от {purchase_date}. "
                        f"Не подтверждает текущую доступность и не является текущим предложением.{human_note}",
                        "Averon Import",
                    )
                    total_cell.comment = Comment(
                        f"Историческая стоимость по предыдущей покупке 1С от {purchase_date}. "
                        f"Не подтверждает текущую доступность и не является текущим предложением.{human_note}",
                        "Averon Import",
                    )
                    target_comment_cells.update((price_cell.coordinate, total_cell.coordinate))
            if historical_count:
                price_header.comment = Comment(
                    "Исторические цены 1С отмечены цветом. Они отражают предыдущие покупки, "
                    "не подтверждают текущую доступность и не являются текущим предложением.",
                    "Averon Import",
                )
                target_comment_cells.add(price_header.coordinate)

            target_letters = (get_column_letter(price_column), get_column_letter(total_column))
            for letter, width in zip(target_letters, (20, 20), strict=True):
                dimension = sheet.column_dimensions[letter]
                dimension.hidden = False
                dimension.width = width
            workbook.save(output_path)
            workbook.close()
            workbook = None
        except TenderWorkspaceError:
            raise
        except Exception as exc:
            raise TenderWorkspaceError("Не удалось безопасно сформировать XLSX.", 409, "TENDER_EXPORT_FAILED") from exc
        finally:
            if workbook is not None:
                workbook.close()
            if normalized_path is not None:
                try:
                    normalized_path.unlink(missing_ok=True)
                except OSError:
                    pass

        if not output_path.is_file() or output_path.stat().st_size > MAX_TENDER_EXPORT_BYTES:
            raise TenderWorkspaceError("Размер готового XLSX превышает допустимый предел.", 413, "TENDER_EXPORT_TOO_LARGE")
        if _sha256_file(source_path) != workspace.get("source_sha256"):
            raise TenderWorkspaceError("Исходная книга изменилась во время экспорта.", 409, "TENDER_SOURCE_CHANGED")
        _restore_source_owned_package_metadata(source_path, output_path)
        _restore_empty_shared_string_cells(source_path, output_path, allowed_cells, target_sheet)
        after = _workbook_snapshot(output_path)
        _verify_semantic_roundtrip(
            before,
            after,
            target_sheet=target_sheet,
            allowed_cells=allowed_cells,
            target_columns={price_column, total_column},
        )
        _verify_package_roundtrip(
            source_path,
            output_path,
            target_sheet=target_sheet,
            allowed_cells=allowed_cells,
            target_comment_cells=target_comment_cells,
            semantic_before=before,
            semantic_after=after,
        )
        return {
            "target_sheet": target_sheet,
            "target_columns": {
                "unit_price": get_column_letter(price_column),
                "total": get_column_letter(total_column),
            },
            "historical_count": historical_count,
            "automatic_historical_count": historical_count - human_confirmed_history_count,
            "human_confirmed_historical_count": human_confirmed_history_count,
        }


class TenderPriceExportRepository:
    """Durable owner-bound export metadata and artifacts under each workspace."""

    def __init__(self, repository: TenderWorkspaceRepository):
        self.repository = repository
        self._lock = threading.RLock()
        self._downloads: dict[tuple[str, str], int] = {}
        self._cleanup_orphan_files()

    def _directory(self, workspace_path: Path, tender_id: str) -> Path:
        path = Path(workspace_path).resolve()
        root = self.repository.workspace_root.resolve()
        safe_id = self.repository._validate_id(tender_id)
        if path.parent != root or path.name != safe_id:
            raise TenderWorkspaceError("Тендер не найден.")
        return path / "exports"

    @staticmethod
    def _atomic_json(path: Path, value: dict[str, Any]) -> None:
        payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{path.stem}.{uuid.uuid4().hex}.", suffix=".tmp", dir=path.parent,
        )
        temporary = Path(temporary_name)
        try:
            with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass

    @staticmethod
    def _read_metadata(path: Path, tender_id: str) -> dict[str, Any]:
        try:
            if path.stat().st_size > 64 * 1024:
                raise ValueError
            value = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(value, dict) or value.get("schema_version") != 1 or value.get("tender_id") != tender_id or value.get("status") != "completed":
                raise ValueError
            return value
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            raise TenderWorkspaceError("Экспорт не найден.", 404, "TENDER_EXPORT_NOT_FOUND") from exc

    def _entries(self, directory: Path, tender_id: str) -> list[tuple[dict[str, Any], Path, Path]]:
        entries = []
        if not directory.is_dir():
            return entries
        for meta_path in directory.glob("[a-f0-9]" * 32 + ".json"):
            try:
                record = self._read_metadata(meta_path, tender_id)
                entries.append((record, meta_path, directory / f"{meta_path.stem}.xlsx"))
            except TenderWorkspaceError:
                continue
        entries.sort(key=lambda item: str(item[0].get("created_at") or ""))
        return entries

    def _cleanup_orphan_files(self) -> None:
        cutoff = datetime.now(timezone.utc).timestamp() - 3600
        for workspace_path in self.repository._workspace_dirs():
            directory = workspace_path / "exports"
            if not directory.is_dir():
                continue
            for candidate in directory.iterdir():
                try:
                    if not candidate.is_file() or candidate.stat().st_mtime >= cutoff:
                        continue
                    if candidate.name.startswith(".") and candidate.suffix == ".tmp":
                        candidate.unlink(missing_ok=True)
                    elif re.fullmatch(r"\.[a-f0-9]{32}\.[a-f0-9]{32}\.tmp\.xlsx", candidate.name):
                        candidate.unlink(missing_ok=True)
                    elif _EXPORT_ID_RE.fullmatch(candidate.stem) and candidate.suffix.casefold() == ".xlsx":
                        if not (directory / f"{candidate.stem}.json").exists():
                            candidate.unlink(missing_ok=True)
                except OSError:
                    continue

    @staticmethod
    def _remove_entry(entry: tuple[dict[str, Any], Path, Path]) -> None:
        _record, meta_path, file_path = entry
        try:
            meta_path.unlink(missing_ok=True)
            file_path.unlink(missing_ok=True)
        except OSError as exc:
            raise TenderWorkspaceError("Не удалось освободить место для нового экспорта.", 409, "TENDER_EXPORT_RETENTION_FAILED") from exc

    def create_completed(
        self,
        workspace_path: Path,
        workspace: dict[str, Any],
        run: dict[str, Any],
        decisions: list[TenderPriceDecision],
        *,
        owner_id: str,
        allow_partial: bool,
        include_historical_prices: bool,
        history_decision_revision: int = 0,
        history_decision_digest: str = "",
        builder: TenderXlsxPriceExporter,
    ) -> dict[str, Any]:
        tender_id = str(workspace.get("tender_id") or "")
        directory = self._directory(workspace_path, tender_id)
        directory.mkdir(mode=0o700, parents=True, exist_ok=True)
        export_id = uuid.uuid4().hex
        candidate = directory / f".{export_id}.{uuid.uuid4().hex}.tmp.xlsx"
        final_path = directory / f"{export_id}.xlsx"
        meta_path = directory / f"{export_id}.json"
        published = False
        try:
            info = builder.write(workspace_path / "source.xlsx", candidate, workspace, decisions)
            if candidate.stat().st_size > MAX_TENDER_EXPORT_BYTES:
                raise TenderWorkspaceError("Размер готового XLSX превышает допустимый предел.", 413, "TENDER_EXPORT_TOO_LARGE")
            if _sha256_file(workspace_path / "source.xlsx") != workspace.get("source_sha256"):
                raise TenderWorkspaceError("Исходная книга изменилась во время экспорта.", 409, "TENDER_SOURCE_CHANGED")
            summary = summarize_decisions(decisions)
            record = {
                "schema_version": 1,
                "status": "completed",
                "export_id": export_id,
                "tender_id": tender_id,
                "owner_id": owner_id,
                "run_id": run["run_id"],
                "source_sha256": workspace["source_sha256"],
                "workspace_revision": int(workspace["revision"]),
                "export_policy_revision": TENDER_EXPORT_POLICY_REVISION,
                "include_historical_prices": bool(include_historical_prices),
                "allow_partial": bool(allow_partial),
                "history_decision_revision": int(history_decision_revision),
                "history_decision_digest": str(history_decision_digest),
                "created_at": datetime.now(timezone.utc).isoformat(),
                "filename": _safe_output_filename(str(workspace.get("filename") or "tender.xlsx")),
                "output_file_sha256": _sha256_file(candidate),
                **info,
                **summary,
                "run_source_mode": run.get("source_mode"),
                "run_completed_at": run.get("completed_at"),
            }
            if len(_json_text(record).encode("utf-8")) > 64 * 1024:
                raise TenderWorkspaceError("Метаданные экспорта превышают допустимый предел.", 413, "TENDER_EXPORT_METADATA_TOO_LARGE")
            with self.repository._lock, self._lock:
                entries = self._entries(directory, tender_id)
                excess = max(0, len(entries) + 1 - MAX_COMPLETED_EXPORTS_PER_WORKSPACE)
                prunable = [
                    item for item in entries
                    if self._downloads.get((tender_id, str(item[0].get("export_id") or "")), 0) == 0
                ]
                if len(prunable) < excess:
                    raise TenderWorkspaceError("Завершённый экспорт сейчас скачивается. Повторите позже.", 409, "TENDER_EXPORT_RETENTION_BUSY")
                prune = prunable[:excess]
                current_bytes = self.repository._workspace_storage_bytes()
                reclaim = sum(
                    (meta.stat().st_size if meta.exists() else 0) + (file.stat().st_size if file.exists() else 0)
                    for _item, meta, file in prune
                )
                metadata_bytes = len(_json_text(record).encode("utf-8"))
                if current_bytes + metadata_bytes - reclaim > MAX_TENDER_STORAGE_BYTES:
                    raise TenderWorkspaceError("Достигнут общий лимит хранилища тендеров.", 409, "TENDER_DISK_QUOTA")
                os.replace(candidate, final_path)
                try:
                    self._atomic_json(meta_path, record)
                    published = True
                except Exception:
                    final_path.unlink(missing_ok=True)
                    raise
                for entry in prune:
                    self._remove_entry(entry)
            return self.public_record(record)
        except TenderWorkspaceError:
            if published:
                try:
                    meta_path.unlink(missing_ok=True)
                    final_path.unlink(missing_ok=True)
                except OSError:
                    pass
            raise
        except Exception as exc:
            if published:
                try:
                    meta_path.unlink(missing_ok=True)
                    final_path.unlink(missing_ok=True)
                except OSError:
                    pass
            raise TenderWorkspaceError("Не удалось безопасно сохранить экспорт.", 409, "TENDER_EXPORT_FAILED") from exc
        finally:
            try:
                candidate.unlink(missing_ok=True)
            except OSError:
                pass

    @staticmethod
    def public_record(record: dict[str, Any]) -> dict[str, Any]:
        return {
            key: record.get(key)
            for key in (
                "export_id", "tender_id", "run_id", "source_sha256",
                "workspace_revision", "export_policy_revision",
                "include_historical_prices", "allow_partial", "created_at",
                "filename", "output_file_sha256", "target_sheet", "target_columns",
                "selected_count", "priced_count", "historical_count", "blank_count",
                "automatic_historical_count", "human_confirmed_historical_count",
                "history_decision_revision", "history_decision_digest",
                "reason_counts", "run_source_mode", "run_completed_at",
            )
        }

    @staticmethod
    def _verify_artifact(file_path: Path, record: dict[str, Any]) -> None:
        try:
            if file_path.stat().st_size <= 0 or file_path.stat().st_size > MAX_TENDER_EXPORT_BYTES:
                raise OSError
            if _sha256_file(file_path) != record.get("output_file_sha256"):
                raise OSError
        except OSError as exc:
            raise TenderWorkspaceError("Файл экспорта повреждён или отсутствует.", 409, "TENDER_EXPORT_CORRUPT") from exc

    def get_public(self, workspace_path: Path, tender_id: str, export_id: str, owner_id: str) -> dict[str, Any]:
        if not _EXPORT_ID_RE.fullmatch(str(export_id or "")):
            raise TenderWorkspaceError("Экспорт не найден.", 404, "TENDER_EXPORT_NOT_FOUND")
        directory = self._directory(workspace_path, tender_id)
        record = self._read_metadata(directory / f"{export_id}.json", tender_id)
        if record.get("owner_id") != owner_id:
            raise TenderWorkspaceError("Экспорт не найден.", 404, "TENDER_EXPORT_NOT_FOUND")
        self._verify_artifact(directory / f"{export_id}.xlsx", record)
        return self.public_record(record)

    def list_public(self, workspace_path: Path, tender_id: str, owner_id: str) -> dict[str, Any]:
        directory = self._directory(workspace_path, tender_id)
        records = []
        with self._lock:
            for record, _meta_path, _file_path in reversed(self._entries(directory, tender_id)):
                if record.get("owner_id") != owner_id:
                    continue
                self._verify_artifact(directory / f"{record['export_id']}.xlsx", record)
                records.append(self.public_record(record))
        return {"exports": records[:MAX_COMPLETED_EXPORTS_PER_WORKSPACE]}

    def acquire_download(self, workspace_path: Path, tender_id: str, export_id: str, owner_id: str) -> tuple[Path, dict[str, Any], Any]:
        if not _EXPORT_ID_RE.fullmatch(str(export_id or "")):
            raise TenderWorkspaceError("Экспорт не найден.", 404, "TENDER_EXPORT_NOT_FOUND")
        directory = self._directory(workspace_path, tender_id)
        with self._lock:
            record = self._read_metadata(directory / f"{export_id}.json", tender_id)
            if record.get("owner_id") != owner_id:
                raise TenderWorkspaceError("Экспорт не найден.", 404, "TENDER_EXPORT_NOT_FOUND")
            file_path = directory / f"{export_id}.xlsx"
            self._verify_artifact(file_path, record)
            key = (tender_id, export_id)
            self._downloads[key] = self._downloads.get(key, 0) + 1
        released = False
        release_lock = threading.Lock()

        def release() -> None:
            nonlocal released
            with release_lock:
                if released:
                    return
                released = True
            with self._lock:
                count = self._downloads.get(key, 0)
                if count <= 1:
                    self._downloads.pop(key, None)
                else:
                    self._downloads[key] = count - 1
                entries = self._entries(directory, tender_id)
                excess = max(0, len(entries) - MAX_COMPLETED_EXPORTS_PER_WORKSPACE)
                if excess:
                    prunable = [
                        item for item in entries
                        if self._downloads.get((tender_id, str(item[0].get("export_id") or "")), 0) == 0
                    ]
                    for item in prunable[:excess]:
                        try:
                            self._remove_entry(item)
                        except TenderWorkspaceError:
                            break

        return file_path, record, release


__all__ = [
    "TENDER_EXPORT_POLICY_REVISION",
    "MAX_TENDER_EXPORT_BYTES",
    "MAX_COMPLETED_EXPORTS_PER_WORKSPACE",
    "XLSX_MIME",
    "TenderPriceDecision",
    "TenderPriceResolver",
    "TenderXlsxPriceExporter",
    "TenderPriceExportRepository",
    "summarize_decisions",
]
