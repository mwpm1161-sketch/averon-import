from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
import zipfile
from dataclasses import asdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from averon_import.services.one_c_history.xlsx_import import OneCImportError, preflight_xlsx as shared_preflight_xlsx
from .template import (
    PREPARED_ROWS, TEMPLATE_HEADERS, TEMPLATE_SCHEMA_NAME,
    TEMPLATE_SCHEMA_VERSION, TEMPLATE_SHEET, TEMPLATE_TABLE,
)
from .models import TenderMapping, TenderSourceManifest, TenderSourceRow, TenderUnitBasis, TenderWorkbookStructure

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_ZIP_MEMBERS = 1000
MAX_COMPRESSION_RATIO = 100
MAX_SHEETS = 12
MAX_ROWS_PER_SHEET = 2000
MAX_COLUMNS_PER_SHEET = 256
MAX_HEADER_SCAN_ROWS = 50
MAX_HEADER_LENGTH = 200
MAX_CELL_TEXT_LENGTH = 20_000
MAX_ACTUAL_ITEMS = 500
MAX_ANALYSIS_DATA_BYTES = 7 * 1024 * 1024
PARSER_VERSION = 1
MAX_PREVIEW_ROWS = 15

FIELD_ALIASES: dict[str, set[str]] = {
    "resource_code": {"кодресурса", "ресурсныйкод", "нсиресурса", "номресурса", "сметкодресурса", "номсметкодресурса", "noсметкодресурса"},
    "name": {"наименование", "название", "наименованиепозиции", "описание", "ресурс", "наименованиересурса", "товар"},
    "unit": {"едизм", "единицаизмерения", "единица", "unit", "измерение"},
    "quantity": {"колво", "количество", "qty", "quantity", "количествоединиц"},
    "article": {"артикул", "sku", "артикулпроизводителя", "кодsku", "manufacturerarticle"},
    "manufacturer": {"производитель", "изготовитель", "бренд", "manufacturer"},
    "model": {"модельтип", "модель", "тип", "марка", "model", "typemark"},
}
REQUIRED_FIELDS = ("name", "quantity")
OFFICIAL_FIELD_BY_HEADER = dict(zip(TEMPLATE_HEADERS, ("resource_code", "name", "unit", "quantity", "article", "manufacturer", "model")))
_NUMBER_RE = re.compile(r"^[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)$")


class TenderParseError(ValueError):
    def __init__(self, message: str, code: str = "TENDER_XLSX_INVALID"):
        super().__init__(message)
        self.code = code


def _header_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", "", text)


def _safe_xml(data: bytes, error_message: str) -> ElementTree.Element:
    upper = data.upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise TenderParseError("Книги с DTD/ENTITY декларациями не поддерживаются.", "TENDER_XLSX_XML_UNSAFE")
    try:
        return ElementTree.fromstring(data)
    except ElementTree.ParseError as exc:
        raise TenderParseError(error_message) from exc


def preflight_tender_xlsx(path: Path) -> None:
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise TenderParseError("Файл не является корректной книгой XLSX.") from exc
    try:
        with archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_ZIP_MEMBERS:
                raise TenderParseError("Архив XLSX содержит недопустимое число компонентов.")
            exact: set[str] = set()
            folded: set[str] = set()
            total = 0
            for info in infos:
                name = info.filename.replace("\\", "/")
                parts = PurePosixPath(name).parts
                if (name.startswith("/") or ".." in parts or "\x00" in name or
                        (parts and ":" in parts[0]) or name in exact):
                    raise TenderParseError("Архив XLSX содержит небезопасные или повторные компоненты.")
                if name.casefold() in folded:
                    raise TenderParseError("Архив XLSX содержит повторные компоненты без учёта регистра.")
                exact.add(name)
                folded.add(name.casefold())
                if info.file_size < 0 or info.compress_size < 0:
                    raise TenderParseError("Архив XLSX повреждён.")
                total += info.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise TenderParseError("Распакованный размер XLSX превышает 20 МиБ.")
                if info.file_size > 1024 and (info.compress_size == 0 or info.file_size / info.compress_size > MAX_COMPRESSION_RATIO):
                    raise TenderParseError("Коэффициент сжатия XLSX превышает допустимый предел.")
                lower = name.casefold()
                if "vbaproject.bin" in lower or lower.startswith("xl/externallinks/") or lower.endswith(".bin") and "vba" in lower:
                    raise TenderParseError("Книги с макросами и внешними ссылками не поддерживаются.")
                if lower.endswith(".rels"):
                    if info.file_size > 1024 * 1024:
                        raise TenderParseError("Метаданные связей XLSX превышают допустимый размер.")
                    root = _safe_xml(archive.read(info), "Метаданные XLSX повреждены.")
                    if any(str(node.attrib.get("TargetMode", "")).casefold() == "external" for node in root.iter()):
                        raise TenderParseError("Внешние связи в XLSX не поддерживаются.")
                elif lower.endswith(".xml"):
                    if info.file_size > 4 * 1024 * 1024:
                        raise TenderParseError("Компонент XML XLSX превышает допустимый размер.")
                    _safe_xml(archive.read(info), "Компонент XML XLSX повреждён.")
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(exact):
                raise TenderParseError("Структура XLSX не содержит обязательных компонентов.")
    except TenderParseError:
        raise
    except Exception as exc:
        raise TenderParseError("Архив XLSX повреждён или не поддерживается.") from exc


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    if len(text) > MAX_CELL_TEXT_LENGTH:
        raise TenderParseError("Текст ячейки превышает допустимую длину.")
    return text


def _decimal_quantity(value: Any, data_type: str) -> tuple[str, str | None, bool]:
    raw = _cell_text(value)
    if not raw:
        return raw, None, False
    if data_type == "f" or data_type == "e" or isinstance(value, (bool, date, datetime)):
        return raw, None, False
    try:
        if isinstance(value, (int, float, Decimal)):
            if isinstance(value, float) and not math.isfinite(value):
                return raw, None, False
            number = Decimal(str(value))
        elif _NUMBER_RE.fullmatch(raw):
            number = Decimal(raw.replace(",", "."))
        else:
            return raw, None, False
    except (InvalidOperation, ValueError):
        return raw, None, False
    if not number.is_finite() or number <= 0:
        return raw, None, False
    return raw, format(number.normalize(), "f"), True


def parse_unit_basis(raw_value: Any) -> dict[str, Any]:
    raw = _cell_text(raw_value)
    normalized = " ".join(unicodedata.normalize("NFKC", raw).casefold().replace("ё", "е").split())
    match = re.fullmatch(r"(?:(\d+)\s*)?(шт|штук|м|метр(?:а|ов)?|кг|килограмм(?:а|ов)?|т|тонн(?:а|ы)?|м2|м²|м3|м³|л|литр(?:а|ов)?|компл|комплект(?:а|ов)?|уп|упак(?:овка|овки|овок)?)\.?", normalized)
    if not match:
        return asdict(TenderUnitBasis(raw, None, None, None, None, False))
    multiplier = int(match.group(1) or 1)
    unit = match.group(2)
    if multiplier not in {1, 10, 100, 1000}:
        return asdict(TenderUnitBasis(raw, None, None, None, None, False))
    if multiplier != 1 and unit not in {"шт", "штук", "м", "метр", "метра", "метров"}:
        return asdict(TenderUnitBasis(raw, None, None, None, None, False))
    if unit in {"шт", "штук"}:
        base, dim, basis = "шт", "count", "each"
    elif unit in {"м", "метр", "метра", "метров"}:
        base, dim, basis = "м", "length", "each"
    elif unit in {"кг", "килограмм", "килограмма", "килограммов"}:
        base, dim, basis = "кг", "mass", "each"
    elif unit in {"т", "тонна", "тонны", "тонн"}:
        base, dim, basis = "кг", "mass", "tonne_to_kg"
        multiplier *= 1000
    elif unit in {"м2", "м²"}:
        base, dim, basis = "м2", "area", "each"
    elif unit in {"м3", "м³"}:
        base, dim, basis = "м3", "volume", "each"
    elif unit in {"л", "литр", "литра", "литров"}:
        base, dim, basis = "л", "volume", "each"
    elif unit in {"компл", "комплект", "комплекта", "комплектов"}:
        base, dim, basis = "компл", "set", "unspecified_set"
    else:
        base, dim, basis = "уп", "package", "unspecified_package"
    return asdict(TenderUnitBasis(raw, base, dim, str(multiplier), basis, True))


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)[:MAX_CELL_TEXT_LENGTH]


def _color_style(color: Any) -> dict[str, Any] | None:
    if color is None:
        return None
    # Some unset openpyxl descriptors return descriptor placeholders from
    # getattr; __dict__ contains only values actually stored in the workbook.
    return {key: value for key, value in vars(color).items() if key in {"type", "rgb", "indexed", "theme", "tint", "auto"}}


def _style_fingerprint(cell: Any) -> str:
    """Hash normalized formatting properties instead of workbook-local style ids."""
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    protection = cell.protection
    sides = {}
    for side_name in ("left", "right", "top", "bottom", "diagonal", "vertical", "horizontal", "start", "end"):
        side = getattr(border, side_name, None)
        sides[side_name] = None if side is None else {
            "style": side.style,
            "color": _color_style(side.color),
        }
    normalized = {
        "number_format": cell.number_format,
        "font": {
            "name": font.name, "size": font.sz, "bold": font.bold,
            "italic": font.italic, "underline": font.underline, "strike": font.strike,
            "vert_align": font.vertAlign, "color": _color_style(font.color),
            "scheme": font.scheme, "charset": font.charset, "family": font.family,
        },
        "fill": {
            "type": fill.fill_type, "foreground": _color_style(fill.fgColor),
            "background": _color_style(fill.bgColor),
        },
        "border": {
            "sides": sides, "diagonal_up": border.diagonalUp,
            "diagonal_down": border.diagonalDown, "outline": border.outline,
        },
        "alignment": {
            key: getattr(alignment, key)
            for key in ("horizontal", "vertical", "textRotation", "wrapText", "shrinkToFit", "indent", "relativeIndent", "justifyLastLine", "readingOrder")
        },
        "protection": {"locked": protection.locked, "hidden": protection.hidden},
    }
    payload = json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def is_semantically_empty_cell(cell: Any) -> bool:
    """Whether a target cell can be reused without replacing workbook meaning.

    Physical empty-string shared-string cells and style-only cells are allowed.
    Formulas, comments, hyperlinks and non-empty constants remain occupied even
    when a formula happens to display an empty result.
    """
    if (
        getattr(cell, "data_type", None) == "f"
        or getattr(cell, "comment", None) is not None
        or getattr(cell, "hyperlink", None) is not None
    ):
        return False
    return getattr(cell, "value", None) in (None, "")


class TenderWorkbookParser:
    def parse(self, path: Path, *, tender_id: str | None = None, mapping_override: dict[str, int | None] | None = None, selected_sheet: str | None = None, selected_header_row: int | None = None) -> dict[str, Any]:
        if path.suffix.casefold() != ".xlsx":
            raise TenderParseError("Загрузите файл в формате .xlsx.")
        if path.stat().st_size > MAX_UPLOAD_BYTES:
            raise TenderParseError("Размер XLSX превышает 5 МиБ.")
        preflight_tender_xlsx(path)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        normalized_path: Path | None = None
        try:
            # Reuse the existing generic OOXML casing normalization after the
            # stricter tender-specific size/security checks have passed.
            normalized_path = shared_preflight_xlsx(path)
            workbook = load_workbook(normalized_path or path, read_only=False, data_only=False, keep_links=False)
        except OneCImportError as exc:
            if normalized_path:
                normalized_path.unlink(missing_ok=True)
            raise TenderParseError("Книга XLSX не прошла проверку безопасности.") from exc
        except Exception as exc:
            if normalized_path:
                normalized_path.unlink(missing_ok=True)
            raise TenderParseError("Не удалось прочитать книгу XLSX. Проверьте формат и структуру файла.") from exc
        try:
            if not workbook.sheetnames or len(workbook.sheetnames) > MAX_SHEETS:
                raise TenderParseError("В книге слишком много листов.")
            for sheet in workbook.worksheets:
                if sheet.max_row > MAX_ROWS_PER_SHEET or sheet.max_column > MAX_COLUMNS_PER_SHEET:
                    raise TenderParseError("Размер листа превышает поддерживаемую границу.")
            table_hits = [(sheet, table) for sheet in workbook.worksheets for table in sheet.tables.values() if table.name == TEMPLATE_TABLE]
            schema_marker = workbook.defined_names.get(TEMPLATE_SCHEMA_NAME)
            if schema_marker is not None and str(schema_marker.attr_text or "").strip() != TEMPLATE_SCHEMA_VERSION:
                raise TenderParseError("Версия схемы официального шаблона Averon не поддерживается.")
            is_official = bool(table_hits)
            if is_official:
                if schema_marker is None:
                    raise TenderParseError("В официальной книге отсутствует версия схемы Averon.")
                if len(table_hits) != 1:
                    raise TenderParseError("Официальная таблица AveronTenderInput должна быть единственной.")
                sheet, table = table_hits[0]
                if sheet.title != TEMPLATE_SHEET:
                    raise TenderParseError("Таблица AveronTenderInput находится на неожиданном листе.")
                min_col, min_row, max_col, max_row = __import__("openpyxl.utils.cell", fromlist=["range_boundaries"]).range_boundaries(table.ref)
                if min_col != 1 or min_row != 1 or max_col != len(TEMPLATE_HEADERS) or not 1 <= max_row <= PREPARED_ROWS + 1:
                    raise TenderParseError("Диапазон официальной таблицы Averon имеет неверную структуру.")
                headers = [_cell_text(sheet.cell(1, column).value) for column in range(1, len(TEMPLATE_HEADERS) + 1)]
                if tuple(headers) != TEMPLATE_HEADERS:
                    raise TenderParseError("Заголовки официальной таблицы Averon повреждены или изменены.")
                table_columns = list(table.tableColumns or [])
                if len(table_columns) != len(TEMPLATE_HEADERS) or tuple(column.name for column in table_columns) != TEMPLATE_HEADERS:
                    raise TenderParseError("Столбцы официальной таблицы Averon повреждены или изменены.")
                if any(
                    (cell.value is not None or cell.comment is not None)
                    and (cell.column > len(TEMPLATE_HEADERS) or cell.row > max_row)
                    for cell in sheet._cells.values()
                ):
                    raise TenderParseError("В официальной книге найдены данные за пределами таблицы A1:G.")
                mapping = {field: index for index, field in enumerate(OFFICIAL_FIELD_BY_HEADER.values(), start=1)}
                header_row = 1
                ambiguous = False
                header_candidates = [{"sheet_name": sheet.title, "header_row": header_row}]
            else:
                if TEMPLATE_SHEET in workbook.sheetnames:
                    official_sheet = workbook[TEMPLATE_SHEET]
                    first_row = tuple(_cell_text(official_sheet.cell(1, column).value) for column in range(1, len(TEMPLATE_HEADERS) + 1))
                    instruction_marker = "Инструкция" in workbook.sheetnames and workbook["Инструкция"]["A1"].value == "Шаблон Averon v1"
                    if schema_marker is not None or first_row == TEMPLATE_HEADERS or any(official_sheet.cell(1, column).value is not None for column in range(1, len(TEMPLATE_HEADERS) + 1)) or instruction_marker:
                        raise TenderParseError("В официальной книге отсутствует корректная таблица AveronTenderInput.")
                if selected_sheet is not None or selected_header_row is not None:
                    sheet = workbook[selected_sheet] if selected_sheet in workbook.sheetnames else None
                    if sheet is None or selected_header_row is None or selected_header_row < 1 or selected_header_row > min(sheet.max_row, MAX_HEADER_SCAN_ROWS):
                        raise TenderParseError("Лист или строка заголовка недоступны для сопоставления.", "TENDER_MAPPING_SCOPE_INVALID")
                    header_row = selected_header_row
                    headers = [_cell_text(sheet.cell(header_row, column).value) for column in range(1, min(sheet.max_column, MAX_COLUMNS_PER_SHEET) + 1)]
                    while headers and not headers[-1]:
                        headers.pop()
                    self._validate_header_lengths(headers)
                    mapping, ambiguous = self._mapping_for_headers(headers)
                    header_candidates = [{"sheet_name": sheet.title, "header_row": header_row}]
                else:
                    sheet, header_row, headers, mapping, ambiguous, header_candidates = self._detect_fallback(workbook)
                if mapping_override is not None:
                    mapping = self.validate_mapping_override(mapping_override, headers)
                    ambiguous = False
            if mapping_override is not None and is_official:
                raise TenderParseError("Для корректного официального шаблона ручное сопоставление не требуется.")
            required_missing = [field for field in REQUIRED_FIELDS if mapping.get(field) is None]
            if required_missing:
                ambiguous = True
            rows: list[dict[str, Any]] = []
            rows_json_bytes = 0
            counts = {kind: 0 for kind in ("item", "section", "total", "ignored", "invalid")}
            for physical_row in range(header_row + 1, sheet.max_row + 1):
                values = {field: self._value(sheet, physical_row, mapping.get(field)) for field in mapping}
                relevant = [value[0] for value in values.values()]
                if not any(value not in (None, "") for value in relevant):
                    continue
                name = values.get("name", (None, None))[0]
                qty, qty_type = values.get("quantity", (None, "n"))
                if not is_official and self._is_column_number_scaffold(values, mapping):
                    kind, invalid_reasons = "ignored", []
                elif not is_official and self._is_merged_section_row(sheet, physical_row, mapping):
                    kind, invalid_reasons = "section", []
                elif not is_official and self._is_numbered_reference_note(values):
                    kind, invalid_reasons = "ignored", []
                else:
                    kind, invalid_reasons = self._classify(name, qty, qty_type, values, official=is_official)
                counts[kind] += 1
                if kind == "item" and counts[kind] > MAX_ACTUAL_ITEMS:
                    raise TenderParseError("В книге больше 500 строк позиций.")
                if kind not in {"item", "section", "total", "invalid", "ignored"}:
                    continue
                sheet_identity = f"{sheet.title}:{sheet.sheet_state}"
                seed = f"{tender_id or digest}:{digest}:{sheet_identity}:{physical_row}"
                source_row_id = hashlib.sha256(seed.encode("utf-8")).hexdigest()[:32]
                row: dict[str, Any] = {
                    "source_row_id": source_row_id,
                    "sheet_name": sheet.title,
                    "excel_row": physical_row,
                    "row_type": kind,
                    "resource_code": self._field_text(values, "resource_code"),
                    "name": self._field_text(values, "name"),
                    "raw_unit": self._field_text(values, "unit"),
                    "quantity_raw": _cell_text(qty),
                    "quantity": None,
                    "quantity_trusted": False,
                    "article": self._field_text(values, "article"),
                    "manufacturer": self._field_text(values, "manufacturer"),
                    "model": self._field_text(values, "model"),
                    "source_cells": {field: f"{get_column_letter(col)}{physical_row}" for field, col in mapping.items() if col is not None},
                    "invalid_reason_codes": invalid_reasons,
                }
                row["quantity_raw"], row["quantity"], row["quantity_trusted"] = _decimal_quantity(qty, qty_type)
                row["unit_basis"] = parse_unit_basis(row["raw_unit"])
                row["warnings"] = []
                if kind == "item" and not row["quantity_trusted"]:
                    row["warnings"].append("Количество не подтверждено: значение отсутствует или неоднозначно.")
                if kind == "item" and row["raw_unit"] and not row["unit_basis"]["trusted"]:
                    row["warnings"].append("Единица измерения не распознана; исходный текст сохранён.")
                serialized_row = asdict(TenderSourceRow(**row))
                rows_json_bytes += len(json.dumps(serialized_row, ensure_ascii=False, separators=(",", ":")).encode("utf-8")) + 1
                if rows_json_bytes > MAX_ANALYSIS_DATA_BYTES:
                    raise TenderParseError("Строки XLSX превышают допустимый объём рабочего пространства.", "TENDER_WORKSPACE_TOO_LARGE")
                rows.append(serialized_row)
            source_edge = max((column for column in mapping.values() if column is not None), default=0)
            logical_edge = max_col if is_official else source_edge
            future_columns = self._future_columns(sheet, logical_edge)
            structure = asdict(TenderWorkbookStructure(
                sheet_name=sheet.title, header_row=header_row, headers=headers,
                mapping=mapping, logical_right_edge=logical_edge,
                logical_right_column=get_column_letter(logical_edge) if logical_edge else None,
                future_output_columns=future_columns,
            ))
            manifest = self._manifest(
                workbook, digest, sheet, headers, header_row,
                max_cell_bytes=MAX_ANALYSIS_DATA_BYTES - rows_json_bytes,
            )
            with zipfile.ZipFile(path) as archive:
                package_names = [item.filename.casefold() for item in archive.infolist()]
            sensitive_parts = [name for name in package_names if any(token in name for token in ("/pivot", "/slicer", "/embedding", "/querytables", "/activex", "customxml", "/drawings/"))]
            manifest["unsupported_preservation_sensitive_objects"] = sorted(set(sensitive_parts))[:200]
            warnings = []
            if not is_official and any("цена" in _header_key(_cell_text(sheet.cell(header_row, column).value)) or "сумма" in _header_key(_cell_text(sheet.cell(header_row, column).value)) for column in range(logical_edge + 1, min(sheet.max_column, logical_edge + 30) + 1)):
                warnings.append({"code": "DISTANT_PRICE_COLUMNS", "message": "В книге обнаружены существующие служебные столбцы с ценовыми заголовками вне основной таблицы. Они не будут изменены."})
            mapping_required = bool(ambiguous or required_missing)
            return {
                "parser_version": PARSER_VERSION,
                "structure": structure,
                "source_sha256": digest,
                "sheet_names": list(workbook.sheetnames),
                "sheet_name": sheet.title,
                "sheet_candidates": sorted({candidate["sheet_name"] for candidate in header_candidates}) or [sheet.title],
                "header_candidates": header_candidates,
                "sheet_visibility": sheet.sheet_state,
                "header_row": header_row,
                "headers": headers,
                "mapping": mapping,
                "mapping_required": mapping_required,
                "official_template": is_official,
                "table_name": TEMPLATE_TABLE if is_official else None,
                "table_ref": table.ref if is_official else None,
                "logical_right_edge": logical_edge,
                "logical_right_column": get_column_letter(logical_edge) if logical_edge else None,
                "future_output_columns": future_columns,
                "counts": counts,
                "item_count": counts["item"],
                "section_count": counts["section"],
                "total_count": counts["total"],
                "invalid_count": counts["invalid"],
                "invalid_examples": [
                    {"excel_row": row["excel_row"], "reasons": row["invalid_reason_codes"]}
                    for row in rows if row["row_type"] == "invalid"
                ][:MAX_PREVIEW_ROWS],
                "warnings": warnings,
                "sample_rows": [row for row in rows if row["row_type"] == "item"][:MAX_PREVIEW_ROWS],
                "rows": rows,
                "manifest": manifest,
            }
        finally:
            workbook.close()
            if normalized_path:
                normalized_path.unlink(missing_ok=True)

    @staticmethod
    def _value(sheet, row: int, column: int | None) -> tuple[Any, str]:
        if column is None:
            return None, "n"
        cell = sheet.cell(row=row, column=column)
        if isinstance(cell.value, str) and len(cell.value) > MAX_CELL_TEXT_LENGTH:
            raise TenderParseError("Текст ячейки превышает допустимую длину.")
        return cell.value, cell.data_type

    @staticmethod
    def _field_text(values: dict[str, tuple[Any, str]], field: str) -> str:
        return _cell_text(values.get(field, (None, "n"))[0])

    @staticmethod
    def _is_column_number_scaffold(
        values: dict[str, tuple[Any, str]], mapping: dict[str, int | None],
    ) -> bool:
        """Recognize a compact row that repeats its mapped physical columns."""
        matched_fields = []
        for field, column in mapping.items():
            if column is None:
                continue
            value = values.get(field, (None, "n"))[0]
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            if isinstance(value, bool):
                return False
            try:
                if isinstance(value, str):
                    text = value.strip()
                    if not re.fullmatch(r"\+?\d+(?:\.0+)?", text):
                        return False
                    number = Decimal(text)
                elif isinstance(value, (int, float, Decimal)):
                    number = Decimal(str(value))
                else:
                    return False
            except (InvalidOperation, ValueError):
                return False
            if number != number.to_integral_value() or not 1 <= number <= MAX_COLUMNS_PER_SHEET:
                return False
            if int(number) != column:
                return False
            matched_fields.append(field)

        # Requiring multiple mapped fields, including both required fields,
        # avoids treating an ordinary numeric code or quantity as scaffolding.
        return len(matched_fields) >= 3 and {"name", "quantity"}.issubset(matched_fields)

    @staticmethod
    def _is_merged_section_row(sheet, physical_row: int, mapping: dict[str, int | None]) -> bool:
        source_columns = [column for column in mapping.values() if column is not None]
        if not source_columns:
            return False
        source_left, source_right = min(source_columns), max(source_columns)
        covering_merges = [
            merged for merged in sheet.merged_cells.ranges
            if merged.min_row == physical_row == merged.max_row
            and merged.min_col <= source_left and merged.max_col >= source_right
        ]
        if not covering_merges:
            return False

        # Only inspect the logical source area. Some workbooks repeat labels in
        # distant print-layout columns that are outside the mapped table.
        row_values = []
        for column in range(source_left, source_right + 1):
            cell = sheet._cells.get((physical_row, column))
            if cell is not None and cell.value not in (None, ""):
                row_values.append(cell)
        if len(row_values) != 1:
            return False
        cell = row_values[0]
        if cell.data_type == "f" or not isinstance(cell.value, str) or not _cell_text(cell.value):
            return False
        return any(
            merged.min_row == physical_row and merged.min_col == cell.column
            and merged.min_col <= cell.column <= merged.max_col
            for merged in covering_merges
        )

    @staticmethod
    def _is_numbered_reference_note(values: dict[str, tuple[Any, str]]) -> bool:
        reference = values.get("resource_code", (None, "n"))[0]
        if not isinstance(reference, str):
            return False
        note_match = re.match(r"^\s*\d{1,3}[.)]\s+(.+)$", reference, re.DOTALL)
        if note_match is None:
            return False
        note_body = note_match.group(1).strip()
        # Require explanatory, multiword content so a short numbered reference
        # code remains an invalid incomplete row instead of being discarded.
        if len(note_body) < 16 or len(note_body.split()) < 3:
            return False
        return not any(
            values.get(field, (None, "n"))[0] not in (None, "")
            for field in ("name", "unit", "quantity", "article", "manufacturer", "model")
        )

    def _detect_fallback(self, workbook):
        candidates = []
        for sheet in workbook.worksheets:
            for row_number in range(1, min(sheet.max_row, MAX_HEADER_SCAN_ROWS) + 1):
                headers = [_cell_text(sheet.cell(row_number, column).value) for column in range(1, min(sheet.max_column, MAX_COLUMNS_PER_SHEET) + 1)]
                if any(len(value) > MAX_HEADER_LENGTH for value in headers):
                    # Long preamble/body cells are valid source content; this
                    # row simply cannot be interpreted as a header candidate.
                    continue
                while headers and not headers[-1]:
                    headers.pop()
                mapping, ambiguous = self._mapping_for_headers(headers)
                required_present = sum(mapping.get(field) is not None for field in REQUIRED_FIELDS)
                nonempty_headers = sum(bool(header) for header in headers)
                if nonempty_headers:
                    score = required_present * 100 + sum(value is not None for value in mapping.values()) * 5 + min(nonempty_headers, 20)
                    candidates.append((score, sheet, row_number, headers, mapping, ambiguous))
        if not candidates:
            raise TenderParseError("Не удалось определить лист и заголовки. Задайте сопоставление вручную.", "TENDER_MAPPING_REQUIRED")
        candidates.sort(key=lambda value: value[0], reverse=True)
        best = candidates[0]
        self._validate_header_lengths(best[3])
        same_score = [candidate for candidate in candidates if candidate[0] == best[0]]
        ambiguous = best[5] or len(same_score) > 1 or any(best[4].get(field) is None for field in REQUIRED_FIELDS)
        bounded_headers = [
            {"sheet_name": candidate[1].title, "header_row": candidate[2], "headers": candidate[3], "mapping": candidate[4]}
            for candidate in candidates[:10]
        ]
        return best[1], best[2], best[3], best[4], ambiguous, bounded_headers

    @staticmethod
    def _validate_header_lengths(headers: list[str]) -> None:
        if any(len(value) > MAX_HEADER_LENGTH for value in headers):
            raise TenderParseError("Заголовок XLSX превышает допустимую длину.")

    @staticmethod
    def _mapping_for_headers(headers: list[str]) -> tuple[dict[str, int | None], bool]:
        mapping: dict[str, int | None] = {field: None for field in FIELD_ALIASES}
        seen: dict[str, list[int]] = {}
        ambiguous = False
        for index, label in enumerate(headers, start=1):
            key = _header_key(label)
            # Generic identifiers stay unresolved; they never become article evidence.
            if key in {"код", "кодтовара", "номер"}:
                ambiguous = True
                continue
            for field, aliases in FIELD_ALIASES.items():
                if key in aliases:
                    seen.setdefault(field, []).append(index)
        for field, columns in seen.items():
            if len(columns) == 1:
                mapping[field] = columns[0]
            else:
                ambiguous = True
        return mapping, ambiguous

    @staticmethod
    def validate_mapping_override(mapping: dict[str, int | None], headers: list[str]) -> dict[str, int | None]:
        valid_fields = set(FIELD_ALIASES)
        result = {field: None for field in FIELD_ALIASES}
        header_count = len(headers)
        for field, column in mapping.items():
            if field not in valid_fields:
                raise TenderParseError("Сопоставление содержит неизвестное поле.", "TENDER_MAPPING_INVALID")
            if column is None:
                continue
            if isinstance(column, bool) or not isinstance(column, int) or column < 1 or column > header_count:
                raise TenderParseError("Сопоставление указывает колонку вне заголовков.", "TENDER_MAPPING_INVALID")
            result[field] = column
        if any(result.get(field) is None for field in REQUIRED_FIELDS):
            raise TenderParseError("Для импорта нужны колонки Наименование и Кол-во.", "TENDER_MAPPING_REQUIRED")
        if result.get("article") is not None and result["article"] == result.get("resource_code"):
            raise TenderParseError("Код ресурса нельзя сопоставить с артикулом.", "TENDER_RESOURCE_CODE_NOT_ARTICLE")
        if result.get("article") is not None and _header_key(headers[result["article"] - 1]) not in FIELD_ALIASES["article"]:
            raise TenderParseError("Артикул можно сопоставить только с колонкой, заголовок которой явно указывает на артикул.", "TENDER_ARTICLE_HEADER_UNSUPPORTED")
        mapped_columns = [column for column in result.values() if column is not None]
        if len(mapped_columns) != len(set(mapped_columns)):
            raise TenderParseError("Одну колонку нельзя сопоставить нескольким полям.", "TENDER_MAPPING_DUPLICATE_COLUMN")
        return asdict(TenderMapping(**result))

    @staticmethod
    def _classify(name: Any, quantity: Any, quantity_type: str, values: dict[str, tuple[Any, str]], *, official: bool = False) -> tuple[str, list[str]]:
        strings = [_cell_text(value).casefold() for value, _kind in values.values() if value not in (None, "")]
        if not official and any(re.search(r"\b(итого|всего|сумма итого|total)\b", value) for value in strings):
            return "total", []
        name_present = name not in (None, "") and bool(_cell_text(name))
        quantity_raw, _normalized_quantity, quantity_trusted = _decimal_quantity(quantity, quantity_type)
        quantity_present = bool(quantity_raw)
        item_signals = any(
            values.get(field, (None, ""))[0] not in (None, "")
            for field in ("resource_code", "unit", "article", "manufacturer", "model")
        )
        if not official and name_present and not quantity_present and not item_signals:
            return "section", []
        if name_present and quantity_trusted:
            return "item", []
        if not strings:
            return "ignored", []
        reasons = []
        if not name_present:
            reasons.append("missing_name")
        if not quantity_present:
            reasons.append("missing_quantity")
        elif not quantity_trusted:
            reasons.append("invalid_quantity")
        return "invalid", reasons

    @staticmethod
    def _future_columns(sheet, logical_edge: int) -> list[dict[str, Any]]:
        result = []
        for offset, column in enumerate((logical_edge + 1, logical_edge + 2)):
            letter = get_column_letter(column)
            dimension = sheet.column_dimensions[letter]
            values = []
            formulas = []
            comments = []
            for row in range(1, min(sheet.max_row, MAX_ROWS_PER_SHEET) + 1):
                cell = sheet.cell(row, column)
                if cell.data_type == "f":
                    formulas.append(cell.coordinate)
                elif cell.value not in (None, ""):
                    values.append(cell.coordinate)
                if cell.comment:
                    comments.append(cell.coordinate)
            merged = [str(rng) for rng in sheet.merged_cells.ranges if rng.min_col <= column <= rng.max_col]
            result.append({
                "column": letter, "planned_header": "Цена за единицу" if offset == 0 else "Общая стоимость",
                "occupied_values": values[:100], "value_count": len(values),
                "formulas": formulas[:100], "formula_count": len(formulas), "comments": comments[:100],
                "merges": merged, "hidden": bool(dimension.hidden),
                "meaningful_objects": bool(sheet._charts or sheet._images),
            })
        return result

    @staticmethod
    def _manifest(workbook, digest: str, selected_sheet, headers: list[str], header_row: int, *, max_cell_bytes: int) -> dict[str, Any]:
        meaningful = []
        meaningful_bytes = 0
        comment_counts: dict[str, int] = {}
        for sheet in workbook.worksheets:
            # Worksheet._cells is sparse: iterating the bounded rectangular
            # range would instantiate millions of empty cells for styled or
            # distant columns in otherwise small workbooks.
            cells = sheet._cells.values()
            comment_counts[sheet.title] = 0
            for cell in cells:
                if cell.value is None and not cell.comment:
                    continue
                if cell.comment:
                    comment_counts[sheet.title] += 1
                item = {"sheet": sheet.title, "cell": cell.coordinate, "value": _json_value(cell.value), "formula": cell.data_type == "f", "type": cell.data_type, "comment": bool(cell.comment), "style_fingerprint": _style_fingerprint(cell)}
                meaningful_bytes += len(json.dumps(item, ensure_ascii=False, separators=(",", ":")).encode("utf-8")) + 1
                if meaningful_bytes > max_cell_bytes:
                    raise TenderParseError("Манифест превышает допустимый объём рабочего пространства.", "TENDER_WORKSPACE_TOO_LARGE")
                meaningful.append(item)
        defined_names = []
        try:
            defined_names = [{"name": name, "attr_text": str(item.attr_text or "")[:1000]} for name, item in workbook.defined_names.items()]
        except AttributeError:
            defined_names = [str(item)[:1000] for item in workbook.defined_names.values()]
        sheets = []
        for sheet in workbook.worksheets:
            tables = [{"name": table.name, "ref": table.ref} for table in sheet.tables.values()]
            sheets.append({
                "name": sheet.title, "state": sheet.sheet_state, "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "tables": tables, "row_dimensions": {str(k): {"hidden": bool(v.hidden), "height": v.height} for k, v in sheet.row_dimensions.items()},
                "column_dimensions": {k: {"hidden": bool(v.hidden), "min": v.min, "max": v.max, "width": v.width} for k, v in sheet.column_dimensions.items()},
                "hidden_columns": [k for k, v in sheet.column_dimensions.items() if v.hidden],
                "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                "print_area": str(sheet.print_area) if sheet.print_area else None,
                "print_title_rows": sheet.print_title_rows, "print_title_cols": sheet.print_title_cols,
                "comments": comment_counts.get(sheet.title, 0),
                "data_validations": len(sheet.data_validations.dataValidation),
                "conditional_formatting_count": len(sheet.conditional_formatting),
                "images": len(sheet._images), "charts": len(sheet._charts),
            })
        return asdict(TenderSourceManifest(
            source_sha256=digest, sheet_names_order=list(workbook.sheetnames), sheets=sheets,
            meaningful_cells=meaningful, defined_names=defined_names,
            unsupported_preservation_sensitive_objects=[], selected_sheet=selected_sheet.title,
            header_row=header_row, detected_headers=headers,
        ))


__all__ = ["TenderWorkbookParser", "TenderParseError", "preflight_tender_xlsx", "parse_unit_basis", "is_semantically_empty_cell", "MAX_UPLOAD_BYTES", "MAX_UNCOMPRESSED_BYTES", "MAX_ACTUAL_ITEMS", "PARSER_VERSION"]
