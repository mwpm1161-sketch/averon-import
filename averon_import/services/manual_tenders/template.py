from __future__ import annotations

from functools import lru_cache
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName


TEMPLATE_SHEET = "Ресурсная ведомость"
TEMPLATE_TABLE = "AveronTenderInput"
TEMPLATE_SCHEMA_NAME = "AveronTenderSchemaVersion"
TEMPLATE_SCHEMA_VERSION = "1"
TEMPLATE_HEADERS = (
    "Код ресурса", "Наименование", "Ед. изм.", "Кол-во",
    "Артикул", "Производитель", "Модель / тип",
)
PREPARED_ROWS = 500


class TenderTemplateService:
    """Build and cache immutable in-memory official template bytes."""

    @staticmethod
    @lru_cache(maxsize=1)
    def _build() -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = TEMPLATE_SHEET
        sheet.append(list(TEMPLATE_HEADERS))
        header_fill = PatternFill(fill_type="solid", fgColor="E9ECEF")
        header_font = Font(name="Calibri", size=11, bold=True, color="263238")
        header_border = Border(bottom=Side(style="thin", color="B0BEC5"))
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = header_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in range(2, PREPARED_ROWS + 2):
            sheet.cell(row=row, column=1).number_format = "@"
            for column in (5, 6, 7):
                sheet.cell(row=row, column=column).number_format = "@"
        table = Table(displayName=TEMPLATE_TABLE, ref=f"A1:G{PREPARED_ROWS + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
        workbook.defined_names.add(DefinedName(TEMPLATE_SCHEMA_NAME, attr_text=TEMPLATE_SCHEMA_VERSION))
        sheet.freeze_panes = "A2"
        for column, width in zip("ABCDEFG", (18, 48, 14, 16, 22, 24, 24)):
            sheet.column_dimensions[column].width = width
        instruction = workbook.create_sheet("Инструкция")
        for line in (
            ("Шаблон Averon v1",),
            ("Заполните Наименование и Кол-во для каждой позиции.",),
            ("Ед. изм. обязательна для автоматического расчёта стоимости; коэффициенты упаковок не предполагаются.",),
            ("Код ресурса — ссылка на исходную ведомость, это не артикул производителя.",),
            ("Артикул, Производитель и Модель / тип заполняйте только при наличии данных.",),
            ("Одна строка таблицы — одна позиция; не объединяйте ячейки внутри таблицы.",),
            ("После подбора Averon в будущем добавит отдельные колонки цены и общей стоимости.",),
            ("Не добавляйте формулы, внешние ссылки или макросы. Подготовлено до 500 строк.",),
        ):
            instruction.append(list(line))
        instruction.column_dimensions["A"].width = 110
        instruction.freeze_panes = "A2"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def bytes(self) -> bytes:
        return self._build()


__all__ = ["TenderTemplateService", "TEMPLATE_SHEET", "TEMPLATE_TABLE", "TEMPLATE_HEADERS", "TEMPLATE_SCHEMA_NAME", "TEMPLATE_SCHEMA_VERSION", "PREPARED_ROWS"]
