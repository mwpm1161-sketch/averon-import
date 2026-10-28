"""Explicitly experimental XLSX for local AI-first estimator pilots.

This exporter is intentionally separate from the production export service.
The ordinary pilot workbook is estimator-facing; the diagnostic workbook is
an administrator-only audit artifact.
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import json
import re
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from averon_import.services.ai_first.contracts import AIDocumentDraft, AIReviewState


MAX_PILOT_EXPORT_PAGES = 3
MAX_EXCEL_CELL_CHARS = 32767
MAX_EXCEL_COLUMNS = 16384
MAX_EXCEL_ROWS = 1_048_576
LONG_TEXT_CHUNK_CHARS = 30000
LONG_TEXT_SHEET_NAME = "Длинные значения"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ROLE_LABELS = {
    "header": "Заголовок",
    "column_numbering": "Нумерация колонок",
    "section": "Раздел",
    "standalone_item": "Самостоятельная позиция",
    "component": "Компонент комплекта",
    "continuation": "Продолжение строки",
    "note": "Примечание",
    "unknown": "Не определена",
}
DISPOSITION_LABELS = {
    "confirmed": "Подтверждена",
    "corrected": "Исправлена и подтверждена",
    "excluded": "Исключена",
}
CANONICAL_FIELDS = (
    ("position_number", "№ позиции"),
    ("name", "Наименование"),
    ("brand", "Марка"),
    ("manufacturer", "Производитель"),
    ("article", "Артикул / обозначение"),
    ("unit", "Единица измерения"),
    ("quantity", "Количество"),
    ("notes", "Примечания"),
)

_HEADER_ALIASES = {
    "position_number": {"номер", "номер позиции", "поз", "позиция", "№ п п", "п п"},
    "name": {"наименование", "название", "наименование продукции", "наименование оборудования", "наименование материала"},
    "brand": {"марка", "марка тип", "марка оборудования", "марка материала"},
    "manufacturer": {"производитель", "изготовитель", "производитель оборудования"},
    "article": {"артикул", "код", "обозначение", "артикул обозначение"},
    "unit": {"ед изм", "единица измерения", "единицы измерения", "ед", "ед изм"},
    "quantity": {"количество", "кол во", "кол", "qty"},
    "notes": {"примечание", "примечания", "комментарий", "особые указания"},
}


class PilotExportError(ValueError):
    """Safe pilot-export rejection code; never contains document data."""


def _header_key(value: object) -> str:
    text = str(value or "").casefold().replace("ё", "е").replace("№", " номер ")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return " ".join(text.split())


def _canonical_field(header: object) -> str | None:
    key = _header_key(header)
    for field, aliases in _HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def _json_text(value: Any) -> str:
    return json.dumps(_redact_secret_fields(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


_DIAGNOSTIC_SECRET_KEY = re.compile(r"(?:api[_-]?key|authorization|password|secret|credential|access[_-]?token|refresh[_-]?token)", re.IGNORECASE)


def _redact_secret_fields(value: Any) -> Any:
    """Redact operational credentials from admin diagnostic JSON projections."""
    if isinstance(value, Mapping):
        return {
            str(key): "[REDACTED]" if _DIAGNOSTIC_SECRET_KEY.search(str(key)) else _redact_secret_fields(child)
            for key, child in value.items()
        }
    if isinstance(value, (list, tuple)):
        return [_redact_secret_fields(child) for child in value]
    return value


def _append_export_row(sheet, values: Sequence[object], long_values: list[dict[str, Any]] | None) -> None:
    """Append text safely, spilling oversized values into an exact-recovery sheet."""
    row_values: list[str] = []
    pending: list[tuple[dict[str, Any], int]] = []
    for column_number, value in enumerate(values, start=1):
        text = "" if value is None else str(value)
        if len(text) <= MAX_EXCEL_CELL_CHARS:
            row_values.append(text)
            continue
        if long_values is None:
            raise PilotExportError("export_cell_limit_exceeded")

        chunks = [text[index:index + LONG_TEXT_CHUNK_CHARS] for index in range(0, len(text), LONG_TEXT_CHUNK_CHARS)]
        first_chunk_row = 5 + sum(len(item["chunks"]) for item in long_values)
        if first_chunk_row + len(chunks) - 1 > MAX_EXCEL_ROWS:
            raise PilotExportError("export_long_text_overflow_exceeded")
        reference = len(long_values) + 1
        item = {
            "reference": reference,
            "source_sheet": sheet.title,
            "source_cell": None,
            "chunks": chunks,
            "first_chunk_row": first_chunk_row,
            "cell": None,
        }
        long_values.append(item)
        pending.append((item, column_number))
        row_values.append(f"Длинный текст · запись {reference} на листе «{LONG_TEXT_SHEET_NAME}»")

    sheet.append(row_values)
    for item, column_number in pending:
        cell = sheet.cell(sheet.max_row, column_number)
        item["source_cell"] = cell.coordinate
        item["cell"] = cell


def _safe_filename(value: object) -> str:
    # The workbook is downloaded by the authenticated local user; never expose a path.
    text = str(value or "document.pdf").replace("\\", "/").split("/")[-1]
    return text[:255] or "document.pdf"


def _recognition_source_label(report: Mapping[str, Any]) -> str:
    adapter = report.get("adapter") if isinstance(report, Mapping) else None
    if not isinstance(adapter, Mapping):
        return "Источник распознавания не определён"
    provider_id = adapter.get("provider_id")
    model_id = adapter.get("model_id")
    if provider_id == "yandex_ai_studio" and model_id == "qwen_vision":
        return "Qwen Vision"
    if provider_id == "offline_synthetic" and model_id == "fake_vision":
        return "Синтетический тест"
    return "Источник распознавания не определён"


def _issue_index(draft: AIDocumentDraft) -> dict[int, list[Any]]:
    by_page: dict[int, list[Any]] = defaultdict(list)
    seen: set[str] = set()
    for issue in [*draft.issues, *(item for page in draft.pages for item in page.issues)]:
        if issue.id in seen:
            continue
        seen.add(issue.id)
        if issue.page_number is not None:
            by_page[issue.page_number].append(issue)
        else:
            for page_number in draft.expected_page_numbers:
                by_page[page_number].append(issue)
    return by_page


def _active_lineage_sources(review: AIReviewState) -> set[str]:
    manual_by_id = {row.id: row for row in review.manual_rows}
    pending = [
        source_id
        for row in review.manual_rows if row.active
        for source_id in row.source_record_ids
    ]
    result: set[str] = set()
    while pending:
        record_id = pending.pop()
        if record_id in result:
            continue
        result.add(record_id)
        ancestor = manual_by_id.get(record_id)
        if ancestor is not None:
            pending.extend(ancestor.source_record_ids)
    return result


def _rows_for_document(draft: AIDocumentDraft, review: AIReviewState) -> list[dict[str, Any]]:
    role_records = review.row_roles
    decisions = review.decisions
    manual_ids = {item.id for item in review.manual_rows}
    superseded = _active_lineage_sources(review)
    result: list[dict[str, Any]] = []
    for page in draft.pages:
        for table in page.tables:
            # Keep the exact (including blank) source header. A display fallback
            # would erase the distinction between an empty header and a real one.
            columns = {item.index: item.raw_header or "" for item in table.columns}
            for row in table.rows:
                cells = sorted(row.raw_cells, key=lambda item: item.column_index)
                role = role_records.get(row.id)
                decision = decisions.get(row.id)
                result.append({
                    "id": row.id,
                    "page": page.page_number,
                    "page_status": page.status,
                    "table_id": table.id,
                    "table_index": table.index + 1,
                    "row_index": row.index + 1,
                    "model_kind": row.model_kind or "",
                    "role": role.role if role and role.active else "unknown",
                    "role_active": bool(role and role.active),
                    "parent_record_id": role.parent_record_id if role and role.active else "",
                    "decision": decision,
                    "active": True,
                    "manual": False,
                    "superseded": row.id in superseded,
                    "row": row,
                    "cells": cells,
                    "columns": columns,
                    "source_values": [
                        {"column_index": cell.column_index, "header": columns.get(cell.column_index, f"Столбец {cell.column_index + 1}"), "value": cell.raw_value or ""}
                        for cell in cells
                    ],
                    "current_values": [
                        {"column_index": cell.column_index, "header": columns.get(cell.column_index, ""), "value": (decision.values.get(f"cell:{cell.column_index}", cell.raw_value or "") if decision and (decision.active or decision.invalidation_reason == "role_change") else cell.raw_value or "")}
                        for cell in cells
                    ],
                    "last_saved_values": dict(decision.values) if decision else {},
                    "evidence": {
                        # Keep the per-row worksheet below Excel's 32,767 character
                        # cell limit. The complete, lossless page payload (including
                        # raw model fields) is chunked in "Исходная структура".
                        "row": {
                            "id": row.id,
                            "index": row.index,
                            "model_kind": row.model_kind,
                            "server_role_candidate": row.server_role_candidate,
                            "role_verified": row.role_verified,
                            "raw_cell_ids": [cell.id for cell in cells],
                            "struck_through": row.struck_through,
                            "handwritten": row.handwritten,
                            "continuation_of": row.continuation_of,
                            "continues_previous_page": row.continues_previous_page,
                            "bbox_candidate": row.bbox_candidate,
                        },
                        "table": {
                            "id": table.id,
                            "index": table.index,
                            "columns": [
                                {
                                    "id": item.id,
                                    "index": item.index,
                                    "raw_header": item.raw_header,
                                    "canonical_field_candidate": item.canonical_field_candidate,
                                }
                                for item in table.columns
                            ],
                            "continues_previous_page": table.continues_previous_page,
                            "bbox_candidate": table.bbox_candidate,
                        },
                        "full_raw_payload_sheet": "Исходная структура",
                    },
                })
    default_headers = ["Поз.", "Наименование", "Артикул", "Ед. изм.", "Количество", "Примечание"]
    for manual in review.manual_rows:
        if manual.page_number not in {page.page_number for page in draft.pages}:
            continue
        decision = decisions.get(manual.id)
        values = [
            {"column_index": index, "header": default_headers[index] if index < len(default_headers) else f"Столбец {index + 1}", "value": value}
            for index, value in sorted(
                ((int(key.split(":", 1)[1]), value) for key, value in manual.values.items() if key.startswith("cell:") and key.split(":", 1)[1].isdigit()),
                key=lambda item: item[0],
            )
        ]
        role = role_records.get(manual.id)
        page = next(item for item in draft.pages if item.page_number == manual.page_number)
        result.append({
            "id": manual.id,
            "page": manual.page_number,
            "page_status": page.status,
            "table_id": "",
            "table_index": "",
            "row_index": "",
            "model_kind": "manual_row",
            "role": role.role if role and role.active else "unknown",
            "role_active": bool(role and role.active),
            "parent_record_id": role.parent_record_id if role and role.active else "",
            "decision": decision,
            "active": manual.active,
            "manual": True,
            "superseded": manual.id in superseded,
            "row": manual,
            "cells": [],
            "columns": {item["column_index"]: item["header"] for item in values},
            "source_values": values,
            "current_values": [
                {**item, "value": (decision.values.get(f"cell:{item['column_index']}", item["value"]) if decision and (decision.active or decision.invalidation_reason == "role_change") else item["value"])}
                for item in values
            ],
            "last_saved_values": dict(decision.values) if decision else {},
            "evidence": {"manual_row": manual.model_dump(mode="json")},
        })
    return result


def _page_report_states(report: Mapping[str, Any]) -> dict[int, str]:
    result: dict[int, str] = {}
    for item in report.get("pages", []) if isinstance(report.get("pages", []), list) else []:
        if isinstance(item, Mapping) and type(item.get("page_number")) is int:
            result[item["page_number"]] = str(item.get("state") or "unknown")
    return result


def _page_report_errors(report: Mapping[str, Any]) -> dict[int, str]:
    result: dict[int, str] = {}
    for item in report.get("pages", []) if isinstance(report.get("pages", []), list) else []:
        if isinstance(item, Mapping) and type(item.get("page_number")) is int:
            attempts = item.get("attempts") if isinstance(item.get("attempts"), list) else []
            last_attempt = attempts[-1] if attempts and isinstance(attempts[-1], Mapping) else {}
            code = last_attempt.get("safe_error_code") or item.get("safe_error_code") or ""
            result[item["page_number"]] = str(code)
    return result


def _row_issue_open(row: Mapping[str, Any], issue_index: Mapping[int, list[Any]], review: AIReviewState) -> bool:
    record_id = str(row["id"])
    evidence_ids = {record_id}
    table_id = row.get("table_id")
    if table_id:
        evidence_ids.add(str(table_id))
    evidence_ids.update(cell.id for cell in row.get("cells", []))
    for issue in issue_index.get(int(row["page"]), []):
        if issue.severity not in {"review", "blocking"} or issue.id in review.resolved_issue_ids:
            continue
        if not issue.record_ids or evidence_ids.intersection(issue.record_ids):
            return True
    for cell in row.get("cells", []):
        uncertain = bool(cell.unreadable or cell.struck_through or cell.handwritten or cell.alternatives or cell.handwritten_replacements)
        if not uncertain:
            continue
        scope_ids = {cell.id, record_id}
        if table_id:
            scope_ids.add(str(table_id))
        related_issue_ids = {
            issue.id for issue in issue_index.get(int(row["page"]), [])
            if issue.severity in {"review", "blocking"}
            and (not issue.record_ids or scope_ids.intersection(issue.record_ids))
        }
        if not related_issue_ids or any(issue_id not in review.resolved_issue_ids for issue_id in related_issue_ids):
            return True
    return False


def _values_by_field(row: Mapping[str, Any]) -> dict[str, list[dict[str, Any]]]:
    labels = row["columns"]
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current = {item["column_index"]: item["value"] for item in row["current_values"]}
    for item in row["source_values"]:
        field = _canonical_field(labels.get(item["column_index"], item["header"]))
        if field:
            grouped[field].append({"raw": item["value"], "current": current.get(item["column_index"], item["value"])})
    return grouped


def _review_status(row: Mapping[str, Any], complete_page: bool, processing_state: str, open_issue: bool) -> str:
    decision = row.get("decision")
    if row.get("superseded"):
        return "Заменена структурным решением"
    if not row.get("active"):
        return "Неактивная ручная запись"
    if not row.get("role_active"):
        return "Роль строки не проверена"
    if not decision or not decision.active:
        return "Значения требуют проверки"
    if decision.disposition == "excluded":
        return "Исключена из проверяемой спецификации"
    if not decision.human_confirmed:
        return "Исправления сохранены, но не подтверждены"
    if not complete_page:
        return "Страница не подтверждена полностью"
    if processing_state != "completed":
        return "Обработка страницы не завершена"
    if open_issue:
        return "Есть неразрешённое замечание"
    if row.get("ambiguous_field_mapping"):
        return "Неоднозначное сопоставление колонок"
    if row["role"] != "standalone_item":
        return "Не является самостоятельной позицией"
    return DISPOSITION_LABELS.get(decision.disposition, "Проверена")


def _prepare_pilot_export_rows(
    draft: AIDocumentDraft,
    review: AIReviewState,
    report: Mapping[str, Any],
    page_numbers: Sequence[int],
    document_name: str,
) -> dict[str, Any]:
    if not page_numbers or len(page_numbers) > MAX_PILOT_EXPORT_PAGES:
        raise PilotExportError("invalid_page_scope")
    if any(type(number) is not int for number in page_numbers) or len(set(page_numbers)) != len(page_numbers):
        raise PilotExportError("invalid_page_scope")
    selected = sorted(page_numbers)
    expected = set(draft.expected_page_numbers)
    if not set(selected) <= expected:
        raise PilotExportError("page_outside_document")
    if draft.source_document_id != review.source_document_id or draft.revision_id != review.revision_id:
        raise PilotExportError("review_state_revision_mismatch")

    page_count = len(draft.expected_page_numbers)
    report_states = _page_report_states(report)
    confirmed_pages = set(review.confirmed_complete_pages)
    pages_by_number = {page.page_number: page for page in draft.pages}
    issue_index = _issue_index(draft)
    all_rows = [row for row in _rows_for_document(draft, review) if row["page"] in selected]
    selected_rows: list[dict[str, Any]] = []
    for row in all_rows:
        grouped = _values_by_field(row)
        row["ambiguous_field_mapping"] = any(len(values) > 1 for values in grouped.values())
        row["processing_state"] = report_states.get(int(row["page"]), "unknown")
        row["page_confirmed"] = int(row["page"]) in confirmed_pages
        row["open_issue"] = _row_issue_open(row, issue_index, review)
        row["review_status"] = _review_status(row, row["page_confirmed"], row["processing_state"], row["open_issue"])
        page = pages_by_number.get(int(row["page"]))
        row["verified_position"] = bool(
            row["active"] and not row["superseded"] and row["role_active"]
            and row["role"] == "standalone_item"
            and row["decision"] is not None and row["decision"].active
            and row["decision"].disposition in {"confirmed", "corrected"}
            and row["decision"].human_confirmed
            and row["page_confirmed"] and row["processing_state"] == "completed"
            and page is not None and page.status == "extracted" and not row["open_issue"]
            and not row["ambiguous_field_mapping"]
        )
        reasons: list[str] = []
        if row["superseded"]:
            reasons.append("superseded")
        if not row["active"]:
            reasons.append("inactive_row")
        if not row["role_active"]:
            reasons.append("role_unconfirmed")
        elif row["role"] != "standalone_item":
            reasons.append("role_not_standalone")
        if row["decision"] is None or not row["decision"].active:
            reasons.append("values_unconfirmed")
        elif not row["decision"].human_confirmed:
            reasons.append("values_not_confirmed")
        if not row["page_confirmed"]:
            reasons.append("page_completeness_unconfirmed")
        if row["processing_state"] != "completed" or page is None or page.status != "extracted":
            reasons.append("page_processing_incomplete")
        if row["open_issue"]:
            reasons.append("open_issue")
        if row["ambiguous_field_mapping"]:
            reasons.append("ambiguous_column_mapping")
        row["export_reason_codes"] = reasons
        if row["verified_position"]:
            selected_rows.append(row)
    return {
        "selected": selected,
        "page_count": page_count,
        "filename": _safe_filename(document_name),
        "report_states": report_states,
        "confirmed_pages": confirmed_pages,
        "pages_by_number": pages_by_number,
        "issue_index": issue_index,
        "all_rows": all_rows,
        "selected_rows": selected_rows,
        "export_reasons": {
            str(row["id"]): list(row["export_reason_codes"])
            for row in all_rows if not row["verified_position"]
        },
    }


def preview_ai_first_pilot_export(
    draft: AIDocumentDraft,
    review: AIReviewState,
    report: Mapping[str, Any],
    *,
    page_numbers: Sequence[int],
    document_name: str,
) -> dict[str, Any]:
    """Return content-free, backend-derived counts for a selected test fragment."""
    prepared = _prepare_pilot_export_rows(draft, review, report, page_numbers, document_name)
    all_rows = prepared["all_rows"]
    selected_rows = prepared["selected_rows"]
    ocr_rows = [row for row in all_rows if not row["manual"]]
    manual_rows = [row for row in all_rows if row["manual"]]
    return {
        "page_numbers": prepared["selected"],
        "document_page_count": prepared["page_count"],
        "ocr_row_count": len(ocr_rows),
        "manual_row_count": len(manual_rows),
        "total_row_count": len(all_rows),
        "verified_position_count": len(selected_rows),
        "preview_checks_admission_only": True,
        "xlsx_technical_limits_checked": False,
        "blocked_rows": [
            {"page_number": row["page"], "source_row_number": row["row_index"],
             "reason_codes": list(row["export_reason_codes"])}
            for row in all_rows if not row["verified_position"]
        ],
        "pages": [
            {
                "page_number": number,
                "ocr_row_count": sum(row["page"] == number and not row["manual"] for row in all_rows),
                "manual_row_count": sum(row["page"] == number and row["manual"] for row in all_rows),
                "total_row_count": sum(row["page"] == number for row in all_rows),
                "verified_position_count": sum(row["page"] == number for row in selected_rows),
            }
            for number in prepared["selected"]
        ],
    }


def _write_sheet(
    sheet, title: str, document_name: str, document_id: str,
    selected_pages: Sequence[int], page_count: int,
    header: Sequence[str], rows: Iterable[Sequence[object]], widths: Sequence[int] | None = None,
    detail_header: Sequence[object] | None = None,
    *, source_label: str = "Источник распознавания не определён",
    long_values: list[dict[str, Any]] | None = None,
) -> None:
    coverage = "Полный охват страниц документа" if set(selected_pages) == set(range(1, page_count + 1)) else "ЧАСТИЧНЫЙ ОХВАТ ДОКУМЕНТА"
    sheet.sheet_view.showGridLines = False
    _append_export_row(sheet, [f"Служебный лист · AI-first · {source_label}"], long_values)
    _append_export_row(sheet, [f"Документ: {document_name} · ID {document_id} · страницы {', '.join(map(str, selected_pages))} из {page_count} · {coverage}"], long_values)
    _append_export_row(sheet, [title], long_values)
    _append_export_row(sheet, header, long_values)
    if detail_header is not None:
        _append_export_row(sheet, detail_header, long_values)
    for values in rows:
        _append_export_row(sheet, values, long_values)
    thin_border = Border(bottom=Side(style="hair", color="D7E0E7"))
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                value = str(cell.value)
                if len(value) > MAX_EXCEL_CELL_CHARS:
                    raise PilotExportError("export_cell_limit_exceeded")
                cell.value = value
                cell.data_type = "s"  # OCR text can never become an Excel formula.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="EEF3F6")
                cell.font = Font(name="Aptos", size=10, bold=True, color="425466")
            elif cell.row == 3:
                cell.font = Font(name="Aptos", size=12, bold=True, color="263746")
            elif cell.row == 4:
                cell.fill = PatternFill("solid", fgColor="1D4ED8")
                cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            elif detail_header is not None and cell.row == 5:
                cell.fill = PatternFill("solid", fgColor="DBEAFE")
                cell.font = Font(name="Aptos", size=9, bold=True, color="1E3A8A")
            else:
                cell.font = Font(name="Aptos", size=9, color="263746")
            if cell.row >= (6 if detail_header is not None else 5):
                cell.border = thin_border
        if row[0].row >= (6 if detail_header is not None else 5):
            sheet.row_dimensions[row[0].row].height = 42
    sheet.freeze_panes = "A6" if detail_header is not None else "A5"
    if detail_header is None and sheet.max_row >= 4 and sheet.max_column:
        sheet.auto_filter.ref = f"A4:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    if widths:
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    else:
        for index in range(1, sheet.max_column + 1):
            heading = sheet.cell(4, index).value or ""
            source_heading = sheet.cell(5, index).value or "" if detail_header is not None else ""
            sheet.column_dimensions[get_column_letter(index)].width = min(42, max(14, len(str(heading)) + 2, min(40, len(str(source_heading)) + 2)))


def _specification_width(header: object, values: Sequence[object]) -> int:
    header_lines = str(header or "").splitlines() or [""]
    longest = max((len(line) for line in header_lines), default=0)
    for value in values[:50]:
        lines = str(value or "").splitlines() or [""]
        longest = max(longest, min(max(map(len, lines), default=0), 52))
    return min(48, max(13, longest + 3))


def _write_specification_sheet(
    sheet,
    *,
    columns: Sequence[tuple[int, str]],
    rows: Sequence[dict[str, Any]],
    title: str,
    metadata: str,
    scope_label: str,
    show_scope_mark: bool,
    long_values: list[dict[str, Any]],
) -> None:
    """Write only admitted rows using their exact source-table columns."""
    if len(columns) > MAX_EXCEL_COLUMNS:
        raise PilotExportError("export_column_limit_exceeded")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "39766A" if show_scope_mark else "6F9B90"
    column_count = max(len(columns), 1)
    _append_export_row(sheet, [scope_label if show_scope_mark else "Дополнительная таблица · тестовый режим"], long_values)
    _append_export_row(sheet, [metadata], long_values)
    _append_export_row(sheet, [title], long_values)
    _append_export_row(sheet, [header for _, header in columns] if columns else [""], long_values)

    for row in rows:
        current = {int(item["column_index"]): item["value"] for item in row["current_values"]}
        _append_export_row(sheet, [current.get(index, "") for index, _ in columns] if columns else [""], long_values)
    if not rows:
        _append_export_row(sheet, ["Нет позиций, допущенных в проверенную спецификацию."] + [""] * (column_count - 1), long_values)

    rule = Side(style="thin", color="DCE4E8")
    body_font = Font(name="Aptos", size=10, color="253746")
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.value = str(cell.value)
                cell.data_type = "s"
            cell.border = Border(bottom=rule)
            cell.alignment = Alignment(
                vertical="center" if cell.row <= 3 else "top",
                wrap_text=cell.row >= 4,
            )
            cell.font = body_font
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="EAF2F0")
                cell.font = Font(name="Aptos", size=10, bold=True, color="275C53")
            elif cell.row == 2:
                cell.font = Font(name="Aptos", size=9, color="5D6B75")
            elif cell.row == 3:
                cell.font = Font(name="Aptos", size=12, bold=True, color="263746")
            elif cell.row == 4:
                cell.fill = PatternFill("solid", fgColor="315F72")
                cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            elif cell.row == 5 and not rows:
                cell.fill = PatternFill("solid", fgColor="F2F5F6")
                cell.font = Font(name="Aptos", size=10, italic=True, color="52616B")
            elif cell.row >= 5:
                cell.font = body_font

    sheet.row_dimensions[1].height = 23
    sheet.row_dimensions[2].height = 20
    sheet.row_dimensions[3].height = 24
    header_lines = max((len(str(header or "").splitlines()) for _, header in columns), default=1)
    sheet.row_dimensions[4].height = min(58, max(32, 16 * header_lines + 8))
    for index, (source_index, header) in enumerate(columns, start=1):
        sample_values = [
            next((item["value"] for item in row["current_values"] if int(item["column_index"]) == source_index), "")
            for row in rows
        ]
        sheet.column_dimensions[get_column_letter(index)].width = _specification_width(header, sample_values)
    if not columns:
        sheet.column_dimensions["A"].width = 24
    if not rows:
        sheet.row_dimensions[5].height = 28
    else:
        for row_number, row in enumerate(rows, start=5):
            current = {int(item["column_index"]): item["value"] for item in row["current_values"]}
            max_lines = 1
            for index, _ in columns:
                width = sheet.column_dimensions[get_column_letter(index + 1)].width or 13
                text = str(current.get(index, "") or "")
                lines = sum(max(1, (len(part) + int(width) - 1) // int(width)) for part in text.splitlines() or [""])
                max_lines = max(max_lines, lines)
            sheet.row_dimensions[row_number].height = min(72, max(28, 14 * max_lines + 6))
    sheet.freeze_panes = "A5"
    sheet.print_title_rows = "1:4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    last_data_row = max(4, 4 + len(rows))
    sheet.auto_filter.ref = f"A4:{get_column_letter(column_count)}{last_data_row}"


def _source_sheet_name(page_number: int, table_index: int | str) -> str:
    suffix = "manual" if table_index == "manual" else f"T{int(table_index):03d}"
    return f"P{int(page_number):03d}-{suffix}"[:31]


def _source_table_sheet(
    workbook: Workbook,
    *,
    name: str,
    title: str,
    filename: str,
    document_id: str,
    selected: Sequence[int],
    page_count: int,
    columns: Sequence[tuple[int, str]],
    rows: Sequence[dict[str, Any]],
    source_label: str,
    long_values: list[dict[str, Any]],
) -> dict[str, int]:
    """Write one source table with exact headers and visible raw/reviewed cells."""
    if len(columns) * 2 + 6 > MAX_EXCEL_COLUMNS:
        raise PilotExportError("export_column_limit_exceeded")
    sheet = workbook.create_sheet(name)
    fixed = ["Строка источника", "ID строки", "Тип по модели", "Роль Human Review", "Статус роли", "Статус проверки"]
    header = list(fixed)
    detail_header: list[object] = ["", "", "", "", "", ""]
    for index, source_header in columns:
        label = source_header if source_header else "[без заголовка]"
        header.extend([f"Индекс {index} · исходно · {label}", f"Индекс {index} · проверено · {label}"])
        # Repeated/blank labels are retained verbatim in the next visible row;
        # column indexes in row 4 keep their identities distinct.
        detail_header.extend([source_header, source_header])
    data: list[list[object]] = []
    location: dict[str, int] = {}
    for offset, row in enumerate(rows, start=6):
        location[str(row["id"])] = offset
        source_values = {int(item["column_index"]): item["value"] for item in row["source_values"]}
        current_values = {int(item["column_index"]): item["value"] for item in row["current_values"]}
        role = ROLE_LABELS.get(row["role"], "Не определена") if row["role_active"] else "Не подтверждена человеком"
        values: list[object] = [
            "Ручная" if row.get("manual") else row.get("row_index", ""),
            row["id"], row["model_kind"], role,
            "Подтверждена человеком" if row["role_active"] else "Не подтверждена",
            row["review_status"],
        ]
        for index, _source_header in columns:
            values.extend([source_values.get(index, ""), current_values.get(index, source_values.get(index, ""))])
        data.append(values)
    _write_sheet(
        sheet, title, filename, document_id, selected, page_count,
        header, data, detail_header=detail_header,
        source_label=source_label, long_values=long_values,
    )
    return location


def build_ai_first_diagnostic_xlsx(
    draft: AIDocumentDraft,
    review: AIReviewState,
    report: Mapping[str, Any],
    *,
    page_numbers: Sequence[int],
    document_name: str,
) -> bytes:
    """Build a local, explicitly test-marked workbook from saved review state."""
    prepared = _prepare_pilot_export_rows(draft, review, report, page_numbers, document_name)
    selected = prepared["selected"]
    page_count = prepared["page_count"]
    expected = set(draft.expected_page_numbers)
    filename = prepared["filename"]
    report_states = prepared["report_states"]
    confirmed_pages = prepared["confirmed_pages"]
    pages_by_number = prepared["pages_by_number"]
    issue_index = prepared["issue_index"]
    all_rows = prepared["all_rows"]
    selected_rows = prepared["selected_rows"]
    source_label = _recognition_source_label(report)
    long_values: list[dict[str, Any]] = []

    workbook = Workbook()
    positions = workbook.active
    positions.title = "Спецификация"
    scope_label = "Тестовая выгрузка · Полный охват" if set(selected) == expected else "Тестовая выгрузка · Частичный охват"
    regions = draft.source_envelope.get("page_regions", {})
    if any(str(number) in regions for number in selected):
        scope_label = "ТЕСТ · Частичный охват · Выбранная область страницы"
    specification_groups: list[dict[str, Any]] = []

    # The user-facing specification keeps each physical table's own schema.
    # Only rows admitted by the existing Human Review checks are shown here.
    for page_number in selected:
        page = pages_by_number.get(page_number)
        if page is None:
            continue
        for table in sorted(page.tables, key=lambda item: (item.index, item.id)):
            table_rows = sorted(
                (row for row in selected_rows if not row["manual"] and row["table_id"] == table.id),
                key=lambda row: row["row_index"],
            )
            if not table_rows:
                continue
            source_headers = {column.index: column.raw_header or "" for column in table.columns}
            for row in table_rows:
                for item in row["current_values"]:
                    source_headers.setdefault(int(item["column_index"]), "")
            specification_groups.append({
                "page": page_number, "table_index": table.index + 1,
                "columns": sorted(source_headers.items()), "rows": table_rows,
            })

    for page_number in selected:
        manual_rows = sorted(
            (row for row in selected_rows if row["manual"] and row["page"] == page_number),
            key=lambda row: row["row_index"],
        )
        if not manual_rows:
            continue
        headers: dict[int, str] = {}
        for row in manual_rows:
            for item in row["source_values"]:
                headers.setdefault(int(item["column_index"]), str(item["header"] or ""))
        specification_groups.append({"page": page_number, "table_index": None, "columns": sorted(headers.items()), "rows": manual_rows})

    if not specification_groups:
        # Keep a useful, clearly empty first sheet without inventing data or
        # promoting unresolved rows into the verified specification.
        for page_number in selected:
            page = pages_by_number.get(page_number)
            if page and page.tables:
                table = sorted(page.tables, key=lambda item: (item.index, item.id))[0]
                headers = {column.index: column.raw_header or "" for column in table.columns}
                for row in all_rows:
                    if not row["manual"] and row["table_id"] == table.id:
                        for item in row["source_values"]:
                            headers.setdefault(int(item["column_index"]), "")
                specification_groups.append({"page": page_number, "table_index": table.index + 1, "columns": sorted(headers.items()), "rows": []})
                break
        if not specification_groups and all_rows:
            row = all_rows[0]
            headers = {int(item["column_index"]): str(item["header"] or "") for item in row["source_values"]}
            specification_groups.append({"page": row["page"], "table_index": None, "columns": sorted(headers.items()), "rows": []})

    for ordinal, group in enumerate(specification_groups, start=1):
        sheet = positions if ordinal == 1 else workbook.create_sheet(f"Спецификация {ordinal}")
        if ordinal == 1:
            metadata = f"Документ: {filename} · страница {group['page']}"
            if group["table_index"] is not None:
                metadata += f" · таблица {group['table_index']}"
            title = "Подтверждённые значения по исходной таблице"
        else:
            metadata = f"Продолжение · {filename} · страница {group['page']}"
            if group["table_index"] is not None:
                metadata += f" · таблица {group['table_index']}"
            title = "Подтверждённые значения по отдельной исходной таблице"
        _write_specification_sheet(
            sheet, columns=group["columns"], rows=group["rows"], title=title,
            metadata=metadata, scope_label=scope_label, show_scope_mark=(ordinal == 1),
            long_values=long_values,
        )

    information = workbook.create_sheet("Информация")
    information_rows = [
        ["Назначение", "Тестовая выгрузка для оценки AI-first распознавания; не коммерческая смета."],
        ["Охват", f"Страницы {', '.join(map(str, selected))} из {page_count}; {'полный' if set(selected) == expected else 'частичный'} охват документа."],
        ["Лист «Спецификация»", f"{len(selected_rows)} строк, прошедших прежние условия Human Review для проверенных самостоятельных позиций."],
        ["Другие строки", "Заголовки, разделы, компоненты, неизвестные роли и неподтверждённые строки сохранены в листах аудита и исходной структуры; роли не переопределялись."],
        ["Исправления", "На листах аудита отдельно сохранены исходные OCR-значения и текущие значения Human Review."],
        ["Ограничения", "Экспорт не меняет допуск строк, подтверждение полноты страниц или ограничения коммерческого экспорта и sourcing."],
    ]
    _write_sheet(
        information, "Информация о тестовой выгрузке", filename, draft.source_document_id,
        selected, page_count, ["Раздел", "Сведения"], information_rows,
        widths=[30, 100], source_label=source_label, long_values=long_values,
    )

    source_locations: dict[str, tuple[str, int]] = {}
    # Keep each physical table on its own worksheet. Heterogeneous tables never
    # share a synthetic set of columns, and source identity is page/table/index.
    for page_number in selected:
        page = pages_by_number.get(page_number)
        if page is None:
            continue
        ordered_tables = sorted(page.tables, key=lambda table: (table.index, table.id))
        for table_ordinal, table in enumerate(ordered_tables, start=1):
            name = _source_sheet_name(page_number, table_ordinal)
            table_rows = sorted(
                (row for row in all_rows if not row["manual"] and row["table_id"] == table.id),
                key=lambda row: row["row_index"],
            )
            source_headers = {column.index: column.raw_header or "" for column in table.columns}
            # Preserve a cell even if a malformed/incomplete extraction omitted
            # its column descriptor; identify it by source index and blank header.
            for row in table_rows:
                for item in row["source_values"]:
                    source_headers.setdefault(int(item["column_index"]), "")
            columns = sorted(source_headers.items())
            locations = _source_table_sheet(
                workbook, name=name, title=f"Исходные и проверенные значения · страница {page_number}, таблица {table.index + 1} · ID {table.id}",
                filename=filename, document_id=draft.source_document_id, selected=selected, page_count=page_count,
                columns=columns, rows=table_rows, source_label=source_label, long_values=long_values,
            )
            source_locations.update({record_id: (name, row_number) for record_id, row_number in locations.items()})

    # Manual rows have their own per-page sheet, separate from model tables.
    for page_number in selected:
        manual_rows = [row for row in all_rows if row["manual"] and row["page"] == page_number]
        if not manual_rows:
            continue
        by_index: dict[int, str] = {}
        for row in manual_rows:
            for item in row["source_values"]:
                by_index.setdefault(int(item["column_index"]), str(item["header"] or ""))
        name = _source_sheet_name(page_number, "manual")
        locations = _source_table_sheet(
            workbook, name=name, title=f"Ручные строки · страница {page_number}",
            filename=filename, document_id=draft.source_document_id, selected=selected, page_count=page_count,
            columns=sorted(by_index.items()), rows=manual_rows,
            source_label=source_label, long_values=long_values,
        )
        source_locations.update({record_id: (name, row_number) for record_id, row_number in locations.items()})

    all_sheet = workbook.create_sheet("Все распознанные строки")
    all_header = [
        "ТЕСТ", "Документ", "ID документа", "Страница", "Таблица", "Строка источника", "Тип по модели",
        "Роль Human Review", "Родитель компонента", "Статус проверки", "Статус обработки страницы",
        "Полнота страницы подтверждена", "Решение по значениям", "Лист исходной таблицы",
        "Строка на листе исходной таблицы", "Последние сохранённые значения Human Review (JSON)",
        "Последнее решение Human Review (JSON)", "Исходная строка и таблица (JSON)", "ID строки",
    ]
    all_data = []
    for row in all_rows:
        decision = row["decision"]
        source_sheet, source_row = source_locations.get(row["id"], ("", ""))
        all_data.append([
            "ТЕСТ", filename, draft.source_document_id, row["page"], row["table_index"], row["row_index"], row["model_kind"],
            ROLE_LABELS.get(row["role"], "Не определена") if row["role_active"] else "Не подтверждена человеком",
            row["parent_record_id"], row["review_status"], row["processing_state"],
            "Да" if row["page_confirmed"] else "Нет", DISPOSITION_LABELS.get(decision.disposition, "") if decision and decision.active else "Нет активного решения",
            source_sheet, source_row, _json_text(row["last_saved_values"]),
            _json_text(decision.model_dump(mode="json") if decision else {}),
            _json_text(row["evidence"]), row["id"],
        ])
    _write_sheet(
        all_sheet, "Все распознанные и добавленные строки выбранных страниц", filename,
        draft.source_document_id, selected, page_count, all_header, all_data,
        source_label=source_label, long_values=long_values,
    )
    all_headers = {all_sheet.cell(4, column).value: column for column in range(1, all_sheet.max_column + 1)}
    for row_number in range(5, all_sheet.max_row + 1):
        target_name = all_sheet.cell(row_number, all_headers["Лист исходной таблицы"]).value
        target_row = all_sheet.cell(row_number, all_headers["Строка на листе исходной таблицы"]).value
        if target_name and target_row:
            link = all_sheet.cell(row_number, all_headers["Лист исходной таблицы"])
            link.hyperlink = f"#'{target_name}'!A{target_row}"
            link.style = "Hyperlink"

    coverage_sheet = workbook.create_sheet("Охват страниц")
    page_errors = _page_report_errors(report)
    coverage_header = ["Страница", "Статус обработки", "Безопасный код ошибки", "Результат извлечения", "Полнота подтверждена", "Строк в результате", "Проверенных самостоятельных позиций", "Неразрешённые замечания", "Охват документа"]
    coverage_data = []
    for page_number in selected:
        page = pages_by_number.get(page_number)
        page_rows = [row for row in all_rows if row["page"] == page_number]
        open_page_issues = [
            issue for issue in issue_index.get(page_number, [])
            if issue.severity in {"review", "blocking"} and issue.id not in review.resolved_issue_ids
        ]
        coverage_data.append([
            page_number, report_states.get(page_number, "не запускалась"), page_errors.get(page_number, ""), page.status if page else "нет результата",
            "Да" if page_number in confirmed_pages else "Нет", len(page_rows),
            sum(bool(row["verified_position"]) for row in page_rows), len(open_page_issues),
            "Только выбранная область" if str(page_number) in regions else "Полный" if set(selected) == expected else "Частичный",
        ])
    _write_sheet(
        coverage_sheet, "Состояние обработки выбранного фрагмента", filename,
        draft.source_document_id, selected, page_count, coverage_header, coverage_data,
        source_label=source_label, long_values=long_values,
    )

    source_sheet = workbook.create_sheet("Исходная структура")
    source_header = ["Страница", "Часть", "Всего частей", "JSON исходного результата (объединить части по порядку без разделителей)"]
    source_rows = []
    for page_number in selected:
        page = pages_by_number.get(page_number)
        if page is None:
            continue
        encoded = _json_text(page.model_dump(mode="json"))
        chunks = [encoded[index:index + 30000] for index in range(0, len(encoded), 30000)] or [""]
        source_rows.extend([
            page_number, index + 1, len(chunks), chunk,
        ] for index, chunk in enumerate(chunks))
    _write_sheet(
        source_sheet, "Полная исходная структура выбранных страниц", filename,
        draft.source_document_id, selected, page_count, source_header, source_rows,
        widths=[12, 12, 14, 100], source_label=source_label, long_values=long_values,
    )

    if regions:
        region_sheet = workbook.create_sheet("Области страниц")
        points = draft.source_envelope.get("page_region_pdf_points", {})
        _write_sheet(region_sheet, "Обработаны только указанные области, не страницы целиком", filename,
            draft.source_document_id, selected, page_count,
            ["Страница", "Область · нормализованные x, y, ширина, высота", "Область PDF · x0, y0, x1, y1"],
            [[number, _json_text(regions[str(number)]), _json_text(points.get(str(number)))]
             for number in selected if str(number) in regions], source_label=source_label, long_values=long_values)
    if report.get("saved_editor_draft"):
        editor_sheet = workbook.create_sheet("Черновик редактора")
        encoded = _json_text(report["saved_editor_draft"])
        chunks = [encoded[i:i + 30000] for i in range(0, len(encoded), 30000)]
        _write_sheet(editor_sheet, "Сохранённые промежуточные правки и история; не подтверждение", filename,
            draft.source_document_id, selected, page_count, ["Часть", "Всего", "JSON · соединить без разделителей"],
            [[i + 1, len(chunks), part] for i, part in enumerate(chunks)], source_label=source_label, long_values=long_values)

    # A separate administrative run sheet uses an explicit allowlist. It never
    # serializes the full report/config/adapter objects or checkpoint hashes.
    adapter = report.get("adapter") if isinstance(report.get("adapter"), Mapping) else {}
    config = report.get("config") if isinstance(report.get("config"), Mapping) else {}
    pricing = config.get("pricing") if isinstance(config.get("pricing"), Mapping) else {}
    run_rows: list[list[object]] = [
        ["Состояние запуска", report.get("status", "Неизвестно")],
        ["Провайдер", adapter.get("provider_id", "Неизвестно")],
        ["Модель", adapter.get("model_id", "Неизвестно")],
        ["Версия модели", adapter.get("model_version", "Неизвестно")],
        ["Версия промпта", adapter.get("prompt_version", "Неизвестно")],
        ["Версия схемы", adapter.get("schema_version", "Неизвестно")],
        ["Целостность исходного PDF", "Подтверждена" if report.get("source_integrity_verified") is True else "Не подтверждена"],
        ["Ожидаемых страниц", report.get("expected_page_count", "Не сообщено")],
        ["Время обработки, с", report.get("elapsed_s", "Не сообщено")],
        ["Фактические входные токены", report.get("actual_input_tokens") if report.get("actual_input_tokens") is not None else "Не сообщено"],
        ["Фактические выходные токены", report.get("actual_output_tokens") if report.get("actual_output_tokens") is not None else "Не сообщено"],
        ["Reasoning tokens", report.get("actual_reasoning_tokens") if report.get("actual_reasoning_tokens") is not None else "Не сообщено"],
        ["Фактическая расчётная стоимость, ₽", report.get("actual_cost_rub") if report.get("actual_cost_rub") is not None else "Не рассчитана"],
        ["Зарезервированная стоимость неизвестных запросов, ₽", report.get("reserved_unknown_cost_rub", "Не сообщено")],
        ["Параллельность", config.get("concurrency", "Не сообщено")],
        ["DPI", config.get("dpi", "Не сообщено")],
        ["Локальный лимит RPM", config.get("requests_per_minute", "Не сообщено")],
        ["Локальный лимит TPM", config.get("tokens_per_minute", "Не сообщено")],
        ["Локальный бюджет, ₽", config.get("max_budget_rub", "Не сообщено")],
        ["Тариф входа за 1000 токенов, ₽", pricing.get("input_rub_per_1k", "Не сообщено")],
        ["Тариф кешированного входа за 1000 токенов, ₽", pricing.get("cached_input_rub_per_1k", "Не сообщено")],
        ["Тариф выхода за 1000 токенов, ₽", pricing.get("output_rub_per_1k", "Не сообщено")],
    ]
    attempt_rows: list[list[object]] = []
    for page_record in report.get("pages", []) if isinstance(report.get("pages"), list) else []:
        if not isinstance(page_record, Mapping):
            continue
        attempts = page_record.get("attempts") if isinstance(page_record.get("attempts"), list) else []
        if not attempts:
            attempt_rows.append([page_record.get("page_number", ""), "", page_record.get("state", ""), "", "", "", "", ""])
            continue
        for attempt_index, attempt in enumerate(attempts, start=1):
            if not isinstance(attempt, Mapping):
                continue
            usage = attempt.get("usage") if isinstance(attempt.get("usage"), Mapping) else {}
            attempt_rows.append([
                page_record.get("page_number", ""), attempt_index, attempt.get("state", ""),
                attempt.get("http_status", ""), attempt.get("finish_reason", ""),
                usage.get("input_tokens", "Не сообщено"), usage.get("output_tokens", "Не сообщено"),
                usage.get("reasoning_tokens", "Не сообщено"),
                attempt.get("elapsed_s", "Не сообщено"), attempt.get("request_latency_ms", "Не сообщено"),
                attempt.get("actual_cost_rub", "Не рассчитана"), attempt.get("safe_error_code", ""),
            ])
    run_sheet = workbook.create_sheet("Диагностика запуска")
    _write_sheet(
        run_sheet, "Разрешённые метаданные запуска", filename, draft.source_document_id,
        selected, page_count, ["Показатель", "Значение"], run_rows,
        widths=[48, 72], source_label=source_label, long_values=long_values,
    )
    attempt_sheet = workbook.create_sheet("Попытки страниц")
    _write_sheet(
        attempt_sheet, "Безопасные статусы попыток по страницам", filename, draft.source_document_id,
        selected, page_count,
        ["Страница", "Попытка", "Статус", "HTTP", "finish_reason", "Входные токены", "Выходные токены", "Reasoning", "Время, с", "Задержка, мс", "Стоимость, ₽", "Код ошибки"],
        attempt_rows, widths=[12, 12, 18, 10, 18, 18, 18, 18, 14, 16, 16, 28],
        source_label=source_label, long_values=long_values,
    )

    if long_values:
        long_sheet = workbook.create_sheet(LONG_TEXT_SHEET_NAME)
        long_rows: list[list[object]] = []
        for item in long_values:
            part_count = len(item["chunks"])
            long_rows.extend([
                item["reference"], item["source_sheet"], item["source_cell"], part_number, part_count, chunk,
            ] for part_number, chunk in enumerate(item["chunks"], start=1))
        _write_sheet(
            long_sheet,
            "Исходный текст по частям; соединяйте части записи без разделителей по порядку",
            filename, draft.source_document_id, selected, page_count,
            ["Запись", "Лист источника", "Ячейка источника", "Часть", "Всего частей", "Исходный текст · часть"],
            long_rows, widths=[12, 24, 18, 12, 16, 100], source_label=source_label,
        )
        for item in long_values:
            source_cell = item["cell"]
            source_cell.hyperlink = f"#'{LONG_TEXT_SHEET_NAME}'!A{item['first_chunk_row']}"
            source_cell.style = "Hyperlink"

    workbook.properties.title = f"ТЕСТ — AI-first {source_label}"
    workbook.properties.subject = "AI-first · административная диагностика выбранного фрагмента"
    stream = BytesIO()
    try:
        workbook.save(stream)
    except (ValueError, OSError):
        raise PilotExportError("pilot_workbook_generation_failed") from None
    return stream.getvalue()


def build_ai_first_pilot_xlsx(
    draft: AIDocumentDraft,
    review: AIReviewState,
    report: Mapping[str, Any],
    *,
    page_numbers: Sequence[int],
    document_name: str,
) -> bytes:
    """Build the estimator-facing workbook without a synthetic column template.

    Each physical source table is a separate block on the one ordinary sheet.
    Only rows admitted by the existing Human Review checks are emitted as
    positions. Notes remain in their own section and are explicitly marked if
    their role/value review is still outstanding.
    """
    prepared = _prepare_pilot_export_rows(draft, review, report, page_numbers, document_name)
    selected = prepared["selected"]
    positions_rows = prepared["selected_rows"]
    if not positions_rows:
        raise PilotExportError("no_verified_positions")

    sheet = Workbook().active
    sheet.title = "Спецификация"
    workbook = sheet.parent
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "39766A"
    long_values: list[dict[str, Any]] = []
    full_page_selection = set(selected) == set(draft.expected_page_numbers)
    all_pages_completed = all(prepared["report_states"].get(number) == "completed" for number in selected)
    marker = "ТЕСТ · Полный охват" if full_page_selection and all_pages_completed else "ТЕСТ · Частичный охват"
    if any(str(number) in draft.source_envelope.get("page_regions", {}) for number in selected):
        marker = "ТЕСТ · Частичный охват области"
    _append_export_row(sheet, [marker], long_values)
    _append_export_row(sheet, [f"{prepared['filename']} · страницы {', '.join(map(str, selected))} из {prepared['page_count']}"], long_values)
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 21

    row_index = 3
    page_lookup = prepared["pages_by_number"]
    table_groups: list[tuple[int, int, list[tuple[int, str]], list[dict[str, Any]]]] = []
    for page_number in selected:
        page = page_lookup.get(page_number)
        if page is None:
            continue
        for table in sorted(page.tables, key=lambda item: (item.index, item.id)):
            rows = sorted(
                (row for row in positions_rows if not row["manual"] and row["table_id"] == table.id),
                key=lambda row: row["row_index"],
            )
            if not rows:
                continue
            headers = {column.index: column.raw_header or "" for column in table.columns}
            for item in rows:
                for value in item["current_values"]:
                    headers.setdefault(int(value["column_index"]), "")
            table_groups.append((page_number, table.index + 1, sorted(headers.items()), rows))

    # Manual rows have no OCR table identity, so keep them grouped by source page.
    for page_number in selected:
        rows = sorted(
            (row for row in positions_rows if row["manual"] and row["page"] == page_number),
            key=lambda row: row["row_index"],
        )
        if rows:
            headers: dict[int, str] = {}
            for item in rows:
                for value in item["source_values"]:
                    headers.setdefault(int(value["column_index"]), str(value["header"] or ""))
            table_groups.append((page_number, 0, sorted(headers.items()), rows))

    if not table_groups:
        raise PilotExportError("no_verified_positions")

    def merged_label(label: str, column_count: int, *, fill: str, color: str, font_size: int = 10) -> None:
        nonlocal row_index
        _append_export_row(sheet, [label] + [""] * max(0, column_count - 1), long_values)
        if column_count > 1:
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=column_count)
        cell = sheet.cell(row_index, 1)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Aptos", size=font_size, bold=True, color=color)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row_index].height = 22
        row_index += 1

    row_number = 0
    for page_number, table_number, columns, rows in table_groups:
        column_count = max(1, len(columns))
        table_label = f"Страница {page_number} · " + (f"Таблица {table_number}" if table_number else "Добавленные позиции")
        merged_label(table_label, column_count, fill="EAF2F0", color="275C53", font_size=10)
        _append_export_row(sheet, [header for _, header in columns] if columns else [""], long_values)
        for cell in sheet[row_index]:
            cell.fill = PatternFill("solid", fgColor="315F72")
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row_index].height = 36
        header_row = row_index
        row_index += 1
        for item in rows:
            current = {int(value["column_index"]): value["value"] for value in item["current_values"]}
            _append_export_row(sheet, [current.get(index, "") for index, _ in columns] or [""], long_values)
            for cell in sheet[row_index]:
                cell.font = Font(name="Aptos", size=10, color="253746")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="thin", color="DCE4E8"))
                if cell.value is not None:
                    cell.data_type = "s"
            sheet.row_dimensions[row_index].height = 42
            row_index += 1
        if len(table_groups) == 1 and rows:
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(column_count)}{row_index - 1}"
        row_index += 1

    note_rows = []
    for item in prepared["all_rows"]:
        kind = str(item["model_kind"] or "").casefold().replace("-", "_")
        candidate = getattr(item["row"], "server_role_candidate", None)
        is_note = (item["role_active"] and item["role"] == "note") or kind in {"note", "document_note", "project_note"} or candidate == "note"
        if is_note and not item["superseded"]:
            note_rows.append(item)
    if note_rows:
        note_groups: dict[tuple[int, str], list[dict[str, Any]]] = defaultdict(list)
        for item in note_rows:
            note_groups[(int(item["page"]), str(item["table_id"]))].append(item)
        # Different note tables keep their own source headings and order.
        for (page_number, _table_id), group in sorted(note_groups.items()):
            source_columns = sorted(group[0]["columns"])
            note_header = ["Страница", "Таблица", "Строка", "Статус проверки"]
            note_header.extend(group[0]["columns"].get(index, "") for index in source_columns)
            note_width = max(1, len(note_header))
            table_label = f"Примечания из документа · страница {page_number}"
            if group[0]["table_index"]:
                table_label += f" · таблица {group[0]['table_index']}"
            merged_label(table_label, note_width, fill="F2F5F6", color="425466")
            _append_export_row(sheet, note_header, long_values)
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="6B7280")
                cell.font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[row_index].height = 34
            row_index += 1
            for item in group:
                decision = item["decision"]
                values = item["current_values"] if decision and decision.active and decision.human_confirmed else item["source_values"]
                by_column = {int(value["column_index"]): value["value"] for value in values}
                table_title = item["table_index"] if item["table_index"] else "—"
                status = "Проверено человеком" if item["role_active"] and decision and decision.active and decision.human_confirmed else "Требует проверки"
                _append_export_row(sheet, [item["page"], table_title, item["row_index"], status] + [by_column.get(index, "") for index in source_columns], long_values)
                for cell in sheet[row_index]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = Border(bottom=Side(style="thin", color="DCE4E8"))
                    if cell.value is not None:
                        cell.data_type = "s"
                sheet.row_dimensions[row_index].height = 36
                row_index += 1

    # Excel limits a cell to 32,767 characters. Only oversized values add a
    # continuation sheet; normal workbooks stay a single readable sheet.
    if long_values:
        overflow = workbook.create_sheet(LONG_TEXT_SHEET_NAME)
        overflow.sheet_view.showGridLines = False
        overflow.append(["Запись", "Ячейка спецификации", "Часть", "Всего частей", "Исходный текст · часть"])
        for cell in overflow[1]:
            cell.fill = PatternFill("solid", fgColor="315F72")
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        overflow.freeze_panes = "A2"
        for item in long_values:
            first_row = overflow.max_row + 1
            item["first_chunk_row"] = first_row
            for part, chunk in enumerate(item["chunks"], start=1):
                overflow.append([item["reference"], item["source_cell"], part, len(item["chunks"]), chunk])
                for cell in overflow[overflow.max_row]:
                    if cell.value is not None:
                        cell.value = str(cell.value)
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            item["cell"].hyperlink = f"#'{LONG_TEXT_SHEET_NAME}'!A{first_row}"
            item["cell"].style = "Hyperlink"
        overflow.column_dimensions["A"].width = 12
        overflow.column_dimensions["B"].width = 22
        overflow.column_dimensions["C"].width = 12
        overflow.column_dimensions["D"].width = 16
        overflow.column_dimensions["E"].width = 100

    rule = Side(style="thin", color="DCE4E8")
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(bottom=rule)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value is not None:
                cell.value = str(cell.value)
                cell.data_type = "s"
    sheet.freeze_panes = "A4"
    sheet.sheet_view.zoomScale = 85
    for column_index in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column_index).value or "") for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(44, max(12, max((len(v) for v in values), default=10) + 2))

    workbook.properties.title = "ТЕСТ · Спецификация AI-first"
    workbook.properties.subject = "Проверенный фрагмент исходной спецификации"
    stream = BytesIO()
    try:
        workbook.save(stream)
    except (ValueError, OSError):
        raise PilotExportError("pilot_workbook_generation_failed") from None
    return stream.getvalue()


__all__ = ["MAX_PILOT_EXPORT_PAGES", "XLSX_MIME", "PilotExportError", "build_ai_first_pilot_xlsx", "build_ai_first_diagnostic_xlsx", "preview_ai_first_pilot_export"]
