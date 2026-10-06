from __future__ import annotations

from collections import Counter
import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Context, Decimal, InvalidOperation, ROUND_HALF_EVEN, localcontext
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import unicodedata
import zipfile
from xml.etree import ElementTree

from openpyxl import load_workbook

from averon_import.core.unit_normalization import normalize_unit_family
from averon_import.services.one_c_history.models import FIELD_NAMES


PARSER_VERSION = 2
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ZIP_MEMBERS = 1500
MAX_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_HEADER_SCAN_ROWS = 50
MAX_HEADER_LENGTH = 512
MAX_PREVIEW_ROWS = 12
MAX_WARNING_EXAMPLES = 5
EFFECTIVE_UNIT_PRICE_PRECISION = 50

FIELD_ALIASES = {
    "item_code": {"кодноменклатуры", "номенклатурныйкод", "номерноменклатуры", "номеркодноменклатуры", "номенклатурныйномер", "itemcode", "nomenclaturenumber", "1citemcode"},
    "item_name": {"номенклатураедизм", "номенклатура", "наименование", "товар", "номенклатуранаименование", "itemname", "item"},
    "unit": {"едизм", "единицаизмерения", "единица", "unit", "measurementunit"},
    "quantity": {"колво", "количество", "количествоединиц", "qty", "quantity"},
    "reported_unit_price_gross": {"ценасндс", "ценасучетомндс", "ценабрутто", "цена", "unitpricegross", "grossunitprice", "reportedunitpricegross"},
    "amount_gross": {"суммасндс", "суммаучетомндс", "суммабрутто", "сумма", "amountgross", "grossamount"},
    "document_date": {"датадокумента", "датапоступления", "дата", "documentdate", "date"},
    "document_type": {"типдокумента", "виддокумента", "типоперации", "documenttype"},
    "document_reference": {"документприхода", "документ", "номердокумента", "основание", "documentreference", "document"},
    "counterparty": {"контрагент", "поставщик", "наименованиеконтрагента", "counterparty", "supplier"},
    "warehouse": {"склад"},
    "contract": {"договор", "контракт", "contract"},
    "article": {"артикул", "кодтовара", "article", "sku"},
    "manufacturer": {"производитель", "бренд", "manufacturer", "brand"},
    "characteristic": {"характеристика", "характеристиканоменклатуры", "characteristic"},
    "supplier_code": {"кодпоставщика", "suppliercode"},
    "supplier_inn": {"инн", "иннпоставщика", "supplierinn"},
    "vat_rate": {"ставкандс", "ндс", "vatrate"},
    "currency": {"валюта", "currency"},
    "organization": {"организация", "organization"},
    "document_stable_reference": {"ссылканадокумент", "идентификатордокумента", "documentid", "documentstablereference"},
    "document_line_number": {"номерстроки", "номерпозиции", "linenumber", "documentlinenumber"},
}

EVENT_FIELDS = (
    "quantity", "reported_unit_price_gross", "amount_gross", "document_date",
    "document_type", "document_reference", "counterparty", "contract",
)


class OneCImportError(ValueError):
    """An import error safe to show to the administrator."""


@dataclass(frozen=True)
class ParsedEvent:
    item_key: str
    item_code: str | None
    item_name: str
    raw_unit: str
    unit_family: str | None
    identity_quality: str
    group_number: int | None
    source_row: int
    document_date: str | None
    document_type: str
    document_reference: str
    counterparty: str
    contract: str
    quantity: str | None
    reported_unit_price_gross: str | None
    effective_unit_price_gross: str | None
    amount_gross: str | None
    price_usable: bool
    optional_facts: dict[str, str | None]
    source_facts: dict[str, str | None]


@dataclass(frozen=True)
class _GroupContext:
    item_code: str | None
    item_name: str
    raw_unit: str
    group_number: int
    raw_label: str
    optional_facts: dict[str, str | None]
    source_facts: dict[str, str | None]


@dataclass
class ParsedWorkbook:
    filename: str
    file_sha256: str
    sheet_name: str
    header_row: int
    headers: list[str]
    header_signature: str
    layout_type: str
    field_mapping: dict[str, int | None]
    group_header_row: int | None
    event_header_row: int | None
    group_headers: list[str]
    group_field_mapping: dict[str, int | None]
    event_field_mapping: dict[str, int | None]
    group_header_signature: str | None
    event_header_signature: str
    item_name_parse_strategy: str
    events: list[ParsedEvent]
    item_count: int
    group_count: int
    physical_row_count: int
    distinct_counterparty_count: int
    unit_vocabulary_count: int
    period_start: str | None
    period_end: str | None
    document_type_counts: dict[str, int]
    document_type_other_event_count: int
    supplier_missing_count: int
    unusable_price_count: int
    missing_code_count: int
    repeated_display_label_count: int
    normalized_display_collision_count: int
    code_conflict_count: int
    skipped_row_count: int
    warnings: list[dict]

    def summary(self) -> dict:
        return {
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "layout_type": self.layout_type,
            "group_header_row": self.group_header_row,
            "event_header_row": self.event_header_row,
            "period_start": self.period_start,
            "period_end": self.period_end,
            "item_count": self.item_count,
            "group_count": self.group_count,
            "physical_row_count": self.physical_row_count,
            "distinct_counterparty_count": self.distinct_counterparty_count,
            "unit_vocabulary_count": self.unit_vocabulary_count,
            "event_count": len(self.events),
            "usable_price_event_count": sum(event.price_usable for event in self.events),
            "supplier_missing_count": self.supplier_missing_count,
            "unusable_price_count": self.unusable_price_count,
            "missing_code_count": self.missing_code_count,
            "repeated_display_label_count": self.repeated_display_label_count,
            "normalized_display_collision_count": self.normalized_display_collision_count,
            "code_conflict_count": self.code_conflict_count,
            "skipped_row_count": self.skipped_row_count,
            "document_type_counts": self.document_type_counts,
            "document_type_other_event_count": self.document_type_other_event_count,
            "warning_count": len(self.warnings),
            "warnings": self.warnings[:MAX_WARNING_EXAMPLES],
        }

    def sample(self) -> list[dict]:
        def short(value: str | None, limit: int = 240) -> str | None:
            return value[:limit] if value is not None else None

        return [
            {
                "source_row": event.source_row,
                "item_name": short(event.item_name),
                "item_code_present": bool(event.item_code),
                "unit": short(event.raw_unit, 80),
                "document_date": event.document_date,
                "document_type": short(event.document_type, 120),
                "quantity": short(event.quantity, 80),
                "reported_unit_price_gross": short(event.reported_unit_price_gross, 80),
                "effective_unit_price_gross": short(event.effective_unit_price_gross, 80),
                "amount_gross": short(event.amount_gross, 80),
                "price_usable": event.price_usable,
            }
            for event in self.events[:MAX_PREVIEW_ROWS]
        ]


def _header_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", "", text)


def _display_repeat_key(value: object) -> str:
    """Normalize case and whitespace while preserving meaningful punctuation."""
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    return " ".join(text.split())


def _effective_unit_price(amount: Decimal | None, quantity: Decimal | None, reported_price: Decimal | None) -> Decimal | None:
    if amount is None or quantity is None or quantity <= 0:
        return reported_price
    # Use an isolated, explicit context so application code cannot change the
    # persisted 50-significant-digit representation by mutating getcontext().
    context = Context(prec=EFFECTIVE_UNIT_PRICE_PRECISION, rounding=ROUND_HALF_EVEN)
    with localcontext(context):
        return amount / quantity


def header_signature(headers: list[str]) -> str:
    normalized = [_header_key(value) for value in headers]
    payload = json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _header_mapping(headers: list[str]) -> dict[str, int | None]:
    result: dict[str, int | None] = {key: None for key in FIELD_NAMES}
    for index, label in enumerate(headers):
        normalized = _header_key(label)
        for field, aliases in FIELD_ALIASES.items():
            if result[field] is None and normalized in aliases:
                result[field] = index
    warehouse_columns = [
        index for index, label in enumerate(headers)
        if _header_key(label) in FIELD_ALIASES["warehouse"]
    ]
    result["warehouse"] = warehouse_columns[0] if len(warehouse_columns) == 1 else None
    return result


def _warehouse_header_ambiguous(headers: list[str]) -> bool:
    return sum(_header_key(label) in FIELD_ALIASES["warehouse"] for label in headers) > 1


def _relationship_mentions_lowercase_shared_strings(archive: zipfile.ZipFile) -> bool:
    member = "xl/_rels/workbook.xml.rels"
    try:
        root = ElementTree.fromstring(archive.read(member))
    except (KeyError, ElementTree.ParseError):
        return False
    for relationship in root:
        rel_type = str(relationship.attrib.get("Type") or "").casefold()
        target = str(relationship.attrib.get("Target") or "")
        if rel_type.endswith("/sharedstrings"):
            return target.replace("\\", "/").endswith("sharedStrings.xml")
    return False


def preflight_xlsx(path: Path) -> Path | None:
    """Validate OOXML bounds; return a temporary narrow shared-string fix if needed."""
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise OneCImportError("Файл не является корректной книгой XLSX.") from exc
    with archive:
        infos = archive.infolist()
        if not infos or len(infos) > MAX_ZIP_MEMBERS:
            raise OneCImportError("Архив XLSX содержит недопустимое число компонентов.")
        total_uncompressed = 0
        names: set[str] = set()
        folded_names: set[str] = set()
        for info in infos:
            name = info.filename.replace("\\", "/")
            parts = Path(name).parts
            if name.startswith("/") or ".." in parts or name in names or (parts and ":" in parts[0]) or "\x00" in name:
                raise OneCImportError("Архив XLSX содержит небезопасные или повторные компоненты.")
            if name.casefold() in folded_names:
                raise OneCImportError("Архив XLSX содержит повторные компоненты без учёта регистра.")
            names.add(name)
            folded_names.add(name.casefold())
            if info.file_size < 0 or info.compress_size < 0:
                raise OneCImportError("Архив XLSX повреждён.")
            total_uncompressed += info.file_size
            if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise OneCImportError("Распакованный размер XLSX превышает допустимый предел.")
            if info.file_size > 1024 and (info.compress_size == 0 or info.file_size / info.compress_size > MAX_COMPRESSION_RATIO):
                raise OneCImportError("Коэффициент сжатия XLSX превышает допустимый предел.")
            lower = name.casefold()
            if "vbaproject.bin" in lower or lower.startswith("xl/externallinks/"):
                raise OneCImportError("Книги с макросами или внешними ссылками не поддерживаются.")
            if name.endswith(".rels"):
                if info.file_size > 1_048_576:
                    raise OneCImportError("Метаданные связей XLSX превышают допустимый размер.")
                try:
                    relationship_root = ElementTree.fromstring(archive.read(info))
                except ElementTree.ParseError as exc:
                    raise OneCImportError("Метаданные XLSX повреждены.") from exc
                if any(str(node.attrib.get("TargetMode", "")).casefold() == "external" for node in relationship_root):
                    raise OneCImportError("Внешние ссылки в XLSX не поддерживаются.")

        canonical = "xl/sharedStrings.xml"
        case_variants = [info for info in infos if info.filename.casefold() == canonical.casefold()]
        if canonical in names or not case_variants or not _relationship_mentions_lowercase_shared_strings(archive):
            return None
        if len(case_variants) != 1 or case_variants[0].filename != "xl/SharedStrings.xml":
            raise OneCImportError("Несовпадение регистра компонента sharedStrings не поддерживается.")
        fd, temp_name = tempfile.mkstemp(prefix="averon-onec-normalized-", suffix=".xlsx", dir=path.parent)
        os.close(fd)
        normalized_path = Path(temp_name)
        try:
            with zipfile.ZipFile(normalized_path, "w") as output:
                output.comment = archive.comment
                for info in infos:
                    replacement = info
                    if info.filename == "xl/SharedStrings.xml":
                        replacement = copy.copy(info)
                        replacement.filename = "xl/sharedStrings.xml"
                        replacement.orig_filename = "xl/sharedStrings.xml"
                    with archive.open(info, "r") as source, output.open(replacement, "w") as destination:
                        while chunk := source.read(1024 * 1024):
                            destination.write(chunk)
            return normalized_path
        except Exception:
            normalized_path.unlink(missing_ok=True)
            raise


def _open_workbook(path: Path):
    normalized = preflight_xlsx(path)
    try:
        return load_workbook(normalized or path, read_only=True, data_only=False, keep_links=False), normalized
    except Exception as exc:
        if normalized:
            normalized.unlink(missing_ok=True)
        raise OneCImportError("Не удалось прочитать книгу XLSX. Проверьте формат и структуру файла.") from exc


def _detect_header(sheet) -> tuple[int, list[str], dict[str, int | None]]:
    best: tuple[int, int, list[str], dict[str, int | None]] | None = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN_ROWS, values_only=True), start=1):
        headers = [str(value).strip() if value is not None else "" for value in row]
        if any(len(value) > MAX_HEADER_LENGTH for value in headers):
            # Long data cells can occur inside the bounded header scan. Such a
            # row cannot be a header candidate, but must not reject the sheet.
            continue
        while headers and not headers[-1]:
            headers.pop()
        mapping = _header_mapping(headers)
        present = {field for field, index in mapping.items() if index is not None}
        score = len(present)
        core = "item_name" in present and bool(present.intersection({"quantity", "reported_unit_price_gross", "amount_gross"}))
        if core and (best is None or score > best[0]):
            best = (score, row_number, headers, mapping)
    if best is None:
        raise OneCImportError("Не удалось определить строку заголовков. Настройте сопоставление вручную.")
    return best[1], best[2], best[3]


def _fallback_header(sheet) -> tuple[int, list[str], dict[str, int | None]]:
    """Find a bounded candidate header row when none of the aliases are known."""
    best: tuple[int, int, list[str]] | None = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN_ROWS, values_only=True), start=1):
        headers = [str(value).strip() if value is not None else "" for value in row]
        if any(len(value) > MAX_HEADER_LENGTH for value in headers):
            continue
        while headers and not headers[-1]:
            headers.pop()
        populated = sum(bool(value) for value in headers)
        if populated >= 2 and (best is None or populated > best[0]):
            best = (populated, row_number, headers)
    if best is None:
        raise OneCImportError("В книге не удалось найти строку заголовков для ручного сопоставления.")
    return best[1], best[2], {key: None for key in FIELD_NAMES}


def _header_values(sheet, row_number: int) -> list[str]:
    values = next(sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True), ())
    headers = [str(value).strip() if value is not None else "" for value in values]
    if any(len(value) > MAX_HEADER_LENGTH for value in headers):
        raise OneCImportError("Заголовок XLSX превышает допустимую длину.")
    return headers[:MAX_COLUMNS]


def _detect_two_level_header(sheet) -> dict | None:
    """Detect an adjacent group-label/event-header pair within the bounded scan."""
    last_row = min(MAX_HEADER_SCAN_ROWS, sheet.max_row or 0)
    candidates = []
    for group_row in range(1, last_row):
        event_row = group_row + 1
        try:
            group_headers = _header_values(sheet, group_row)
            event_headers = _header_values(sheet, event_row)
        except OneCImportError:
            # An overlong row is not a candidate header pair. Explicitly
            # selected rows are still rejected by inspect_sheet_mapping().
            continue
        group_mapping = _header_mapping(group_headers)
        group_mapping["warehouse"] = None
        event_mapping = _header_mapping(event_headers)
        core_count = sum(event_mapping.get(field) is not None for field in (
            "quantity", "reported_unit_price_gross", "amount_gross",
        ))
        context_count = sum(event_mapping.get(field) is not None for field in (
            "document_date", "document_reference", "counterparty", "contract",
        ))
        if group_mapping.get("item_name") is None or core_count < 2 or context_count < 1:
            continue
        # A group label row and the event header row have separate roles. Do not
        # mistake a normal one-row flat header for a pair with item_name repeated.
        if event_mapping.get("item_name") is not None:
            continue
        score = sum(index is not None for index in event_mapping.values()) + core_count * 4 + context_count
        candidates.append((score, -group_row, {
            "layout_type": "hierarchical_grouped",
            "header_row": event_row,
            "group_header_row": group_row,
            "event_header_row": event_row,
            "headers": event_headers,
            "group_headers": group_headers,
            "field_mapping": event_mapping,
            "group_field_mapping": group_mapping,
            "event_field_mapping": event_mapping,
            "warehouse_header_ambiguous": _warehouse_header_ambiguous(event_headers),
            "item_name_parse_strategy": "comma_suffix_unit" if group_mapping.get("unit") is None else "none",
        }))
    if not candidates:
        return None
    return max(candidates, key=lambda item: (item[0], item[1]))[2]


def _detect_sheet_structure(sheet) -> dict:
    two_level = _detect_two_level_header(sheet)
    if two_level is not None:
        group_headers = two_level["group_headers"]
        event_headers = two_level["headers"]
        two_level["group_header_signature"] = header_signature(group_headers)
        two_level["event_header_signature"] = header_signature(event_headers)
        combined = json.dumps(
            [two_level["group_header_signature"], two_level["event_header_signature"]], separators=(",", ":")
        )
        two_level["header_signature"] = hashlib.sha256(combined.encode("utf-8")).hexdigest()
        two_level["warehouse_header_ambiguous"] = _warehouse_header_ambiguous(event_headers)
        return two_level

    try:
        header_row, headers, mapping = _detect_header(sheet)
    except OneCImportError:
        header_row, headers, mapping = _fallback_header(sheet)
    layout_type = _infer_layout(sheet, header_row, mapping)
    if layout_type == "hierarchical_grouped":
        group_row = event_row = header_row
        group_headers = headers
        group_mapping = dict(mapping)
        group_mapping["warehouse"] = None
        event_mapping = dict(mapping)
        group_sig = header_signature(group_headers)
    else:
        group_row = event_row = None
        group_headers = []
        group_mapping = {key: None for key in FIELD_NAMES}
        event_mapping = dict(mapping)
        group_sig = None
    return {
        "layout_type": layout_type,
        "header_row": header_row,
        "group_header_row": group_row,
        "event_header_row": event_row,
        "headers": headers,
        "group_headers": group_headers,
        "field_mapping": mapping,
        "group_field_mapping": group_mapping,
        "event_field_mapping": event_mapping,
        "group_header_signature": group_sig,
        "event_header_signature": header_signature(headers),
        "header_signature": header_signature(headers),
        "item_name_parse_strategy": "comma_suffix_unit" if mapping.get("unit") is None else "none",
        "warehouse_header_ambiguous": _warehouse_header_ambiguous(headers),
    }


def inspect_sheet_mapping(
    path: Path,
    sheet_name: str,
    *,
    group_header_row: int | None = None,
    event_header_row: int | None = None,
) -> dict:
    """Return role-specific, bounded headers and mappings for a selected sheet."""
    workbook, normalized = _open_workbook(path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise OneCImportError("Выбранный лист отсутствует в книге.")
        sheet = workbook[sheet_name]
        if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + MAX_HEADER_SCAN_ROWS:
            raise OneCImportError("Размер листа XLSX превышает допустимые пределы.")
        if group_header_row is None and event_header_row is None:
            return {**_detect_sheet_structure(sheet), "sheet_names": list(workbook.sheetnames), "sheet_name": sheet_name}

        selected_event_row = event_header_row or group_header_row
        if selected_event_row is None or selected_event_row > min(MAX_HEADER_SCAN_ROWS, sheet.max_row or 0):
            raise OneCImportError("Строка заголовков выходит за допустимый диапазон.")
        automatic = _detect_two_level_header(sheet)
        if automatic and selected_event_row == automatic["event_header_row"] and group_header_row in (None, automatic["group_header_row"]):
            structure = automatic
        elif group_header_row is not None and group_header_row != selected_event_row:
            if group_header_row > min(MAX_HEADER_SCAN_ROWS, sheet.max_row or 0):
                raise OneCImportError("Строка заголовков группы выходит за допустимый диапазон.")
            group_headers = _header_values(sheet, group_header_row)
            event_headers = _header_values(sheet, selected_event_row)
            group_mapping = _header_mapping(group_headers)
            group_mapping["warehouse"] = None
            event_mapping = _header_mapping(event_headers)
            group_signature = header_signature(group_headers)
            event_signature = header_signature(event_headers)
            combined = hashlib.sha256(json.dumps([group_signature, event_signature], separators=(",", ":")).encode()).hexdigest()
            structure = {
                "layout_type": "hierarchical_grouped", "header_row": selected_event_row,
                "group_header_row": group_header_row, "event_header_row": selected_event_row,
                "headers": event_headers, "group_headers": group_headers,
                "field_mapping": event_mapping, "group_field_mapping": group_mapping,
                "event_field_mapping": event_mapping, "group_header_signature": group_signature,
                "event_header_signature": event_signature, "header_signature": combined,
                "item_name_parse_strategy": "comma_suffix_unit" if group_mapping.get("unit") is None else "none",
                "warehouse_header_ambiguous": _warehouse_header_ambiguous(event_headers),
            }
        else:
            headers = _header_values(sheet, selected_event_row)
            while headers and not headers[-1]:
                headers.pop()
            mapping = _header_mapping(headers)
            layout = _infer_layout(sheet, selected_event_row, mapping)
            if layout == "hierarchical_grouped":
                group_row = event_row = selected_event_row
                group_headers = headers
                group_mapping = dict(mapping)
                group_mapping["warehouse"] = None
                event_mapping = dict(mapping)
                group_signature = header_signature(group_headers)
            else:
                group_row = event_row = None
                group_headers = []
                group_mapping = {key: None for key in FIELD_NAMES}
                event_mapping = dict(mapping)
                group_signature = None
            structure = {
                "layout_type": layout, "header_row": selected_event_row,
                "group_header_row": group_row, "event_header_row": event_row,
                "headers": headers, "group_headers": group_headers,
                "field_mapping": mapping, "group_field_mapping": group_mapping,
                "event_field_mapping": event_mapping, "group_header_signature": group_signature,
                "event_header_signature": header_signature(headers), "header_signature": header_signature(headers),
                "item_name_parse_strategy": "comma_suffix_unit" if mapping.get("unit") is None else "none",
                "warehouse_header_ambiguous": _warehouse_header_ambiguous(headers),
            }
        return {**structure, "sheet_names": list(workbook.sheetnames), "sheet_name": sheet_name}
    finally:
        workbook.close()
        if normalized:
            normalized.unlink(missing_ok=True)


def inspect_sheet(path: Path, sheet_name: str, header_row: int | None = None) -> tuple[int, list[str], dict[str, int | None], str]:
    """Compatibility wrapper for callers that only need a single row mapping."""
    structure = inspect_sheet_mapping(path, sheet_name, event_header_row=header_row)
    return (
        structure["header_row"], structure["headers"], structure["field_mapping"], structure["layout_type"],
    )


def _cell_value(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    return text if text else None


def _has_mapped_value(row: tuple, mapping: dict[str, int | None], fields: tuple[str, ...]) -> bool:
    return any(
        (index := mapping.get(field)) is not None and index < len(row) and _cell_value(row[index]) is not None
        for field in fields
    )


def _row_has_event_evidence(
    row: tuple,
    group_mapping: dict[str, int | None],
    event_mapping: dict[str, int | None],
) -> bool:
    present_fields = [
        field for field in EVENT_FIELDS
        if (index := event_mapping.get(field)) is not None
        and index < len(row) and _cell_value(row[index]) is not None
    ]
    if not present_fields:
        return False
    group_identity_columns = {
        index for field in ("item_name", "item_code", "unit", "article", "manufacturer", "characteristic")
        if (index := group_mapping.get(field)) is not None
    }
    if any(event_mapping[field] not in group_identity_columns for field in present_fields):
        return True
    price_index = event_mapping.get("reported_unit_price_gross")
    if price_index is None or price_index >= len(row):
        return False
    value = row[price_index]
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return _parse_decimal(value) is not None
    text = _cell_value(value)
    if text is None:
        return False
    scalar = text.replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    return bool(re.fullmatch(r"[+-]?(?:(?:\d+(?:[.,]\d*)?)|(?:[.,]\d+))(?:[eE][+-]?\d+)?", scalar))


def _infer_layout(sheet, header_row: int, mapping: dict[str, int | None]) -> str:
    populated_names = 0
    detail_rows = 0
    name_index = mapping.get("item_name")
    if name_index is None:
        return "flat"
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=min(header_row + 101, MAX_ROWS), values_only=True):
        if _has_mapped_value(row, mapping, EVENT_FIELDS):
            detail_rows += 1
            if name_index < len(row) and _cell_value(row[name_index]):
                populated_names += 1
    if detail_rows and populated_names / detail_rows < 0.65:
        return "hierarchical_grouped"
    return "flat"


def detect_workbook(path: Path) -> dict:
    workbook, normalized = _open_workbook(path)
    try:
        sheets = list(workbook.sheetnames)
        candidates = []
        for sheet_name in sheets:
            sheet = workbook[sheet_name]
            if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + MAX_HEADER_SCAN_ROWS:
                continue
            try:
                structure = _detect_sheet_structure(sheet)
            except OneCImportError:
                continue
            score = sum(index is not None for index in structure["event_field_mapping"].values())
            score += sum(index is not None for index in structure["group_field_mapping"].values())
            candidates.append((sheet_name, structure, sheet, score))
        if not candidates:
            raise OneCImportError("В книге нет листа с заголовками для ручного сопоставления.")
        tdsheet = next((candidate for candidate in candidates if candidate[0].casefold() == "tdsheet"), None)
        recognized = [candidate for candidate in candidates if candidate[3] > 0]
        selected = tdsheet or (max(recognized, key=lambda candidate: candidate[3]) if recognized else candidates[0])
        sheet_name, structure, sheet, _score = selected
        if len(structure["headers"]) > MAX_COLUMNS or (sheet.max_column or 0) > MAX_COLUMNS:
            raise OneCImportError("В листе XLSX слишком много столбцов.")
        if (sheet.max_row or 0) > MAX_ROWS + structure["header_row"]:
            raise OneCImportError("В листе XLSX слишком много строк.")
        return {
            "sheet_names": sheets,
            "sheet_name": sheet_name,
            **structure,
        }
    finally:
        workbook.close()
        if normalized:
            normalized.unlink(missing_ok=True)


def _parse_decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        parsed = value
    elif isinstance(value, (int, float)):
        parsed = Decimal(str(value))
    else:
        raw = str(value).strip().replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
        if not raw:
            return None
        if re.fullmatch(r"[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)[eE][+-]?\d+", raw):
            try:
                parsed = Decimal(raw.replace(",", "."))
            except InvalidOperation:
                return None
            if not parsed.is_finite():
                return None
            parts = parsed.as_tuple()
            return parsed if len(parts.digits) <= 30 and parts.exponent >= -18 and abs(parsed.adjusted()) <= 30 else None
        negative = raw.startswith("(") and raw.endswith(")")
        raw = raw.strip("()")
        raw = re.sub(r"[^0-9,\.\-+]", "", raw)
        if not raw or raw in {"-", "+", ".", ","}:
            return None
        if "," in raw and "." in raw:
            decimal_separator = "," if raw.rfind(",") > raw.rfind(".") else "."
            group_separator = "." if decimal_separator == "," else ","
            raw = raw.replace(group_separator, "").replace(decimal_separator, ".")
        elif "," in raw:
            raw = raw.replace(".", "").replace(",", ".")
        elif raw.count(".") > 1:
            head, tail = raw.rsplit(".", 1)
            raw = head.replace(".", "") + ("." + tail if len(tail) <= 2 else tail)
        try:
            parsed = Decimal(raw)
        except InvalidOperation:
            return None
        if negative:
            parsed = -abs(parsed)
    if not parsed.is_finite():
        return None
    parts = parsed.as_tuple()
    if len(parts.digits) > 30 or parts.exponent < -18 or abs(parsed.adjusted()) > 30:
        return None
    return parsed


def _decimal_text(value: Decimal | None) -> str | None:
    if value is None:
        return None
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return "0" if text in {"", "-0"} else text


def _date_value(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    text = str(value).strip()
    for pattern in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    return None


def _document_type_value(explicit_value: object | None, document_reference: str) -> str:
    explicit = _cell_value(explicit_value)
    if explicit:
        return explicit
    normalized = unicodedata.normalize("NFKC", document_reference).casefold().replace("ё", "е").strip()
    known_prefixes = (
        "поступление товаров и услуг",
        "авансовый отчет",
        "корректировка поступления",
    )
    for prefix in known_prefixes:
        if normalized.startswith(prefix):
            return prefix[0].upper() + prefix[1:]
    return ""


def _mapped(row: tuple, mapping: dict[str, int | None], field: str) -> object | None:
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def _warehouse_value(row: tuple, mapping: dict[str, int | None]) -> str | None:
    """Accept only a bounded text value from this physical event row."""
    value = _mapped(row, mapping, "warehouse")
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text or len(text) > 32_767:
        return None
    try:
        text.encode("utf-8")
    except UnicodeEncodeError:
        return None
    return text


def _split_name_unit(name: str, raw_unit: str, strategy: str) -> tuple[str, str]:
    if raw_unit:
        if raw_unit != name and normalize_unit_family(raw_unit):
            return name, raw_unit
        if normalize_unit_family(raw_unit):
            return name, raw_unit
    if strategy in {"comma_suffix_unit", "comma_or_parentheses"} and "," in name:
        base, suffix = name.rsplit(",", 1)
        suffix = suffix.strip()
        if base.strip() and suffix and normalize_unit_family(suffix):
            return base.strip(), suffix
    if strategy == "comma_or_parentheses":
        match = re.match(r"^(.*?)[(]\s*([^()]+?)\s*[)]$", name)
        if match and normalize_unit_family(match.group(2)):
            return match.group(1).strip(), match.group(2).strip()
    return name.strip(), raw_unit


def parse_workbook(
    path: Path,
    *,
    filename: str,
    file_sha256: str,
    sheet_name: str,
    header_row: int,
    headers: list[str],
    layout_type: str,
    field_mapping: dict[str, int | None],
    item_name_parse_strategy: str,
    group_header_row: int | None = None,
    event_header_row: int | None = None,
    group_headers: list[str] | None = None,
    group_field_mapping: dict[str, int | None] | None = None,
    event_field_mapping: dict[str, int | None] | None = None,
    group_header_signature: str | None = None,
    event_header_signature: str | None = None,
) -> ParsedWorkbook:
    if layout_type not in {"hierarchical_grouped", "flat"}:
        raise OneCImportError("Неизвестный тип структуры отчёта.")
    if not headers or len(headers) > MAX_COLUMNS:
        raise OneCImportError("В XLSX нет допустимых заголовков.")
    event_mapping = {
        key: (event_field_mapping or field_mapping).get(key)
        for key in FIELD_NAMES
    }
    group_mapping = {
        key: (group_field_mapping or field_mapping).get(key)
        for key in FIELD_NAMES
    } if layout_type == "hierarchical_grouped" else {key: None for key in FIELD_NAMES}
    if layout_type == "hierarchical_grouped":
        # Warehouses belong to event rows only, even when a one-row header is
        # used for a grouped report.
        group_mapping["warehouse"] = None
    group_header_row = group_header_row or (header_row if layout_type == "hierarchical_grouped" else None)
    event_header_row = event_header_row or header_row
    group_headers = group_headers or (headers if layout_type == "hierarchical_grouped" else [])
    if any(index is not None and index >= len(headers) for index in event_mapping.values()):
        raise OneCImportError("Сопоставление содержит столбец вне заголовков.")
    if layout_type == "hierarchical_grouped" and any(
        index is not None and index >= len(group_headers) for index in group_mapping.values()
    ):
        raise OneCImportError("Сопоставление полей группы содержит столбец вне заголовков.")
    if (group_mapping if layout_type == "hierarchical_grouped" else event_mapping).get("item_name") is None:
        raise OneCImportError("Сопоставьте поле «Наименование номенклатуры».")
    if not any(event_mapping.get(name) is not None for name in ("quantity", "reported_unit_price_gross", "amount_gross")):
        raise OneCImportError("Сопоставьте хотя бы одно поле: количество, цена или сумма.")

    workbook, normalized = _open_workbook(path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise OneCImportError("Выбранный лист отсутствует в книге.")
        sheet = workbook[sheet_name]
        if (sheet.max_column or 0) > MAX_COLUMNS or len(headers) > MAX_COLUMNS:
            raise OneCImportError("В листе XLSX слишком много столбцов.")
        first_data_row = event_header_row + 1
        if (sheet.max_row or 0) > MAX_ROWS + first_data_row:
            raise OneCImportError("В листе XLSX слишком много строк.")
        physical_row_count = sheet.max_row or 0

        events: list[ParsedEvent] = []
        current_item: _GroupContext | None = None
        item_keys: set[str] = set()
        labels_by_group: Counter[tuple[str, str]] = Counter()
        normalized_labels_by_group: Counter[tuple[str, str]] = Counter()
        codes_by_key: dict[str, dict[str, str | None]] = {}
        unit_values: set[str] = set()
        identity_issues: list[dict] = []
        type_counts: Counter[str] = Counter()
        event_dates: list[str] = []
        group_count = 0
        observed_rows = 0
        skipped_row_count = 0

        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=first_data_row, values_only=False), start=first_data_row
        ):
            observed_rows += 1
            if observed_rows > MAX_ROWS:
                raise OneCImportError("В листе XLSX слишком много строк.")
            row = tuple(cell.value for cell in cells)
            if not any(_cell_value(value) is not None for value in row):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise OneCImportError("Формулы в исходном отчёте не поддерживаются. Сохраните значения без формул.")

            if layout_type == "hierarchical_grouped":
                # In grouped reports, group and event headers have independent
                # meanings. In particular, the same physical column can hold
                # an item label on a group row and a price on event rows.
                event_present = _row_has_event_evidence(row, group_mapping, event_mapping)
                if not event_present:
                    raw_group_name = _cell_value(_mapped(row, group_mapping, "item_name")) or ""
                    raw_group_code = _cell_value(_mapped(row, group_mapping, "item_code"))
                    if raw_group_code and re.fullmatch(r"[-+]?\d+\.0+", raw_group_code):
                        raw_group_code = raw_group_code.split(".", 1)[0]
                    if raw_group_name or raw_group_code:
                        raw_group_unit = _cell_value(_mapped(row, group_mapping, "unit")) or ""
                        name, unit = _split_name_unit(raw_group_name, raw_group_unit, item_name_parse_strategy)
                        group_count += 1
                        if unit:
                            unit_values.add(unit)
                        if not raw_group_code:
                            if name:
                                labels_by_group[(_display_repeat_key(name), _display_repeat_key(unit))] += 1
                                normalized_labels_by_group[(_header_key(name), _header_key(unit))] += 1
                        group_optional = {
                            field: _cell_value(_mapped(row, group_mapping, field))
                            for field in ("article", "manufacturer", "characteristic", "supplier_code", "supplier_inn", "vat_rate", "currency", "organization")
                        }
                        group_source = {
                            f"group_{field}": _cell_value(_mapped(row, group_mapping, field))
                            for field in FIELD_NAMES if field != "warehouse" and group_mapping.get(field) is not None
                        }
                        current_item = _GroupContext(
                            item_code=raw_group_code,
                            item_name=name,
                            raw_unit=unit,
                            group_number=group_count,
                            raw_label=raw_group_name,
                            optional_facts=group_optional,
                            source_facts=group_source,
                        )
                    else:
                        skipped_row_count += 1
                    continue

                if current_item is None or not (current_item.item_name or current_item.item_code):
                    raise OneCImportError(
                        f"Строка {row_number}: найдено событие закупки без предшествующей группы номенклатуры."
                    )
                raw_event_code = _cell_value(_mapped(row, event_mapping, "item_code"))
                if raw_event_code and re.fullmatch(r"[-+]?\d+\.0+", raw_event_code):
                    raw_event_code = raw_event_code.split(".", 1)[0]
                if current_item.item_code and raw_event_code and current_item.item_code != raw_event_code:
                    identity_issues.append({"code": "group_event_code_conflict", "source_row": row_number})
                # Item identity comes only from the group row. An event-row
                # code is retained below as source provenance, never promoted
                # into a new canonical item identity.
                code = current_item.item_code or None
                name = current_item.item_name
                unit = current_item.raw_unit or (_cell_value(_mapped(row, event_mapping, "unit")) or "")
                group_num: int | None = current_item.group_number
                raw_group_label = current_item.raw_label
                row_optional = {
                    field: _cell_value(_mapped(row, event_mapping, field))
                    for field in ("article", "manufacturer", "characteristic", "supplier_code", "supplier_inn", "vat_rate", "currency", "organization", "document_stable_reference", "document_line_number")
                }
                row_optional["warehouse"] = _warehouse_value(row, event_mapping)
                optional = {
                    field: row_optional.get(field) or current_item.optional_facts.get(field)
                    for field in row_optional
                }
                source_facts = dict(current_item.source_facts)
                source_facts.update({
                    field: _cell_value(_mapped(row, event_mapping, field))
                    for field in FIELD_NAMES if event_mapping.get(field) is not None
                })
                if event_mapping.get("warehouse") is not None:
                    source_facts["warehouse"] = _warehouse_value(row, event_mapping)
            else:
                event_present = _has_mapped_value(row, event_mapping, EVENT_FIELDS)
                raw_name = _cell_value(_mapped(row, event_mapping, "item_name")) or ""
                raw_unit_cell = _cell_value(_mapped(row, event_mapping, "unit")) or ""
                item_code_value = _cell_value(_mapped(row, event_mapping, "item_code"))
                if item_code_value and re.fullmatch(r"[-+]?\d+\.0+", item_code_value):
                    item_code_value = item_code_value.split(".", 1)[0]
                if not event_present:
                    skipped_row_count += 1
                    continue
                name, unit = _split_name_unit(raw_name, raw_unit_cell, item_name_parse_strategy)
                if unit:
                    unit_values.add(unit)
                code = item_code_value or None
                group_count += 1
                group_num = group_count
                raw_group_label = raw_name
                if not code:
                    if name:
                        labels_by_group[(_display_repeat_key(name), _display_repeat_key(unit))] += 1
                        normalized_labels_by_group[(_header_key(name), _header_key(unit))] += 1
                optional = {
                    field: _cell_value(_mapped(row, event_mapping, field))
                    for field in ("article", "manufacturer", "characteristic", "supplier_code", "supplier_inn", "vat_rate", "currency", "organization", "document_stable_reference", "document_line_number")
                }
                optional["warehouse"] = _warehouse_value(row, event_mapping)
                source_facts = {
                    field: _cell_value(_mapped(row, event_mapping, field))
                    for field in FIELD_NAMES if event_mapping.get(field) is not None
                }
                if event_mapping.get("warehouse") is not None:
                    source_facts["warehouse"] = _warehouse_value(row, event_mapping)

            if layout_type == "hierarchical_grouped":
                raw_unit_cell = _cell_value(_mapped(row, event_mapping, "unit")) or ""
                if name and not unit and raw_unit_cell:
                    name, unit = _split_name_unit(name, raw_unit_cell, item_name_parse_strategy)
                if unit:
                    unit_values.add(unit)
            if raw_group_label:
                source_facts["raw_item_label"] = raw_group_label

            if name and not unit and item_name_parse_strategy != "none":
                name, unit = _split_name_unit(name, unit, item_name_parse_strategy)
            code_key = unicodedata.normalize("NFKC", code).strip() if code else ""
            if code_key:
                item_key = "1c:" + hashlib.sha256(code_key.encode("utf-8")).hexdigest()
                quality = "stable_1c_code"
                current_description = {
                    "unit": unit,
                    **{
                        field: optional.get(field)
                        for field in ("article", "manufacturer", "characteristic")
                    },
                }
                previous_description = codes_by_key.get(item_key)
                if previous_description is None:
                    codes_by_key[item_key] = current_description
                else:
                    conflict = False
                    for field, current_fact in current_description.items():
                        previous_fact = previous_description.get(field)
                        if not previous_fact and current_fact:
                            previous_description[field] = current_fact
                        elif previous_fact and current_fact:
                            if field == "unit":
                                previous_family = normalize_unit_family(previous_fact)
                                current_family = normalize_unit_family(current_fact)
                                differs = previous_family != current_family if previous_family and current_family else _header_key(previous_fact) != _header_key(current_fact)
                            else:
                                differs = _header_key(previous_fact) != _header_key(current_fact)
                            conflict = conflict or differs
                    if conflict:
                        identity_issues.append({"code": "stable_code_description_conflict", "source_row": row_number})
            else:
                identity_material = f"{_header_key(name)}|{_header_key(unit)}"
                if layout_type == "hierarchical_grouped":
                    item_key = "weak:" + hashlib.sha256(f"{identity_material}|group:{group_num}".encode("utf-8")).hexdigest()
                else:
                    item_key = "weak:" + hashlib.sha256(f"{identity_material}|row:{row_number}".encode("utf-8")).hexdigest()
                quality = "missing_stable_code" if name else "missing_item_identity"

            raw_quantity = _mapped(row, event_mapping, "quantity")
            raw_price = _mapped(row, event_mapping, "reported_unit_price_gross")
            raw_amount = _mapped(row, event_mapping, "amount_gross")
            quantity = _parse_decimal(raw_quantity)
            reported_price = _parse_decimal(raw_price)
            amount = _parse_decimal(raw_amount)
            effective_price = _effective_unit_price(amount, quantity, reported_price)
            document_reference = _cell_value(_mapped(row, event_mapping, "document_reference")) or ""
            document_type = _document_type_value(_mapped(row, event_mapping, "document_type"), document_reference)
            date_value = _date_value(_mapped(row, event_mapping, "document_date"))
            if date_value:
                event_dates.append(date_value)
            if document_type:
                type_counts[document_type] += 1
            item_keys.add(item_key)
            events.append(ParsedEvent(
                item_key=item_key,
                item_code=code,
                item_name=name,
                raw_unit=unit,
                unit_family=normalize_unit_family(unit),
                identity_quality=quality,
                group_number=group_num,
                source_row=row_number,
                document_date=date_value,
                document_type=document_type,
                document_reference=document_reference,
                counterparty=_cell_value(_mapped(row, event_mapping, "counterparty")) or "",
                contract=_cell_value(_mapped(row, event_mapping, "contract")) or "",
                quantity=_decimal_text(quantity),
                reported_unit_price_gross=_decimal_text(reported_price),
                effective_unit_price_gross=_decimal_text(effective_price),
                amount_gross=_decimal_text(amount),
                price_usable=effective_price is not None,
                optional_facts=optional,
                source_facts=source_facts,
            ))

        workbook.close()
        repeated_display_labels = sum(1 for count in labels_by_group.values() if count > 1)
        normalized_display_collisions = sum(1 for count in normalized_labels_by_group.values() if count > 1)
        missing_code_count = sum(1 for event in events if event.item_code is None)
        supplier_missing_count = sum(1 for event in events if not event.counterparty)
        distinct_counterparty_count = len({event.counterparty for event in events if event.counterparty})
        unit_vocabulary_count = len(unit_values)
        unusable_price_count = sum(1 for event in events if not event.price_usable)
        code_conflict_count = len(identity_issues)
        type_count_pairs = type_counts.most_common(50)
        bounded_type_counts = dict(type_count_pairs)
        other_type_event_count = sum(type_counts.values()) - sum(bounded_type_counts.values())
        warnings: list[dict] = []
        if missing_code_count:
            warnings.append({"code": "missing_stable_code", "count": missing_code_count, "message": "В отчёте не выгружен код номенклатуры. История будет импортирована, но идентификация одинаковых позиций будет менее надёжной."})
        if repeated_display_labels:
            warnings.append({"code": "repeated_display_label_groups", "count": repeated_display_labels, "message": "Обнаружены повторяющиеся подписи групп без кода 1С; их каноническая идентичность неизвестна, группы сохранены раздельно без объединения."})
        if supplier_missing_count:
            warnings.append({"code": "events_without_supplier", "count": supplier_missing_count, "message": "Часть событий не содержит контрагента; такие события сохраняются."})
        if unusable_price_count:
            warnings.append({"code": "events_without_usable_price", "count": unusable_price_count, "message": "События без цены или суммы сохраняются без пригодной цены."})
        if code_conflict_count:
            warnings.append({"code": "stable_code_identity_conflict", "count": code_conflict_count, "message": "У одного кода обнаружены различные описательные данные или единицы; факты сохранены для проверки."})
        group_event_code_conflicts = sum(issue["code"] == "group_event_code_conflict" for issue in identity_issues)
        if group_event_code_conflicts:
            warnings.append({"code": "group_event_code_conflict", "count": group_event_code_conflicts, "message": "Код из строки события отличается от кода группы; для идентичности сохранён код группы, оба исходных значения оставлены в provenance."})
        if not events:
            raise OneCImportError("В выбранной структуре не найдены строки событий закупки.")
        return ParsedWorkbook(
            filename=Path(filename).name[:255], file_sha256=file_sha256, sheet_name=sheet_name,
            header_row=header_row, headers=headers,
            header_signature=(
                hashlib.sha256(json.dumps([
                    group_header_signature or header_signature(group_headers),
                    event_header_signature or header_signature(headers),
                ], separators=(",", ":")).encode()).hexdigest()
                if layout_type == "hierarchical_grouped" and group_header_row != event_header_row
                else header_signature(headers)
            ),
            layout_type=layout_type, field_mapping=event_mapping,
            group_header_row=group_header_row, event_header_row=event_header_row,
            group_headers=list(group_headers), group_field_mapping=group_mapping,
            event_field_mapping=event_mapping,
            group_header_signature=group_header_signature or (header_signature(group_headers) if group_headers else None),
            event_header_signature=event_header_signature or header_signature(headers),
            item_name_parse_strategy=item_name_parse_strategy, events=events,
            item_count=len(item_keys), group_count=group_count,
            physical_row_count=physical_row_count,
            distinct_counterparty_count=distinct_counterparty_count,
            unit_vocabulary_count=unit_vocabulary_count,
            period_start=min(event_dates) if event_dates else None,
            period_end=max(event_dates) if event_dates else None,
            document_type_counts=bounded_type_counts, document_type_other_event_count=other_type_event_count,
            supplier_missing_count=supplier_missing_count,
            unusable_price_count=unusable_price_count, missing_code_count=missing_code_count,
            repeated_display_label_count=repeated_display_labels,
            normalized_display_collision_count=normalized_display_collisions,
            code_conflict_count=code_conflict_count,
            skipped_row_count=skipped_row_count,
            warnings=warnings,
        )
    except OneCImportError:
        raise
    except Exception as exc:
        raise OneCImportError("Не удалось разобрать выбранный лист XLSX.") from exc
    finally:
        try:
            workbook.close()
        finally:
            if normalized:
                normalized.unlink(missing_ok=True)


def summarize_rows(parsed: ParsedWorkbook) -> list[dict]:
    return parsed.sample()


__all__ = [
    "FIELD_ALIASES", "MAX_PREVIEW_ROWS", "PARSER_VERSION", "ParsedWorkbook",
    "OneCImportError", "detect_workbook", "header_signature", "inspect_sheet", "parse_workbook",
    "preflight_xlsx",
]
