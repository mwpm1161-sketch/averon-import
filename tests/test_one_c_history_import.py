from __future__ import annotations

import asyncio
from datetime import datetime
from decimal import Context, Decimal, ROUND_DOWN, ROUND_HALF_UP, ROUND_UP, getcontext, localcontext, setcontext
import io
import json
from pathlib import Path
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from threading import Event
import zipfile
from xml.etree import ElementTree as ET

import pytest
from fastapi import UploadFile
from openpyxl import Workbook

from averon_import.core.unit_normalization import normalize_unit_family
from averon_import.services.one_c_history.models import ImportMappingRequest, PreviewMappingRequest
from averon_import.services.one_c_history.activity import OneCHistoryActivityConflict, OneCHistoryActivityRegistry
from averon_import.services.one_c_history.repository import OneCHistoryRepository
from averon_import.services.one_c_history.service import OneCHistoryImportService
from averon_import.services.one_c_history.xlsx_import import (
    OneCImportError,
    EFFECTIVE_UNIT_PRICE_PRECISION,
    _parse_decimal,
    detect_workbook,
    parse_workbook,
    preflight_xlsx,
)


HIERARCHICAL_HEADERS = [
    "Номенклатура, ед. изм.", "Цена с НДС", "Кол-во", "Сумма с НДС",
    "Дата документа", "Документ прихода", "Контрагент", "Договор", "Код номенклатуры",
]


def _xlsx_bytes(path: Path, *, headers=HIERARCHICAL_HEADERS, rows=(), sheet_name="TDSheet") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path.read_bytes()


def _flat_xlsx_bytes(path: Path, *, shift=0, sheet_name="TDSheet", name="Шайба") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for _ in range(shift):
        sheet.append(["Отчёт о закупках"])
    sheet.append(["Код номенклатуры", "Наименование", "Ед. изм.", "Количество", "Цена с НДС", "Сумма с НДС", "Контрагент"])
    sheet.append(["001", name, "шт", 2, 3, 6, "Поставщик"])
    workbook.save(path)
    return path.read_bytes()


def _two_level_xlsx_bytes(path: Path, *, shift=0) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TDSheet"
    sheet.cell(row=1, column=1, value="Отчёт о закупках")
    group_row = 9 + shift
    for row_number, values in (
        (group_row, ["Номенклатура, ед. изм.", "Код номенклатуры", "Ед. изм.", None, None, None, None]),
        (group_row + 1, ["Цена с НДС", "Количество", "Сумма с НДС", "Дата документа", "Документ прихода", "Контрагент", "Договор"]),
        (group_row + 2, ["Шайба", "0001", "шт", None, None, None, None]),
        (group_row + 3, [3, 2, 6, datetime(2026, 9, 1), "Поступление №1", "Поставщик", "Договор"]),
        (group_row + 4, [None, None, None, None, None, None, None]),
        (group_row + 5, [4, 3, 12, datetime(2026, 9, 2), "Поступление №2", "Поставщик", "Договор"]),
    ):
        for column, value in enumerate(values, start=1):
            if value is not None:
                sheet.cell(row=row_number, column=column, value=value)
    workbook.save(path)
    return path.read_bytes()


def _hierarchical_rows(*, second_quantity=1):
    return [
        ["Прокладка, кг", None, None, None, None, None, None, None, None],
        [None, "12,345", 2, "25.00", datetime(2026, 6, 23), "Поступление товаров и услуг №1", "Поставщик А", "Договор А", None],
        ["Прокладка, кг", None, None, None, None, None, None, None, None],
        [None, 2, second_quantity, 2, datetime(2026, 9, 22), "Авансовый отчет", None, None, None],
        ["Корректировка, шт", None, None, None, None, None, None, None, None],
        [None, None, 1, None, datetime(2026, 9, 22), "Корректировка поступления", None, None, None],
    ]


def _upload(payload: bytes, filename="history.xlsx"):
    return UploadFile(filename=filename, file=io.BytesIO(payload))


def _make_service(tmp_path: Path) -> OneCHistoryImportService:
    repository = OneCHistoryRepository(tmp_path / "data")
    service = OneCHistoryImportService(repository)
    service.temp_root = tmp_path / "temp"
    service.temp_root.mkdir()
    return service


async def _import_preview(service: OneCHistoryImportService, payload: bytes, filename="history.xlsx", *, save_profile=False, profile_name=None):
    preview = await service.create_preview(_upload(payload, filename))
    request = ImportMappingRequest(
        preview_id=preview["preview_id"],
        sheet_name=preview["sheet_name"],
        header_row=preview["header_row"],
        layout_type=preview["layout_type"],
        field_mapping=preview["field_mapping"],
        item_name_parse_strategy=preview["item_name_parse_strategy"],
        save_profile=save_profile,
        profile_name=profile_name,
    )
    return preview, await service.import_confirmed(request)


def test_hierarchical_report_preserves_duplicate_groups_events_and_prices(tmp_path):
    path = tmp_path / "hierarchical.xlsx"
    payload = _xlsx_bytes(path, rows=_hierarchical_rows())
    detected = detect_workbook(path)
    parsed = parse_workbook(
        path,
        filename="history.xlsx",
        file_sha256="a" * 64,
        sheet_name=detected["sheet_name"],
        header_row=detected["header_row"],
        headers=detected["headers"],
        layout_type=detected["layout_type"],
        field_mapping=detected["field_mapping"],
        item_name_parse_strategy=detected["item_name_parse_strategy"],
    )

    assert payload
    assert detected["sheet_name"] == "TDSheet"
    assert detected["layout_type"] == "hierarchical_grouped"
    assert parsed.group_count == 3
    assert parsed.item_count == 3
    assert parsed.repeated_display_label_count == 1
    assert parsed.period_start == "2026-06-23"
    assert parsed.period_end == "2026-09-22"
    assert parsed.physical_row_count == 7
    assert parsed.distinct_counterparty_count == 1
    assert parsed.unit_vocabulary_count == 2
    assert parsed.document_type_counts == {
        "Поступление товаров и услуг": 1,
        "Авансовый отчет": 1,
        "Корректировка поступления": 1,
    }
    assert parsed.events[0].reported_unit_price_gross == "12.345"
    assert parsed.events[0].effective_unit_price_gross == "12.5"
    assert parsed.events[0].raw_unit == "кг"
    assert parsed.events[0].source_facts["raw_item_label"] == "Прокладка, кг"
    assert parsed.events[0].unit_family == "kilogram"
    assert parsed.events[0].counterparty == "Поставщик А"
    assert parsed.events[0].document_type == "Поступление товаров и услуг"
    assert parsed.events[1].document_type == "Авансовый отчет"
    assert parsed.events[1].counterparty == ""
    assert parsed.events[1].contract == ""
    assert parsed.events[2].document_type == "Корректировка поступления"
    assert parsed.events[2].price_usable is False
    assert parsed.events[2].effective_unit_price_gross is None
    assert parsed.events[0].identity_quality == "missing_stable_code"
    repeated_warning = next(warning for warning in parsed.warnings if warning["code"] == "repeated_display_label_groups")
    assert "каноническая идентичность неизвестна" in repeated_warning["message"]
    assert "без объединения" in repeated_warning["message"]
    missing_code_warning = next(warning for warning in parsed.warnings if warning["code"] == "missing_stable_code")
    assert missing_code_warning["message"] == "В отчёте не выгружен код номенклатуры. История будет импортирована, но идентификация одинаковых позиций будет менее надёжной."


def test_flat_report_uses_stable_codes_and_keeps_same_named_codes_distinct(tmp_path):
    headers = ["Код номенклатуры", "Наименование", "Ед. изм.", "Кол-во", "Цена с НДС", "Сумма с НДС", "Дата документа", "Документ", "Контрагент", "Договор", "Артикул"]
    rows = [
        ["C-1", "Одинаковое имя", "шт", 1, "10,25", 10.25, datetime(2026, 7, 1), "Поступление 1", "", "", "A-1"],
        ["C-1", "Новое отображаемое имя", "шт", 2, 20, 40, datetime(2026, 8, 1), "Поступление 2", "", "", "A-1"],
        ["C-1", "Новое отображаемое имя", "кг", 1, 20, 20, datetime(2026, 8, 2), "Поступление 2а", "", "", "A-1"],
        ["C-2", "Одинаковое имя", "шт", 1, 10, 10, datetime(2026, 9, 1), "Поступление 3", "", "", "A-2"],
    ]
    path = tmp_path / "flat.xlsx"
    _xlsx_bytes(path, headers=headers, rows=rows)
    detected = detect_workbook(path)
    parsed = parse_workbook(
        path, filename="flat.xlsx", file_sha256="b" * 64,
        sheet_name=detected["sheet_name"], header_row=detected["header_row"], headers=detected["headers"],
        layout_type=detected["layout_type"], field_mapping=detected["field_mapping"],
        item_name_parse_strategy=detected["item_name_parse_strategy"],
    )
    assert detected["layout_type"] == "flat"
    assert parsed.item_count == 2
    assert parsed.events[0].item_key == parsed.events[1].item_key
    assert parsed.events[0].item_key != parsed.events[3].item_key
    assert parsed.events[0].identity_quality == "stable_1c_code"
    assert parsed.events[1].item_name == "Новое отображаемое имя"
    assert parsed.events[0].effective_unit_price_gross == "10.25"
    assert parsed.events[0].reported_unit_price_gross == "10.25"
    assert parsed.supplier_missing_count == 4
    assert parsed.code_conflict_count == 1
    service = _make_service(tmp_path / "store")
    asyncio.run(_import_preview(service, path.read_bytes()))
    with sqlite3.connect(service.repository.database_path) as connection:
        variants = connection.execute(
            "SELECT item_name, raw_unit FROM nomenclature_descriptive_variants WHERE item_id = (SELECT item_id FROM nomenclature_items WHERE source_item_code = 'C-1') ORDER BY first_source_row"
        ).fetchall()
    assert variants == [("Одинаковое имя", "шт"), ("Новое отображаемое имя", "шт"), ("Новое отображаемое имя", "кг")]


@pytest.mark.parametrize(
    ("source", "expected"),
    [("1 234,56", "1234.56"), ("1,234.56", "1234.56"), ("12.375", "12.375"), (12.5, "12.5"), (Decimal("9.125"), "9.125")],
)
def test_decimal_parser_handles_comma_dot_and_numeric_cells(source, expected):
    assert format(_parse_decimal(source), "f") == expected


def test_decimal_parser_rejects_unbounded_exponents():
    assert _parse_decimal("1e999999") is None


def test_effective_unit_price_is_context_independent_and_preserves_source_values(tmp_path):
    path = tmp_path / "decimal-context.xlsx"
    _xlsx_bytes(
        path,
        rows=[
            ["Расходный материал, кг", None, None, None, None, None, None, None, None],
            [None, "10.12", 3, "30.35", datetime(2026, 7, 1), "Поступление №1", "Контрагент", None, None],
        ],
    )
    detected = detect_workbook(path)
    original_context = getcontext().copy()
    serialized_results = []
    try:
        for context in (Context(prec=4, rounding=ROUND_DOWN), Context(prec=80, rounding=ROUND_UP)):
            setcontext(context)
            parsed = parse_workbook(
                path, filename="decimal-context.xlsx", file_sha256="1" * 64,
                sheet_name=detected["sheet_name"], header_row=detected["header_row"],
                headers=detected["headers"], layout_type=detected["layout_type"],
                field_mapping=detected["field_mapping"],
                item_name_parse_strategy=detected["item_name_parse_strategy"],
            )
            event = parsed.events[0]
            serialized_results.append(event.effective_unit_price_gross)
            assert event.quantity == "3"
            assert event.amount_gross == "30.35"
            assert event.reported_unit_price_gross == "10.12"
            with localcontext(Context(prec=EFFECTIVE_UNIT_PRICE_PRECISION)):
                expected = Decimal("30.35") / Decimal("3")
            assert Decimal(event.effective_unit_price_gross) == expected
            with localcontext(Context(prec=60)):
                assert Decimal(event.effective_unit_price_gross).quantize(
                    Decimal("0.01"), rounding=ROUND_HALF_UP
                ) == Decimal("10.12")
        assert serialized_results[0] == serialized_results[1]
        assert len(Decimal(serialized_results[0]).as_tuple().digits) == EFFECTIVE_UNIT_PRICE_PRECISION
    finally:
        setcontext(original_context)


def test_effective_unit_price_handles_terminating_quotient(tmp_path):
    path = tmp_path / "terminating-price.xlsx"
    _xlsx_bytes(
        path,
        rows=[
            ["Кабель, м", None, None, None, None, None, None, None, None],
            [None, "0.13", 8, "1", datetime(2026, 7, 1), "Поступление №1", "Контрагент", None, None],
        ],
    )
    detected = detect_workbook(path)
    parsed = parse_workbook(
        path, filename="terminating-price.xlsx", file_sha256="2" * 64,
        sheet_name=detected["sheet_name"], header_row=detected["header_row"],
        headers=detected["headers"], layout_type=detected["layout_type"],
        field_mapping=detected["field_mapping"],
        item_name_parse_strategy=detected["item_name_parse_strategy"],
    )
    assert parsed.events[0].effective_unit_price_gross == "0.125"


def test_display_repeat_metric_preserves_punctuation_and_normalizes_case_whitespace(tmp_path):
    path = tmp_path / "display-repeat-keys.xlsx"
    _xlsx_bytes(
        path,
        rows=[
            ["Датчик-А, шт", None, None, None, None, None, None, None, None],
            [None, 1, 1, 1, datetime(2026, 7, 1), "Поступление №1", None, None, None],
            ["Датчик А, шт", None, None, None, None, None, None, None, None],
            [None, 1, 1, 1, datetime(2026, 7, 2), "Поступление №2", None, None, None],
            ["  Крепёж B, шт", None, None, None, None, None, None, None, None],
            [None, 1, 1, 1, datetime(2026, 7, 3), "Поступление №3", None, None, None],
            ["крепёж   b, шт", None, None, None, None, None, None, None, None],
            [None, 1, 1, 1, datetime(2026, 7, 4), "Поступление №4", None, None, None],
        ],
    )
    detected = detect_workbook(path)
    parsed = parse_workbook(
        path, filename="display-repeat-keys.xlsx", file_sha256="3" * 64,
        sheet_name=detected["sheet_name"], header_row=detected["header_row"],
        headers=detected["headers"], layout_type=detected["layout_type"],
        field_mapping=detected["field_mapping"],
        item_name_parse_strategy=detected["item_name_parse_strategy"],
    )
    assert parsed.repeated_display_label_count == 1
    assert parsed.normalized_display_collision_count == 2
    assert parsed.group_count == parsed.item_count == 4


def test_current_unit_vocabulary_has_conservative_normalization():
    expected = {"шт": "piece", "кг": "kilogram", "м": "meter", "м2": "square_meter", "л": "litre", "Пар": "pair", "м3": "cubic_meter", "компл": "set", "пог. м": "meter", "упак": "pack", "г": "gram", "боб": "bobbin", "т": "tonne", "мл": "millilitre"}
    assert {unit: normalize_unit_family(unit) for unit in expected} == expected
    assert normalize_unit_family("неизвестная ед.") is None


def test_synthetic_tdsheet_acceptance_profile_and_price_rounding(tmp_path):
    units = ["шт", "кг", "м", "м2", "л", "Пар", "м3", "компл", "пог. м", "упак", "г", "боб", "т", "мл"]
    path = tmp_path / "accepted-profile.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TDSheet"
    for index in range(1, 9):
        sheet.append([f"Параметр отчёта {index}"])
    sheet.append(["Номенклатура, ед. изм."])
    sheet.merge_cells("A9:K9")
    sheet.append(["Цена с НДС", None, None, "Количество", "Сумма с НДС", "Дата", "Документ прихода", None, None, "Контрагент", "Договор"])
    sheet.merge_cells("A10:C10")
    sheet.merge_cells("G10:I10")
    event_index = 0

    for group_index in range(3066):
        if group_index < 52:
            label_key = group_index % 26
            item_label = f"Повторяющаяся группа {label_key}"
            unit = units[label_key % len(units)]
        else:
            item_label = f"Номенклатура {group_index}"
            unit = units[group_index % len(units)]
        group_row = sheet.max_row + 1
        sheet.append([f"{item_label}, {unit}"] + [None] * 10)
        if group_index < 3:
            sheet.merge_cells(start_row=group_row, start_column=1, end_row=group_row, end_column=11)
        event_count_for_group = 1458 if group_index == 3065 else 1
        for _ in range(event_count_for_group):
            if event_index < 4081:
                reference = f"Поступление товаров и услуг №{event_index + 1}"
                counterparty = f"Контрагент {event_index % 264 + 1}"
            elif event_index < 4521:
                reference = f"Авансовый отчет №{event_index - 4080}"
                counterparty = None
            else:
                reference = f"Корректировка поступления №{event_index - 4520}"
                counterparty = None
            event_date = datetime(2026, 6, 23) if event_index == 0 else (
                datetime(2026, 9, 22) if event_index == 4522 else datetime(2026, 7, 1)
            )
            is_correction = event_index >= 4521
            if is_correction:
                price = amount = None
                quantity = 1
            elif event_index < 1190:
                price, quantity, amount = "10.12", 3, "30.35"
            else:
                price, quantity, amount = "10.12", 1, "10.12"
            detail_row = sheet.max_row + 1
            sheet.append([
                price, None, None, quantity, amount, event_date, reference,
                None, None, counterparty, None,
            ])
            if event_index < 5:
                sheet.merge_cells(start_row=detail_row, start_column=1, end_row=detail_row, end_column=3)
                sheet.merge_cells(start_row=detail_row, start_column=7, end_row=detail_row, end_column=9)
            event_index += 1
    workbook.save(path)

    detected = detect_workbook(path)
    parsed = parse_workbook(
        path, filename="accepted-profile.xlsx", file_sha256="e" * 64,
        sheet_name=detected["sheet_name"], header_row=detected["header_row"], headers=detected["headers"],
        layout_type=detected["layout_type"], field_mapping=detected["field_mapping"],
        item_name_parse_strategy=detected["item_name_parse_strategy"],
        group_header_row=detected["group_header_row"], event_header_row=detected["event_header_row"],
        group_headers=detected["group_headers"], group_field_mapping=detected["group_field_mapping"],
        event_field_mapping=detected["event_field_mapping"],
        group_header_signature=detected["group_header_signature"],
        event_header_signature=detected["event_header_signature"],
    )

    assert parsed.physical_row_count == 7599
    assert detected["group_header_row"] == 9 and detected["event_header_row"] == 10
    assert detected["group_field_mapping"]["item_name"] == 0
    assert detected["event_field_mapping"]["reported_unit_price_gross"] == 0
    assert detected["event_field_mapping"]["quantity"] == 3
    assert parsed.group_count == 3066
    assert len(parsed.events) == 4523
    assert parsed.document_type_counts == {
        "Поступление товаров и услуг": 4081,
        "Авансовый отчет": 440,
        "Корректировка поступления": 2,
    }
    assert parsed.distinct_counterparty_count == 264
    assert parsed.unit_vocabulary_count == 14
    assert parsed.period_start == "2026-06-23"
    assert parsed.period_end == "2026-09-22"
    assert parsed.repeated_display_label_count == 26
    assert parsed.normalized_display_collision_count == 26
    assert parsed.events[0].item_name.startswith("Повторяющаяся группа")
    assert parsed.events[0].reported_unit_price_gross == "10.12"

    advances = [event for event in parsed.events if event.document_type == "Авансовый отчет"]
    corrections = [event for event in parsed.events if event.document_type == "Корректировка поступления"]
    assert len(advances) == 440 and all(not event.counterparty for event in advances)
    assert len(corrections) == 2
    assert all(event.quantity and event.document_date and event.document_reference for event in corrections)
    assert all(event.reported_unit_price_gross is None and event.amount_gross is None for event in corrections)

    rounded_price_mismatches = [
        event for event in parsed.events
        if event.reported_unit_price_gross and event.amount_gross
        and Decimal(event.reported_unit_price_gross) * Decimal(event.quantity) != Decimal(event.amount_gross)
    ]
    assert len(rounded_price_mismatches) == 1190
    for event in rounded_price_mismatches:
        reported = Decimal(event.reported_unit_price_gross)
        effective = Decimal(event.effective_unit_price_gross)
        amount = Decimal(event.amount_gross)
        quantity = Decimal(event.quantity)
        with localcontext(Context(prec=EFFECTIVE_UNIT_PRICE_PRECISION)):
            assert effective == amount / quantity
        with localcontext(Context(prec=60)):
            assert effective.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) == reported
            assert abs(effective - reported) <= Decimal("0.005")


def test_group_level_identity_and_descriptive_facts_are_inherited_by_detail_events(tmp_path):
    path = tmp_path / "group-facts.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TDSheet"
    sheet.append(["Номенклатура", "Код номенклатуры", "Артикул", "Цена", "Количество", "Сумма", "Дата", "Документ"])
    sheet.append(["Цена с НДС", "Код номенклатуры", "Артикул", "Количество", "Сумма с НДС", "Дата", "Документ прихода"])
    sheet.append(["Позиция, м", "00123", "ART-9", None, None, None, None, None])
    sheet.append(["4.5", None, None, 2, "9", datetime(2026, 6, 1), "Поступление товаров и услуг №1"])
    workbook.save(path)

    detected = detect_workbook(path)
    parsed = parse_workbook(
        path, filename="group-facts.xlsx", file_sha256="f" * 64,
        sheet_name=detected["sheet_name"], header_row=detected["header_row"],
        headers=detected["headers"], layout_type=detected["layout_type"],
        field_mapping=detected["field_mapping"],
        item_name_parse_strategy=detected["item_name_parse_strategy"],
        group_header_row=detected["group_header_row"], event_header_row=detected["event_header_row"],
        group_headers=detected["group_headers"], group_field_mapping=detected["group_field_mapping"],
        event_field_mapping=detected["event_field_mapping"],
    )

    event = parsed.events[0]
    assert event.item_name == "Позиция"
    assert event.item_code == "00123"
    assert event.optional_facts["article"] == "ART-9"
    assert event.source_facts["group_item_code"] == "00123"
    assert event.source_facts["group_article"] == "ART-9"
    assert event.reported_unit_price_gross == "4.5"

    service = _make_service(tmp_path / "service")
    preview = asyncio.run(service.create_preview(_upload(path.read_bytes())))
    assert preview["mapping_required"] is False
    request = ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=preview["sheet_name"],
        header_row=preview["header_row"], group_header_row=preview["group_header_row"],
        event_header_row=preview["event_header_row"], layout_type=preview["layout_type"],
        field_mapping=preview["field_mapping"],
        group_field_mapping=preview["group_field_mapping"],
        event_field_mapping=preview["event_field_mapping"],
        item_name_parse_strategy=preview["item_name_parse_strategy"],
    )
    result = asyncio.run(service.import_confirmed(request))
    assert result["status"] == "succeeded"
    with sqlite3.connect(service.repository.database_path) as connection:
        stored = connection.execute(
            "SELECT item_code,item_name,reported_unit_price_gross_decimal,optional_facts_json,source_facts_json FROM purchase_events"
        ).fetchone()
    assert stored[0] == "00123" and stored[1] == "Позиция"
    assert stored[2] == "4.5"
    assert json.loads(stored[3])["article"] == "ART-9"
    assert json.loads(stored[4])["group_item_code"] == "00123"


def test_hierarchical_event_without_group_fails_closed(tmp_path):
    path = tmp_path / "orphan-event.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TDSheet"
    sheet.append(["Номенклатура, ед. изм."])
    sheet.merge_cells("A1:K1")
    sheet.append(["Цена с НДС", None, None, "Количество", "Сумма с НДС", "Дата", "Документ прихода", None, None, "Контрагент", "Договор"])
    sheet.merge_cells("A2:C2")
    sheet.merge_cells("G2:I2")
    sheet.append(["12", None, None, 1, "12", datetime(2026, 6, 1), "Поступление №1", None, None, "Поставщик", None])
    sheet.merge_cells("A3:C3")
    sheet.merge_cells("G3:I3")
    workbook.save(path)
    detected = detect_workbook(path)

    with pytest.raises(OneCImportError, match="без предшествующей группы"):
        parse_workbook(
            path, filename="orphan-event.xlsx", file_sha256="0" * 64,
            sheet_name=detected["sheet_name"], header_row=detected["header_row"],
            headers=detected["headers"], layout_type=detected["layout_type"],
            field_mapping=detected["field_mapping"],
            item_name_parse_strategy=detected["item_name_parse_strategy"],
            group_header_row=detected["group_header_row"], event_header_row=detected["event_header_row"],
            group_headers=detected["group_headers"], group_field_mapping=detected["group_field_mapping"],
            event_field_mapping=detected["event_field_mapping"],
        )


def test_active_snapshot_keeps_immutable_mapping_profile_revision(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "profiled.xlsx", rows=_hierarchical_rows())
    _preview, result = asyncio.run(_import_preview(service, payload, save_profile=True, profile_name="Исходный профиль"))
    profile = service.repository.profile(result["profile_id"])
    active = service.repository.active_metadata()
    provenance = active["mapping_provenance"]

    assert profile.revision == 1
    assert provenance["profile_id"] == profile.profile_id
    assert provenance["profile_revision"] == 1
    assert provenance["fingerprint"]
    assert provenance["layout_type"] == "hierarchical_grouped"
    assert provenance["group_field_mapping"]
    assert provenance["event_field_mapping"]

    updated = service.repository.save_profile(profile.model_copy(update={"name": "Переименованный профиль"}))
    assert updated.revision == 2
    assert service.repository.delete_profile(profile.profile_id) is True
    assert service.repository.active_metadata()["mapping_provenance"] == provenance


def test_sourcing_status_is_an_allowlisted_aggregate_projection(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "private-source-name.xlsx", rows=_hierarchical_rows())
    asyncio.run(_import_preview(service, payload, filename="private-source-name.xlsx"))

    status = service.repository.sourcing_status()

    assert set(status) == {
        "available", "catalog_version", "period_start", "period_end", "imported_at",
        "item_count", "event_count", "usable_price_event_count",
    }
    assert status["available"] is True
    assert status["period_start"] == "2026-06-23"
    assert status["period_end"] == "2026-09-22"
    assert status["event_count"] == 3
    assert status["item_count"] == 3
    assert status["catalog_version"]
    assert status["imported_at"]
    assert all("private-source-name" not in str(value) for value in status.values())


def test_sourcing_status_without_snapshot_has_stable_unavailable_shape(tmp_path):
    repository = _make_service(tmp_path).repository
    status = repository.sourcing_status()

    assert status == {
        "available": False,
        "catalog_version": None,
        "period_start": None,
        "period_end": None,
        "imported_at": None,
        "item_count": 0,
        "event_count": 0,
        "usable_price_event_count": 0,
    }

    repository.database_path.write_bytes(b"not a valid sqlite snapshot")
    assert repository.sourcing_status() == status


def test_manually_changed_mapping_does_not_claim_an_unmodified_saved_profile(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "profiled.xlsx", rows=_hierarchical_rows())
    asyncio.run(_import_preview(service, payload, save_profile=True, profile_name="Профиль"))
    changed_rows = _hierarchical_rows()
    changed_rows[1][1] = "11.00"
    changed_payload = _xlsx_bytes(tmp_path / "profiled-updated.xlsx", rows=changed_rows)
    preview = asyncio.run(service.create_preview(_upload(changed_payload)))
    changed_mapping = dict(preview["field_mapping"])
    changed_mapping["counterparty"] = None
    analyzed = asyncio.run(service.analyze_preview(preview["preview_id"], PreviewMappingRequest(
        sheet_name=preview["sheet_name"], header_row=preview["header_row"],
        layout_type=preview["layout_type"], field_mapping=changed_mapping,
        item_name_parse_strategy=preview["item_name_parse_strategy"],
    )))
    result = asyncio.run(service.import_confirmed(ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=analyzed["sheet_name"],
        header_row=analyzed["header_row"], layout_type=analyzed["layout_type"],
        field_mapping=analyzed["field_mapping"],
        item_name_parse_strategy=analyzed["item_name_parse_strategy"],
    )))

    provenance = result["active_import"]["mapping_provenance"]
    assert provenance["profile_id"] is None
    assert provenance["profile_revision"] is None
    assert provenance["event_field_mapping"]["counterparty"] is None


def test_activation_replace_failure_keeps_previous_snapshot(tmp_path, monkeypatch):
    service = _make_service(tmp_path)
    first_payload = _xlsx_bytes(tmp_path / "first.xlsx", rows=_hierarchical_rows())
    _first_preview, first = asyncio.run(_import_preview(service, first_payload))
    old_bytes = service.repository.database_path.read_bytes()
    second_rows = _hierarchical_rows()
    second_rows[1][1] = "99.00"
    second_payload = _xlsx_bytes(tmp_path / "second.xlsx", rows=second_rows)
    original_replace = __import__("os").replace

    def fail_snapshot_replace(source, destination):
        if Path(destination) == service.repository.database_path:
            raise OSError("injected replace failure")
        return original_replace(source, destination)

    monkeypatch.setattr("averon_import.services.one_c_history.repository.os.replace", fail_snapshot_replace)
    with pytest.raises(OneCImportError, match="Текущий снимок не изменён"):
        asyncio.run(_import_preview(service, second_payload))

    assert service.repository.database_path.read_bytes() == old_bytes
    assert service.repository.active_metadata()["sha256"] == first["active_import"]["sha256"]


def test_post_replace_directory_sync_failure_reports_committed_snapshot(tmp_path, monkeypatch):
    service = _make_service(tmp_path)
    first_payload = _xlsx_bytes(tmp_path / "first.xlsx", rows=_hierarchical_rows())
    _first_preview, first = asyncio.run(_import_preview(service, first_payload))
    second_rows = _hierarchical_rows()
    second_rows[1][1] = "77.00"
    second_payload = _xlsx_bytes(tmp_path / "second.xlsx", rows=second_rows)

    def fail_directory_sync():
        raise OSError("injected directory fsync failure")

    monkeypatch.setattr(service.repository, "_fsync_directory", fail_directory_sync)
    _preview, result = asyncio.run(_import_preview(service, second_payload))

    assert result["status"] == "succeeded"
    assert result["active_import"]["sha256"] != first["active_import"]["sha256"]
    assert result["warnings"]
    assert service.repository.active_metadata()["sha256"] == result["active_import"]["sha256"]


def test_status_bookkeeping_failure_after_replace_does_not_claim_old_snapshot(tmp_path, monkeypatch):
    service = _make_service(tmp_path)
    first_payload = _xlsx_bytes(tmp_path / "first.xlsx", rows=_hierarchical_rows())
    _first_preview, first = asyncio.run(_import_preview(service, first_payload))
    second_rows = _hierarchical_rows()
    second_rows[1][1] = "88.00"
    second_payload = _xlsx_bytes(tmp_path / "second.xlsx", rows=second_rows)
    original_record = service.repository.record_attempt

    def fail_success_status(status, **kwargs):
        if status == "succeeded":
            raise OSError("injected status bookkeeping failure")
        return original_record(status, **kwargs)

    monkeypatch.setattr(service.repository, "record_attempt", fail_success_status)
    _preview, result = asyncio.run(_import_preview(service, second_payload))

    assert result["status"] == "succeeded"
    assert result["active_import"]["sha256"] != first["active_import"]["sha256"]
    assert result["active_import"]["sha256"] == service.repository.active_metadata()["sha256"]
    assert result["warnings"]


def test_preview_is_bounded_path_free_and_deletes_temp_after_success(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "source.xlsx", rows=_hierarchical_rows())
    preview, result = asyncio.run(_import_preview(service, payload, "C:\\private\\report.xlsx"))
    assert result["status"] == "succeeded"
    assert len(preview["sample"]) <= 12
    assert str(service.temp_root) not in json.dumps(preview)
    assert "C:\\private" not in json.dumps(preview)
    assert not list(service.temp_root.glob("pending-*.xlsx"))
    active = service.repository.active_metadata()
    assert active["price_basis"] == "gross_including_vat"
    assert active["event_count"] == 3
    assert active["physical_row_count"] == 7
    assert active["distinct_counterparty_count"] == 1
    assert active["identity_quality"] == "degraded_missing_stable_code"
    with sqlite3.connect(service.repository.database_path) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("SELECT COUNT(*) FROM purchase_events").fetchone()[0] == 3


def test_preview_ownership_busy_retry_profile_permission_and_actor_audit(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "owned.xlsx", rows=_hierarchical_rows())
    registry = OneCHistoryActivityRegistry()
    service.activity_registry = registry

    async def exercise():
        preview = await service.create_preview(_upload(payload), owner_key="user:a")
        other_mapping = PreviewMappingRequest(
            sheet_name=preview["sheet_name"],
            header_row=preview["header_row"],
            layout_type=preview["layout_type"],
            field_mapping=preview["field_mapping"],
            group_header_row=preview.get("group_header_row"),
            event_header_row=preview.get("event_header_row"),
            group_field_mapping=preview.get("group_field_mapping") or {},
            event_field_mapping=preview.get("event_field_mapping") or {},
            item_name_parse_strategy=preview["item_name_parse_strategy"],
        )
        with pytest.raises(OneCImportError):
            await service.inspect_preview_sheet(preview["preview_id"], preview["sheet_name"], owner_key="user:b")
        with pytest.raises(OneCImportError):
            await service.analyze_preview(preview["preview_id"], other_mapping, owner_key="user:b")

        inspected = await service.inspect_preview_sheet(
            preview["preview_id"], preview["sheet_name"], owner_key="user:a",
        )
        assert inspected["sheet_name"] == preview["sheet_name"]
        analyzed = await service.analyze_preview(preview["preview_id"], other_mapping, owner_key="user:a")
        request = ImportMappingRequest(
            preview_id=preview["preview_id"],
            sheet_name=analyzed["sheet_name"],
            header_row=analyzed["header_row"],
            group_header_row=analyzed.get("group_header_row"),
            event_header_row=analyzed.get("event_header_row"),
            layout_type=analyzed["layout_type"],
            field_mapping=analyzed["field_mapping"],
            group_field_mapping=analyzed.get("group_field_mapping") or {},
            event_field_mapping=analyzed.get("event_field_mapping") or {},
            item_name_parse_strategy=analyzed["item_name_parse_strategy"],
        )
        with pytest.raises(OneCImportError):
            await service.import_confirmed(request, owner_key="user:b")
        assert preview["preview_id"] in service._pending

        with pytest.raises(PermissionError):
            await service.import_confirmed(
                request.model_copy(update={"save_profile": True, "profile_name": "Глобальный"}),
                owner_key="user:a", can_manage_profiles=False,
            )
        assert preview["preview_id"] in service._pending

        reader = registry.acquire_sourcing()
        with pytest.raises(OneCHistoryActivityConflict) as busy:
            await service.import_confirmed(request, owner_key="user:a")
        assert busy.value.code == "ONE_C_HISTORY_IN_USE"
        record = service._pending[preview["preview_id"]]
        assert record.path.is_file()
        assert record.analysis_key == service._analysis_key(record.detected)
        reader.release()

        actor = {"user_id": "user-a-id", "username": "colleague", "role": "user"}
        result = await service.import_confirmed(request, owner_key="user:a", actor=actor, can_manage_profiles=False)
        second_preview = await service.create_preview(_upload(payload), owner_key="user:b")
        second_result = await service.import_confirmed(
            request.model_copy(update={"preview_id": second_preview["preview_id"]}),
            owner_key="user:b",
            actor={"user_id": "user-b-id", "username": "another-user", "role": "user"},
            can_manage_profiles=False,
        )
        return preview, result, second_result

    preview, result, second_result = asyncio.run(exercise())
    assert result["status"] == "succeeded"
    assert second_result["status"] == "already_active" and second_result["idempotent"] is True
    assert result["active_import"]["semantic_import_fingerprint"] == second_result["active_import"]["semantic_import_fingerprint"]
    assert preview["preview_id"] not in service._pending
    audit = service.repository.public_status()["recent_audit"]
    assert audit[-2]["user_id"] == "user-a-id"
    assert audit[-2]["username"] == "colleague"
    assert audit[-2]["role"] == "user"
    assert audit[-2]["outcome"] == "succeeded"
    assert audit[-2]["catalog_version"] == service.repository.catalog_version()
    assert audit[-1]["user_id"] == "user-b-id"
    assert audit[-1]["outcome"] == "already_active"
    active = service.repository.active_metadata()
    assert "user_id" not in active and "username" not in active


def test_two_concurrent_final_imports_fail_fast_and_second_preview_retries_safely(tmp_path):
    service = _make_service(tmp_path)
    registry = OneCHistoryActivityRegistry()
    service.activity_registry = registry
    payload = _flat_xlsx_bytes(tmp_path / "shared-import.xlsx")

    async def prepare():
        await _import_preview(
            service,
            _flat_xlsx_bytes(tmp_path / "previous-active.xlsx", name="Предыдущая история"),
        )
        previous_active = service.repository.active_metadata()
        first = await service.create_preview(_upload(payload), owner_key="user:first")
        second = await service.create_preview(_upload(payload), owner_key="user:second")

        def request_for(preview):
            return ImportMappingRequest(
                preview_id=preview["preview_id"],
                sheet_name=preview["sheet_name"],
                header_row=preview["header_row"],
                group_header_row=preview.get("group_header_row"),
                event_header_row=preview.get("event_header_row"),
                layout_type=preview["layout_type"],
                field_mapping=preview["field_mapping"],
                group_field_mapping=preview.get("group_field_mapping") or {},
                event_field_mapping=preview.get("event_field_mapping") or {},
                item_name_parse_strategy=preview["item_name_parse_strategy"],
            )

        return previous_active, first, second, request_for(first), request_for(second)

    previous_active, first_preview, second_preview, first_request, second_request = asyncio.run(prepare())
    build_entered = Event()
    allow_build = Event()
    original_build = service.repository.build_staging_snapshot

    def blocked_build(*args, **kwargs):
        build_entered.set()
        if not allow_build.wait(timeout=10):
            raise TimeoutError("test did not release the staged import")
        return original_build(*args, **kwargs)

    service.repository.build_staging_snapshot = blocked_build
    with ThreadPoolExecutor(max_workers=1) as pool:
        first_future = pool.submit(
            lambda: asyncio.run(service.import_confirmed(first_request, owner_key="user:first"))
        )
        try:
            assert build_entered.wait(timeout=10), "first import did not enter staging while holding its writer lease"
            assert registry.status()["update_in_progress"] is True
            still_active = service.repository.active_metadata()
            assert still_active["sha256"] == previous_active["sha256"]
            with pytest.raises(OneCHistoryActivityConflict) as competing:
                asyncio.run(service.import_confirmed(second_request, owner_key="user:second"))
            assert competing.value.code == "ONE_C_HISTORY_UPDATE_IN_PROGRESS"
            assert second_preview["preview_id"] in service._pending
            second_pending = service._pending[second_preview["preview_id"]]
            assert second_pending.path.is_file()
            assert second_pending.analysis_key == service._analysis_key(second_pending.detected)
        finally:
            allow_build.set()
        first_result = first_future.result(timeout=20)

    assert first_result["status"] == "succeeded"
    assert registry.status()["replacement_allowed"] is True
    assert service.repository.active_metadata()["sha256"] != previous_active["sha256"]
    active_version = service.repository.catalog_version()
    retry_result = asyncio.run(service.import_confirmed(second_request, owner_key="user:second"))
    assert retry_result["status"] == "already_active"
    assert retry_result["idempotent"] is True
    assert second_preview["preview_id"] not in service._pending
    assert service.repository.catalog_version() == active_version
    assert first_preview["preview_id"] not in service._pending
    with sqlite3.connect(service.repository.database_path) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("SELECT COUNT(*) FROM purchase_events").fetchone()[0] == 1


def test_preview_bounds_are_per_owner_and_global_without_cross_user_eviction(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "bounded.xlsx", rows=_hierarchical_rows())

    async def exercise():
        first_user = [await service.create_preview(_upload(payload), owner_key="user:a") for _ in range(3)]
        other_user = await service.create_preview(_upload(payload), owner_key="user:b")
        await service.create_preview(_upload(payload), owner_key="user:a")
        assert other_user["preview_id"] in service._pending
        assert first_user[0]["preview_id"] not in service._pending
        assert sum(item.owner_key == "user:a" for item in service._pending.values()) == 3

        for index in range(2, 18):
            await service.create_preview(_upload(payload), owner_key=f"user:{index}")
        assert len(service._pending) == 20
        with pytest.raises(OneCImportError, match="Слишком много подготовленных отчётов"):
            await service.create_preview(_upload(payload), owner_key="user:20")
        assert len(service._pending) == 20
        assert other_user["preview_id"] in service._pending

    asyncio.run(exercise())


def test_expired_preview_is_removed_on_next_capacity_check(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "expiry.xlsx", rows=_hierarchical_rows())

    async def exercise():
        expired = await service.create_preview(_upload(payload), owner_key="user:a")
        record = service._pending[expired["preview_id"]]
        path = record.path
        record.expires_at = 0
        await service.create_preview(_upload(payload), owner_key="user:b")
        assert expired["preview_id"] not in service._pending
        assert not path.exists()

    asyncio.run(exercise())


def test_same_sha_is_idempotent_and_profiles_survive_snapshot_replacement(tmp_path):
    service = _make_service(tmp_path)
    payload = _xlsx_bytes(tmp_path / "source.xlsx", rows=_hierarchical_rows())
    first_preview, first = asyncio.run(_import_preview(service, payload, save_profile=True, profile_name="Основной отчёт"))
    profile_id = first["profile_id"]
    assert profile_id and len(service.repository.profiles()) == 1
    second_preview = asyncio.run(service.create_preview(_upload(payload)))
    assert second_preview["mapping_profile_id"] == profile_id
    assert second_preview["mapping_required"] is False
    second_request = ImportMappingRequest(
        preview_id=second_preview["preview_id"], sheet_name=second_preview["sheet_name"],
        header_row=second_preview["header_row"], layout_type=second_preview["layout_type"],
        field_mapping=second_preview["field_mapping"], item_name_parse_strategy=second_preview["item_name_parse_strategy"],
    )
    second = asyncio.run(service.import_confirmed(second_request))
    assert second["idempotent"] is True
    assert first["active_import"]["semantic_import_fingerprint"] == second["active_import"]["semantic_import_fingerprint"]
    service.repository.database_path.unlink()
    assert service.repository.profile(profile_id).name == "Основной отчёт"
    assert first_preview["sha256"] == second_preview["sha256"]


def test_same_sha_with_changed_event_mapping_rebuilds_snapshot(tmp_path):
    service = _make_service(tmp_path)
    payload = _flat_xlsx_bytes(tmp_path / "flat.xlsx")
    _first_preview, first = asyncio.run(_import_preview(service, payload))
    first_fingerprint = first["active_import"]["semantic_import_fingerprint"]

    preview = asyncio.run(service.create_preview(_upload(payload)))
    changed_mapping = dict(preview["event_field_mapping"])
    changed_mapping["counterparty"] = None
    analyzed = asyncio.run(service.analyze_preview(preview["preview_id"], PreviewMappingRequest(
        sheet_name=preview["sheet_name"], header_row=preview["header_row"],
        layout_type="flat", field_mapping=changed_mapping,
        item_name_parse_strategy=preview["item_name_parse_strategy"],
    )))
    second = asyncio.run(service.import_confirmed(ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=analyzed["sheet_name"],
        header_row=analyzed["header_row"], layout_type=analyzed["layout_type"],
        field_mapping=analyzed["field_mapping"],
        item_name_parse_strategy=analyzed["item_name_parse_strategy"],
    )))

    assert second["status"] == "succeeded" and second["idempotent"] is False
    assert second["active_import"]["sha256"] == first["active_import"]["sha256"]
    assert second["active_import"]["semantic_import_fingerprint"] != first_fingerprint
    assert second["active_import"]["mapping_provenance"]["event_field_mapping"]["counterparty"] is None


def test_same_sha_with_changed_group_mapping_rebuilds_snapshot(tmp_path):
    service = _make_service(tmp_path)
    payload = _two_level_xlsx_bytes(tmp_path / "grouped.xlsx")
    _first_preview, first = asyncio.run(_import_preview(service, payload))
    first_fingerprint = first["active_import"]["semantic_import_fingerprint"]

    preview = asyncio.run(service.create_preview(_upload(payload)))
    changed_group_mapping = dict(preview["group_field_mapping"])
    changed_group_mapping["characteristic"] = 0
    analyzed = asyncio.run(service.analyze_preview(preview["preview_id"], PreviewMappingRequest(
        sheet_name=preview["sheet_name"], header_row=preview["header_row"],
        group_header_row=preview["group_header_row"], event_header_row=preview["event_header_row"],
        layout_type="hierarchical_grouped", field_mapping=preview["event_field_mapping"],
        group_field_mapping=changed_group_mapping, event_field_mapping=preview["event_field_mapping"],
        item_name_parse_strategy=preview["item_name_parse_strategy"],
    )))
    second = asyncio.run(service.import_confirmed(ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=analyzed["sheet_name"],
        header_row=analyzed["header_row"], group_header_row=analyzed["group_header_row"],
        event_header_row=analyzed["event_header_row"], layout_type=analyzed["layout_type"],
        field_mapping=analyzed["field_mapping"], group_field_mapping=analyzed["group_field_mapping"],
        event_field_mapping=analyzed["event_field_mapping"],
        item_name_parse_strategy=analyzed["item_name_parse_strategy"],
    )))

    assert second["status"] == "succeeded" and second["idempotent"] is False
    assert second["active_import"]["sha256"] == first["active_import"]["sha256"]
    assert second["active_import"]["semantic_import_fingerprint"] != first_fingerprint
    assert second["active_import"]["mapping_provenance"]["group_field_mapping"]["characteristic"] == 0


def test_same_sha_with_another_selected_sheet_rebuilds_snapshot(tmp_path):
    path = tmp_path / "two-flat-sheets.xlsx"
    workbook = Workbook()
    first_sheet = workbook.active
    first_sheet.title = "TDSheet"
    headers = ["Код номенклатуры", "Наименование", "Ед. изм.", "Количество", "Цена с НДС", "Сумма с НДС", "Контрагент"]
    first_sheet.append(headers)
    first_sheet.append(["001", "Шайба", "шт", 2, 3, 6, "Поставщик"])
    second_sheet = workbook.create_sheet("Другой лист")
    second_sheet.append(headers)
    second_sheet.append(["002", "Гайка", "шт", 4, 5, 20, "Поставщик 2"])
    workbook.save(path)
    payload = path.read_bytes()
    service = _make_service(tmp_path / "store")
    _first_preview, first = asyncio.run(_import_preview(service, payload))
    first_fingerprint = first["active_import"]["semantic_import_fingerprint"]

    preview = asyncio.run(service.create_preview(_upload(payload)))
    inspected = asyncio.run(service.inspect_preview_sheet(preview["preview_id"], "Другой лист"))
    analyzed = asyncio.run(service.analyze_preview(preview["preview_id"], PreviewMappingRequest(
        sheet_name=inspected["sheet_name"], header_row=inspected["header_row"],
        layout_type="flat", field_mapping=inspected["field_mapping"],
        item_name_parse_strategy=inspected["item_name_parse_strategy"],
    )))
    second = asyncio.run(service.import_confirmed(ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=analyzed["sheet_name"],
        header_row=analyzed["header_row"], layout_type=analyzed["layout_type"],
        field_mapping=analyzed["field_mapping"],
        item_name_parse_strategy=analyzed["item_name_parse_strategy"],
    )))

    assert second["status"] == "succeeded" and second["idempotent"] is False
    assert second["active_import"]["sha256"] == first["active_import"]["sha256"]
    assert second["active_import"]["sheet_name"] == "Другой лист"
    assert second["active_import"]["semantic_import_fingerprint"] != first_fingerprint


def test_profile_rename_does_not_change_semantic_idempotency(tmp_path):
    service = _make_service(tmp_path)
    payload = _flat_xlsx_bytes(tmp_path / "profiled.xlsx")
    _preview, first = asyncio.run(_import_preview(service, payload, save_profile=True, profile_name="Профиль"))
    profile = service.repository.profile(first["profile_id"])
    renamed = service.repository.save_profile(profile.model_copy(update={"name": "Новое имя"}))
    assert renamed.revision == profile.revision + 1

    preview = asyncio.run(service.create_preview(_upload(payload)))
    second = asyncio.run(service.import_confirmed(ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=preview["sheet_name"],
        header_row=preview["header_row"], layout_type=preview["layout_type"],
        field_mapping=preview["field_mapping"], item_name_parse_strategy=preview["item_name_parse_strategy"],
    )))

    assert preview["mapping_profile_id"] == profile.profile_id
    assert second["idempotent"] is True
    assert second["active_import"]["semantic_import_fingerprint"] == first["active_import"]["semantic_import_fingerprint"]


def test_saved_two_level_profile_uses_shifted_current_header_rows(tmp_path):
    service = _make_service(tmp_path)
    original = _two_level_xlsx_bytes(tmp_path / "original.xlsx")
    original_preview, first = asyncio.run(_import_preview(service, original, save_profile=True, profile_name="Двухуровневый"))
    assert original_preview["group_header_row"] == 9
    assert original_preview["event_header_row"] == 10

    shifted = _two_level_xlsx_bytes(tmp_path / "shifted.xlsx", shift=1)
    preview = asyncio.run(service.create_preview(_upload(shifted)))

    assert preview["mapping_profile_id"] == first["profile_id"]
    assert preview["mapping_required"] is False
    assert preview["group_header_row"] == 10
    assert preview["event_header_row"] == 11
    assert preview["summary"]["group_count"] == 1
    assert preview["summary"]["event_count"] == 2
    assert preview["sample"][0]["source_row"] == 13
    assert preview["sample"][0]["item_name"] == "Шайба"
    assert preview["sample"][0]["item_code_present"] is True


def test_saved_flat_profile_uses_shifted_current_header_row(tmp_path):
    service = _make_service(tmp_path)
    original = _flat_xlsx_bytes(tmp_path / "flat-original.xlsx")
    _preview, first = asyncio.run(_import_preview(service, original, save_profile=True, profile_name="Плоский"))
    shifted = _flat_xlsx_bytes(tmp_path / "flat-shifted.xlsx", shift=1)

    preview = asyncio.run(service.create_preview(_upload(shifted)))

    assert preview["mapping_profile_id"] == first["profile_id"]
    assert preview["mapping_required"] is False
    assert preview["header_row"] == 2
    assert preview["summary"]["event_count"] == 1
    assert preview["sample"][0]["source_row"] == 3


def test_conflicting_compatible_profiles_require_explicit_selection(tmp_path):
    service = _make_service(tmp_path)
    payload = _flat_xlsx_bytes(tmp_path / "flat.xlsx")
    _preview, imported = asyncio.run(_import_preview(service, payload, save_profile=True, profile_name="Первый"))
    original_profile = service.repository.profile(imported["profile_id"])
    changed_mapping = {key: value for key, value in original_profile.field_mapping.items() if key != "counterparty"}
    conflicting_profile = original_profile.model_copy(update={
        "profile_id": "a" * 32,
        "name": "Другой вариант",
        "field_mapping": changed_mapping,
        "event_field_mapping": changed_mapping,
        "revision": 1,
    })
    service.repository.save_profile(conflicting_profile)

    preview = asyncio.run(service.create_preview(_upload(payload)))

    assert preview["mapping_required"] is True
    assert preview["mapping_profile_id"] is None
    assert preview["summary"]["warnings"][0]["code"] == "multiple_compatible_profiles"
    assert preview["summary"]["warnings"][0]["message"] == "Найдено несколько совместимых профилей. Выберите профиль импорта."

    explicitly_selected = asyncio.run(service.create_preview(_upload(payload), profile_id=conflicting_profile.profile_id))
    assert explicitly_selected["mapping_required"] is False
    assert explicitly_selected["mapping_profile_id"] == conflicting_profile.profile_id


def test_identical_compatible_profiles_do_not_invent_profile_provenance(tmp_path):
    service = _make_service(tmp_path)
    payload = _flat_xlsx_bytes(tmp_path / "flat.xlsx")
    _preview, imported = asyncio.run(_import_preview(service, payload, save_profile=True, profile_name="Первый"))
    original_profile = service.repository.profile(imported["profile_id"])
    duplicate = original_profile.model_copy(update={
        "profile_id": "b" * 32,
        "name": "Дубликат",
        "revision": 1,
    })
    service.repository.save_profile(duplicate)

    changed_payload = _flat_xlsx_bytes(tmp_path / "flat-updated.xlsx", name="Гайка")
    preview = asyncio.run(service.create_preview(_upload(changed_payload)))
    result = asyncio.run(service.import_confirmed(ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=preview["sheet_name"],
        header_row=preview["header_row"], layout_type=preview["layout_type"],
        field_mapping=preview["field_mapping"], item_name_parse_strategy=preview["item_name_parse_strategy"],
    )))

    assert preview["mapping_required"] is False
    assert preview["mapping_profile_id"] is None
    assert result["profile_id"] is None
    assert result["active_import"]["mapping_provenance"]["profile_id"] is None


def test_changed_headers_require_remapping_and_preview_can_analyze_new_mapping(tmp_path):
    service = _make_service(tmp_path)
    first = _xlsx_bytes(tmp_path / "first.xlsx", rows=_hierarchical_rows())
    asyncio.run(_import_preview(service, first, save_profile=True, profile_name="Профиль"))
    changed_headers = ["Номенклатура / единица", "Цена брутто операции", "Объём поставки", "К оплате", "Момент операции", "Основание документа", "Организация-контрагент", "Условия договора", "Идентификатор позиции"]
    changed = _xlsx_bytes(tmp_path / "changed.xlsx", headers=changed_headers, rows=_hierarchical_rows())
    preview = asyncio.run(service.create_preview(_upload(changed)))
    assert preview["mapping_required"] is True
    analysis = PreviewMappingRequest(
        sheet_name=preview["sheet_name"], header_row=preview["header_row"], layout_type="hierarchical_grouped",
        field_mapping={
            "item_name": 0, "reported_unit_price_gross": 1, "quantity": 2, "amount_gross": 3,
            "document_date": 4, "document_reference": 5, "counterparty": 6, "contract": 7,
            "item_code": 8,
        }, item_name_parse_strategy="comma_suffix_unit",
    )
    analyzed = asyncio.run(service.analyze_preview(preview["preview_id"], analysis))
    assert analyzed["mapping_required"] is False
    assert analyzed["summary"]["event_count"] == 3
    request = ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=analyzed["sheet_name"], header_row=analyzed["header_row"],
        layout_type=analyzed["layout_type"], field_mapping=analyzed["field_mapping"],
        item_name_parse_strategy=analyzed["item_name_parse_strategy"],
    )
    assert asyncio.run(service.import_confirmed(request))["status"] == "succeeded"


def test_preview_can_select_another_sheet_and_rebuild_its_mapping(tmp_path):
    path = tmp_path / "two-sheets.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "TDSheet"
    first.append(HIERARCHICAL_HEADERS)
    for row in _hierarchical_rows():
        first.append(row)
    second = workbook.create_sheet("Purchase rows")
    headers = ["Код номенклатуры", "Наименование", "Ед. изм.", "Кол-во", "Цена с НДС", "Сумма с НДС", "Дата документа"]
    second.append(headers)
    second.append(["C-9", "Шайба", "шт", 2, 3, 6, datetime(2026, 9, 1)])
    workbook.save(path)
    service = _make_service(tmp_path / "store")
    preview = asyncio.run(service.create_preview(_upload(path.read_bytes())))
    assert preview["sheet_names"] == ["TDSheet", "Purchase rows"]
    inspected = asyncio.run(service.inspect_preview_sheet(preview["preview_id"], "Purchase rows"))
    assert inspected["header_row"] == 1
    assert inspected["field_mapping"]["item_code"] == 0
    analyzed = asyncio.run(service.analyze_preview(preview["preview_id"], PreviewMappingRequest(
        sheet_name="Purchase rows", header_row=inspected["header_row"], layout_type="flat",
        field_mapping=inspected["field_mapping"], item_name_parse_strategy="none",
    )))
    assert analyzed["summary"]["event_count"] == 1


def test_failed_staging_keeps_previous_active_snapshot_and_deletes_raw_file(tmp_path, monkeypatch):
    service = _make_service(tmp_path)
    first = _xlsx_bytes(tmp_path / "first.xlsx", rows=_hierarchical_rows())
    asyncio.run(_import_preview(service, first))
    old_sha = service.repository.active_metadata()["sha256"]
    second = _xlsx_bytes(tmp_path / "second.xlsx", rows=_hierarchical_rows(second_quantity=3))
    preview = asyncio.run(service.create_preview(_upload(second)))
    pending_path = service._pending[preview["preview_id"]].path
    monkeypatch.setattr(service.repository, "build_staging_snapshot", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("synthetic staging failure")))
    request = ImportMappingRequest(
        preview_id=preview["preview_id"], sheet_name=preview["sheet_name"], header_row=preview["header_row"],
        layout_type=preview["layout_type"], field_mapping=preview["field_mapping"],
        item_name_parse_strategy=preview["item_name_parse_strategy"],
    )
    with pytest.raises(OneCImportError):
        asyncio.run(service.import_confirmed(request))
    assert service.repository.active_metadata()["sha256"] == old_sha
    assert not pending_path.exists()


def test_preview_failure_deletes_raw_upload_and_unsafe_archive_is_rejected(tmp_path):
    service = _make_service(tmp_path)
    with pytest.raises(OneCImportError):
        asyncio.run(service.create_preview(_upload(b"not an xlsx")))
    assert not list(service.temp_root.glob("pending-*.xlsx"))
    archive_path = tmp_path / "zipbomb.xlsx"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/large.xml", b"0" * 300_000)
    with pytest.raises(OneCImportError, match="Коэффициент сжатия"):
        preflight_xlsx(archive_path)
    traversal_path = tmp_path / "traversal.xlsx"
    with zipfile.ZipFile(traversal_path, "w") as archive:
        archive.writestr("../outside.xml", b"not extracted")
    with pytest.raises(OneCImportError, match="небезопасные"):
        preflight_xlsx(traversal_path)


def test_shared_strings_case_fallback_is_narrow_and_does_not_mutate_source(tmp_path):
    source = tmp_path / "casing.xlsx"
    _xlsx_bytes(source, rows=_hierarchical_rows())
    original_detection = detect_workbook(source)
    original_parsed = parse_workbook(
        source, filename="original.xlsx", file_sha256="f" * 64,
        sheet_name=original_detection["sheet_name"], header_row=original_detection["header_row"],
        headers=original_detection["headers"], layout_type=original_detection["layout_type"],
        field_mapping=original_detection["field_mapping"],
        item_name_parse_strategy=original_detection["item_name_parse_strategy"],
    )
    original_logical_values = [
        (event.item_name, event.raw_unit, event.document_date, event.document_type, event.document_reference,
         event.quantity, event.reported_unit_price_gross, event.effective_unit_price_gross, event.amount_gross)
        for event in original_parsed.events
    ]
    replacements = {}
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    ET.register_namespace("", main_ns)
    ET.register_namespace("r", office_rel_ns)
    with zipfile.ZipFile(source) as original:
        members = {name: original.read(name) for name in original.namelist()}
    strings = []
    string_ids = {}
    sheet_root = ET.fromstring(members["xl/worksheets/sheet1.xml"])
    for cell in sheet_root.findall(f".//{{{main_ns}}}c"):
        inline = cell.find(f"{{{main_ns}}}is/{{{main_ns}}}t")
        if inline is None:
            continue
        value = inline.text or ""
        index = string_ids.setdefault(value, len(strings))
        if index == len(strings):
            strings.append(value)
        cell.attrib["t"] = "s"
        for child in list(cell):
            cell.remove(child)
        ET.SubElement(cell, f"{{{main_ns}}}v").text = str(index)
    members["xl/worksheets/sheet1.xml"] = ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
    rel_root = ET.fromstring(members["xl/_rels/workbook.xml.rels"])
    ET.SubElement(rel_root, f"{{{rel_ns}}}Relationship", {
        "Id": "rIdSharedStrings", "Type": f"{office_rel_ns}/sharedStrings", "Target": "sharedStrings.xml",
    })
    members["xl/_rels/workbook.xml.rels"] = ET.tostring(rel_root, encoding="utf-8", xml_declaration=True)
    types_root = ET.fromstring(members["[Content_Types].xml"])
    ET.SubElement(types_root, f"{{{content_ns}}}Override", {
        "PartName": "/xl/sharedStrings.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    })
    members["[Content_Types].xml"] = ET.tostring(types_root, encoding="utf-8", xml_declaration=True)
    shared_root = ET.Element(f"{{{main_ns}}}sst", {"count": str(len(strings)), "uniqueCount": str(len(strings))})
    for value in strings:
        item = ET.SubElement(shared_root, f"{{{main_ns}}}si")
        ET.SubElement(item, f"{{{main_ns}}}t").text = value
    members["xl/SharedStrings.xml"] = ET.tostring(shared_root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    source_before_preflight = source.read_bytes()
    normalized = preflight_xlsx(source)
    try:
        assert normalized is not None
        with zipfile.ZipFile(source) as original:
            assert "xl/SharedStrings.xml" in original.namelist()
            assert "xl/sharedStrings.xml" not in original.namelist()
            original_members = set(original.namelist())
            original_contents = {name: original.read(name) for name in original_members}
        with zipfile.ZipFile(normalized) as fixed:
            assert "xl/sharedStrings.xml" in fixed.namelist()
            assert "xl/SharedStrings.xml" not in fixed.namelist()
            assert set(fixed.namelist()) == (original_members - {"xl/SharedStrings.xml"}) | {"xl/sharedStrings.xml"}
            for name in original_members - {"xl/SharedStrings.xml"}:
                assert fixed.read(name) == original_contents[name]
            assert fixed.read("xl/sharedStrings.xml") == original_contents["xl/SharedStrings.xml"]
            assert fixed.read("[Content_Types].xml") == original_contents["[Content_Types].xml"]
        assert source.read_bytes() == source_before_preflight
        detected = detect_workbook(source)
        parsed = parse_workbook(
            source, filename="case.xlsx", file_sha256="d" * 64,
            sheet_name=detected["sheet_name"], header_row=detected["header_row"], headers=detected["headers"],
            layout_type=detected["layout_type"], field_mapping=detected["field_mapping"],
            item_name_parse_strategy=detected["item_name_parse_strategy"],
        )
        assert parsed.events[0].item_name == "Прокладка"
        assert parsed.events[0].document_reference == "Поступление товаров и услуг №1"
        fixed_logical_values = [
            (event.item_name, event.raw_unit, event.document_date, event.document_type, event.document_reference,
             event.quantity, event.reported_unit_price_gross, event.effective_unit_price_gross, event.amount_gross)
            for event in parsed.events
        ]
        assert fixed_logical_values == original_logical_values

        alternate = tmp_path / "alternate-case.xlsx"
        alternate_members = dict(members)
        alternate_members["xl/sharedstrings.xml"] = alternate_members.pop("xl/SharedStrings.xml")
        with zipfile.ZipFile(alternate, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, content in alternate_members.items():
                archive.writestr(name, content)
        with pytest.raises(OneCImportError, match="Несовпадение регистра"):
            preflight_xlsx(alternate)

        with zipfile.ZipFile(source, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/sharedStrings.xml", original_contents["xl/SharedStrings.xml"])
        with pytest.raises(OneCImportError, match="повторные компоненты без учёта регистра"):
            preflight_xlsx(source)
    finally:
        if normalized:
            normalized.unlink(missing_ok=True)
    assert not normalized.exists()


def test_formula_cells_are_rejected_without_execution(tmp_path):
    path = tmp_path / "formula.xlsx"
    _xlsx_bytes(path, rows=[["Заглушка, шт", None, None, None, None, None, None, None, None], [None, "=1+1", 1, 1, datetime(2026, 1, 1), "Документ", None, None, None]])
    detected = detect_workbook(path)
    with pytest.raises(OneCImportError, match="Формулы"):
        parse_workbook(
            path, filename="formula.xlsx", file_sha256="c" * 64, sheet_name=detected["sheet_name"],
            header_row=detected["header_row"], headers=detected["headers"], layout_type=detected["layout_type"],
            field_mapping=detected["field_mapping"], item_name_parse_strategy=detected["item_name_parse_strategy"],
        )
