from __future__ import annotations

import asyncio
import hashlib
import json as json_module
import threading
import time
from io import BytesIO
from urllib.parse import urlsplit

import fitz
import pytest
from openpyxl import load_workbook

from averon_import.services.ai_first.local_workflow import SyntheticPageAdapter
from averon_import.services.ai_first.private_acl import verify_private_tree
from averon_import.services.auth import CurrentUser, Role
from averon_import.services.workspace import WorkspaceService


class _Response:
    def __init__(self, status_code, body):
        self.status_code = status_code
        self.content = body
        self.text = body.decode("utf-8", errors="replace")

    def json(self):
        return json_module.loads(self.content)


class _ApiClient:
    def __init__(self, app):
        self.app = app

    def get(self, path):
        return self._request("GET", path)

    def post(self, path, *, json=None):
        return self._request("POST", path, json=json)

    def _request(self, method, target, *, json=None):
        parsed = urlsplit(target)
        body = b"" if json is None else json_module.dumps(json, ensure_ascii=False).encode("utf-8")
        headers = [(b"host", b"localhost")]
        if json is not None:
            headers.append((b"content-type", b"application/json"))
        messages = []
        sent = False

        async def receive():
            nonlocal sent
            if sent:
                return {"type": "http.disconnect"}
            sent = True
            return {"type": "http.request", "body": body, "more_body": False}

        async def send(message):
            messages.append(message)

        scope = {
            "type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1",
            "method": method, "scheme": "http", "path": parsed.path,
            "raw_path": parsed.path.encode(), "query_string": parsed.query.encode(),
            "headers": headers, "client": ("127.0.0.1", 12345), "server": ("localhost", 80),
        }
        asyncio.run(self.app(scope, receive, send))
        start = next(message for message in messages if message["type"] == "http.response.start")
        content = b"".join(message.get("body", b"") for message in messages if message["type"] == "http.response.body")
        return _Response(start["status"], content)


@pytest.fixture()
def ai_first_client(monkeypatch, tmp_path):
    monkeypatch.setenv("AVERON_DATA_DIR", str(tmp_path / "test-data"))
    from averon_import.services.ocr.qwen_vision_client import _HttpsTransport
    from averon_import import main

    http_attempts = []
    def forbid_http(*args, **kwargs):
        http_attempts.append(True)
        raise AssertionError("V2D offline workflow must not send HTTP")
    monkeypatch.setattr(_HttpsTransport, "post", forbid_http)
    monkeypatch.setattr(main, "workspace_service", WorkspaceService(tmp_path / "test-data"))
    monkeypatch.setattr(main, "AI_FIRST_UI_ENABLED", True)
    monkeypatch.setattr(main, "AI_FIRST_FAKE_OCR_ENABLED", True)
    monkeypatch.setattr(main, "AI_FIRST_LIVE_ENABLED", False)
    main.app.dependency_overrides[main.require_authenticated] = lambda: CurrentUser(
        username="synthetic-reviewer", role=Role.USER, user_id="synthetic-reviewer-id",
    )
    yield main, _ApiClient(main.app)
    main.app.dependency_overrides.clear()
    assert http_attempts == []


def _make_document(main, page_count: int = 2):
    pdf = fitz.open()
    for number in range(page_count):
        page = pdf.new_page(width=612, height=792)
        page.insert_text((72, 72), f"Synthetic test page {number + 1}")
    source = main.workspace_service.data_dir / "synthetic.pdf"
    source.parent.mkdir(parents=True, exist_ok=True)
    pdf.save(source)
    pdf.close()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workspace = main.workspace_service.create(
        source,
        {"filename": "synthetic.pdf", "page_count": page_count, "size": source.stat().st_size},
        source_sha256=digest,
    )
    return workspace


def _make_sized_document(main, dimensions):
    pdf = fitz.open()
    for number, (width, height) in enumerate(dimensions, 1):
        page = pdf.new_page(width=width, height=height)
        page.insert_text((36, 48), f"Synthetic sized page {number}")
    source = main.workspace_service.data_dir / "synthetic-sized.pdf"
    source.parent.mkdir(parents=True, exist_ok=True)
    pdf.save(source)
    pdf.close()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    return main.workspace_service.create(
        source,
        {"filename":"synthetic-sized.pdf","page_count":len(dimensions),"size":source.stat().st_size},
        source_sha256=digest,
    )


def _wait_for_job(client, job_id, timeout_s=15):
    deadline = time.monotonic() + timeout_s
    last_job = None
    while time.monotonic() < deadline:
        response = client.get(f"/api/jobs/{job_id}")
        assert response.status_code == 200, response.text
        job = response.json()
        last_job = job
        if job["status"] in {"completed", "failed", "expired"}:
            return job
        time.sleep(0.025)
    raise AssertionError(f"AI-first job did not finish: {json_module.dumps(last_job, ensure_ascii=False)}")


def _start_fake(client, document_id, *, pages=None, context=None):
    response = client.post(
        f"/api/documents/{document_id}/ai-first/recognize",
        json={"mode": "fake", "pages": pages, "user_context": context},
    )
    assert response.status_code == 200, response.text
    job = _wait_for_job(client, response.json()["id"])
    assert job["status"] == "completed", job
    return job["result"]


def test_fake_ai_first_runs_all_pdf_pages_and_preserves_private_local_state(ai_first_client, tmp_path):
    main, client = ai_first_client
    workspace = _make_document(main, 2)

    private_hint = "Синтетическая подсказка конфиденциальна"
    result = _start_fake(client, workspace.document_id, context=private_hint)

    assert result["mode"] == "fake"
    assert result["status"] == "completed"
    assert result["completed_pages"] == 2
    assert result["counts"] == {"tables": 2, "rows": 4, "cells": 24}
    snapshot = client.get(f"/api/documents/{workspace.document_id}/ai-first")
    assert snapshot.status_code == 200, snapshot.text
    body = snapshot.json()
    assert body["draft"]["expected_page_numbers"] == [1, 2]
    assert [page["page_number"] for page in body["draft"]["pages"]] == [1, 2]
    assert body["mode"] == "fake_vision"
    assert body["pages"] and all(page["state"] == "completed" for page in body["pages"])
    assert body["review_ready_for_guard"]["ready"] is False
    assert not workspace.result_path.exists()
    assert not workspace.review_decisions_path.exists()
    for name in ("transcription.json", "review.json", "report.json", "progress.json"):
        assert (workspace.root / "ai_first" / name).is_file()
    assert verify_private_tree(workspace.root / "ai_first") >= 5
    assert not list((workspace.root / "ai_first").glob("*.png"))
    assert all(private_hint not in path.read_text(encoding="utf-8") for path in (workspace.root / "ai_first").rglob("*.json"))
    report_text = (workspace.root / "ai_first" / "report.json").read_text(encoding="utf-8")
    transcription_text = (workspace.root / "ai_first" / "transcription.json").read_text(encoding="utf-8")
    assert "Демо-позиция A" not in report_text
    assert "Демо-позиция A" in transcription_text


def test_fake_ai_first_completes_28_pages_including_large_format_and_persists_checkpoint(ai_first_client):
    main, client = ai_first_client
    dimensions = [(612, 792) if number % 2 else (1190, 1684) for number in range(1, 29)]
    workspace = _make_sized_document(main, dimensions)

    result = _start_fake(client, workspace.document_id)

    assert result["status"] == "completed"
    assert result["completed_pages"] == result["total_pages"] == 28
    assert result["counts"] == {"tables":28,"rows":56,"cells":336}
    state = client.get(f"/api/documents/{workspace.document_id}/ai-first").json()
    assert len(state["draft"]["pages"]) == 28
    assert all(page["status"] == "extracted" for page in state["draft"]["pages"])
    assert all(page["state"] == "completed" for page in state["pages"])
    assert len(list((workspace.root / "ai_first" / "checkpoint").glob("page-*.json"))) == 28


def test_oversized_page_fails_safely_and_processing_issue_cannot_be_resolved(ai_first_client):
    main, client = ai_first_client
    workspace = _make_sized_document(main, [(3000, 4000)])

    result = _start_fake(client, workspace.document_id)
    state = client.get(f"/api/documents/{workspace.document_id}/ai-first").json()
    assert result["status"] == "completed_with_errors"
    assert result["completed_pages"] == 0
    assert state["draft"]["pages"][0]["status"] == "failed"
    assert state["pages"][0]["safe_error_code"] == "image_pixel_limit"
    issue = next(item for item in state["draft"]["issues"] if item["code"] == "page_extraction_failed")
    response = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/review/resolve-issue",
        json={"issue_id":issue["id"],"detail":"Synthetic check"},
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "processing_issue_requires_retry"
    assert client.get(f"/api/documents/{workspace.document_id}/ai-first").json()["review"]["resolved_issue_ids"] == []


def test_review_edits_adds_excludes_and_bulk_confirms_without_changing_source_draft(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    _start_fake(client, workspace.document_id)
    path = f"/api/documents/{workspace.document_id}/ai-first"
    original = client.get(path).json()
    rows = original["draft"]["pages"][0]["tables"][0]["rows"]
    first_id = rows[0]["id"]
    raw_name = rows[0]["raw_cells"][1]["raw_value"]

    edited = client.post(
        f"{path}/review/row",
        json={"record_id": first_id, "disposition": "corrected", "values": {"cell:1": "Проверенное название"}},
    )
    assert edited.status_code == 200, edited.text
    state = client.get(path).json()
    assert state["draft"]["pages"][0]["tables"][0]["rows"][0]["raw_cells"][1]["raw_value"] == raw_name
    assert state["review"]["decisions"][first_id]["values"]["cell:1"] == "Проверенное название"

    added = client.post(
        f"{path}/review/add-row",
        json={"page_number": 1, "values": {"cell:0": "3", "cell:1": "Ручная синтетическая позиция"}},
    )
    assert added.status_code == 200, added.text
    state = client.get(path).json()
    manual = next(row for row in state["review"]["manual_rows"] if row["active"])
    assert manual["id"] in state["review"]["decisions"]

    excluded = client.post(
        f"{path}/review/row",
        json={"record_id": rows[1]["id"], "disposition": "excluded", "values": {}},
    )
    assert excluded.status_code == 200, excluded.text
    assert client.get(path).json()["review"]["decisions"][rows[1]["id"]]["disposition"] == "excluded"

    bulk = client.post(
        f"{path}/review/bulk-confirm",
        json={"record_ids": [first_id]},
    )
    assert bulk.status_code == 200, bulk.text
    final = client.get(path).json()
    assert final["review"]["decisions"][first_id]["disposition"] == "confirmed"
    assert [event["action"] for event in final["review"]["history"]].count("edit_value") == 1
    assert "add_row" in [event["action"] for event in final["review"]["history"]]
    assert "exclude_row" in [event["action"] for event in final["review"]["history"]]
    assert client.post(f"{path}/review/save").json()["saved"] is True


def test_pilot_xlsx_exports_saved_review_fragment_without_production_export(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main, 2)
    _start_fake(client, workspace.document_id, pages=[1])
    path = f"/api/documents/{workspace.document_id}/ai-first"
    snapshot = client.get(path).json()
    rows = snapshot["draft"]["pages"][0]["tables"][0]["rows"]
    first_id, second_id = (row["id"] for row in rows)

    assert client.post(f"{path}/review/row-role", json={"record_id":first_id,"role":"standalone_item"}).status_code == 200
    assert client.post(f"{path}/review/row-role", json={"record_id":second_id,"role":"section"}).status_code == 200
    assert client.post(f"{path}/review/row", json={
        "record_id":first_id,"disposition":"corrected","values":{"cell:1":"Synthetic corrected item"},
    }).status_code == 200
    assert client.post(f"{path}/review/row", json={"record_id":second_id,"disposition":"confirmed"}).status_code == 200
    assert client.post(f"{path}/review/complete-page", json={"page_number":1}).status_code == 200

    response = client.post(f"{path}/pilot-export", json={"page_numbers":[1]})
    assert response.status_code == 200, response.text
    assert not workspace.result_path.exists()
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert {"Проверенные позиции", "Все распознанные строки", "Охват страниц", "Исходная структура"} <= set(workbook.sheetnames)
    assert workbook["Проверенные позиции"].max_row == 5
    assert workbook["Все распознанные строки"].max_row == 6
    assert "ЧАСТИЧНЫЙ ОХВАТ ДОКУМЕНТА" in workbook["Проверенные позиции"]["A2"].value


def test_ai_first_locator_returns_only_safe_no_match_status_for_fake_page(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    _start_fake(client, workspace.document_id)
    state = client.get(f"/api/documents/{workspace.document_id}/ai-first").json()
    row = state["draft"]["pages"][0]["tables"][0]["rows"][0]

    response = client.get(
        f"/api/documents/{workspace.document_id}/ai-first/locator?record_id={row['id']}"
    )
    assert response.status_code == 200, response.text
    assert response.json()["matched"] is False
    assert "bbox" not in response.json()
    assert "raw_value" not in response.text


def test_corrected_quantity_survives_confirmation_save_and_reopen(ai_first_client):
    _main, client = ai_first_client
    workspace = _make_document(_main, 1)
    _start_fake(client, workspace.document_id)
    path = f"/api/documents/{workspace.document_id}/ai-first"
    before = client.get(path).json()
    row = before["draft"]["pages"][0]["tables"][0]["rows"][0]

    correction = client.post(
        f"{path}/review/row",
        json={"record_id":row["id"],"disposition":"corrected","values":{"cell:4":"9"}},
    )
    assert correction.status_code == 200, correction.text
    confirmed = client.post(
        f"{path}/review/row",
        json={"record_id":row["id"],"disposition":"confirmed","values":{}},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert client.post(f"{path}/review/save").json()["saved"] is True

    reopened = client.get(path).json()
    decision = reopened["review"]["decisions"][row["id"]]
    assert decision["disposition"] == "confirmed"
    assert decision["values"]["cell:4"] == "9"
    confirm_event = next(event for event in reversed(reopened["review"]["history"]) if event["action"] == "confirm_row")
    assert confirm_event["payload"]["values"]["cell:4"] == "9"


def test_review_change_during_result_publication_is_serialized_and_audited(ai_first_client, monkeypatch):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    _start_fake(client, workspace.document_id)
    path = f"/api/documents/{workspace.document_id}/ai-first"
    row = client.get(path).json()["draft"]["pages"][0]["tables"][0]["rows"][0]

    from averon_import.services.ai_first.local_workflow import AIWorkspaceStore
    original_write = AIWorkspaceStore.write
    publication_blocked = threading.Event()
    release_publication = threading.Event()

    def pause_publication(store, name, payload):
        if name == "transcription.json" and not publication_blocked.is_set():
            publication_blocked.set()
            assert release_publication.wait(5), "test did not release AI-first publication"
        return original_write(store, name, payload)

    monkeypatch.setattr(AIWorkspaceStore, "write", pause_publication)
    start = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/recognize",
        json={"mode":"fake","pages":[1]},
    )
    assert start.status_code == 200, start.text
    job_id = start.json()["id"]
    assert publication_blocked.wait(5), "recognition should reach atomic publication"

    review_started = threading.Event()
    review_finished = threading.Event()
    review_result = {}

    def update_review():
        review_started.set()
        review_result["response"] = client.post(
            f"{path}/review/row",
            json={"record_id":row["id"],"disposition":"corrected","values":{"cell:4":"11"}},
        )
        review_finished.set()

    reviewer = threading.Thread(target=update_review)
    reviewer.start()
    assert review_started.wait(2)
    assert not review_finished.wait(0.15), "review mutation must wait while publication holds the document lock"
    release_publication.set()
    job = _wait_for_job(client, job_id)
    reviewer.join(timeout=3)
    assert not reviewer.is_alive()
    assert job["status"] == "completed", job
    assert review_result["response"].status_code == 200, review_result["response"].text

    reopened = client.get(path).json()
    decision = reopened["review"]["decisions"][row["id"]]
    assert decision["disposition"] == "corrected"
    assert decision["values"]["cell:4"] == "11"
    assert any(
        event["action"] == "edit_value" and event["payload"]["values"].get("cell:4") == "11"
        for event in reopened["review"]["history"]
    )


def test_ai_first_server_routes_are_hidden_when_experiment_flag_is_disabled(ai_first_client, monkeypatch):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    monkeypatch.setattr(main, "AI_FIRST_UI_ENABLED", False)
    base = f"/api/documents/{workspace.document_id}/ai-first"
    blocked = [
        client.post(f"{base}/recognize", json={"mode":"fake","pages":[1]}),
        client.get(base),
        client.get(f"{base}/progress"),
        client.post(f"{base}/review/row", json={"record_id":"row","disposition":"confirmed","values":{}}),
        client.post(f"{base}/review/row-role", json={"record_id":"row","role":"standalone_item"}),
        client.post(f"{base}/review/bulk-confirm", json={"record_ids":["row"]}),
        client.post(f"{base}/review/add-row", json={"page_number":1,"values":{"cell:0":"1"}}),
        client.post(f"{base}/review/complete-page", json={"page_number":1}),
        client.post(f"{base}/review/resolve-issue", json={"issue_id":"issue","detail":"checked"}),
        client.post(f"{base}/review/save"),
        client.post(f"{base}/pilot-export", json={"page_numbers":[1]}),
        client.get(f"{base}/locator?record_id=row"),
        client.post("/api/jobs/not-a-job/cancel"),
    ]
    assert [response.status_code for response in blocked] == [404] * len(blocked)
    assert all("Экспериментальный режим выключен" in response.text for response in blocked)
    assert not (workspace.root / "ai_first").exists()


def test_bulk_confirmation_does_not_confirm_page_or_resolve_remarks(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    _start_fake(client, workspace.document_id)
    path = f"/api/documents/{workspace.document_id}/ai-first"
    initial = client.get(path).json()
    assert initial["review_ready_for_guard"]["ready"] is False
    rows = initial["draft"]["pages"][0]["tables"][0]["rows"]
    assert client.post(f"{path}/review/bulk-confirm", json={"record_ids":[row["id"] for row in rows]}).status_code == 200
    after_rows = client.get(path).json()
    assert after_rows["review"]["confirmed_complete_pages"] == []
    assert after_rows["review_ready_for_guard"]["ready"] is False

    page = after_rows["draft"]["pages"][0]
    assert client.post(f"{path}/review/complete-page", json={"page_number":1}).status_code == 200
    still_blocked = client.get(path).json()
    assert still_blocked["review_ready_for_guard"]["ready"] is False
    for issue in page["issues"]:
        resolved = client.post(
            f"{path}/review/resolve-issue",
            json={"issue_id":issue["id"],"detail":"Сверено с синтетической страницей."},
        )
        assert resolved.status_code == 200, resolved.text
    assert client.get(path).json()["review_ready_for_guard"]["ready"] is False
    for row in rows:
        classified = client.post(
            f"{path}/review/row-role",
            json={"record_id": row["id"], "role": "standalone_item"},
        )
        assert classified.status_code == 200, classified.text
    assert client.get(path).json()["review_ready_for_guard"]["ready"] is True


def test_context_limit_is_fail_closed_and_context_is_not_echoed(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    hint = "SYNTHETIC-PRIVATE-HINT-" + ("x" * 1000)
    response = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/recognize",
        json={"mode": "fake", "pages": [1], "user_context": hint},
    )
    assert response.status_code == 422
    assert hint not in response.text
    assert not (workspace.root / "ai_first").exists()


def test_qwen_requires_server_flag_and_explicit_per_run_ack(ai_first_client, monkeypatch):
    main, client = ai_first_client
    workspace = _make_document(main, 1)
    missing_server_flag = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/recognize",
        json={"mode": "qwen", "pages": [1], "live_acknowledged": True},
    )
    assert missing_server_flag.status_code == 403

    monkeypatch.setattr(main, "AI_FIRST_LIVE_ENABLED", True)
    missing_ack = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/recognize",
        json={"mode": "qwen", "pages": [1], "live_acknowledged": False},
    )
    assert missing_ack.status_code == 400
    assert not (workspace.root / "ai_first" / "checkpoint").exists()


def test_cancel_endpoint_only_cancels_ai_first_job(ai_first_client, monkeypatch):
    main, client = ai_first_client
    workspace = _make_document(main, 3)
    original = SyntheticPageAdapter.extract
    second_page_started = threading.Event()
    calls = {}

    async def slow_extract(self, request):
        calls[request.page_number] = calls.get(request.page_number, 0) + 1
        if request.page_number == 2 and not second_page_started.is_set():
            second_page_started.set()
            await asyncio.sleep(0.5)
        return await original(self, request)

    monkeypatch.setattr(SyntheticPageAdapter, "extract", slow_extract)
    start = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/recognize",
        json={"mode": "fake", "pages": None},
    )
    assert start.status_code == 200, start.text
    job_id = start.json()["id"]
    assert second_page_started.wait(3), "first page should complete before the second starts"
    cancel = client.post(f"/api/jobs/{job_id}/cancel")
    assert cancel.status_code == 200, cancel.text
    job = _wait_for_job(client, job_id)
    assert job["status"] == "completed"
    assert job["result"]["status"] == "cancelled"
    assert job["result"]["completed_pages"] == 1
    stored = client.get(f"/api/documents/{workspace.document_id}/ai-first").json()
    assert stored["draft"]["pages"][0]["status"] == "extracted"
    assert stored["draft"]["pages"][1]["status"] == "failed"

    resume = client.post(
        f"/api/documents/{workspace.document_id}/ai-first/recognize",
        json={"mode":"fake","pages":None,"retry_pages":[2]},
    )
    assert resume.status_code == 200, resume.text
    resumed = _wait_for_job(client, resume.json()["id"])
    assert resumed["status"] == "completed", json_module.dumps(resumed, ensure_ascii=False)
    assert resumed["result"]["status"] == "completed"
    assert calls == {1:1, 2:2, 3:1}


def test_existing_averon_template_exposes_ai_first_review_only_behind_server_config():
    from pathlib import Path

    root = Path(__file__).parents[1]
    template = (root / "averon_import" / "templates" / "index.html").read_text(encoding="utf-8")
    app_js = (root / "averon_import" / "static" / "app.js").read_text(encoding="utf-8")
    ai_js = (root / "averon_import" / "static" / "ai_first_ui.js").read_text(encoding="utf-8")
    assert 'id="ai-first-local-controls" hidden' in template
    assert 'id="ai-first-context" maxlength="1000"' in template
    assert 'id="ai-first-review-view"' in template
    assert 'id="ai-first-pdf-image"' in template and 'id="ai-first-table-scroll"' in template
    assert 'id="ai-first-confirm-selected"' in template and 'id="ai-first-confirm-page"' in template
    assert 'if (state.aiFirst.mode !== "standard") return startAIFirstRecognition();' in app_js
    assert 'postAIFirstReview("bulk-confirm"' in ai_js
    assert 'postAIFirstReview("complete-page"' in ai_js
    assert 'postAIFirstReview("resolve-issue"' in ai_js
    assert 'postAIFirstReview("add-row"' in ai_js
    assert 'postAIFirstReview("row"' in ai_js
    assert 'ai-first/review/save' in ai_js
    assert "aiFirstRowInputSnapshot" in ai_js
    assert "Сохранить несохранённые изменения и подтвердить эту строку?" in ai_js
    assert "несохранённых правок нет" in ai_js
    assert 'if (mode === "qwen") {' in ai_js and 'updateAIFirstModeControls();' in ai_js
    assert 'page.safe_error_code' in ai_js
