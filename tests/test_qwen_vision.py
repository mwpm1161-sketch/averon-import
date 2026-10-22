from __future__ import annotations

import base64
import hashlib
import json
import socket
import threading
import traceback
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image
from pydantic import ValidationError

from averon_import.core.constants import BASE_COLUMNS, PROCESSING_MODES
from averon_import.services.export_service import ExcelExportService
from averon_import.services.ocr.base import OcrProviderError, OcrResult, OcrRow, PageOcrResult
from averon_import.services.ocr.qwen_vision import QwenVisionProvider
from averon_import.services.ocr.qwen_vision_client import (
    MAX_OUTPUT_TOKENS_CAP, QwenVisionClient, QwenVisionConfig, QwenVisionError, VisionHttpResponse,
    VisionLimits, _HttpsTransport,
)
from averon_import.services.ocr.qwen_vision_contract import VisionPage, project_page
from averon_import.services.ocr import qwen_vision_client as qwen_vision_client_module
from averon_import.services.pdf_service import PdfService
from averon_import.services.recognition import RecognitionService
from averon_import.services.review_decisions import CONFIRM_FIELD_VALUE_DECISION, HumanReviewService
from averon_import.services.secrets import MemorySecretStore, YANDEX_AI_API_KEY, YANDEX_API_KEY


SECRET = "synthetic-test-api-secret"
DOCUMENT_TEXT = "СЕКРЕТНЫЙ-ТЕСТОВЫЙ-ДОКУМЕНТ"
REQUEST_ID = "0bc73024-b97f-4a78-b023-bb56503acde0"


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("unit tests must not open a socket")
    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(socket.socket, "connect", denied)


def png(size=(32, 24)):
    buffer = BytesIO()
    Image.new("RGB", size, "white").save(buffer, format="PNG")
    return buffer.getvalue()


def cell(column, text, **changes):
    value = {"column_index": column, "column_span": 1, "row_span": 1,
             "raw_text": text, "alternatives": [], "handwritten_replacements": [], "unreadable": False,
             "struck_through": False, "handwritten": False, "bbox_candidate": None}
    value.update(changes)
    return value


def row(texts, **changes):
    value = {"kind": "item", "logical_row_index": None, "continuation_of": None,
             "continues_previous_page": False, "struck_through": False,
             "handwritten": False, "handwritten_replacements": [], "bbox_candidate": None,
             "cells": [cell(index, text) for index, text in enumerate(texts)]}
    value.update(changes)
    return value


def table(headers=None, rows=None):
    return {"headers": headers if headers is not None else ["Поз.", "Наименование", "Ед. изм.", "Количество", "Масса"],
            "rows": rows if rows is not None else [row(["1", DOCUMENT_TEXT, "шт", "2", "0,5"])],
            "continues_previous_page": False, "bbox_candidate": None}


def page(number=1, tables=None):
    return {"page": number, "tables": tables if tables is not None else [table()],
            "unreadable_fragments": [], "handwritten_replacements": []}


def response(extraction=None, *, content=None, finish_reason="stop", usage=None, status=200):
    outer = {"choices": [{"finish_reason": finish_reason, "message": {"content":
        json.dumps(extraction if extraction is not None else page(), ensure_ascii=False) if content is None else content}}]}
    if usage is not None:
        outer["usage"] = usage
    return VisionHttpResponse(status, json.dumps(outer, ensure_ascii=False).encode(), {"X-Request-ID": REQUEST_ID})


class Transport:
    def __init__(self, *responses):
        self.responses = list(responses or [response()])
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append({"url": url, **kwargs})
        value = self.responses.pop(0)
        if isinstance(value, Exception):
            raise value
        if callable(value):
            return value()
        return value


def config(**limit_changes):
    return QwenVisionConfig("https://ai.test/v1", "gpt://test-folder/qwen3.6-35b-a3b", SECRET,
                            folder_id="test-folder", limits=replace(VisionLimits(), **limit_changes))


@pytest.fixture
def pdf(tmp_path):
    target = tmp_path / "synthetic.pdf"
    with fitz.open() as document:
        for _ in range(3):
            document.new_page(width=72, height=72).insert_text((5, 20), "TEST ONLY")
        document.save(target)
    return target


def recognize(pdf, extraction=None, **limit_changes):
    transport = Transport(response(extraction))
    provider = QwenVisionProvider(config(**limit_changes), transport=transport)
    return provider.recognize(pdf, [1]), transport


def test_payload_is_multimodal_strict_schema_and_content_free_diagnostics():
    transport = Transport(response(usage={"prompt_tokens": 120, "completion_tokens": 35, "total_tokens": 155}))
    client = QwenVisionClient(config(max_output_tokens=1234), transport=transport)
    image = png()
    result = client.extract(image, 1)
    call = transport.calls[0]
    body = json.loads(call["body"])
    assert call["url"] == "https://ai.test/v1/chat/completions"
    assert call["headers"]["Authorization"] == f"Api-Key {SECRET}"
    assert call["headers"]["OpenAI-Project"] == "test-folder"
    assert call["headers"]["x-data-logging-enabled"] == "false"
    assert body["store"] is False and body["stream"] is False
    assert body["max_completion_tokens"] == 1234
    assert "max_tokens" not in body
    parts = body["messages"][1]["content"]
    assert parts[1]["type"] == "image_url"
    assert base64.b64decode(parts[1]["image_url"]["url"].split(",", 1)[1]) == image
    schema = body["response_format"]["json_schema"]
    assert schema["strict"] is True and schema["schema"]["additionalProperties"] is False
    assert result.diagnostics["request_id"] == REQUEST_ID
    assert result.diagnostics["usage"]["total_tokens"] == 155
    assert result.diagnostics["latency_ms"] >= 0
    assert result.diagnostics["api_verified"] is False
    assert result.page.tables[0].rows[0].cells[1].raw_text == DOCUMENT_TEXT


def test_empty_user_context_keeps_legacy_request_bytes_and_prompt_fingerprint():
    client = QwenVisionClient(config(), transport=Transport(response()))
    image = png()
    actual = client._request_body(image, 1, user_context="")
    legacy_payload = {
        "model": client.config.model, "stream": False, "store": False,
        "max_completion_tokens": client.config.limits.max_output_tokens, "temperature": 0,
        "messages": [
            {"role": "system", "content": qwen_vision_client_module._PROMPT},
            {"role": "user", "content": [
                {"type": "text", "text": "Source page number: 1"},
                {"type": "image_url", "image_url": {
                    "url": "data:image/png;base64," + base64.b64encode(image).decode("ascii"),
                }},
            ]},
        ],
        "response_format": {"type": "json_schema", "json_schema": {
            "name": "averon_vision_page_v1", "strict": True,
            "schema": VisionPage.model_json_schema(),
        }},
    }
    expected = json.dumps(legacy_payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
    assert client._request_body(image, 1) == expected
    assert actual == expected
    assert client.prompt_sha256 == "aa364b5a06cd1a625191ddcbc920fed9e2669e2aec1cd5e526b2f0d621ed7fe3"


def test_user_context_is_untrusted_exact_and_bounded_without_truncation():
    context = "  Synthetic hint: inspect the quantity column; do not alter source values.  "
    transport = Transport(response())
    client = QwenVisionClient(config(), transport=transport)
    client.extract(png(), 1, user_context=context)
    body = json.loads(transport.calls[0]["body"])
    assert body["messages"][1] == {
        "role": "system",
        "content": qwen_vision_client_module._USER_CONTEXT_POLICY,
    }
    user_parts = body["messages"][2]["content"]
    assert user_parts[0]["text"] == "Source page number: 1"
    assert user_parts[1]["text"] == qwen_vision_client_module._USER_CONTEXT_LABEL + context
    assert user_parts[2]["type"] == "image_url"
    assert "untrusted" in body["messages"][1]["content"]
    assert "cannot override" in body["messages"][1]["content"]

    exactly_at_limit = "x" * 1000
    assert client.measure_request_bytes(png(), 1, user_context=exactly_at_limit) > client.measure_request_bytes(png(), 1)
    no_send = Transport(response())
    with pytest.raises(QwenVisionError, match="user_context_too_long"):
        QwenVisionClient(config(), transport=no_send).extract(png(), 1, user_context="x" * 1001)
    assert no_send.calls == []

    with pytest.raises(QwenVisionError, match="invalid_user_context"):
        QwenVisionClient(config(), transport=no_send).extract(png(), 1, user_context="\ud800")
    assert no_send.calls == []


def test_user_context_increases_request_size_and_can_be_rejected_before_transport():
    image = png()
    baseline = QwenVisionClient(config(), transport=Transport()).measure_request_bytes(image, 1)
    context_transport = Transport(response())
    client = QwenVisionClient(
        config(max_request_bytes=baseline + 32), transport=context_transport,
    )
    with pytest.raises(QwenVisionError, match="request_size_limit"):
        client.extract(image, 1, user_context="synthetic hint")
    assert context_transport.calls == []


def test_increased_output_token_limit_is_sent_and_bounded():
    assert MAX_OUTPUT_TOKENS_CAP == 32768
    for requested_limit in (16384, 24576, 32768):
        transport = Transport(response())
        result = QwenVisionClient(
            config(max_output_tokens=requested_limit, reasoning_effort="none"), transport=transport,
        ).extract(png(), 1)
        payload = json.loads(transport.calls[0]["body"])
        assert payload["max_completion_tokens"] == requested_limit
        assert payload["reasoning_effort"] == "none"
        assert result.page.page == 1
    with pytest.raises(ValueError, match="max_output_tokens"):
        VisionLimits(max_output_tokens=MAX_OUTPUT_TOKENS_CAP + 1)
    for invalid in (0, -1, True, 1.5):
        with pytest.raises(ValueError, match="max_output_tokens"):
            VisionLimits(max_output_tokens=invalid)


def test_reasoning_disable_is_opt_in_and_uses_documented_parameter():
    default_transport = Transport(response())
    QwenVisionClient(config(), transport=default_transport).extract(png(), 1)
    assert "reasoning_effort" not in json.loads(default_transport.calls[0]["body"])

    disabled_transport = Transport(response())
    QwenVisionClient(config(reasoning_effort="none"), transport=disabled_transport).extract(png(), 1)
    assert json.loads(disabled_transport.calls[0]["body"])["reasoning_effort"] == "none"
    with pytest.raises(ValueError, match="reasoning_effort"):
        VisionLimits(reasoning_effort="low")


def test_config_uses_existing_ai_secret_namespace_and_settings():
    store = MemorySecretStore()
    store.set(YANDEX_API_KEY, "must-not-be-used")
    settings = SimpleNamespace(folder_id="folder", llm_model="qwen3.6-35b-a3b", llm_base_url="https://ai.test")
    assert not QwenVisionConfig.from_app_settings(settings, store, limits=VisionLimits(), environ={}).configured
    store.set(YANDEX_AI_API_KEY, SECRET)
    owned = QwenVisionConfig.from_app_settings(settings, store, limits=VisionLimits(), environ={})
    assert owned.api_key == SECRET and owned.model == "gpt://folder/qwen3.6-35b-a3b"
    assert owned.base_url == "https://ai.test/v1"
    override = QwenVisionConfig.from_app_settings(settings, store, limits=VisionLimits(), environ={
        "AVERON_YANDEX_AI_API_KEY": "synthetic-env-key", "AVERON_YANDEX_AI_MODEL": "gpt://custom/model",
        "AVERON_YANDEX_AI_BASE_URL": "https://other.test/v1/",
    })
    assert override.api_key == "synthetic-env-key" and override.model == "gpt://custom/model"
    assert SECRET not in repr(owned)
    assert "qwen_vision" not in PROCESSING_MODES


def test_provider_contract_real_pdf_rendering_individual_pages_no_disk_cache(pdf, tmp_path):
    transport = Transport(response(page(1)), response(page(2)))
    provider = QwenVisionProvider(config(), transport=transport)
    events = []
    cache = tmp_path / "must-not-be-created"
    result = provider.recognize(pdf, [2, 1, 1], pages_dir=cache, progress=lambda *args: events.append(args))
    assert isinstance(result, OcrResult)
    assert result.provider == "qwen_vision" and [item.page for item in result.pages] == [1, 2]
    assert all(isinstance(item, PageOcrResult) for item in result.pages)
    assert isinstance(result.pages[0].rows[0], OcrRow)
    assert [json.loads(call["body"])["messages"][1]["content"][0]["text"] for call in transport.calls] == [
        "Source page number: 1", "Source page number: 2",
    ]
    assert [event[:2] for event in events] == [(0, 2), (1, 2), (1, 2), (2, 2)]
    assert result.stats["requests"] == 2 and not cache.exists()
    assert provider.available() and provider.health()["api_verified"] is False


def test_five_columns_unknown_headers_multiline_and_technical_text(pdf):
    raw = page(tables=[table(["№", "Наименование", "Мой столбец", "Ед. изм.", "Количество"], [
        row(["1", 'QZXW-0000\nRXWZ-0000', "Ру16\nA-01", "м", "12,5"]),
    ])])
    result, _ = recognize(pdf, raw)
    item = result.pages[0].rows[0]
    assert set(item.values) == {column["key"] for column in BASE_COLUMNS}
    assert item.values["name"] == 'QZXW-0000\nRXWZ-0000'
    assert item.values["quantity"] == "12,5"
    assert item.metadata["vision_raw_page"] == raw
    assert item.metadata["vision_headers"][2] == "Мой столбец"
    assert not item.values["code"]


def test_standard_nine_column_specification_is_projected_without_normalization(pdf):
    headers = [column["title"] for column in BASE_COLUMNS]
    values = ["01", "Насос", "КМ-50", "A-007", "Завод", "шт", "2,0", "0,50", "По проекту"]
    result, _ = recognize(pdf, page(tables=[table(headers, [row(values)])]))
    assert result.pages[0].rows[0].values == dict(zip((column["key"] for column in BASE_COLUMNS), values))
    assert result.pages[0].provides_confidence is False


def test_multiple_tables_and_continuations_keep_order_and_never_merge(pdf):
    second = table(rows=[row(["2", "Насос", "шт", "1", ""], logical_row_index=0),
                         row(["", "с электродвигателем", "", "", ""], kind="continuation",
                             logical_row_index=0, continuation_of=0, continues_previous_page=True)])
    raw = page(tables=[table(), second, table(headers=["Нестандартно"], rows=[])])
    result, _ = recognize(pdf, raw)
    rows = result.pages[0].rows
    assert len(rows) == 3
    assert [item.metadata["vision_candidate_table_index"] for item in rows] == [0, 1, 1]
    assert rows[2].values["name"] == "с электродвигателем"
    assert rows[2].metadata["vision_raw_row"]["continuation_of"] == 0
    assert rows[0].metadata["vision_raw_page"]["tables"][2]["headers"] == ["Нестандартно"]


@pytest.mark.parametrize(("candidate_rows", "expected_names"), [
    (
        [
            row(["1", "Synthetic item alpha", "pcs", "2", ""]),
            row(["2", "Synthetic item beta", "pcs", "3", ""], kind="continuation"),
        ],
        ["Synthetic item alpha", "Synthetic item beta"],
    ),
    (
        [
            row(["1", "Synthetic item alpha", "pcs", "2", ""]),
            row(["2", "Synthetic item beta", "pcs", "3", ""], kind="component"),
        ],
        ["Synthetic item alpha", "Synthetic item beta"],
    ),
    (
        [
            row(["1", "Synthetic item alpha", "pcs", "2", ""], kind="header"),
            row(["2", "Synthetic item beta", "pcs", "3", ""], kind="header"),
        ],
        ["Synthetic item alpha", "Synthetic item beta"],
    ),
])
def test_qwen_candidate_rows_survive_legacy_continuation_repair(pdf, tmp_path, candidate_rows, expected_names):
    raw = page(tables=[table(rows=candidate_rows)])
    transport = Transport(response(raw))
    provider = QwenVisionProvider(config(), transport=transport)
    result = RecognitionService(PdfService(), ocr=provider).recognize(
        pdf, tmp_path / "pages", [1], None, 150, lambda *args: None,
    )
    assert [item["name"] for item in result["rows"]] == expected_names
    assert [item["source_row"] for item in result["rows"]] == [0, 1]
    assert all(item["provider_has_explicit_rows"] is True for item in result["rows"])
    assert all(item["ocr_metadata"]["vision_rows_are_candidate_observations"] is True
               for item in result["rows"])
    assert all(item["ocr_metadata"]["physical_row_refs"] == [] for item in result["rows"])
    assert all(not item.get("physical_row_refs") for item in result["rows"])


@pytest.mark.parametrize("change", ["row_strike", "cell_strike", "handwritten_quantity", "missing_quantity", "unreadable_code"])
def test_uncertain_readings_preserve_originals_without_replacement(pdf, change):
    data = page()
    item = data["tables"][0]["rows"][0]
    if change == "row_strike":
        item["struck_through"] = True
    elif change == "cell_strike":
        item["cells"][3]["struck_through"] = True
    elif change == "handwritten_quantity":
        item["cells"][3].update(handwritten=True, alternatives=["3"])
    elif change == "missing_quantity":
        item["cells"][3]["raw_text"] = ""
    else:
        data["tables"][0]["headers"][1] = "Артикул"
        item["cells"][1].update(raw_text="A?7", unreadable=True, alternatives=["A17", "A77"])
    result, _ = recognize(pdf, data)
    target = result.pages[0]
    assert len(target.rows) == 1
    assert target.rows[0].metadata["vision_raw_page"] == data
    assert target.rows[0].values["quantity"] == ("2" if change == "unreadable_code" else "")
    if change == "unreadable_code":
        assert target.rows[0].values["code"] == ""
    assert target.page_status["output_status"] == "REVIEW_REQUIRED"


def test_duplicate_header_mapping_and_spanning_cells_do_not_guess():
    data = page(tables=[table(["Количество", "Кол-во", "Наименование"], [
        row(["2", "3", "Насос"]),
    ])])
    item = project_page(VisionPage.model_validate(data))[0]
    assert item.values["quantity"] == "" and item.metadata["vision_ambiguous_fields"] == ["quantity"]
    data = page(tables=[table(["Наименование", "Количество"], [
        row([], cells=[cell(0, "Насос 2", column_span=2)]),
    ])])
    item = project_page(VisionPage.model_validate(data))[0]
    assert not any(item.values.values())
    assert item.metadata["vision_raw_row"]["cells"][0]["raw_text"] == "Насос 2"


def test_row_and_column_spans_are_validated_across_the_whole_table():
    raw = page(tables=[table(["A", "B", "\u041f\u0440\u0438\u043c\u0435\u0447\u0430\u043d\u0438\u0435"], [
        row([], cells=[
            cell(0, "merged", row_span=2, column_span=2),
            cell(2, "upper note"),
        ]),
        row([], cells=[cell(2, "lower note")]),
    ])])
    parsed = VisionPage.model_validate(raw)
    projected = project_page(parsed)
    assert len(projected) == 2
    assert projected[0].values["note"] == "upper note"
    assert projected[1].values["note"] == "lower note"
    assert parsed.tables[0].rows[0].cells[0].row_span == 2
    assert parsed.tables[0].rows[0].cells[0].column_span == 2

@pytest.mark.parametrize(("row_span", "column_span", "overlap_column"), [
    (2, 1, 0), (2, 2, 1),
])
def test_row_span_overlap_in_later_row_is_rejected(row_span, column_span, overlap_column):
    raw = page(tables=[table(["A", "B", "C"], [
        row([], cells=[cell(0, "merged", row_span=row_span, column_span=column_span)]),
        row([], cells=[cell(overlap_column, "contradiction")]),
    ])])
    with pytest.raises(ValidationError):
        VisionPage.model_validate(raw)


@pytest.mark.parametrize("rows,headers", [
    ([row([], cells=[cell(1, "outside", column_span=2)])], ["A", "B"]),
    ([row([], cells=[cell(0, "outside", row_span=2)])], ["A", "B"]),
])
def test_row_or_column_span_outside_table_is_rejected(rows, headers):
    with pytest.raises(ValidationError):
        VisionPage.model_validate(page(tables=[table(headers, rows)]))


def test_overlapping_row_span_discards_the_whole_page_without_partial_projection(pdf):
    raw = page(tables=[table(rows=[
        row(["1", "Do not leak", "шт", "2", ""]),
        row(["2", "Contradictory row", "шт", "3", ""]),
    ])])
    raw["tables"][0]["rows"][0]["cells"][1]["row_span"] = 2
    result, _ = recognize(pdf, raw)
    assert result.pages[0].rows == []
    assert result.pages[0].page_status["output_status"] == "NO_SPEC_OUTPUT"
    assert "schema_validation_failed" in result.pages[0].errors[0]


def test_raw_transcription_ocr_alternative_and_handwritten_replacement_keep_distinct_roles(pdf):
    raw = page()
    quantity = raw["tables"][0]["rows"][0]["cells"][3]
    quantity.update(
        raw_text="5", alternatives=["6"], struck_through=True,
        handwritten_replacements=[{
            "text": "8", "bbox_candidate": {"x": .7, "y": .2, "width": .1, "height": .1},
        }],
    )
    result, _ = recognize(pdf, raw)
    item = result.pages[0].rows[0]
    retained = item.metadata["vision_raw_page"]["tables"][0]["rows"][0]["cells"][3]
    assert retained["raw_text"] == "5"
    assert retained["alternatives"] == ["6"]
    assert retained["handwritten_replacements"] == [{
        "text": "8", "bbox_candidate": {"x": .7, "y": .2, "width": .1, "height": .1},
    }]
    assert item.values["quantity"] == ""


def test_localized_handwriting_does_not_clear_other_printed_cells():
    raw = page()
    raw["tables"][0]["rows"][0]["cells"][3]["handwritten_replacements"] = [
        {"text": "8", "bbox_candidate": None},
    ]
    item = project_page(VisionPage.model_validate(raw))[0]
    assert item.values["position"] == "1"
    assert item.values["name"] == DOCUMENT_TEXT
    assert item.values["unit"] == "шт"
    assert item.values["mass"] == "0,5"
    assert item.values["quantity"] == ""


def test_unlocalized_handwritten_replacement_conservatively_clears_row_values():
    raw = page()
    raw["tables"][0]["rows"][0]["handwritten_replacements"] = [
        {"text": "неразборчиво", "bbox_candidate": None},
    ]
    item = project_page(VisionPage.model_validate(raw))[0]
    assert not any(item.values.values())
    assert "vision_handwriting_unlocalized" in item.metadata["review_reasons"]


def test_page_scoped_handwriting_without_row_association_clears_page_projection():
    raw = page(tables=[table(rows=[
        row(["1", "Насос", "шт", "2", ""]),
        row(["2", "Клапан", "шт", "4", ""]),
    ])])
    raw["handwritten_replacements"] = [{"text": "замена вне строк", "bbox_candidate": None}]
    projected = project_page(VisionPage.model_validate(raw))
    assert len(projected) == 2
    assert all(not any(item.values.values()) for item in projected)
    assert projected[0].metadata["vision_raw_page"]["handwritten_replacements"] == raw["handwritten_replacements"]


@pytest.mark.parametrize("mutate", [
    lambda data: data.update(verified=True),
    lambda data: data.update(physical_row_refs=[{"row_index": 1}]),
    lambda data: data["tables"][0]["rows"][0].update(confidence=0.99),
    lambda data: data["tables"][0]["rows"][0].update(semantic_authoritative=True),
    lambda data: data["tables"][0]["rows"][0]["cells"][0].update(raw_text=123),
    lambda data: data["tables"][0]["rows"][0]["cells"][0].update(column_index=1000),
    lambda data: data["tables"][0]["rows"][0]["cells"][0].update(column_span=10**30),
    lambda data: data["tables"][0]["rows"][0]["cells"].append(cell(0, "duplicate")),
    lambda data: data["tables"][0]["rows"][0].update(continuation_of=0),
    lambda data: data["tables"][0]["rows"][0].update(bbox_candidate={"x": .9, "y": 0, "width": .2, "height": .2}),
    lambda data: data["tables"][0]["rows"][0]["cells"][0].pop("alternatives"),
])
def test_strict_schema_rejects_extra_safety_fields_wrong_types_and_invalid_addresses(mutate):
    raw = page()
    mutate(raw)
    transport = Transport(response(raw))
    with pytest.raises(QwenVisionError, match="schema_validation_failed"):
        QwenVisionClient(config(), transport=transport).extract(png(), 1)
    assert len(transport.calls) == 1


@pytest.mark.parametrize("raw,expected", [
    (response(content="not JSON " + DOCUMENT_TEXT), "invalid_json_response"),
    (response(content='{"page":1,"page":2}'), "invalid_json_response"),
    (response(content='{"page":NaN}'), "invalid_json_response"),
    (response(content=""), "empty_response"),
    (response(content="  "), "empty_response"),
    (response(finish_reason="length"), "incomplete_response"),
    (response(finish_reason="content_filter"), "incomplete_response"),
    (response(finish_reason=None), "incomplete_response"),
    (response(finish_reason=SECRET), "incomplete_response"),
    (response(page(2)), "page_mismatch"),
    (VisionHttpResponse(200, b"not JSON", {}), "invalid_json_response"),
    (VisionHttpResponse(200, b'{"choices":[]}', {}), "invalid_json_response"),
    (VisionHttpResponse(200, b'\xff', {}), "invalid_json_response"),
    (VisionHttpResponse(200, b'{"choices":[],"choices":[]}', {}), "invalid_json_response"),
])
def test_invalid_or_truncated_responses_are_not_repaired_or_retried(raw, expected):
    transport = Transport(raw)
    with pytest.raises(QwenVisionError) as error:
        QwenVisionClient(config(), transport=transport).extract(png(), 1)
    assert error.value.code == expected and len(transport.calls) == 1
    assert error.value.diagnostics["request_attempted"] is True


def test_length_response_records_only_safe_size_diagnostics_and_never_retries():
    content = json.dumps(page(), ensure_ascii=False) + " " + DOCUMENT_TEXT
    transport = Transport(response(content=content, finish_reason="length"))
    with pytest.raises(QwenVisionError) as error:
        QwenVisionClient(config(max_output_tokens=16384), transport=transport).extract(png(), 1)
    assert error.value.code == "incomplete_response"
    assert error.value.diagnostics["response_content_chars"] == len(content)
    assert error.value.diagnostics["response_content_bytes"] == len(content.encode("utf-8"))
    assert DOCUMENT_TEXT not in repr(error.value.diagnostics)
    assert len(transport.calls) == 1


@pytest.mark.parametrize("status", [301, 400, 401, 403, 429, 500, 503])
def test_http_errors_are_controlled_and_never_retry(status):
    transport = Transport(VisionHttpResponse(status, (SECRET + DOCUMENT_TEXT).encode(), {}))
    with pytest.raises(QwenVisionError) as error:
        QwenVisionClient(config(), transport=transport).extract(png(), 1)
    assert error.value.code == "http_error" and error.value.diagnostics["http_status"] == status
    assert len(transport.calls) == 1


@pytest.mark.parametrize("exc,expected", [(TimeoutError(SECRET), "timeout"), (OSError(DOCUMENT_TEXT), "transport_error")])
def test_timeout_and_network_errors_never_retry_or_expose_exception(exc, expected):
    transport = Transport(exc)
    with pytest.raises(QwenVisionError) as error:
        QwenVisionClient(config(), transport=transport).extract(png(), 1)
    assert error.value.code == expected and len(transport.calls) == 1
    trace = "".join(traceback.format_exception(error.value))
    assert SECRET not in trace and DOCUMENT_TEXT not in trace


@pytest.mark.parametrize("usage,expected", [
    (None, {}), ({}, {}), ({"prompt_tokens": 8}, {"prompt_tokens": 8}),
    ({"prompt_tokens": True, "completion_tokens": -1, "total_tokens": SECRET}, {}),
    ({"prompt_tokens": 8, "completion_tokens": 3, "total_tokens": 11,
      "prompt_tokens_details": {"cached_tokens": 2, "secret": SECRET},
      "completion_tokens_details": {"reasoning_tokens": 1}, "unknown": DOCUMENT_TEXT},
     {"prompt_tokens": 8, "completion_tokens": 3, "total_tokens": 11,
      "prompt_tokens_details": {"cached_tokens": 2}, "completion_tokens_details": {"reasoning_tokens": 1}}),
])
def test_usage_is_whitelisted_and_missing_counts_are_not_zero_filled(usage, expected):
    result = QwenVisionClient(config(), transport=Transport(response(usage=usage))).extract(png(), 1)
    assert result.diagnostics["usage"] == expected


def test_results_diagnostics_health_errors_and_logs_do_not_leak_payloads(pdf, caplog):
    image = png()
    bad = VisionHttpResponse(401, (SECRET + DOCUMENT_TEXT).encode(), {"x-request-id": SECRET})
    client = QwenVisionClient(config(), transport=Transport(bad))
    with pytest.raises(QwenVisionError) as error:
        client.extract(image, 1)
    result, _ = recognize(pdf)
    public = json.dumps({"error": str(error.value), "diagnostics": error.value.diagnostics,
                         "page_stats": result.pages[0].stats, "status": result.pages[0].page_status,
                         "stats": result.stats, "health": QwenVisionProvider(config()).health()}, ensure_ascii=False)
    for marker in (SECRET, DOCUMENT_TEXT, base64.b64encode(image).decode()):
        assert marker not in public and marker not in caplog.text
    assert SECRET not in repr(config())
    assert SECRET not in repr(bad) and DOCUMENT_TEXT not in repr(bad)


@pytest.mark.parametrize("limits,image,expected", [
    ({"max_image_bytes": 4}, png(), "image_size_limit"),
    ({"max_image_pixels": 10}, png(), "image_pixel_limit"),
    ({"max_request_bytes": 30}, png(), "request_size_limit"),
    ({}, b"invalid image", "invalid_image"),
])
def test_preflight_image_and_request_limits_make_no_request(limits, image, expected):
    transport = Transport()
    with pytest.raises(QwenVisionError) as error:
        QwenVisionClient(config(**limits), transport=transport).extract(image, 1)
    assert error.value.code == expected and not transport.calls
    assert error.value.diagnostics["request_attempted"] is False


def test_response_size_limit_even_with_custom_transport():
    transport = Transport(response())
    with pytest.raises(QwenVisionError, match="response_too_large"):
        QwenVisionClient(config(max_response_bytes=12), transport=transport).extract(png(), 1)


@pytest.mark.parametrize("limit", ["max_tables", "max_rows_per_page", "max_columns_per_table", "max_cells_per_page"])
def test_output_shape_is_bounded_before_schema_validation_and_projection(limit):
    raw = page(tables=[table(rows=[row(["1", "Насос", "шт", "2", ""]),
                                   row(["2", "Клапан", "шт", "3", ""])]), table()])
    transport = Transport(response(raw))
    with pytest.raises(QwenVisionError, match="output_structure_limit"):
        QwenVisionClient(config(max_output_tokens=16384, **{limit: 1}), transport=transport).extract(png(), 1)
    assert len(transport.calls) == 1


def test_refusal_is_rejected_even_if_finish_reason_is_stop():
    raw = json.loads(response().body)
    raw["choices"][0]["message"]["refusal"] = DOCUMENT_TEXT
    transport = Transport(VisionHttpResponse(200, json.dumps(raw).encode(), {}))
    with pytest.raises(QwenVisionError, match="model_refusal"):
        QwenVisionClient(config(), transport=transport).extract(png(), 1)


def test_request_budget_limits_calls_and_does_not_render_remaining_pages(pdf):
    class CountingPdf(PdfService):
        def __init__(self):
            self.pages = []
        def render_page(self, pdf_path, page_number, dpi=110):
            self.pages.append(page_number)
            return super().render_page(pdf_path, page_number, dpi)
    rendered = CountingPdf()
    transport = Transport(response())
    result = QwenVisionProvider(config(max_requests=1), transport=transport, pdf_service=rendered).recognize(pdf, [1, 2, 3])
    assert len(transport.calls) == 1 and rendered.pages == [1]
    assert result.stats["requests"] == 1
    assert all("request_budget_exhausted" in item.errors[0] for item in result.pages[1:])
    assert all(item.page_status["output_status"] == "NO_SPEC_OUTPUT" for item in result.pages[1:])


def test_error_pages_are_atomic_and_later_pages_can_succeed(pdf):
    transport = Transport(response(content="broken"), response(page(2)))
    result = QwenVisionProvider(config(), transport=transport).recognize(pdf, [1, 2])
    assert result.pages[0].rows == [] and result.pages[0].page_status["output_status"] == "NO_SPEC_OUTPUT"
    assert result.pages[1].rows and result.pages[1].page_status["output_status"] == "REVIEW_REQUIRED"


def test_cancel_before_request_between_pages_and_after_inflight_response(pdf):
    cancel = threading.Event()
    cancel.set()
    transport = Transport()
    provider = QwenVisionProvider(config(), transport=transport)
    with pytest.raises(OcrProviderError, match="cancelled"):
        provider.recognize(pdf, [1], cancel=cancel)
    assert not transport.calls
    cancel.clear()
    def progress(current, total, message):
        if current == 1:
            cancel.set()
    with pytest.raises(OcrProviderError, match="cancelled"):
        provider.recognize(pdf, [1, 2], cancel=cancel, progress=progress)
    assert len(transport.calls) == 1
    cancel.clear()
    def cancel_during_request():
        cancel.set()
        return response()
    transport = Transport(cancel_during_request)
    with pytest.raises(OcrProviderError, match="cancelled"):
        QwenVisionProvider(config(), transport=transport).recognize(pdf, [1, 2], cancel=cancel)
    assert len(transport.calls) == 1


def test_request_timeout_300_seconds_is_passed_to_http_transport(pdf):
    transport = Transport(response())
    provider = QwenVisionProvider(config(request_timeout_s=300), transport=transport, clock=lambda: 10.0)
    result = provider.recognize(pdf, [1])

    assert result.pages[0].stats["http_status"] == 200
    assert len(transport.calls) == 1
    assert transport.calls[0]["timeout_s"] == 300
    assert config().limits.total_timeout_s == 600


def test_elapsed_deadline_and_request_timeout_reject_late_results(pdf):
    now = [0.0]
    def late_response():
        now[0] = 7.0
        return response()
    transport = Transport(late_response)
    result = QwenVisionProvider(config(request_timeout_s=5, total_timeout_s=6), transport=transport,
                                clock=lambda: now[0]).recognize(pdf, [1, 2])
    assert len(transport.calls) == 1
    assert "timeout" in result.pages[0].errors[0]
    assert "total_timeout" in result.pages[1].errors[0]
    assert result.stats["elapsed_ms"] == 7000


def test_slow_pdf_rendering_exhausts_total_budget_without_request(pdf):
    now = [0.0]
    class SlowPdf(PdfService):
        def render_page(self, *args, **kwargs):
            now[0] = 5.0
            return png()
    transport = Transport()
    result = QwenVisionProvider(config(total_timeout_s=4), transport=transport, pdf_service=SlowPdf(),
                                clock=lambda: now[0]).recognize(pdf, [1])
    assert not transport.calls and "total_timeout" in result.pages[0].errors[0]


def test_large_pdf_page_is_rejected_before_pixmap_allocation(tmp_path):
    path = tmp_path / "large.pdf"
    with fitz.open() as document:
        document.new_page(width=5000, height=5000)
        document.save(path)
    class MustNotRender(PdfService):
        def render_page(self, *args, **kwargs):
            pytest.fail("oversize page must be rejected before render")
    transport = Transport()
    result = QwenVisionProvider(config(max_image_pixels=100), pdf_service=MustNotRender(), transport=transport).recognize(path, [1])
    assert "image_pixel_limit" in result.pages[0].errors[0] and not transport.calls


def test_crop_keeps_model_boxes_as_candidates_in_submitted_image_space(pdf):
    raw = page()
    raw["tables"][0]["rows"][0]["bbox_candidate"] = {"x": .1, "y": .2, "width": .5, "height": .1}
    transport = Transport(response(raw))
    crop = {"x": .25, "y": .25, "width": .5, "height": .5}
    result = QwenVisionProvider(config(dpi=72), transport=transport).recognize(pdf, [1], crop=crop)
    submitted = json.loads(transport.calls[0]["body"])["messages"][1]["content"][1]["image_url"]["url"]
    with Image.open(BytesIO(base64.b64decode(submitted.split(",", 1)[1]))) as opened:
        assert opened.size == (36, 36)
    item = result.pages[0].rows[0]
    assert item.bbox == {} and item.metadata["vision_image_crop"] == crop
    assert item.metadata["vision_raw_row"]["bbox_candidate"]["x"] == .1


@pytest.mark.parametrize("kwargs,code", [
    ({"pages": []}, "invalid_pages"), ({"pages": [True]}, "invalid_pages"),
    ({"pages": [0]}, "invalid_pages"), ({"pages": [1, 2], "max_pages": 1}, "page_limit"),
    ({"dpi": 9999}, "dpi_limit"), ({"crop": {"x": 1}}, "invalid_crop"),
    ({"crop": {"x": .9, "y": 0, "width": .5, "height": 1}}, "invalid_crop"),
    ({"mode": "yandex"}, "unsupported_mode"),
])
def test_invalid_options_do_not_call_api(pdf, kwargs, code):
    kwargs = dict(kwargs)
    limits = {"max_pages": kwargs.pop("max_pages")} if "max_pages" in kwargs else {}
    pages = kwargs.pop("pages", [1])
    transport = Transport()
    with pytest.raises(QwenVisionError, match=code):
        QwenVisionProvider(config(max_output_tokens=16384, **limits), transport=transport).recognize(pdf, pages, **kwargs)
    assert not transport.calls


@pytest.mark.parametrize("raw", [page(tables=[]), page(tables=[table(rows=[])]),
                                  page(tables=[table(rows=[row(["", "", "", "", ""])])])])
def test_no_valid_rows_is_never_successful(pdf, raw):
    result, _ = recognize(pdf, raw)
    assert all(item.metadata.get("vision_evidence_only") for item in result.pages[0].rows)
    assert result.pages[0].page_status["output_status"] == "NO_SPEC_OUTPUT"
    assert "empty_extraction" in result.pages[0].errors[0]


def test_header_only_and_unreadable_page_evidence_is_preserved_but_not_an_output_row(pdf):
    raw = page(tables=[table(headers=["Нечитаемый столбец"], rows=[])])
    raw["unreadable_fragments"] = ["неразборчиво: A?7"]
    result, _ = recognize(pdf, raw)
    target = result.pages[0]
    assert target.rows[0].metadata["vision_evidence_only"] is True
    assert target.rows[0].metadata["vision_raw_page"] == raw
    assert target.page_status["output_status"] == "NO_SPEC_OUTPUT"
    assert not any(target.rows[0].values.values())


def test_evidence_only_record_is_skip_and_never_an_exported_position(pdf, tmp_path):
    raw = page(tables=[table(headers=["Нечитаемый столбец"], rows=[])])
    provider = QwenVisionProvider(config(), transport=Transport(response(raw)))
    result = RecognitionService(PdfService(), ocr=provider).recognize(
        pdf, tmp_path / "pages", [1], None, 150, lambda *args: None,
    )
    assert len(result["rows"]) == 1
    assert result["rows"][0]["ocr_metadata"]["vision_evidence_only"] is True
    assert result["rows"][0]["row_type"] == "skip"
    assert result["page_statuses"]["1"]["output_status"] == "NO_SPEC_OUTPUT"

    output = tmp_path / "evidence-only.xlsx"
    ExcelExportService().export(
        result["rows"], ["name", "quantity"], output,
        page_statuses=result["page_statuses"], enforce_safety=False,
    )
    workbook = load_workbook(output, read_only=True)
    assert workbook["Спецификация"].max_row == 1
    workbook.close()


def test_never_usable_no_confidence_no_physical_refs_and_existing_export_fails_closed(pdf, tmp_path):
    raw = page()
    raw["tables"][0]["rows"][0]["bbox_candidate"] = {"x": .1, "y": .1, "width": .8, "height": .2}
    provider = QwenVisionProvider(config(max_output_tokens=16384), transport=Transport(response(raw)))
    # Exercise the current consumer through injection, without registering a mode.
    result = RecognitionService(PdfService(), ocr=provider).recognize(
        pdf, tmp_path / "pages", [1], None, 150, lambda *args: None,
    )
    assert result["page_statuses"]["1"]["output_status"] == "REVIEW_REQUIRED"
    assert result["page_statuses"]["1"]["blockers"]
    item = result["rows"][0]
    assert not item.get("physical_row_refs") and not item.get("confidences")
    assert not item.get("semantic_authoritative")
    with pytest.raises(ValueError):
        ExcelExportService().export(result["rows"], ["name", "quantity"], tmp_path / "blocked.xlsx",
                                    page_statuses=result["page_statuses"])
    with pytest.raises(ValueError, match="Физическая строка"):
        HumanReviewService().create_decision(result, document_fingerprint="synthetic", page=1,
                                            physical_refs=[], decision=CONFIRM_FIELD_VALUE_DECISION,
                                            field="quantity", confirmed_value="2")


def test_usage_aggregation_reports_partial_counts_explicitly(pdf):
    transport = Transport(response(page(1), usage={"prompt_tokens": 5}), response(page(2)))
    result = QwenVisionProvider(config(), transport=transport).recognize(pdf, [1, 2])
    assert result.stats["requests"] == 2 and result.stats["usage_complete_pages"] == 0
    assert result.stats["reported_usage"] == {"prompt_tokens": 5}


def test_unreadable_pdf_error_never_exposes_source_path_or_document_bytes(tmp_path):
    path = tmp_path / f"{SECRET}.pdf"
    path.write_text(DOCUMENT_TEXT)
    transport = Transport()
    result = QwenVisionProvider(config(), transport=transport).recognize(path, [1])
    assert not transport.calls and result.pages[0].rows == []
    safe = json.dumps(result.as_dict(), ensure_ascii=False)
    assert SECRET not in safe and DOCUMENT_TEXT not in safe
    assert "page_extraction_failed" in result.pages[0].errors[0]


def test_call_diagnostics_and_budget_are_local_to_each_recognition(pdf):
    transport = Transport(response(usage={"prompt_tokens": 5}), response(usage={"completion_tokens": 2}))
    provider = QwenVisionProvider(config(max_requests=1), transport=transport)
    first = provider.recognize(pdf, [1])
    second = provider.recognize(pdf, [1])
    assert first.stats["reported_usage"] == {"prompt_tokens": 5}
    assert second.stats["reported_usage"] == {"completion_tokens": 2}
    assert first.stats["requests"] == second.stats["requests"] == 1


def test_default_https_transport_reads_with_cap_and_never_follows_redirects(monkeypatch):
    closed = []
    class Sock:
        def settimeout(self, value):
            assert value > 0
    class Reply:
        status = 200
        def getheader(self, key, default):
            return REQUEST_ID
        def read1(self, size):
            return b"x" * size
    class Connection:
        def __init__(self, *args, **kwargs):
            self.sock = Sock()
        def set_debuglevel(self, value):
            assert value == 0
        def connect(self):
            pass
        def request(self, *args, **kwargs):
            pass
        def getresponse(self):
            return Reply()
        def close(self):
            closed.append(True)
    monkeypatch.setattr("http.client.HTTPSConnection", Connection)
    transport = _HttpsTransport()
    with pytest.raises(QwenVisionError, match="response_too_large"):
        transport.post("https://ai.test/v1/chat/completions", body=b"test", headers={}, timeout_s=5,
                       max_response_bytes=10, cancel=None)
    assert closed == [True]
    Reply.status = 302
    assert transport.post("https://ai.test/v1/chat/completions", body=b"test", headers={}, timeout_s=5,
                          max_response_bytes=10, cancel=None).status == 302
    assert len(closed) == 2
