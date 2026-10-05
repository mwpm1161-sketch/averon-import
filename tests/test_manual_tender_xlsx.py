from __future__ import annotations

import hashlib
import json
import math
import threading
import asyncio
import os
import re
import subprocess
import shutil
import time
import uuid
import zipfile
from copy import copy, deepcopy
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Side

from averon_import.services.manual_tenders import TenderActivityRegistry, TenderTemplateService, TenderWorkbookParser, TenderWorkspaceRepository
from averon_import.services.manual_tenders.sourcing import TenderSourcingRowAdapter, TenderSourcingRunStore, canonical_tender_projection
from averon_import.services.jobs import JobService
from averon_import.services.manual_tenders.parser import (
    MAX_ACTUAL_ITEMS,
    MAX_HEADER_LENGTH,
    MAX_UNCOMPRESSED_BYTES,
    MAX_UPLOAD_BYTES,
    TenderParseError,
    parse_unit_basis,
    preflight_tender_xlsx,
)
from averon_import.services.manual_tenders.repository import (
    ABSOLUTE_TTL_SECONDS,
    IDLE_TTL_SECONDS,
    MAX_TENDER_RUN_BYTES,
    MAX_WORKSPACES_GLOBAL,
    MAX_WORKSPACES_PER_USER,
    PREVIEW_TTL_SECONDS,
    TenderWorkspaceError,
)
from averon_import.services.manual_tenders.template import PREPARED_ROWS, TEMPLATE_HEADERS, TEMPLATE_SCHEMA_NAME, TEMPLATE_SCHEMA_VERSION, TEMPLATE_SHEET, TEMPLATE_TABLE


def _official(path: Path, *, include_required=True) -> bytes:
    content = TenderTemplateService().bytes()
    path.write_bytes(content)
    if include_required:
        workbook = load_workbook(path)
        sheet = workbook[TEMPLATE_SHEET]
        sheet["A2"] = "R-001"
        sheet["B2"] = "Синтетический ресурс"
        sheet["C2"] = "шт"
        sheet["D2"] = 2
        sheet["E2"] = "SKU-1"
        sheet["F2"] = "Тестовый производитель"
        sheet["G2"] = "Модель X"
        workbook.save(path)
    return path.read_bytes()


def _fallback(path: Path, *, count=370) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ресурсная ведомость1 - СУВР(О)_"
    for coordinate, value in {
        "A12": "№ смет/Код ресурса", "C12": "Наименование", "D12": "Ед. изм.", "E12": "Кол-во",
        "AI12": "Цена за единицу", "AJ12": "Общая стоимость", "CF12": "Итого",
    }.items():
        sheet[coordinate] = value
    sheet["F12"] = "Служебное скрытое поле F"
    sheet["G12"] = "Служебное скрытое поле G"
    sheet.column_dimensions["F"].hidden = True
    sheet.column_dimensions["G"].hidden = True
    sheet.column_dimensions["AI"].width = 15
    sheet["AI12"].number_format = "0.00"
    sheet["AJ12"].number_format = "0.00"
    sheet["B5"] = "Сводная часть"
    sheet.merge_cells("B5:C5")
    sheet.print_area = "A1:E391"
    row_number = 13
    sheet.cell(row_number, 3, "Синтетический раздел А")
    row_number += 1
    scaled_units = ["10 шт", "100 шт", "1000 шт", "10 м", "1000 м"] * 3 + ["т", "10 шт"]
    item_index = 0
    for _ in range(count // 2):
        sheet.cell(row_number, 1, "REP-1")
        sheet.cell(row_number, 3, f"Синтетическая позиция {item_index + 1}")
        sheet.cell(row_number, 4, scaled_units[item_index] if item_index < len(scaled_units) else "шт")
        sheet.cell(row_number, 5, 1 + (item_index % 3))
        item_index += 1
        row_number += 1
    sheet.cell(row_number, 3, "Синтетический раздел Б")
    row_number += 1
    while item_index < count:
        sheet.cell(row_number, 1, "REP-1")
        sheet.cell(row_number, 3, f"Синтетическая позиция {item_index + 1}")
        sheet.cell(row_number, 4, scaled_units[item_index] if item_index < len(scaled_units) else "кг")
        sheet.cell(row_number, 5, 1 + (item_index % 3))
        item_index += 1
        row_number += 1
    for label in ("ИТОГО по разделу", "Всего по ведомости"):
        sheet.cell(row_number, 3, label)
        sheet.cell(row_number, 5, 740)
        row_number += 1
    for index in range(1, 7):
        sheet.cell(row_number + index, 1, f"FORMAL-CF-{index}")
    workbook.save(path)
    return path.read_bytes()


def test_official_template_contract_and_parser_fast_path(tmp_path):
    path = tmp_path / "official.xlsx"
    _official(path)
    workbook = load_workbook(path, data_only=False, keep_links=False)
    assert workbook.sheetnames == [TEMPLATE_SHEET, "Инструкция"]
    sheet = workbook[TEMPLATE_SHEET]
    assert tuple(sheet.cell(1, col).value for col in range(1, 8)) == TEMPLATE_HEADERS
    assert list(sheet.tables) == [TEMPLATE_TABLE]
    assert sheet.tables[TEMPLATE_TABLE].ref == f"A1:G{PREPARED_ROWS + 1}"
    assert workbook.defined_names[TEMPLATE_SCHEMA_NAME].attr_text == TEMPLATE_SCHEMA_VERSION
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].fill.fgColor.rgb.endswith("E9ECEF")
    assert sheet.column_dimensions["B"].width >= 40
    assert not any(cell.data_type == "f" for ws in workbook.worksheets for row in ws.iter_rows() for cell in row)
    instructions = [str(workbook["Инструкция"].cell(row, 1).value) for row in range(1, workbook["Инструкция"].max_row + 1)]
    assert any("Одна строка таблицы" in line for line in instructions)
    assert any("не объединяйте ячейки" in line for line in instructions)
    assert any("будущем добавит" in line for line in instructions)
    workbook.close()

    parsed = TenderWorkbookParser().parse(path, tender_id="a" * 32)
    assert parsed["official_template"] is True
    assert parsed["mapping_required"] is False
    assert parsed["item_count"] == 1
    row = parsed["rows"][0]
    assert row["resource_code"] == "R-001"
    assert row["article"] == "SKU-1"
    assert row["resource_code"] != row["article"]
    assert parsed["logical_right_column"] == "G"
    assert [item["column"] for item in parsed["future_output_columns"]] == ["H", "I"]


def test_official_template_blank_prepared_rows_are_ignored(tmp_path):
    path = tmp_path / "blank.xlsx"
    _official(path, include_required=False)
    parsed = TenderWorkbookParser().parse(path)
    assert parsed["official_template"] is True
    assert parsed["mapping_required"] is False
    assert parsed["item_count"] == 0
    assert parsed["counts"]["ignored"] == 0


def test_official_template_rejects_changed_headers_and_duplicate_table(tmp_path):
    path = tmp_path / "bad-template.xlsx"
    _official(path, include_required=False)
    workbook = load_workbook(path)
    workbook[TEMPLATE_SHEET]["B1"] = "Имя"
    workbook.save(path)
    with pytest.raises(TenderParseError, match="Заголовки официальной таблицы"):
        TenderWorkbookParser().parse(path)

    stripped_path = tmp_path / "stripped-template.xlsx"
    _official(stripped_path, include_required=False)
    stripped = load_workbook(stripped_path)
    for column in range(1, len(TEMPLATE_HEADERS) + 1):
        stripped[TEMPLATE_SHEET].cell(1, column).value = None
    stripped[TEMPLATE_SHEET].tables.clear()
    stripped.save(stripped_path)
    with pytest.raises(TenderParseError, match="отсутствует корректная таблица AveronTenderInput"):
        TenderWorkbookParser().parse(stripped_path)


def test_official_template_accepts_resized_371_row_table_and_requires_schema_marker(tmp_path):
    path = tmp_path / "resized-official.xlsx"
    _official(path, include_required=False)
    workbook = load_workbook(path)
    sheet = workbook[TEMPLATE_SHEET]
    sheet.tables[TEMPLATE_TABLE].ref = "A1:G371"
    sheet["A371"] = "R-371"
    sheet["B371"] = "Синтетическая позиция 371"
    sheet["C371"] = "шт"
    sheet["D371"] = 1
    workbook.save(path)
    workbook.close()
    parsed = TenderWorkbookParser().parse(path)
    assert parsed["table_ref"] == "A1:G371"
    assert parsed["item_count"] == 1 and parsed["invalid_count"] == 0
    assert parsed["rows"][0]["excel_row"] == 371

    missing_marker = tmp_path / "missing-schema.xlsx"
    _official(missing_marker, include_required=False)
    workbook = load_workbook(missing_marker)
    del workbook.defined_names[TEMPLATE_SCHEMA_NAME]
    workbook.save(missing_marker)
    workbook.close()
    with pytest.raises(TenderParseError, match="отсутствует версия схемы"):
        TenderWorkbookParser().parse(missing_marker)

    wrong_marker = tmp_path / "wrong-schema.xlsx"
    _official(wrong_marker, include_required=False)
    workbook = load_workbook(wrong_marker)
    workbook.defined_names[TEMPLATE_SCHEMA_NAME].attr_text = "2"
    workbook.save(wrong_marker)
    workbook.close()
    with pytest.raises(TenderParseError, match="Версия схемы официального шаблона"):
        TenderWorkbookParser().parse(wrong_marker)


def test_manifest_format_fingerprint_is_stable_and_tracks_style_properties(tmp_path):
    original = tmp_path / "format-base.xlsx"
    _official(original)
    parser = TenderWorkbookParser()

    def fingerprint(path):
        parsed = parser.parse(path)
        cell = next(item for item in parsed["manifest"]["meaningful_cells"] if item["cell"] == "B2")
        return cell["style_fingerprint"]

    baseline = fingerprint(original)
    assert fingerprint(original) == baseline
    mutations = ("number_format", "font", "fill", "border", "alignment", "protection")
    for property_name in mutations:
        path = tmp_path / f"format-{property_name}.xlsx"
        path.write_bytes(original.read_bytes())
        workbook = load_workbook(path)
        cell = workbook[TEMPLATE_SHEET]["B2"]
        if property_name == "number_format":
            cell.number_format = "0.000"
        elif property_name == "font":
            style = copy(cell.font)
            style.bold = True
            cell.font = style
        elif property_name == "fill":
            cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        elif property_name == "border":
            style = copy(cell.border)
            style.left = Side(style="thin", color="FF000000")
            cell.border = style
        elif property_name == "alignment":
            style = copy(cell.alignment)
            style.wrap_text = True
            cell.alignment = style
        else:
            style = copy(cell.protection)
            style.hidden = True
            cell.protection = style
        workbook.save(path)
        workbook.close()
        assert fingerprint(path) != baseline, property_name


def test_invalid_item_candidates_are_exposed_safely_and_block_confirmation(tmp_path):
    source = tmp_path / "incomplete.xlsx"
    _official(source, include_required=False)
    workbook = load_workbook(source)
    sheet = workbook[TEMPLATE_SHEET]
    sheet["A2"] = "R-2"
    sheet["B2"] = "Синтетическая позиция"
    sheet["C2"] = "шт"
    sheet["D2"] = None
    workbook.save(source)
    workbook.close()
    payload = source.read_bytes()
    parser = TenderWorkbookParser()
    analysis = parser.parse(source)
    assert analysis["counts"]["invalid"] == analysis["invalid_count"] == 1
    assert analysis["item_count"] == 0
    assert analysis["invalid_examples"] == [{"excel_row": 2, "reasons": ["missing_quantity"]}]
    assert "Синтетическая позиция" not in json.dumps(analysis["invalid_examples"], ensure_ascii=False)

    repository = TenderWorkspaceRepository(tmp_path / "repo-data")
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner", source.name, len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    repository.update_preview(preview["preview_id"], "owner", status="ready", analysis=analysis, parser_version=analysis["parser_version"], mapping=analysis["mapping"])
    with pytest.raises(TenderWorkspaceError) as rejected:
        repository.confirm(preview["preview_id"], "owner")
    assert rejected.value.status_code == 409
    assert rejected.value.code == "TENDER_INVALID_ROWS"


def test_369_valid_items_plus_one_invalid_candidate_cannot_be_confirmed(tmp_path):
    source = tmp_path / "large-incomplete.xlsx"
    _official(source, include_required=False)
    workbook = load_workbook(source)
    sheet = workbook[TEMPLATE_SHEET]
    for row in range(2, 371):
        sheet.cell(row, 1, f"R-{row}")
        sheet.cell(row, 2, f"Синтетическая позиция {row}")
        sheet.cell(row, 3, "шт")
        sheet.cell(row, 4, 1)
    sheet["A371"] = "R-371"
    sheet["B371"] = "Синтетическая позиция 371"
    sheet["C371"] = "шт"
    sheet["D371"] = None
    workbook.save(source)
    workbook.close()
    payload = source.read_bytes()
    analysis = TenderWorkbookParser().parse(source)
    assert analysis["item_count"] == 369 and analysis["invalid_count"] == 1
    repository = TenderWorkspaceRepository(tmp_path / "large-repo-data")
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner-large", source.name, len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    repository.update_preview(preview["preview_id"], "owner-large", status="ready", analysis=analysis, parser_version=analysis["parser_version"], mapping=analysis["mapping"])
    with pytest.raises(TenderWorkspaceError) as rejected:
        repository.confirm(preview["preview_id"], "owner-large")
    assert rejected.value.code == "TENDER_INVALID_ROWS"


def test_fallback_name_only_may_be_section_but_item_signals_without_quantity_are_invalid(tmp_path):
    path = tmp_path / "row-classes.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Код ресурса", "Наименование", "Ед. изм.", "Количество", "Артикул"])
    sheet.append([None, "Синтетический раздел", None, None, None])
    sheet.append(["R-1", "Синтетическая позиция", "шт", None, None])
    sheet.append([None, "Ещё одна позиция", None, "не число", None])
    workbook.save(path)
    parsed = TenderWorkbookParser().parse(path)
    assert [row["row_type"] for row in parsed["rows"]] == ["section", "invalid", "invalid"]
    assert parsed["invalid_count"] == 2
    assert parsed["invalid_examples"] == [
        {"excel_row": 3, "reasons": ["missing_quantity"]},
        {"excel_row": 4, "reasons": ["invalid_quantity"]},
    ]


def test_fallback_ignores_column_scaffold_and_numbered_note_and_recognizes_merged_sections(tmp_path):
    path = tmp_path / "structural-row-classes.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Код ресурса"
    sheet["C1"] = "Наименование"
    sheet["D1"] = "Ед. изм."
    sheet["E1"] = "Кол-во"
    # A/C/D/E repeat their physical column indexes, a source-layout scaffold.
    sheet["A2"] = 1
    sheet["C2"] = 3
    sheet["D2"] = 4
    sheet["E2"] = "5"
    # The section label is anchored outside the mapped name column.
    sheet.merge_cells("A3:E3")
    sheet["A3"] = "Синтетический раздел"
    # A resource code plus item name and unit without quantity remains invalid.
    sheet["A4"] = "REF-4"
    sheet["C4"] = "Синтетическая позиция без количества"
    sheet["D4"] = "шт"
    # A numbered note in the resource/reference column is not an item.
    sheet["A5"] = "1. Синтетическое дополнительное примечание"
    # Unknown unit vocabulary alone does not invalidate a complete item.
    sheet["A6"] = "REF-6"
    sheet["C6"] = "Синтетическая позиция с неизвестной единицей"
    sheet["D6"] = "единица-неизвестна"
    sheet["E6"] = 2
    # Short numbered or ordinary resource codes alone remain incomplete rows.
    sheet["A7"] = "1. X"
    sheet["A8"] = "REF-8"
    workbook.save(path)
    workbook.close()

    parsed = TenderWorkbookParser().parse(path)
    rows_by_excel_row = {row["excel_row"]: row for row in parsed["rows"]}
    assert rows_by_excel_row[2]["row_type"] == "ignored"
    assert rows_by_excel_row[3]["row_type"] == "section"
    assert rows_by_excel_row[4]["row_type"] == "invalid"
    assert rows_by_excel_row[5]["row_type"] == "ignored"
    assert rows_by_excel_row[6]["row_type"] == "item"
    assert rows_by_excel_row[7]["row_type"] == "invalid"
    assert rows_by_excel_row[8]["row_type"] == "invalid"
    assert rows_by_excel_row[6]["unit_basis"]["trusted"] is False
    assert rows_by_excel_row[6]["warnings"]
    assert parsed["counts"] == {"item": 1, "section": 1, "total": 0, "ignored": 2, "invalid": 3}


def test_parser_deletes_normalized_temporary_and_preserves_uploaded_source(tmp_path, monkeypatch):
    import shutil
    import averon_import.services.manual_tenders.parser as parser_module
    path = tmp_path / "source.xlsx"
    _official(path)
    original = path.read_bytes()
    normalized = tmp_path / "normalized-temporary.xlsx"
    def normalize(source):
        shutil.copyfile(source, normalized)
        return normalized
    monkeypatch.setattr(parser_module, "shared_preflight_xlsx", normalize)
    assert TenderWorkbookParser().parse(path)["item_count"] == 1
    assert path.read_bytes() == original
    assert not normalized.exists()


def test_fallback_header_12_a_c_d_e_boundary_hidden_columns_and_row_identity(tmp_path):
    path = tmp_path / "synthetic-legacy.xlsx"
    _fallback(path)
    parsed = TenderWorkbookParser().parse(path, tender_id="tender-id")
    assert parsed["sheet_name"] == "Ресурсная ведомость1 - СУВР(О)_"
    assert parsed["header_row"] == 12
    assert parsed["mapping"] == {"resource_code": 1, "name": 3, "unit": 4, "quantity": 5, "article": None, "manufacturer": None, "model": None}
    assert parsed["counts"]["item"] == 370
    assert parsed["counts"]["section"] == 2
    assert parsed["counts"]["total"] == 2
    assert parsed["logical_right_column"] == "E"
    assert parsed["future_output_columns"][0]["column"] == "F"
    assert parsed["future_output_columns"][0]["hidden"] is True
    assert parsed["future_output_columns"][0]["value_count"] == 1
    assert parsed["future_output_columns"][1]["hidden"] is True
    assert parsed["future_output_columns"][1]["value_count"] == 1
    assert "DISTANT_PRICE_COLUMNS" in {warning["code"] for warning in parsed["warnings"]}
    items = [row for row in parsed["rows"] if row["row_type"] == "item"]
    assert len({row["source_row_id"] for row in items}) == 370
    assert all(row["resource_code"] == "REP-1" and row["article"] == "" for row in items)
    assert len([row for row in items if row["unit_basis"]["trusted"] and row["raw_unit"] not in {"шт", "кг"}]) == 17
    manifest = parsed["manifest"]
    assert manifest["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert manifest["sheets"][0]["print_area"]
    assert manifest["sheets"][0]["merged_ranges"]
    assert manifest["sheets"][0]["hidden_columns"] == ["F", "G"]


def test_fallback_skips_long_data_rows_as_header_candidates_without_truncating_source(tmp_path):
    path = tmp_path / "long-preamble-and-data.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ресурсная ведомость"
    for coordinate, value in {
        "A12": "№ смет/Код ресурса", "C12": "Наименование", "D12": "Ед. изм.", "E12": "Кол-во",
    }.items():
        sheet[coordinate] = value
    long_449 = "Д" * 449
    long_393 = "Е" * 393
    for row, code, name in ((21, "R-21", long_449), (24, "R-24", long_393)):
        sheet.cell(row, 1, code)
        sheet.cell(row, 3, name)
        sheet.cell(row, 4, "шт")
        sheet.cell(row, 5, 1)
    workbook.save(path)
    workbook.close()

    parsed = TenderWorkbookParser().parse(path)
    assert parsed["header_row"] == 12
    assert parsed["mapping"] == {"resource_code": 1, "name": 3, "unit": 4, "quantity": 5, "article": None, "manufacturer": None, "model": None}
    assert parsed["item_count"] == 2 and parsed["invalid_count"] == 0
    parsed_names = {row["excel_row"]: row["name"] for row in parsed["rows"] if row["row_type"] == "item"}
    assert parsed_names == {21: long_449, 24: long_393}
    manifest_values = {cell["cell"]: cell["value"] for cell in parsed["manifest"]["meaningful_cells"]}
    assert manifest_values["C21"] == long_449 and manifest_values["C24"] == long_393


def test_overlong_header_candidate_and_explicitly_selected_header_are_rejected(tmp_path):
    path = tmp_path / "overlong-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Long header"
    sheet.cell(12, 1, "Наименование" + "X" * (MAX_HEADER_LENGTH + 1))
    sheet.cell(12, 2, "Количество")
    workbook.save(path)
    workbook.close()

    with pytest.raises(TenderParseError, match="Не удалось определить лист и заголовки"):
        TenderWorkbookParser().parse(path)
    with pytest.raises(TenderParseError, match="Заголовок XLSX превышает допустимую длину"):
        TenderWorkbookParser().parse(path, selected_sheet="Long header", selected_header_row=12)


def test_generic_code_header_is_not_guessed_as_article_and_mapping_can_be_explicit(tmp_path):
    path = tmp_path / "ambiguous-code.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Код", "Наименование", "Количество"])
    sheet.append(["001", "Синтетическая позиция", 1])
    workbook.save(path)
    parser = TenderWorkbookParser()
    preview = parser.parse(path)
    assert preview["mapping"]["article"] is None
    assert preview["mapping"]["resource_code"] is None
    assert preview["mapping_required"] is True
    assert preview["rows"][0]["article"] == ""
    mapped = parser.parse(path, mapping_override={"resource_code": 1, "name": 2, "quantity": 3}, selected_sheet=sheet.title, selected_header_row=1)
    assert mapped["rows"][0]["resource_code"] == "001"
    assert mapped["rows"][0]["article"] == ""
    with pytest.raises(TenderParseError, match="Код ресурса нельзя"):
        parser.parse(path, mapping_override={"resource_code": 1, "name": 2, "quantity": 3, "article": 1}, selected_sheet=sheet.title, selected_header_row=1)
    with pytest.raises(TenderParseError, match="явно указывает на артикул"):
        parser.parse(path, mapping_override={"name": 2, "quantity": 3, "article": 1}, selected_sheet=sheet.title, selected_header_row=1)

    explicit_article = tmp_path / "explicit-article.xlsx"
    explicit = Workbook()
    explicit_sheet = explicit.active
    explicit_sheet.append(["Артикул производителя", "Наименование", "Количество"])
    explicit_sheet.append(["SKU-001", "Синтетическая позиция", 1])
    explicit.save(explicit_article)
    explicit_result = parser.parse(explicit_article)
    assert explicit_result["rows"][0]["article"] == "SKU-001"


@pytest.mark.parametrize("value", [True, datetime(2026, 1, 1), float("nan"), float("inf"), 0, -1, "1 шт", "1,2.3"])
def test_ambiguous_quantity_values_are_not_trusted(tmp_path, value):
    path = tmp_path / "quantity.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Наименование", "Количество"])
    sheet.append(["Синтетическая позиция", value])
    workbook.save(path)
    parsed = TenderWorkbookParser().parse(path)
    assert parsed["rows"][0]["quantity_trusted"] is False


def test_formula_quantity_never_becomes_authoritative(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Наименование", "Количество"])
    sheet.append(["Синтетическая позиция", "=1+2"])
    workbook.save(path)
    parsed = TenderWorkbookParser().parse(path)
    assert parsed["rows"][0]["quantity_raw"] == "=1+2"
    assert parsed["rows"][0]["quantity_trusted"] is False


@pytest.mark.parametrize(("unit", "base", "scale", "dimension"), [
    ("шт", "шт", "1", "count"), ("10 шт", "шт", "10", "count"),
    ("100 шт", "шт", "100", "count"), ("1000 шт", "шт", "1000", "count"),
    ("м", "м", "1", "length"), ("10 м", "м", "10", "length"),
    ("1000 м", "м", "1000", "length"), ("кг", "кг", "1", "mass"),
    ("т", "кг", "1000", "mass"), ("м2", "м2", "1", "area"),
    ("м3", "м3", "1", "volume"), ("л", "л", "1", "volume"),
    ("компл", "компл", "1", "set"), ("уп", "уп", "1", "package"),
])
def test_tender_unit_basis_supported_units(unit, base, scale, dimension):
    result = parse_unit_basis(unit)
    assert result["trusted"] is True
    assert (result["base_unit"], result["scale"], result["dimension"]) == (base, scale, dimension)


def test_unknown_and_unapproved_pack_scales_remain_unknown():
    assert parse_unit_basis("коробка 12 шт")["trusted"] is False
    assert parse_unit_basis("12 уп")["trusted"] is False
    assert parse_unit_basis("10 кг")["trusted"] is False
    assert parse_unit_basis("10 компл")["trusted"] is False


def test_xlsx_preflight_rejects_non_xlsx_and_unsafe_archive_members(tmp_path):
    not_xlsx = tmp_path / "book.xlsm"
    not_xlsx.write_bytes(b"fake")
    with pytest.raises(TenderParseError, match=r"\.xlsx"):
        TenderWorkbookParser().parse(not_xlsx)
    malformed = tmp_path / "malformed.xlsx"
    malformed.write_bytes(b"not a zip archive")
    with pytest.raises(TenderParseError, match="корректной книгой XLSX"):
        preflight_tender_xlsx(malformed)

    malicious = tmp_path / "traversal.xlsx"
    with zipfile.ZipFile(malicious, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("../escape", "x")
    with pytest.raises(TenderParseError, match="небезопасные"):
        preflight_tender_xlsx(malicious)


def test_xlsx_preflight_rejects_case_collisions_external_relationships_and_entities(tmp_path):
    collision = tmp_path / "collision.xlsx"
    with zipfile.ZipFile(collision, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("XL/WORKBOOK.XML", "<workbook/>")
    with pytest.raises(TenderParseError, match="без учёта регистра"):
        preflight_tender_xlsx(collision)

    external = tmp_path / "external.xlsx"
    with zipfile.ZipFile(external, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Target="https://example.invalid" TargetMode="External"/></Relationships>')
    with pytest.raises(TenderParseError, match="Внешние связи"):
        preflight_tender_xlsx(external)

    entities = tmp_path / "entities.xlsx"
    with zipfile.ZipFile(entities, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", '<!DOCTYPE workbook [<!ENTITY x "unsafe">]><workbook/>')
    with pytest.raises(TenderParseError, match="DTD/ENTITY"):
        preflight_tender_xlsx(entities)


def test_xlsx_preflight_rejects_vba_and_compressed_size_bounds(tmp_path):
    macro = tmp_path / "macro.xlsx"
    with zipfile.ZipFile(macro, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/vbaProject.bin", b"not a macro")
    with pytest.raises(TenderParseError, match="макросами"):
        preflight_tender_xlsx(macro)


def test_xlsx_preflight_rejects_duplicate_members_member_count_uncompressed_size_and_ratio(tmp_path, monkeypatch):
    import averon_import.services.manual_tenders.parser as parser_module
    duplicate = tmp_path / "duplicate.xlsx"
    with zipfile.ZipFile(duplicate, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(TenderParseError, match="повторные компоненты"):
        preflight_tender_xlsx(duplicate)

    too_many = tmp_path / "too-many.xlsx"
    with zipfile.ZipFile(too_many, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/third.xml", "x")
    monkeypatch.setattr(parser_module, "MAX_ZIP_MEMBERS", 2)
    with pytest.raises(TenderParseError, match="число компонентов"):
        preflight_tender_xlsx(too_many)
    monkeypatch.setattr(parser_module, "MAX_ZIP_MEMBERS", 1000)

    oversized = tmp_path / "oversized.xlsx"
    with zipfile.ZipFile(oversized, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/data.bin", b"x" * 80)
    monkeypatch.setattr(parser_module, "MAX_UNCOMPRESSED_BYTES", 40)
    with pytest.raises(TenderParseError, match="Распакованный размер"):
        preflight_tender_xlsx(oversized)
    monkeypatch.setattr(parser_module, "MAX_UNCOMPRESSED_BYTES", 20 * 1024 * 1024)

    high_ratio = tmp_path / "high-ratio.xlsx"
    with zipfile.ZipFile(high_ratio, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/large.xml", "A" * 2048)
    monkeypatch.setattr(parser_module, "MAX_COMPRESSION_RATIO", 1)
    with pytest.raises(TenderParseError, match="Коэффициент сжатия"):
        preflight_tender_xlsx(high_ratio)


def test_fallback_ambiguous_columns_require_manual_mapping_and_item_limit_is_enforced(tmp_path):
    ambiguous = tmp_path / "duplicate-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Наименование", "Наименование", "Количество"])
    sheet.append(["Первое имя", "Второе имя", 1])
    workbook.save(ambiguous)
    parsed = TenderWorkbookParser().parse(ambiguous)
    assert parsed["mapping_required"] is True
    mapped = TenderWorkbookParser().parse(
        ambiguous,
        mapping_override={"name": 2, "quantity": 3},
        selected_sheet=sheet.title,
        selected_header_row=1,
    )
    assert mapped["mapping_required"] is False
    assert mapped["rows"][0]["name"] == "Второе имя"

    excessive = tmp_path / "over-500.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Наименование", "Количество"])
    for index in range(MAX_ACTUAL_ITEMS + 1):
        sheet.append([f"Синтетическая позиция {index}", 1])
    workbook.save(excessive)
    with pytest.raises(TenderParseError, match="больше 500"):
        TenderWorkbookParser().parse(excessive)


def test_preview_and_confirmed_workspaces_are_owner_bound_persistent_and_path_free(tmp_path):
    activity = TenderActivityRegistry()
    repository = TenderWorkspaceRepository(tmp_path / "data", activity)
    source_path = tmp_path / "sample.xlsx"
    payload = _official(source_path)
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner-a", "sample.xlsx", len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    analysis = TenderWorkbookParser().parse(repository.preview_root / preview["preview_id"] / "source.xlsx", tender_id=preview["preview_id"])
    repository.update_preview(preview["preview_id"], "owner-a", status="ready", parser_version=analysis["parser_version"], mapping=analysis["mapping"], analysis=analysis)
    with pytest.raises(TenderWorkspaceError) as denied:
        repository.public_preview(preview["preview_id"], "owner-b")
    assert denied.value.status_code == 404
    public_preview = repository.public_preview(preview["preview_id"], "owner-a")
    assert "source.xlsx" not in json.dumps(public_preview)
    assert str(tmp_path) not in json.dumps(public_preview)
    workspace = repository.confirm(preview["preview_id"], "owner-a")
    assert workspace["counts"]["item"] == 1
    assert workspace["rows"][0]["resource_code"] == "R-001"
    assert workspace["rows"][0]["article"] == "SKU-1"
    assert "owner_id" not in workspace and str(tmp_path) not in json.dumps(workspace)
    original_source = repository.workspace_root / workspace["tender_id"] / "source.xlsx"
    assert original_source.read_bytes() == payload
    assert original_source.stat().st_mode & 0o200 == 0
    recovered = TenderWorkspaceRepository(tmp_path / "data", activity).public_workspace(workspace["tender_id"], "owner-a")
    assert recovered["source_sha256"] == digest


def test_preview_and_confirmed_ttl_absolute_lifetime_and_restart_job_semantics(tmp_path):
    repository = TenderWorkspaceRepository(tmp_path / "data")
    payload_path = tmp_path / "source.xlsx"
    payload = _official(payload_path)
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner", "source.xlsx", len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    metadata_path = repository.preview_root / preview["preview_id"] / "workspace.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert datetime.fromisoformat(metadata["expires_at"]) - datetime.fromisoformat(metadata["created_at"]) == timedelta(seconds=PREVIEW_TTL_SECONDS)
    metadata["status"] = "queued"
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
    restarted = TenderWorkspaceRepository(tmp_path / "data")
    assert restarted.public_preview(preview["preview_id"], "owner")["status"] == "interrupted"
    assert "перезапуском" in restarted.public_preview(preview["preview_id"], "owner")["error"]


def test_confirmed_ttl_does_not_extend_absolute_deadline(tmp_path):
    repository = TenderWorkspaceRepository(tmp_path / "data")
    payload_path = tmp_path / "source.xlsx"
    payload = _official(payload_path)
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner", "source.xlsx", len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    analysis = TenderWorkbookParser().parse(repository.preview_root / preview["preview_id"] / "source.xlsx")
    repository.update_preview(preview["preview_id"], "owner", status="ready", analysis=analysis, parser_version=analysis["parser_version"], mapping=analysis["mapping"])
    workspace = repository.confirm(preview["preview_id"], "owner")
    metadata_path = repository.workspace_root / workspace["tender_id"] / "workspace.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert datetime.fromisoformat(metadata["absolute_expires_at"]) - datetime.fromisoformat(metadata["created_at"]) == timedelta(seconds=ABSOLUTE_TTL_SECONDS)
    before = metadata["absolute_expires_at"]
    repository.public_workspace(workspace["tender_id"], "owner")
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert metadata["absolute_expires_at"] == before
    assert datetime.fromisoformat(metadata["last_access_at"]) - datetime.fromisoformat(metadata["created_at"]) < timedelta(seconds=IDLE_TTL_SECONDS)


def test_concurrent_workspace_touches_keep_metadata_and_source_coherent(tmp_path):
    repository = TenderWorkspaceRepository(tmp_path / "concurrent-data")
    source_path = tmp_path / "concurrent-source.xlsx"
    payload = _official(source_path)
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner-concurrent", source_path.name, len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    analysis = TenderWorkbookParser().parse(repository.preview_root / preview["preview_id"] / "source.xlsx")
    repository.update_preview(preview["preview_id"], "owner-concurrent", status="ready", analysis=analysis, parser_version=analysis["parser_version"], mapping=analysis["mapping"])
    workspace = repository.confirm(preview["preview_id"], "owner-concurrent")
    workspace_dir = repository.workspace_root / workspace["tender_id"]
    metadata_path = workspace_dir / "workspace.json"
    original = json.loads(metadata_path.read_text(encoding="utf-8"))
    barrier = threading.Barrier(12)

    def touch_workspace(_):
        barrier.wait()
        return repository.public_workspace(workspace["tender_id"], "owner-concurrent", touch=True)

    with ThreadPoolExecutor(max_workers=12) as executor:
        results = list(executor.map(touch_workspace, range(12)))
    persisted = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert len(results) == 12
    assert persisted["owner_id"] == original["owner_id"]
    assert persisted["filename"] == original["filename"]
    assert persisted["source_sha256"] == original["source_sha256"] == digest
    assert persisted["absolute_expires_at"] == original["absolute_expires_at"]
    assert hashlib.sha256((workspace_dir / "source.xlsx").read_bytes()).hexdigest() == digest
    assert datetime.fromisoformat(persisted["last_access_at"])
    assert not list(workspace_dir.glob(".workspace.json.*.tmp"))
    orphan = workspace_dir / ".workspace.json.orphan.tmp"
    orphan.write_text("incomplete", encoding="utf-8")
    os.utime(orphan, (1, 1))
    repository.cleanup()
    assert not orphan.exists()


def test_workspace_quotas_are_atomic_and_enforced(tmp_path, monkeypatch):
    import averon_import.services.manual_tenders.repository as repository_module
    monkeypatch.setattr(repository_module, "MAX_WORKSPACES_GLOBAL", 20)
    repository = TenderWorkspaceRepository(tmp_path / "data")

    def prepare(target_repository, owner, marker):
        source = tmp_path / f"{marker}.xlsx"
        payload = _official(source)
        workbook = load_workbook(source)
        workbook[TEMPLATE_SHEET]["A2"] = marker
        workbook.save(source)
        payload = source.read_bytes()
        digest = hashlib.sha256(payload).hexdigest()
        preview = target_repository.reserve_preview(owner, f"{marker}.xlsx", len(payload), digest)
        target_repository.write_preview_source(preview["preview_id"], payload)
        analysis = TenderWorkbookParser().parse(target_repository.preview_root / preview["preview_id"] / "source.xlsx", tender_id=preview["preview_id"])
        target_repository.update_preview(preview["preview_id"], owner, status="ready", analysis=analysis, parser_version=analysis["parser_version"], mapping=analysis["mapping"])
        return preview["preview_id"]

    for index in range(MAX_WORKSPACES_PER_USER):
        preview_id = prepare(repository, "same-owner", f"quota-{index}")
        assert repository.confirm(preview_id, "same-owner")
    over_user = prepare(repository, "same-owner", "quota-over-user")
    with pytest.raises(TenderWorkspaceError) as user_limit:
        repository.confirm(over_user, "same-owner")
    assert user_limit.value.code == "TENDER_WORKSPACE_QUOTA"
    assert len(repository._workspace_dirs()) == MAX_WORKSPACES_PER_USER

    global_repository = TenderWorkspaceRepository(tmp_path / "global-data")
    monkeypatch.setattr(repository_module, "MAX_WORKSPACES_GLOBAL", 2)
    previews = []
    for index in range(5):
        owner = f"global-owner-{index}"
        previews.append((prepare(global_repository, owner, f"global-{index}"), owner))
    barrier = threading.Barrier(len(previews))
    def confirm_global(pair):
        preview_id, owner = pair
        barrier.wait()
        try:
            return global_repository.confirm(preview_id, owner)["tender_id"]
        except TenderWorkspaceError as exc:
            return exc.code
    with ThreadPoolExecutor(max_workers=len(previews)) as executor:
        results = list(executor.map(confirm_global, previews))
    assert sum(not result.startswith("TENDER_") for result in results) == 2
    assert results.count("TENDER_WORKSPACE_QUOTA") == 3
    assert len(global_repository._workspace_dirs()) == 2
    assert MAX_UNCOMPRESSED_BYTES == 20 * 1024 * 1024 and MAX_UPLOAD_BYTES == 5 * 1024 * 1024 and MAX_ACTUAL_ITEMS == 500


def test_activity_lease_prevents_delete_and_lazy_cleanup(tmp_path):
    activity = TenderActivityRegistry()
    repository = TenderWorkspaceRepository(tmp_path / "data", activity)
    payload_path = tmp_path / "source.xlsx"
    payload = _official(payload_path)
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("owner", "source.xlsx", len(payload), digest)
    repository.write_preview_source(preview["preview_id"], payload)
    analysis = TenderWorkbookParser().parse(repository.preview_root / preview["preview_id"] / "source.xlsx")
    repository.update_preview(preview["preview_id"], "owner", status="ready", analysis=analysis, parser_version=analysis["parser_version"], mapping=analysis["mapping"])
    workspace = repository.confirm(preview["preview_id"], "owner")
    lease = activity.acquire(workspace["tender_id"])
    with pytest.raises(TenderWorkspaceError) as busy:
        repository.delete(workspace["tender_id"], "owner")
    assert busy.value.code == "TENDER_WORKSPACE_BUSY"
    metadata_path = repository.workspace_root / workspace["tender_id"] / "workspace.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    expired = datetime.now(timezone.utc) - timedelta(days=3)
    metadata["last_access_at"] = expired.isoformat()
    metadata["absolute_expires_at"] = expired.isoformat()
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
    repository.cleanup()
    assert metadata_path.exists()
    lease.release()
    lease.release()
    repository.cleanup()
    assert not metadata_path.exists()


class _ApiResponse:
    def __init__(self, status, headers, body):
        self.status_code = status
        self.headers = {key.decode().lower(): value.decode() for key, value in headers}
        self.content = body
        self.text = body.decode("utf-8", errors="replace")

    def json(self):
        return json.loads(self.content.decode("utf-8"))


def _api_request(app, method, path, *, headers=None, body=b""):
    sent = False
    messages = []
    route_path, separator, query_string = path.partition("?")
    scope = {
        "type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1",
        "method": method, "scheme": "http", "path": route_path, "raw_path": route_path.encode(),
        "query_string": query_string.encode() if separator else b"", "headers": [(k.lower().encode(), v.encode()) for k, v in (headers or {}).items()],
        "client": ("127.0.0.1", 12345), "server": ("127.0.0.1", 8765),
    }
    async def receive():
        nonlocal sent
        if sent:
            return {"type": "http.disconnect"}
        sent = True
        return {"type": "http.request", "body": body, "more_body": False}
    async def send(message):
        messages.append(message)
    asyncio.run(app(scope, receive, send))
    start = next(item for item in messages if item["type"] == "http.response.start")
    response_body = b"".join(item.get("body", b"") for item in messages if item["type"] == "http.response.body")
    return _ApiResponse(start["status"], start.get("headers", []), response_body)


def _multipart_file(name: str, filename: str, payload: bytes) -> tuple[str, bytes]:
    boundary = f"----Tender-{uuid.uuid4().hex}"
    body = (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"; filename=\"{filename}\"\r\n"
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + payload + f"\r\n--{boundary}--\r\n".encode()
    return boundary, body


@pytest.fixture
def tender_api(monkeypatch, tmp_path):
    monkeypatch.setenv("AVERON_AUTH_MODE", "trusted_proxy")
    monkeypatch.setenv("AVERON_PROXY_SECRET", "tender-test-proxy")
    monkeypatch.setenv("AVERON_ADMIN_USERS", "tender-admin")
    monkeypatch.setenv("AVERON_DATA_DIR", str(tmp_path / "app-data"))
    from averon_import import main
    activity = TenderActivityRegistry()
    repository = TenderWorkspaceRepository(tmp_path / "api-data", activity)
    monkeypatch.setattr(main, "tender_activity", activity)
    monkeypatch.setattr(main, "tender_repository", repository)
    monkeypatch.setattr(main, "tender_sourcing_runs", TenderSourcingRunStore(repository))
    from averon_import.services.manual_tenders.history_decisions import TenderHistoryDecisionStore
    monkeypatch.setattr(main, "tender_history_decisions", TenderHistoryDecisionStore(repository))
    from averon_import.services.manual_tenders.price_export import TenderPriceExportRepository
    monkeypatch.setattr(main, "tender_price_exports", TenderPriceExportRepository(repository))
    isolated_jobs = JobService(max_workers=1)
    monkeypatch.setattr(main, "job_service", isolated_jobs)
    return main, repository


def _auth_headers(username="tender-user"):
    return {"X-Averon-User": username, "X-Averon-Proxy": "tender-test-proxy"}


def test_api_auth_template_preview_confirm_owner_isolation_and_delete(tender_api, tmp_path):
    main, repository = tender_api
    unauthenticated = _api_request(main.app, "GET", "/api/manual-tenders/template")
    assert unauthenticated.status_code == 401
    template = _api_request(main.app, "GET", "/api/manual-tenders/template", headers=_auth_headers())
    assert template.status_code == 200
    assert "attachment" in template.headers.get("content-disposition", "")
    workbook_path = tmp_path / "downloaded-template.xlsx"
    workbook_path.write_bytes(template.content)
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    assert TEMPLATE_SHEET in workbook.sheetnames and TEMPLATE_TABLE in workbook[TEMPLATE_SHEET].tables
    workbook.close()

    payload = _official(tmp_path / "api-upload.xlsx")
    boundary, body = _multipart_file("file", "test.xlsx", payload)
    headers = _auth_headers()
    headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    uploaded = _api_request(main.app, "POST", "/api/manual-tenders/previews", headers=headers, body=body)
    assert uploaded.status_code == 202, uploaded.text
    preview_id = uploaded.json()["preview_id"]
    job_id = uploaded.json().get("job_id")
    assert repository.preview_path(preview_id, "username:tender-user")[1]["owner_id"] == "username:tender-user"
    preview = None
    for _ in range(100):
        preview = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers())
        assert "status" in preview.json(), f"preview={preview_id} http={preview.status_code} body={preview.text}"
        if preview.json()["status"] in {"ready", "failed"}:
            break
        threading.Event().wait(0.02)
    job_traceback = main.job_service.get(job_id, owner_id="username:tender-user").traceback if job_id else None
    assert preview.json()["status"] == "ready", f"{preview.text}\n{job_traceback or ''}"
    same_boundary, same_body = _multipart_file("file", "renamed.xlsx", payload)
    same_headers = {**_auth_headers(), "Content-Type": f"multipart/form-data; boundary={same_boundary}"}
    duplicate = _api_request(main.app, "POST", "/api/manual-tenders/previews", headers=same_headers, body=same_body)
    assert duplicate.status_code == 202
    assert duplicate.json()["preview_id"] == preview_id and duplicate.json()["deduplicated"] is True
    confirmed = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/confirm", headers=_auth_headers())
    assert confirmed.status_code == 201, confirmed.text
    workspace = confirmed.json()
    assert workspace["counts"]["item"] == 1
    encoded = json.dumps(workspace)
    assert str(repository.root) not in encoded
    assert "owner_id" not in workspace
    forbidden = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers("other-user"))
    assert forbidden.status_code == 404
    fetched = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers())
    assert fetched.status_code == 200 and fetched.json()["tender_id"] == workspace["tender_id"]
    deleted = _api_request(main.app, "DELETE", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers())
    assert deleted.status_code == 200 and deleted.json()["deleted"] is True
    main.job_service.executor.shutdown(wait=True)


def test_recent_tender_collection_is_owner_scoped_safe_non_touching_and_cleans_expired(tender_api, tmp_path):
    main, repository = tender_api
    owner_a_old = _confirm_tender_for_owner(main, repository, tmp_path, "username:user-a", "a-old.xlsx")
    owner_a_new = _confirm_tender_for_owner(main, repository, tmp_path, "username:user-a", "a-new.xlsx")
    owner_a_latest = _confirm_tender_for_owner(main, repository, tmp_path, "username:user-a", "a-latest.xlsx")
    owner_b = _confirm_tender_for_owner(main, repository, tmp_path, "username:user-b", "b-only.xlsx")

    old_access = datetime.now(timezone.utc) - timedelta(hours=2)
    new_access = datetime.now(timezone.utc) - timedelta(hours=1)
    for workspace, last_access in ((owner_a_old, old_access), (owner_a_new, new_access), (owner_a_latest, new_access)):
        metadata_path = repository.workspace_root / workspace["tender_id"] / "workspace.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["last_access_at"] = last_access.isoformat()
        repository._atomic_json(metadata_path, metadata)

    unauthenticated = _api_request(main.app, "GET", "/api/manual-tenders?limit=10")
    assert unauthenticated.status_code == 401
    invalid_limit = _api_request(main.app, "GET", "/api/manual-tenders?limit=21", headers=_auth_headers("user-a"))
    assert invalid_limit.status_code == 422

    before = {
        workspace["tender_id"]: json.loads((repository.workspace_root / workspace["tender_id"] / "workspace.json").read_text(encoding="utf-8"))["last_access_at"]
        for workspace in (owner_a_old, owner_a_new, owner_a_latest, owner_b)
    }
    listed_a = _api_request(main.app, "GET", "/api/manual-tenders?limit=10", headers=_auth_headers("user-a"))
    assert listed_a.status_code == 200
    payload_a = listed_a.json()
    assert payload_a["active_count"] == 3
    assert payload_a["limit"] == MAX_WORKSPACES_PER_USER
    assert [item["tender_id"] for item in payload_a["tenders"]] == [owner_a_latest["tender_id"], owner_a_new["tender_id"], owner_a_old["tender_id"]]
    assert {item["filename"] for item in payload_a["tenders"]} == {"a-old.xlsx", "a-new.xlsx", "a-latest.xlsx"}
    assert all(item["item_count"] == 1 for item in payload_a["tenders"])
    limited = _api_request(main.app, "GET", "/api/manual-tenders?limit=1", headers=_auth_headers("user-a"))
    assert limited.status_code == 200
    assert limited.json()["active_count"] == 3 and len(limited.json()["tenders"]) == 1
    assert all(set(item) <= {"tender_id", "filename", "sheet_name", "item_count", "created_at", "last_access_at", "absolute_expires_at", "revision"} for item in payload_a["tenders"])
    encoded_a = json.dumps(payload_a)
    assert str(repository.root) not in encoded_a
    for forbidden in ("owner_id", "source_manifest", "rows", "run", "decisions", "exports", "global_count"):
        assert forbidden not in encoded_a
    after_list = {
        workspace["tender_id"]: json.loads((repository.workspace_root / workspace["tender_id"] / "workspace.json").read_text(encoding="utf-8"))["last_access_at"]
        for workspace in (owner_a_old, owner_a_new, owner_a_latest, owner_b)
    }
    assert after_list == before, "listing summaries must not refresh idle TTL"

    listed_b = _api_request(main.app, "GET", "/api/manual-tenders", headers=_auth_headers("user-b"))
    assert listed_b.status_code == 200
    assert listed_b.json()["active_count"] == 1
    assert [item["tender_id"] for item in listed_b.json()["tenders"]] == [owner_b["tender_id"]]
    assert owner_a_old["tender_id"] not in json.dumps(listed_b.json())

    guessed_open = _api_request(main.app, "GET", f"/api/manual-tenders/{owner_a_old['tender_id']}", headers=_auth_headers("user-b"))
    guessed_delete = _api_request(main.app, "DELETE", f"/api/manual-tenders/{owner_a_old['tender_id']}", headers=_auth_headers("user-b"))
    assert guessed_open.status_code == guessed_delete.status_code == 404
    assert (repository.workspace_root / owner_a_old["tender_id"]).exists()

    opened = _api_request(main.app, "GET", f"/api/manual-tenders/{owner_a_old['tender_id']}", headers=_auth_headers("user-a"))
    assert opened.status_code == 200
    opened_access = opened.json()["last_access_at"]
    assert datetime.fromisoformat(opened_access) > datetime.fromisoformat(before[owner_a_old["tender_id"]])

    lease = main.tender_activity.acquire(owner_a_new["tender_id"])
    try:
        busy_delete = _api_request(main.app, "DELETE", f"/api/manual-tenders/{owner_a_new['tender_id']}", headers=_auth_headers("user-a"))
        assert busy_delete.status_code == 409
        assert busy_delete.json()["detail"]["code"] == "TENDER_WORKSPACE_BUSY"
        assert (repository.workspace_root / owner_a_new["tender_id"]).exists()
    finally:
        lease.release()

    expired_path = repository.workspace_root / owner_a_old["tender_id"] / "workspace.json"
    expired_metadata = json.loads(expired_path.read_text(encoding="utf-8"))
    expired_metadata["last_access_at"] = (datetime.now(timezone.utc) - timedelta(hours=25)).isoformat()
    repository._atomic_json(expired_path, expired_metadata)
    after_cleanup = _api_request(main.app, "GET", "/api/manual-tenders", headers=_auth_headers("user-a"))
    assert after_cleanup.status_code == 200
    assert after_cleanup.json()["active_count"] == 2
    assert [item["tender_id"] for item in after_cleanup.json()["tenders"]] == [owner_a_latest["tender_id"], owner_a_new["tender_id"]]
    assert not (repository.workspace_root / owner_a_old["tender_id"]).exists()


def test_api_confirmation_returns_typed_invalid_rows_error(tender_api, tmp_path):
    main, _repository = tender_api
    path = tmp_path / "api-incomplete.xlsx"
    _official(path, include_required=False)
    workbook = load_workbook(path)
    workbook[TEMPLATE_SHEET]["A2"] = "R-2"
    workbook[TEMPLATE_SHEET]["B2"] = "Синтетическая позиция"
    workbook[TEMPLATE_SHEET]["C2"] = "шт"
    workbook.save(path)
    workbook.close()
    boundary, body = _multipart_file("file", path.name, path.read_bytes())
    uploaded = _api_request(main.app, "POST", "/api/manual-tenders/previews", headers={**_auth_headers(), "Content-Type": f"multipart/form-data; boundary={boundary}"}, body=body)
    assert uploaded.status_code == 202
    preview_id = uploaded.json()["preview_id"]
    for _ in range(100):
        preview = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers()).json()
        if preview["status"] in {"ready", "failed"}:
            break
        threading.Event().wait(0.02)
    assert preview["status"] == "ready"
    assert preview["analysis"]["invalid_count"] == 1
    assert preview["analysis"]["invalid_examples"] == [{"excel_row": 2, "reasons": ["missing_quantity"]}]
    confirmed = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/confirm", headers=_auth_headers())
    assert confirmed.status_code == 409
    assert confirmed.json()["detail"]["code"] == "TENDER_INVALID_ROWS"
    assert "требующие исправления" in confirmed.json()["detail"]["message"]
    main.job_service.executor.shutdown(wait=True)


def test_api_mapping_resolves_generic_code_without_article_leak(tender_api, tmp_path):
    main, _repository = tender_api
    path = tmp_path / "manual-mapping.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Код", "Наименование", "Количество"])
    sheet.append(["R-09", "Синтетическая позиция", 2])
    workbook.save(path)
    boundary, body = _multipart_file("file", path.name, path.read_bytes())
    headers = {**_auth_headers(), "Content-Type": f"multipart/form-data; boundary={boundary}"}
    uploaded = _api_request(main.app, "POST", "/api/manual-tenders/previews", headers=headers, body=body)
    preview_id = uploaded.json()["preview_id"]
    for _ in range(100):
        preview = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers()).json()
        assert "status" in preview, preview
        if preview["status"] in {"ready", "failed"}:
            break
        threading.Event().wait(0.02)
    assert preview["analysis"]["mapping_required"] is True
    analysis = preview["analysis"]
    mapping = {"resource_code": 1, "name": 2, "unit": None, "quantity": 3, "article": None, "manufacturer": None, "model": None}
    unsafe_article_mapping = {**mapping, "resource_code": None, "article": 1}
    unsafe_request = json.dumps({"sheet_name": analysis["sheet_name"], "header_row": analysis["header_row"], "mapping": unsafe_article_mapping}, ensure_ascii=False).encode("utf-8")
    unsafe = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/mapping", headers={**_auth_headers(), "Content-Type": "application/json"}, body=unsafe_request)
    assert unsafe.status_code == 400
    assert _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers()).json()["status"] == "ready"
    request = json.dumps({"sheet_name": analysis["sheet_name"], "header_row": analysis["header_row"], "mapping": mapping}, ensure_ascii=False).encode("utf-8")
    mapped = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/mapping", headers={**_auth_headers(), "Content-Type": "application/json"}, body=request)
    assert mapped.status_code == 202, mapped.text
    for _ in range(100):
        preview_response = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers())
        preview = preview_response.json()
        if preview["status"] in {"ready", "failed"}:
            break
        threading.Event().wait(0.02)
    assert preview["status"] == "ready" and preview["analysis"]["mapping_required"] is False
    confirmed = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/confirm", headers=_auth_headers())
    assert confirmed.status_code == 201, confirmed.text
    row = confirmed.json()["rows"][0]
    assert row["resource_code"] == "R-09"
    assert row["article"] == ""
    main.job_service.executor.shutdown(wait=True)


def test_api_mapping_accepts_only_a_detected_alternate_header_candidate(tender_api, tmp_path):
    main, _repository = tender_api
    path = tmp_path / "multiple-header-candidates.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Первый лист"
    first.append(["Наименование", "Количество"])
    first.append(["Не выбранная позиция", 1])
    second = workbook.create_sheet("Второй лист")
    second.append(["Наименование", "Количество"])
    second.append(["Выбранная позиция", 2])
    workbook.save(path)
    boundary, body = _multipart_file("file", path.name, path.read_bytes())
    headers = {**_auth_headers(), "Content-Type": f"multipart/form-data; boundary={boundary}"}
    uploaded = _api_request(main.app, "POST", "/api/manual-tenders/previews", headers=headers, body=body)
    assert uploaded.status_code == 202, uploaded.text
    preview_id = uploaded.json()["preview_id"]
    for _ in range(100):
        preview_response = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers())
        preview = preview_response.json()
        if preview["status"] in {"ready", "failed"}:
            break
        threading.Event().wait(0.02)
    assert preview["status"] == "ready", preview_response.text
    analysis = preview["analysis"]
    alternate = next(candidate for candidate in analysis["header_candidates"] if candidate["sheet_name"] == "Второй лист")
    mapping = {"resource_code": None, "name": 1, "unit": None, "quantity": 2, "article": None, "manufacturer": None, "model": None}

    invalid = json.dumps({"sheet_name": "Незаявленный лист", "header_row": 1, "mapping": mapping}, ensure_ascii=False).encode("utf-8")
    rejected = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/mapping", headers={**_auth_headers(), "Content-Type": "application/json"}, body=invalid)
    assert rejected.status_code == 409
    invalid_mapping = {**mapping, "resource_code": 1, "article": 1}
    invalid = json.dumps({"sheet_name": alternate["sheet_name"], "header_row": alternate["header_row"], "mapping": invalid_mapping}, ensure_ascii=False).encode("utf-8")
    rejected = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/mapping", headers={**_auth_headers(), "Content-Type": "application/json"}, body=invalid)
    assert rejected.status_code == 400
    still_ready = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers()).json()
    assert still_ready["status"] == "ready"

    request = json.dumps({"sheet_name": alternate["sheet_name"], "header_row": alternate["header_row"], "mapping": mapping}, ensure_ascii=False).encode("utf-8")
    mapped = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/mapping", headers={**_auth_headers(), "Content-Type": "application/json"}, body=request)
    assert mapped.status_code == 202, mapped.text
    for _ in range(100):
        preview_response = _api_request(main.app, "GET", f"/api/manual-tenders/previews/{preview_id}", headers=_auth_headers())
        preview = preview_response.json()
        if preview["status"] in {"ready", "failed"}:
            break
        threading.Event().wait(0.02)
    assert preview["status"] == "ready", preview_response.text
    confirmed = _api_request(main.app, "POST", f"/api/manual-tenders/previews/{preview_id}/confirm", headers=_auth_headers())
    assert confirmed.status_code == 201, confirmed.text
    assert confirmed.json()["sheet_name"] == "Второй лист"
    assert confirmed.json()["rows"][0]["name"] == "Выбранная позиция"
    main.job_service.executor.shutdown(wait=True)


def test_upload_returns_normal_busy_response_when_document_processing_lane_is_saturated(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    from averon_import.services.jobs import DOCUMENT_PROCESSING, SOURCING
    isolated = JobService(capacities={DOCUMENT_PROCESSING: (1, 0), SOURCING: (1, 2)})
    monkeypatch.setattr(main, "job_service", isolated)
    started = threading.Event()
    release = threading.Event()
    isolated.submit(lambda progress: (started.set(), release.wait(5))[1], lane=DOCUMENT_PROCESSING)
    assert started.wait(2)
    payload = _official(tmp_path / "busy.xlsx")
    boundary, body = _multipart_file("file", "busy.xlsx", payload)
    response = _api_request(main.app, "POST", "/api/manual-tenders/previews", headers={**_auth_headers(), "Content-Type": f"multipart/form-data; boundary={boundary}"}, body=body)
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "JOB_LANE_BUSY"
    release.set()
    isolated.executor.shutdown(wait=True)


def test_ui_exposes_excel_path_without_changing_manual_sourcing_and_clears_on_logout():
    root = Path(__file__).resolve().parents[1]
    template = (root / "averon_import" / "templates" / "index.html").read_text(encoding="utf-8")
    script = (root / "averon_import" / "static" / "app.js").read_text(encoding="utf-8")
    assert 'id="open-manual-entry">Ввести позиции вручную' in template
    assert 'id="open-tender-xlsx">Загрузить Excel' in template
    assert 'id="download-tender-template">Скачать шаблон Excel' in template
    assert "async function downloadExcelTenderTemplate()" in script
    assert "async function uploadExcelTender(file)" in script
    assert "async function waitForExcelTenderPreview(generation)" in script
    assert "invalidCount > 0" in script
    assert "Требуют исправления:" in script
    assert "TENDER_INVALID_ROWS" not in script
    assert "function renderExcelTenderRows()" in script
    assert "function clearExcelTenderState()" in script
    assert 'state.manual.rows.filter((row) => row.selected !== false)' in script
    assert "sessionStorage.setItem(EXCEL_TENDER_VIEW_KEY, JSON.stringify({tender_id:tenderId,user_id:userId}))" in script
    assert "sessionStorage.setItem(EXCEL_TENDER_VIEW_KEY, JSON.stringify({rows" not in script
    assert "localStorage.setItem(EXCEL_TENDER_VIEW_KEY" not in script
    assert "const generation = ++state.excelTender.pollGeneration" in script
    assert 'run.status === "completed"' in script
    assert "/runs/${encodeURIComponent(runId)}/export" in script
    assert "include_historical_prices:false, allow_partial:false" in script
    assert "TENDER_EXPORT_HISTORICAL_CONFIRMATION_REQUIRED" in script
    assert "TENDER_EXPORT_PARTIAL_CONFIRMATION_REQUIRED" in script
    export_action = script[script.index("async function startExcelTenderPriceExport()") : script.index("async function pollExcelTenderPriceExport(")]
    assert 'id="tender-history-confirmation-modal"' in template
    assert 'value="include">Включить исторические цены' in template
    assert 'value="exclude">Продолжить без исторических цен' in template
    assert 'value="cancel">Отмена' in template
    assert "historical_decision_confirmed" in export_action
    assert "chooseTenderHistoricalPricePolicy(summary, runId)" in export_action
    assert "body:JSON.stringify(options)" in export_action
    assert '"price"' not in export_action and "localStorage" not in export_action and "sessionStorage" not in export_action
    assert 'id="tender-export-run"' in template and 'id="tender-export-download"' in template


def test_read_only_excel_workspace_renders_370_rows_with_filtering():
    node = shutil.which("node")
    assert node, "Node.js is required for tender UI lifecycle regression"
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [node, str(root / "tests" / "js" / "manual_tender_lifecycle.cjs"), str(root / "averon_import" / "static" / "app.js")],
        capture_output=True, text=True, timeout=20, check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "PASS: 370-row workspace, default tender selection, and generation-safe adaptive sourcing polling" in result.stdout


def _confirm_synthetic_tender(main, repository, tmp_path, *, count=1, official=False):
    path = tmp_path / f"phase-b-{count}-{uuid.uuid4().hex}.xlsx"
    if official:
        payload = _official(path)
    else:
        _official(path, include_required=False)
        workbook = load_workbook(path)
        sheet = workbook[TEMPLATE_SHEET]
        units = ["100 шт", "1000 шт", "10 м", "затрата", "шт", "м", "кг"]
        for index in range(count):
            row_number = index + 2
            sheet.cell(row_number, 1, f"R-{index + 1:04d}")
            sheet.cell(row_number, 2, f"Синтетическая позиция {index + 1}")
            sheet.cell(row_number, 3, units[index % len(units)])
            sheet.cell(row_number, 4, index % 5 + 1)
            sheet.cell(row_number, 5, "")
        workbook.save(path)
        workbook.close()
        payload = path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("username:tender-user", path.name, len(payload), digest)
    source_path = repository.write_preview_source(preview["preview_id"], payload)
    analysis = main.tender_parser.parse(source_path, tender_id=preview["preview_id"])
    repository.update_preview(
        preview["preview_id"], "username:tender-user", status="ready",
        parser_version=analysis["parser_version"], mapping=analysis["mapping"], analysis=analysis,
    )
    return repository.confirm(preview["preview_id"], "username:tender-user")


def _confirm_tender_for_owner(main, repository, tmp_path, owner_id, filename):
    path = tmp_path / f"owner-bound-{uuid.uuid4().hex}.xlsx"
    payload = _official(path)
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview(owner_id, filename, len(payload), digest)
    source_path = repository.write_preview_source(preview["preview_id"], payload)
    analysis = main.tender_parser.parse(source_path, tender_id=preview["preview_id"])
    repository.update_preview(
        preview["preview_id"], owner_id, status="ready",
        parser_version=analysis["parser_version"], mapping=analysis["mapping"], analysis=analysis,
    )
    return repository.confirm(preview["preview_id"], owner_id)


def _fake_tender_sourcing_service(monkeypatch, main, *, catalog_version="catalog-v1"):
    from types import SimpleNamespace
    from averon_import.services.sourcing.models import (
        ProjectSourcingResult, SourcingResult, SourcingRouteMetadata, SourcingSourceMode,
    )
    from averon_import.services.sourcing.product_understanding import build_fallback_intent

    class FakeService:
        def __init__(self):
            self.received = []
            self.calls = 0
            self.modes = []

        def provider(self, provider_key=None):
            if provider_key == "invalid-provider":
                raise ValueError("unknown provider")
            return SimpleNamespace(key=provider_key or "fake-provider", label="Fake provider")

        def _project(self, rows, source_mode, progress):
            self.calls += 1
            self.received = [dict(row) for row in rows]
            self.modes.append(SourcingSourceMode(source_mode))
            if progress:
                progress(0, len(rows), "test")
            results = []
            for row in rows:
                route = None
                if source_mode != SourcingSourceMode.PROVIDER_ONLY:
                    route = SourcingRouteMetadata(
                        source_mode=source_mode, final_source_kind="none",
                        history_catalog_version=catalog_version,
                    )
                results.append(SourcingResult(intent=build_fallback_intent(row), route=route))
            if progress:
                progress(len(rows), len(rows), "test")
            return ProjectSourcingResult(
                positions_total=len(rows), positions_processed=len(rows),
                positions_without_offers=len(rows), source_mode=source_mode,
                catalog_version=catalog_version, results=results,
            )

        def search_project(self, rows, *, source_mode=SourcingSourceMode.PROVIDER_ONLY, progress=None, **_kwargs):
            return self._project(rows, SourcingSourceMode.PROVIDER_ONLY, progress)

        def search_project_routed(self, rows, *, source_mode, progress=None, **_kwargs):
            return self._project(rows, source_mode, progress)

    service = FakeService()
    monkeypatch.setattr(main, "sourcing_service", service)
    monkeypatch.setattr(main, "sourcing_runtime", SimpleNamespace(service=service))
    return service


def _post_tender_sourcing(main, tender_id, payload, *, username="tender-user"):
    body = json.dumps(payload).encode("utf-8")
    return _api_request(
        main.app, "POST", f"/api/manual-tenders/{tender_id}/sourcing",
        headers={**_auth_headers(username), "Content-Type": "application/json"}, body=body,
    )


def _wait_tender_job(main, job_id, *, timeout=3):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            job = main.job_service.get_public(job_id, owner_id="tender-user")
        except KeyError:
            job = None
        if job and job["status"] in {"completed", "failed", "expired"}:
            return job
        time.sleep(0.01)
    raise AssertionError(f"tender sourcing job did not finish: {job_id}")


def test_tender_sourcing_is_server_authoritative_durable_and_owner_bound(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=370)
    service = _fake_tender_sourcing_service(monkeypatch, main)
    before_sha = hashlib.sha256((repository.workspace_root / workspace["tender_id"] / "source.xlsx").read_bytes()).hexdigest()
    items = [row for row in workspace["rows"] if row["row_type"] == "item"]
    selected_ids = [row["source_row_id"] for row in items]
    assert len(selected_ids) == len(set(selected_ids)) == 370

    response = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": selected_ids, "source_mode": "provider_only", "limit": 20,
    })
    assert response.status_code == 202, response.text
    job = _wait_tender_job(main, response.json()["id"])
    assert job["status"] == "completed", json.dumps(job, ensure_ascii=False, default=str)
    assert job["result"]["tender_id"] == workspace["tender_id"]
    assert job["result"]["run_id"]
    assert len(service.received) == 370
    assert all("resource_code" not in row and "code" not in row and "source_text" not in row for row in service.received)
    assert all(row["quantity_trusted"] is True for row in service.received)
    assert {row["unit"] for row in service.received} >= {"100 шт", "1000 шт", "10 м", "затрата"}
    for source, adapted in zip(items, service.received, strict=True):
        assert adapted["source_row_id"] == source["source_row_id"]
        assert adapted["article"] == source["article"]
        assert adapted["quantity"] == source["quantity"]
        assert adapted["unit"] == source["raw_unit"]
    no_article = next(row for row in items if not row["article"] and row["resource_code"])
    adapted = next(row for row in service.received if row["source_row_id"] == no_article["source_row_id"])
    from averon_import.services.sourcing.product_understanding import build_fallback_intent
    assert adapted["article"] == "" and build_fallback_intent(adapted).article == ""

    run_id = job["result"]["run_id"]
    owner_detail = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}", headers=_auth_headers())
    assert owner_detail.status_code == 200
    run = owner_detail.json()
    assert run["status"] == "completed"
    assert run["source_sha256"] == workspace["source_sha256"]
    assert run["workspace_revision"] == workspace["revision"]
    assert len(run["rows"]) == 370
    assert {item["source_row_id"] for item in run["rows"]} == set(selected_ids)
    foreign_list = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs", headers=_auth_headers("other-user"))
    foreign_detail = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}", headers=_auth_headers("other-user"))
    assert foreign_list.status_code == foreign_detail.status_code == 404
    foreign_sourcing = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": selected_ids[:1], "source_mode": "provider_only",
    }, username="other-user")
    assert foreign_sourcing.status_code == 404
    listing = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs", headers=_auth_headers())
    assert listing.status_code == 200 and listing.json()["runs"][0]["run_id"] == run_id
    after_sha = hashlib.sha256((repository.workspace_root / workspace["tender_id"] / "source.xlsx").read_bytes()).hexdigest()
    assert before_sha == after_sha == workspace["source_sha256"]


def test_tender_sourcing_rejects_browser_facts_unknown_duplicate_nonitems_and_oversize(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    service = _fake_tender_sourcing_service(monkeypatch, main)
    item = next(row for row in workspace["rows"] if row["row_type"] == "item")
    section = {"source_row_id":"d"*32,"row_type":"section","excel_row":1}
    workspace_path = repository.workspace_root / workspace["tender_id"]
    metadata = repository._metadata(workspace_path)
    metadata["rows"].append(section)
    repository._atomic_json(workspace_path / "workspace.json", metadata)

    browser_facts = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": [item["source_row_id"]], "source_mode": "provider_only",
        "name": "browser override", "quantity": "999", "resource_code": "browser code", "rows": [{"article": "FAKE"}],
    })
    assert browser_facts.status_code == 422
    duplicate = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": [item["source_row_id"], item["source_row_id"]], "source_mode": "provider_only",
    })
    assert duplicate.status_code == 422
    unknown = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": ["f" * 32], "source_mode": "provider_only",
    })
    assert unknown.status_code == 400 and unknown.json()["detail"]["code"] == "TENDER_SOURCE_ROW_UNKNOWN"
    non_item = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": [section["source_row_id"]], "source_mode": "provider_only",
    })
    assert non_item.status_code == 400 and non_item.json()["detail"]["code"] == "TENDER_SOURCE_ROW_NOT_ITEM"
    oversized = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids": [f"{index:032x}" for index in range(501)], "source_mode": "provider_only",
    })
    assert oversized.status_code == 422
    empty = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[],"source_mode":"provider_only"})
    malformed_mode = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[item["source_row_id"]],"source_mode":"history_first"})
    malformed_limit = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[item["source_row_id"]],"source_mode":"provider_only","limit":0})
    malformed_limit_type = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[item["source_row_id"]],"source_mode":"provider_only","limit":"20"})
    assert empty.status_code == malformed_mode.status_code == malformed_limit.status_code == malformed_limit_type.status_code == 422
    invalid_provider = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[item["source_row_id"]],"source_mode":"provider_only","provider":"invalid-provider"})
    assert invalid_provider.status_code == 400
    assert service.calls == 0


def test_tender_sourcing_accepts_500_and_true_article_without_resource_code_leak(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    official = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    service = _fake_tender_sourcing_service(monkeypatch, main)
    source = next(row for row in official["rows"] if row["row_type"] == "item")
    assert source["article"] == "SKU-1" and source["resource_code"] == "R-001"
    accepted = _post_tender_sourcing(main, official["tender_id"], {
        "source_row_ids": [source["source_row_id"]], "source_mode": "provider_only",
    })
    assert accepted.status_code == 202
    assert _wait_tender_job(main, accepted.json()["id"])["status"] == "completed"
    from averon_import.services.sourcing.product_understanding import build_fallback_intent
    intent = build_fallback_intent(service.received[0])
    assert intent.article == "SKU-1"
    assert "R-001" not in intent.source_text

    large = _confirm_synthetic_tender(main, repository, tmp_path, count=500)
    large_items = [row for row in large["rows"] if row["row_type"] == "item"]
    assert len(large_items) == 500
    accepted_large = _post_tender_sourcing(main, large["tender_id"], {
        "source_row_ids": [row["source_row_id"] for row in large_items], "source_mode": "provider_only",
    })
    assert accepted_large.status_code == 202, accepted_large.text
    assert _wait_tender_job(main, accepted_large.json()["id"], timeout=8)["status"] == "completed"
    assert len(service.received) == 500


def test_tender_activity_lease_spans_queue_coalesces_and_releases(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    _fake_tender_sourcing_service(monkeypatch, main)
    started = threading.Event()
    release = threading.Event()
    blocker = main.job_service.submit(lambda progress: (started.set(), release.wait(5))[1], lane="sourcing", owner_id="blocker")
    assert started.wait(2)
    row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    request = {"source_row_ids": [row["source_row_id"]], "source_mode": "provider_only"}
    first = _post_tender_sourcing(main, workspace["tender_id"], request)
    second = _post_tender_sourcing(main, workspace["tender_id"], request)
    assert first.status_code == second.status_code == 202
    assert first.json()["id"] == second.json()["id"]
    assert repository.activity.active(workspace["tender_id"])
    tender_sourcing_runs = main.tender_sourcing_runs.list_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"])
    assert tender_sourcing_runs["runs"] == []
    deleted = _api_request(main.app, "DELETE", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers())
    assert deleted.status_code == 409
    release.set()
    _wait_tender_job(main, first.json()["id"])
    main.job_service.get_public(blocker.id)
    assert not repository.activity.active(workspace["tender_id"])
    final = _api_request(main.app, "DELETE", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers())
    assert final.status_code == 200


def test_tender_admission_rejection_releases_lease_without_ghost_run(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    _fake_tender_sourcing_service(monkeypatch, main)
    from averon_import.services.jobs import DOCUMENT_PROCESSING, SOURCING, JobService
    jobs = JobService(capacities={DOCUMENT_PROCESSING:(1,2), SOURCING:(1,0)})
    monkeypatch.setattr(main, "job_service", jobs)
    started = threading.Event()
    release = threading.Event()
    jobs.submit(lambda progress: (started.set(), release.wait(5))[1], lane=SOURCING, owner_id="blocker")
    assert started.wait(2)
    row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[row["source_row_id"]],"source_mode":"provider_only"})
    assert response.status_code == 409 and not repository.activity.active(workspace["tender_id"])
    release.set()
    jobs.executor.shutdown(wait=True)


def test_tender_history_modes_hold_captured_history_lease_and_use_existing_router(tender_api, tmp_path, monkeypatch):
    main, _repository = tender_api
    workspace = _confirm_synthetic_tender(main, _repository, tmp_path, count=1)
    service = _fake_tender_sourcing_service(monkeypatch, main, catalog_version="history-v1")
    monkeypatch.setattr(main.one_c_history_repository, "catalog_version", lambda: "history-v1")
    row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    from averon_import.services.jobs import SOURCING
    started = threading.Event()
    release = threading.Event()
    blocker = main.job_service.submit(lambda progress: (started.set(), release.wait(5))[1], lane=SOURCING, owner_id="history-blocker")
    assert started.wait(2)
    for index, mode in enumerate(("one_c_only", "one_c_then_provider")):
        response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[row["source_row_id"]],"source_mode":mode})
        assert response.status_code == 202, response.text
        if index == 0:
            assert main.one_c_history_activity.status()["active_sourcing_count"] == 1
            release.set()
        assert _wait_tender_job(main, response.json()["id"])["status"] == "completed"
        assert service.modes[-1].value == mode
        detail = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs/{_wait_tender_job(main, response.json()['id'])['result']['run_id']}", headers=_auth_headers()).json()
        assert detail["history_catalog_version"] == "history-v1"
        assert detail["rows"][0]["route"]["history_catalog_version"] == "history-v1"
        assert main.one_c_history_activity.status()["active_sourcing_count"] == 0
    response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[row["source_row_id"]],"source_mode":"provider_only"})
    assert response.status_code == 202
    _wait_tender_job(main, response.json()["id"])
    assert main.one_c_history_activity.status()["active_sourcing_count"] == 0
    main.job_service.get_public(blocker.id)


def test_tender_runs_recover_running_as_interrupted_are_immutable_and_bounded(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    path = repository.workspace_root / workspace["tender_id"]
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    running = main.tender_sourcing_runs.create_running(
        path, {"tender_id":workspace["tender_id"],"source_sha256":workspace["source_sha256"],"revision":workspace["revision"]},
        source_mode="provider_only",provider=None,selected_ids=[source["source_row_id"]],history_catalog_version=None,
    )
    complete_before_restart = main.tender_sourcing_runs.create_running(
        path, {"tender_id":workspace["tender_id"],"source_sha256":workspace["source_sha256"],"revision":workspace["revision"]},
        source_mode="provider_only",provider=None,selected_ids=[source["source_row_id"]],history_catalog_version=None,
    )
    main.tender_sourcing_runs.complete(path, complete_before_restart["run_id"],summary={"positions_total":1,"positions_processed":1,"positions_matched":1,"positions_review":0,"positions_without_offers":0},catalog_version="v",history_catalog_version=None,rows=[])
    recovered = TenderSourcingRunStore(repository)
    interrupted = recovered.get_public(path, workspace["tender_id"], running["run_id"])
    assert interrupted["status"] == "interrupted"
    assert "Подбор был прерван перезапуском сервера" in interrupted["failure"]["message"]
    assert recovered.get_public(path, workspace["tender_id"], complete_before_restart["run_id"])["status"] == "completed"
    completed_ids = []
    for _ in range(6):
        run = recovered.create_running(
            path, {"tender_id":workspace["tender_id"],"source_sha256":workspace["source_sha256"],"revision":workspace["revision"]},
            source_mode="provider_only",provider=None,selected_ids=[source["source_row_id"]],history_catalog_version=None,
        )
        completed_ids.append(run["run_id"])
        recovered.complete(path, run["run_id"],summary={"positions_total":1,"positions_processed":1,"positions_matched":0,"positions_review":0,"positions_without_offers":1},catalog_version="v",history_catalog_version=None,rows=[])
    newest = recovered.get_public(path, workspace["tender_id"], completed_ids[-1])
    with pytest.raises(TenderWorkspaceError):
        recovered.complete(path, newest["run_id"],summary={},catalog_version=None,history_catalog_version=None,rows=[])
    files = list((path / "runs").glob("*.json"))
    assert len(files) == 5
    active = recovered.create_running(
        path, {"tender_id":workspace["tender_id"],"source_sha256":workspace["source_sha256"],"revision":workspace["revision"]},
        source_mode="provider_only",provider=None,selected_ids=[source["source_row_id"]],history_catalog_version=None,
    )
    assert recovered.get_public(path, workspace["tender_id"], active["run_id"])["status"] == "running"
    assert len(list((path / "runs").glob("*.json"))) == 6
    oversized = path / "runs" / f"{'e' * 32}.json"
    with pytest.raises(TenderWorkspaceError, match="превышает допустимый объём"):
        recovered._atomic_write(oversized, {"payload":"x" * (1024 * 1024)})


def test_tender_result_correlation_fence_rejects_missing_duplicate_and_unknown_ids():
    source_rows = [{"source_row_id":"a"*32,"excel_row":12},{"source_row_id":"b"*32,"excel_row":13}]
    def result(*ids):
        return {"results":[{"intent":{"source_row_id":row_id}} for row_id in ids]}
    with pytest.raises(TenderWorkspaceError, match="соответствия строк"):
        canonical_tender_projection(result("a"*32),source_rows,["a"*32,"b"*32])
    with pytest.raises(TenderWorkspaceError, match="соответствия строк"):
        canonical_tender_projection(result("a"*32,"a"*32),source_rows,["a"*32,"b"*32])
    with pytest.raises(TenderWorkspaceError, match="соответствия строк"):
        canonical_tender_projection(result("a"*32,"c"*32),source_rows,["a"*32,"b"*32])


def test_tender_runner_failure_releases_lease_and_records_safe_failure(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    service = _fake_tender_sourcing_service(monkeypatch, main)
    def fail_search(*_args, **_kwargs):
        raise RuntimeError("provider secret/token must not be persisted")
    monkeypatch.setattr(service, "search_project", fail_search)
    row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[row["source_row_id"]],"source_mode":"provider_only"})
    assert response.status_code == 202
    job = _wait_tender_job(main, response.json()["id"])
    assert job["status"] == "failed"
    assert not repository.activity.active(workspace["tender_id"])
    run = main.tender_sourcing_runs.list_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"])["runs"][0]
    encoded = json.dumps(run, ensure_ascii=False).casefold()
    assert run["status"] == "failed"
    assert run["failure"]["code"] == "SOURCING_FAILED"
    assert "secret" not in encoded and "token" not in encoded


def test_tender_queue_expiry_has_no_durable_run_and_releases_lease(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    _fake_tender_sourcing_service(monkeypatch, main)
    from averon_import.services.jobs import DOCUMENT_PROCESSING, SOURCING, JobService
    jobs = JobService(capacities={DOCUMENT_PROCESSING:(1,2), SOURCING:(1,2)}, queue_wait_seconds=5)
    monkeypatch.setattr(main, "job_service", jobs)
    started = threading.Event()
    release = threading.Event()
    blocker = jobs.submit(lambda progress: (started.set(), release.wait(5))[1], lane=SOURCING, owner_id="blocker")
    assert started.wait(2)
    row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[row["source_row_id"]],"source_mode":"provider_only"})
    assert response.status_code == 202
    queued = jobs.get(response.json()["id"], owner_id="tender-user")
    queued._created_monotonic = time.monotonic() - 6
    job = jobs.get_public(response.json()["id"], owner_id="tender-user")
    assert job["status"] == "expired"
    assert not repository.activity.active(workspace["tender_id"])
    assert main.tender_sourcing_runs.list_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"])["runs"] == []
    release.set()
    jobs.get_public(blocker.id)
    jobs.executor.shutdown(wait=True)


def test_tender_queue_revalidates_source_and_canonical_rows_before_provider_call(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=1)
    service = _fake_tender_sourcing_service(monkeypatch, main)
    from averon_import.services.jobs import SOURCING
    started = threading.Event()
    release = threading.Event()
    blocker = main.job_service.submit(lambda progress: (started.set(), release.wait(5))[1], lane=SOURCING, owner_id="integrity-blocker")
    assert started.wait(2)
    item = next(row for row in workspace["rows"] if row["row_type"] == "item")
    response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[item["source_row_id"]],"source_mode":"provider_only"})
    assert response.status_code == 202
    path = repository.workspace_root / workspace["tender_id"]
    metadata = repository._metadata(path)
    metadata["rows"][0]["name"] = "synthetic tamper before execution"
    repository._atomic_json(path / "workspace.json", metadata)
    release.set()
    job = _wait_tender_job(main, response.json()["id"])
    assert job["status"] == "failed"
    assert service.calls == 0
    run = main.tender_sourcing_runs.list_public(path, workspace["tender_id"])["runs"][0]
    detail = main.tender_sourcing_runs.get_public(path, workspace["tender_id"], run["run_id"])
    assert run["status"] == "failed" and detail["source_sha256"] == workspace["source_sha256"]
    main.job_service.get_public(blocker.id)


def test_tender_runtime_is_captured_at_admission_and_job_owners_do_not_coalesce(tender_api, tmp_path, monkeypatch):
    main, _repository = tender_api
    workspace = _confirm_synthetic_tender(main, _repository, tmp_path, count=1)
    admitted_service = _fake_tender_sourcing_service(monkeypatch, main)
    from averon_import.services.jobs import SOURCING, JobService
    started = threading.Event()
    release = threading.Event()
    main.job_service.submit(lambda progress: (started.set(), release.wait(5))[1], lane=SOURCING, owner_id="runtime-blocker")
    assert started.wait(2)
    row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    response = _post_tender_sourcing(main, workspace["tender_id"], {"source_row_ids":[row["source_row_id"]],"source_mode":"provider_only"})
    assert response.status_code == 202
    replacement = _fake_tender_sourcing_service(monkeypatch, main)
    release.set()
    job = _wait_tender_job(main, response.json()["id"])
    assert job["status"] == "completed"
    assert admitted_service.calls == 1 and replacement.calls == 0

    jobs = JobService()
    hold = threading.Event()
    entered = threading.Event()
    first = jobs.submit(lambda progress: (entered.set(), hold.wait(5))[1], lane=SOURCING, owner_id="owner-a", dedupe_key="same")
    assert entered.wait(2)
    second = jobs.submit(lambda progress: "second", lane=SOURCING, owner_id="owner-b", dedupe_key="same")
    assert first.id != second.id
    hold.set()
    jobs.executor.shutdown(wait=True)


def _project_result_with_offer(source_row, *, provider, provenance, source_mode="provider_only"):
    from averon_import.services.sourcing.models import (
        MatchDecision, MatchResult, Offer, ProjectSourcingResult,
        SourcingResult, SourcingRouteMetadata, SourcingSourceMode,
    )
    from averon_import.services.sourcing.product_understanding import build_fallback_intent

    adapted = TenderSourcingRowAdapter.convert(source_row)
    intent = build_fallback_intent(adapted)
    source_item_id = str(provenance.get("history_item_id") or "synthetic-source-item") if provider == "one_c_history" else "synthetic-source-item"
    offer = Offer(
        offer_id=f"{provider}:synthetic-offer",
        provider=provider,
        source_item_id=source_item_id,
        title="Synthetic replacement assembly with bounded provider title",
        article="SYN-100",
        manufacturer="Synthetic manufacturer",
        brand="Synthetic brand",
        price=Decimal("1234.5600"),
        currency="RUB",
        price_unit="шт",
        availability=True,
        availability_text="Synthetic stock status",
        data_provenance=provenance,
    )
    mode = SourcingSourceMode(source_mode)
    if mode == SourcingSourceMode.ONE_C_ONLY:
        route = SourcingRouteMetadata(
            source_mode=mode,
            final_source_kind="historical_purchase",
            history_outcome="SAFE_MATCH",
            history_catalog_version="history-snapshot-v9",
            history_selected_event_id=str(provenance.get("selected_event_id") or ""),
            history_purchase_date=str(provenance.get("purchase_date") or ""),
        )
    else:
        route = SourcingRouteMetadata(source_mode=mode, final_source_kind="provider")
    match = MatchResult(offer=offer, decision=MatchDecision.MATCH, rank=1, matched_attributes=["name"])
    item = SourcingResult(intent=intent, recommended_offer=offer, match_results=[match], route=route)
    return ProjectSourcingResult(
        positions_total=1,
        positions_processed=1,
        positions_matched=1,
        source_mode=mode,
        results=[item],
    )


def _persist_project_result(store, repository, workspace, result, *, source_row, source_mode):
    workspace_path = repository.workspace_root / workspace["tender_id"]
    run = store.create_running(
        workspace_path,
        workspace,
        source_mode=source_mode,
        provider=None,
        selected_ids=[source_row["source_row_id"]],
        history_catalog_version="history-snapshot-v9" if source_mode == "one_c_only" else None,
    )
    payload = result.model_dump(mode="json")
    rows = canonical_tender_projection(payload, [source_row], [source_row["source_row_id"]])
    store.complete(
        workspace_path,
        run["run_id"],
        summary={
            "positions_total": 1, "positions_processed": 1, "positions_matched": 1,
            "positions_review": 0, "positions_without_offers": 0,
        },
        catalog_version="synthetic-provider-catalog-v4",
        history_catalog_version="history-snapshot-v9" if source_mode == "one_c_only" else None,
        rows=rows,
    )
    return run["run_id"]


def _persist_price_export_run(main, repository, workspace, source_row, *, provider="etm_ipro", provenance=None, history=False):
    if provenance is None:
        provenance = {
            "source": provider,
            "catalog_version": "etm-catalog-v4",
            "source_item_id": "synthetic-source-item",
            "price_field": "pricewnds",
        }
    mode = "one_c_only" if history else "provider_only"
    provider_key = "one_c_history" if history else provider
    result = _project_result_with_offer(source_row, provider=provider_key, provenance=provenance, source_mode=mode)
    if history:
        from averon_import.services.sourcing.models import HistorySafeMatchBasis
        item = result.results[0]
        route = item.route.model_copy(update={
            "history_safe_basis": HistorySafeMatchBasis.EXACT_ARTICLE,
            "history_catalog_version": "history-snapshot-v9",
            "history_selected_event_id": str(provenance["selected_event_id"]),
            "history_purchase_date": str(provenance["purchase_date"]),
            "history_outcome": "SAFE_MATCH",
        })
        result = result.model_copy(update={"results": [item.model_copy(update={"route": route})]})
    return _persist_project_result(
        main.tender_sourcing_runs, repository, workspace, result,
        source_row=source_row, source_mode=mode,
    )


def _persist_mixed_price_export_run(main, repository, workspace, source_rows):
    from averon_import.services.sourcing.models import (
        HistorySafeMatchBasis, ProjectSourcingResult, SourcingSourceMode,
    )

    history_provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase",
        "snapshot_version":"history-snapshot-v9", "history_item_id":"history-item-1",
        "selected_event_id":"history-event-1", "purchase_date":"2025-01-24",
        "price_basis":"gross_including_vat", "effective_unit_price_gross":"1234.5600",
        "currency_basis":"RUB", "unit_family":"count",
    }
    historical = _project_result_with_offer(
        source_rows[0], provider="one_c_history", provenance=history_provenance,
        source_mode="one_c_only",
    ).results[0]
    historical_route = historical.route.model_copy(update={
        "history_safe_basis": HistorySafeMatchBasis.EXACT_ARTICLE,
        "history_catalog_version":"history-snapshot-v9",
        "history_selected_event_id":"history-event-1",
        "history_purchase_date":"2025-01-24",
        "history_outcome":"SAFE_MATCH",
    })
    historical = historical.model_copy(update={"route":historical_route})
    provider_result = _project_result_with_offer(
        source_rows[1], provider="etm_ipro", provenance={
            "source":"etm_ipro", "catalog_version":"etm-catalog-v4",
            "source_item_id":"synthetic-source-item", "price_field":"pricewnds",
        }, source_mode="one_c_then_provider",
    ).results[0]
    result = ProjectSourcingResult(
        positions_total=2, positions_processed=2, positions_matched=2,
        source_mode=SourcingSourceMode.ONE_C_THEN_PROVIDER,
        results=[historical, provider_result],
    )
    selected_ids = [row["source_row_id"] for row in source_rows]
    workspace_path = repository.workspace_root / workspace["tender_id"]
    running = main.tender_sourcing_runs.create_running(
        workspace_path, workspace, source_mode="one_c_then_provider", provider="etm_ipro",
        selected_ids=selected_ids, history_catalog_version="history-snapshot-v9",
    )
    projected = canonical_tender_projection(result.model_dump(mode="json"), source_rows, selected_ids)
    main.tender_sourcing_runs.complete(
        workspace_path, running["run_id"], summary={
            "positions_total":2, "positions_processed":2, "positions_matched":2,
            "positions_review":0, "positions_without_offers":0,
        }, catalog_version="etm-catalog-v4", history_catalog_version="history-snapshot-v9",
        rows=projected,
    )
    return running["run_id"]


def test_provider_only_real_unrouted_sourcing_shape_is_exportable_through_api(tender_api, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from averon_import.services.sourcing.models import (
        MatchDecision, MatchResult, Offer, ProjectSourcingResult, SourcingResult, SourcingSourceMode,
    )
    from averon_import.services.sourcing.product_understanding import build_fallback_intent

    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    provenance = {
        "source":"etm_ipro", "catalog_version":"fake-etm-catalog-v1",
        "source_item_id":"fake-etm-item-1", "price_field":"pricewnds",
    }

    class ProviderOnlyService:
        def provider(self, provider_key=None):
            assert provider_key == "etm_ipro"
            return SimpleNamespace(key="etm_ipro", label="Fake ETM")

        def search_project(self, rows, *, progress=None, **_kwargs):
            results = []
            for row in rows:
                offer = Offer(
                    offer_id="etm_ipro:fake-offer-1", provider="etm_ipro",
                    source_item_id="fake-etm-item-1", title="Synthetic provider offer",
                    price=Decimal("125.50"), currency="RUB", price_unit="шт",
                    data_provenance=provenance,
                )
                match = MatchResult(offer=offer, decision=MatchDecision.MATCH, rank=1)
                # This is the actual provider_only shape: the normal sourcing path
                # has no routed OneC metadata to persist for provider results.
                results.append(SourcingResult(
                    intent=build_fallback_intent(row), recommended_offer=offer,
                    match_results=[match],
                ))
            if progress:
                progress(len(rows), len(rows), "test")
            return ProjectSourcingResult(
                positions_total=len(rows), positions_processed=len(rows),
                positions_matched=len(rows), source_mode=SourcingSourceMode.PROVIDER_ONLY,
                provider_key="etm_ipro", provider_label="Fake ETM", results=results,
            )

    service = ProviderOnlyService()
    monkeypatch.setattr(main, "sourcing_service", service)
    monkeypatch.setattr(main, "sourcing_runtime", SimpleNamespace(service=service))
    started = _post_tender_sourcing(main, workspace["tender_id"], {
        "source_row_ids":[source["source_row_id"]], "source_mode":"provider_only", "provider":"etm_ipro",
    })
    assert started.status_code == 202, started.text
    sourcing_job = _wait_tender_job(main, started.json()["id"])
    assert sourcing_job["status"] == "completed", sourcing_job
    run_id = sourcing_job["result"]["run_id"]
    workspace_path = repository.workspace_root / workspace["tender_id"]
    durable_run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    canonical = durable_run["rows"][0]
    assert canonical["route"] == {}
    assert canonical["provider_source"]["kind"] is None

    export = _api_request(
        main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":false,"allow_partial":false}',
    )
    assert export.status_code == 202, export.text
    export_job = _wait_tender_job(main, export.json()["id"])
    assert export_job["status"] == "completed", export_job.get("error") or export_job
    assert export_job["result"]["priced_count"] == 1
    file_path, _record, release = main.tender_price_exports.acquire_download(
        workspace_path, workspace["tender_id"], export_job["result"]["export_id"], "username:tender-user",
    )
    try:
        workbook = load_workbook(file_path, data_only=False)
        sheet = workbook[TEMPLATE_SHEET]
        assert Decimal(str(sheet["H2"].value)) == Decimal("125.500000")
        assert Decimal(str(sheet["I2"].value)) == Decimal("251.00")
        workbook.close()
    finally:
        release()

    bad_provenance_run = deepcopy(durable_run)
    bad_provenance_run["rows"][0]["price_provenance"]["price_field"] = "price"
    bad_provenance = main.tender_price_resolver.resolve_run(
        workspace, bad_provenance_run, tender_id=workspace["tender_id"], run_id=run_id,
        include_historical_prices=False,
    )[0]
    assert not bad_provenance.eligible and bad_provenance.reason_code == "PRICE_BASIS_UNPROVEN"

    missing_route_mixed_run = deepcopy(durable_run)
    missing_route_mixed_run["source_mode"] = "one_c_then_provider"
    missing_route_decision = main.tender_price_resolver.resolve_run(
        workspace, missing_route_mixed_run, tender_id=workspace["tender_id"], run_id=run_id,
        include_historical_prices=False,
    )[0]
    assert not missing_route_decision.eligible
    assert missing_route_decision.reason_code == "MATCH_NOT_EXPORTABLE"


def test_mixed_history_export_can_exclude_history_and_continue_with_provider(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=2)
    source_rows = [row for row in workspace["rows"] if row["row_type"] == "item"]
    assert len(source_rows) == 2
    run_id = _persist_mixed_price_export_run(main, repository, workspace, source_rows)
    endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export"

    def post(payload):
        return _api_request(
            main.app, "POST", endpoint,
            headers={**_auth_headers(), "Content-Type":"application/json"},
            body=json.dumps(payload).encode("utf-8"),
        )

    not_decided = post({"include_historical_prices":False, "allow_partial":False})
    assert not_decided.status_code == 409
    assert not_decided.json()["detail"]["code"] == "TENDER_EXPORT_HISTORICAL_CONFIRMATION_REQUIRED"

    excluded_but_not_partial = post({
        "include_historical_prices":False,
        "historical_decision_confirmed":True,
        "allow_partial":False,
    })
    assert excluded_but_not_partial.status_code == 409
    detail = excluded_but_not_partial.json()["detail"]
    assert detail["code"] == "TENDER_EXPORT_PARTIAL_CONFIRMATION_REQUIRED"
    assert detail["summary"]["priced_count"] == 1
    assert detail["summary"]["reason_counts"]["HISTORICAL_PRICE_NOT_INCLUDED"] == 1

    excluded = post({
        "include_historical_prices":False,
        "historical_decision_confirmed":True,
        "allow_partial":True,
    })
    assert excluded.status_code == 202, excluded.text
    excluded_job = _wait_tender_job(main, excluded.json()["id"])
    assert excluded_job["status"] == "completed", excluded_job.get("error") or excluded_job
    excluded_record = excluded_job["result"]
    assert excluded_record["priced_count"] == 1
    assert excluded_record["historical_count"] == 0
    assert excluded_record["reason_counts"]["HISTORICAL_PRICE_NOT_INCLUDED"] == 1
    workspace_path = repository.workspace_root / workspace["tender_id"]
    file_path, _record, release = main.tender_price_exports.acquire_download(
        workspace_path, workspace["tender_id"], excluded_record["export_id"], "username:tender-user",
    )
    try:
        workbook = load_workbook(file_path, data_only=False)
        sheet = workbook[TEMPLATE_SHEET]
        assert sheet["H2"].value is None and sheet["I2"].value is None
        assert isinstance(sheet["H3"].value, (int, float, Decimal))
        assert isinstance(sheet["I3"].value, (int, float, Decimal))
        workbook.close()
    finally:
        release()

    included = post({"include_historical_prices":True, "allow_partial":False})
    assert included.status_code == 202, included.text
    included_job = _wait_tender_job(main, included.json()["id"])
    assert included_job["status"] == "completed", included_job
    included_record = included_job["result"]
    assert included_record["priced_count"] == 2
    assert included_record["historical_count"] == 1
    included_path, _record, release = main.tender_price_exports.acquire_download(
        workspace_path, workspace["tender_id"], included_record["export_id"], "username:tender-user",
    )
    try:
        workbook = load_workbook(included_path, data_only=False)
        sheet = workbook[TEMPLATE_SHEET]
        assert isinstance(sheet["H2"].value, (int, float, Decimal))
        assert isinstance(sheet["I2"].value, (int, float, Decimal))
        assert isinstance(sheet["H3"].value, (int, float, Decimal))
        assert isinstance(sheet["I3"].value, (int, float, Decimal))
        assert sheet["H2"].comment and "2025-01-24" in sheet["H2"].comment.text
        assert sheet["I2"].comment and "2025-01-24" in sheet["I2"].comment.text
        workbook.close()
    finally:
        release()


def test_tender_price_export_target_collision_uses_semantic_empty_cells(tmp_path):
    from averon_import.services.manual_tenders.price_export import _check_target_collision
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation

    def make_book(path, mutate=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet["F1"] = ""
        sheet["G1"] = ""
        sheet["F2"].fill = PatternFill(fill_type="solid", fgColor="FFEEEEEE")
        if mutate:
            mutate(sheet)
        workbook.save(path)
        workbook.close()
        return path

    workspace = {"logical_right_edge": 5, "source_manifest": {}}
    clean = make_book(tmp_path / "semantic-empty.xlsx")
    with zipfile.ZipFile(clean) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    assert b'r="F1"' in sheet_xml and b'r="G1"' in sheet_xml
    assert _check_target_collision(clean, workspace, "Sheet") == (6, 7)

    for name, mutate in (
        ("constant", lambda sheet: setattr(sheet["F3"], "value", "occupied")),
        ("formula-empty", lambda sheet: setattr(sheet["F3"], "value", '=""')),
        ("comment", lambda sheet: setattr(sheet["G3"], "comment", Comment("note", "tester"))),
        ("hyperlink", lambda sheet: setattr(sheet["F3"], "hyperlink", "https://example.invalid/item")),
        ("merge", lambda sheet: sheet.merge_cells("F1:G1")),
        ("filter", lambda sheet: setattr(sheet.auto_filter, "ref", "A1:G5")),
        ("validation", lambda sheet: (lambda validation: (validation.add("F3"), sheet.add_data_validation(validation)))(DataValidation(type="whole", operator="greaterThan", formula1="0"))),
        ("conditional-format", lambda sheet: sheet.conditional_formatting.add("G3", CellIsRule(operator="greaterThan", formula=["0"]))),
    ):
        path = make_book(tmp_path / f"occupied-{name}.xlsx", mutate)
        with pytest.raises(TenderWorkspaceError) as error:
            _check_target_collision(path, workspace, "Sheet")
        assert error.value.code == "TENDER_EXPORT_TARGET_OCCUPIED"


def _confirm_ivy_history_tender(main, repository, tmp_path, *, source_article="", source_name="Плющ искусственный", source_model="", source_manufacturer="", source_unit="шт"):
    path = tmp_path / f"ivy-{uuid.uuid4().hex}.xlsx"
    _official(path, include_required=False)
    workbook = load_workbook(path)
    sheet = workbook[TEMPLATE_SHEET]
    sheet["A2"] = "R-IVY"
    sheet["B2"] = source_name
    sheet["C2"] = source_unit
    sheet["D2"] = 2
    sheet["E2"] = source_article or None
    sheet["F2"] = source_manufacturer or None
    sheet["G2"] = source_model or None
    workbook.save(path)
    workbook.close()
    payload = path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    preview = repository.reserve_preview("username:tender-user", path.name, len(payload), digest)
    source_path = repository.write_preview_source(preview["preview_id"], payload)
    analysis = main.tender_parser.parse(source_path, tender_id=preview["preview_id"])
    repository.update_preview(
        preview["preview_id"], "username:tender-user", status="ready",
        parser_version=analysis["parser_version"], mapping=analysis["mapping"], analysis=analysis,
    )
    return repository.confirm(preview["preview_id"], "username:tender-user")


def _persist_history_review_run(
    main, repository, workspace, *, candidates=1, classification="EXACT_NAME_UNIT",
    offer_overrides=None, provenance_overrides=None, offer_attributes=None, candidate_specs=None,
):
    from averon_import.services.sourcing.models import (
        HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent,
        ProjectSourcingResult, SourcingResult, SourcingRouteMetadata, SourcingSourceMode,
    )

    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    version = "history-snapshot-d3"
    intent = ProductIntent(
        source_row_id=source["source_row_id"], source_text=source["name"],
        normalized_name=source["name"], unit=source["raw_unit"], quantity=str(source["quantity"]),
        article=source.get("article", ""), manufacturer=source.get("manufacturer", ""),
        model=source.get("model", ""),
    )
    offers = []
    matches = []
    for index in range(candidates):
        spec = candidate_specs[index] if candidate_specs and index < len(candidate_specs) else {}
        item_id = f"history-item-{index + 1}"
        event_id = f"history-event-{index + 1}"
        provenance = {
            "source":"one_c_history", "source_kind":"historical_purchase",
            "snapshot_version":version, "history_item_id":item_id,
            "selected_event_id":event_id, "purchase_date":"2025-04-16",
            "price_basis":"gross_including_vat", "effective_unit_price_gross":Decimal("123.4500"),
            "currency_basis":"company_default", "unit_family":"piece",
        }
        provenance.update(provenance_overrides or {})
        provenance.update(spec.get("provenance_overrides") or {})
        offer_values = {
            "title":source["name"], "article":"", "manufacturer":"", "brand":"",
            "price":Decimal("123.4500"), "currency":"RUB", "price_unit":"шт",
        }
        offer_values.update(offer_overrides or {})
        offer_values.update(spec.get("offer_overrides") or {})
        offer = Offer(
            offer_id=f"one_c_history:{item_id}", provider="one_c_history", source_item_id=item_id,
            **offer_values,
            attributes={**(offer_attributes or {}), **(spec.get("offer_attributes") or {})},
            data_provenance=provenance,
            history_retrieval_classification=classification,
        )
        offers.append(offer)
        matches.append(MatchResult(
            offer=offer, decision=MatchDecision(spec.get("decision", "REVIEW")), rank=index + 1,
            matched_attributes=spec.get("matched_attributes", ["name", "unit"]),
            supporting_attributes=spec.get("supporting_attributes", []),
            conflicting_attributes=spec.get("conflicting_attributes", []),
            missing_attributes=spec.get("missing_attributes", []),
        ))
    route = SourcingRouteMetadata(
        source_mode=SourcingSourceMode.ONE_C_ONLY, final_source_kind="history_review",
        history_outcome="REVIEW", history_catalog_version=version,
        history_reason_code="ambiguous_exact_name_identity" if candidates > 1 else "exact_name_unit_needs_review",
        history_candidate_count=candidates,
    )
    item = SourcingResult(intent=intent, offers=offers, match_results=matches, route=route)
    result = ProjectSourcingResult(
        positions_total=1, positions_processed=1, positions_review=1,
        source_mode=SourcingSourceMode.ONE_C_ONLY, results=[item],
    )
    workspace_path = repository.workspace_root / workspace["tender_id"]
    selected_ids = [source["source_row_id"]]
    running = main.tender_sourcing_runs.create_running(
        workspace_path, workspace, source_mode="one_c_only", provider=None,
        selected_ids=selected_ids, history_catalog_version=version,
    )
    projection = canonical_tender_projection(result.model_dump(mode="json"), [source], selected_ids)
    main.tender_sourcing_runs.complete(
        workspace_path, running["run_id"],
        summary={"positions_total":1, "positions_processed":1, "positions_matched":0, "positions_review":1, "positions_without_offers":0},
        catalog_version=version, history_catalog_version=version, rows=projection,
    )
    return running["run_id"]


def test_d4a_normalized_history_confirmation_survives_restart_export_and_revoke(tender_api, tmp_path, monkeypatch):
    from averon_import.services.sourcing.history_identity import (
        HISTORY_IDENTITY_NORMALIZER_REVISION,
        history_name_signature_digest,
    )
    from averon_import.services.manual_tenders.history_decisions import TenderHistoryDecisionStore

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace,
        classification="NORMALIZED_NAME_UNIT",
        offer_overrides={"title":"Искусственный плющ"},
        provenance_overrides={
            "normalizer_revision":HISTORY_IDENTITY_NORMALIZER_REVISION,
            "normalized_name_signature":history_name_signature_digest(source["name"]),
        },
    )
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    assert candidate["offer"]["retrieval_classification"] == "NORMALIZED_NAME_UNIT"
    assert candidate["confirmation_basis"] == "NORMALIZED_CONFIRMATION"
    assert candidate["confirmable"] is True

    run = main.tender_sourcing_runs.get_public(
        repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id,
    )
    assert run["rows"][0]["route"]["history_outcome"] == "REVIEW"
    assert run["rows"][0]["route"]["history_safe_basis"] is None

    confirmed_response = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE",
            "source_row_id":snapshot["rows"][0]["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"],
            "expected_revision":0,
        }).encode(),
    )
    assert confirmed_response.status_code == 200, confirmed_response.text
    confirmed = confirmed_response.json()
    assert confirmed["events"][0]["confirmation_basis"] == "NORMALIZED_CONFIRMATION"
    assert confirmed["effective"][source["source_row_id"]]["confirmation_basis"] == "NORMALIZED_CONFIRMATION"

    monkeypatch.setattr(main, "tender_history_decisions", TenderHistoryDecisionStore(repository))
    restored = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    assert restored["effective"][source["source_row_id"]]["confirmation_basis"] == "NORMALIZED_CONFIRMATION"

    export_endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export"
    included = _api_request(
        main.app, "POST", export_endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert included.status_code == 202, included.text
    exported = _wait_tender_job(main, included.json()["id"])
    assert exported["status"] == "completed", exported
    assert exported["result"]["automatic_historical_count"] == 0
    assert exported["result"]["human_confirmed_historical_count"] == 1

    decision_id = restored["effective"][source["source_row_id"]]["decision_id"]
    revoked = _api_request(
        main.app, "POST", f"{base}/{decision_id}/revoke",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"decision":"REVOKE_HISTORY_CONFIRMATION","expected_revision":1}',
    )
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["effective"] == {}


@pytest.mark.parametrize("provenance_override", [
    {"normalizer_revision":"normalized-name-unit-old"},
    {"normalized_name_signature":"f" * 64},
])
def test_d4a_normalized_history_candidate_recomputes_server_evidence(tender_api, tmp_path, provenance_override):
    from averon_import.services.sourcing.history_identity import (
        HISTORY_IDENTITY_NORMALIZER_REVISION,
        history_name_signature_digest,
    )

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    provenance = {
        "normalizer_revision":HISTORY_IDENTITY_NORMALIZER_REVISION,
        "normalized_name_signature":history_name_signature_digest(source["name"]),
    }
    provenance.update(provenance_override)
    run_id = _persist_history_review_run(
        main, repository, workspace,
        classification="NORMALIZED_NAME_UNIT",
        offer_overrides={"title":"Искусственный плющ"},
        provenance_overrides=provenance,
    )
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    assert snapshot["rows"][0]["candidates"][0]["confirmable"] is False
    response = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE",
            "source_row_id":source["source_row_id"],
            "candidate_offer_id":"one_c_history:history-item-1",
            "expected_revision":0,
        }).encode(),
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_CONFIRMABLE"


@pytest.mark.parametrize(("classification", "source_article", "history_characteristic", "expected"), [
    ("NORMALIZED_NAME_UNIT", "", "25-40", True),
    ("NORMALIZED_NAME_UNIT", "", "25-60", False),
    ("EXACT_NAME_UNIT", "", "25-60", False),
    ("EXACT_ARTICLE", "SKU-42", "25-60", False),
    ("NORMALIZED_NAME_UNIT", "", "", True),
])
def test_d4a_model_characteristic_candidate_gate_is_consistent(tender_api, tmp_path, monkeypatch, classification, source_article, history_characteristic, expected):
    from averon_import.services.sourcing.history_identity import (
        HISTORY_IDENTITY_NORMALIZER_REVISION,
        history_name_signature_digest,
    )

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(
        main, repository, tmp_path, source_article=source_article,
        source_name="Насос циркуляционный", source_model="25-40",
    )
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    normalized = classification == "NORMALIZED_NAME_UNIT"
    provenance = ({
        "normalizer_revision":HISTORY_IDENTITY_NORMALIZER_REVISION,
        "normalized_name_signature":history_name_signature_digest(source["name"]),
    } if normalized else {})
    offer_overrides = {
        "title":"Циркуляционный насос" if normalized else source["name"],
        "article":source_article,
    }
    run_id = _persist_history_review_run(
        main, repository, workspace, classification=classification,
        offer_overrides=offer_overrides, provenance_overrides=provenance,
        offer_attributes={"characteristic":history_characteristic} if history_characteristic else {},
    )
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    assert candidate["offer"]["history_characteristic"] == history_characteristic
    assert candidate["confirmable"] is expected
    if not expected:
        assert candidate["reason_code"] == "HISTORY_CANDIDATE_SOURCE_CONFLICT"
        assert candidate["evidence_fingerprint"] is None
        response = _api_request(
            main.app, "POST", base,
            headers={**_auth_headers(), "Content-Type":"application/json"},
            body=json.dumps({
                "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
                "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
            }).encode(),
        )
        assert response.status_code == 409
        assert _api_request(main.app, "GET", base, headers=_auth_headers()).json()["events"] == []
    elif classification == "NORMALIZED_NAME_UNIT" and history_characteristic:
        from averon_import.services.manual_tenders.history_decisions import TenderHistoryDecisionStore
        monkeypatch.setattr(main, "tender_history_decisions", TenderHistoryDecisionStore(repository))
        restarted = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
        assert restarted["rows"][0]["candidates"][0]["offer"]["history_characteristic"] == "25-40"
        assert restarted["rows"][0]["candidates"][0]["confirmable"] is True
        assert restarted["rows"][0]["candidates"][0]["evidence_fingerprint"] == candidate["evidence_fingerprint"]


def test_d4a_normalized_fingerprint_binds_history_characteristic(tender_api, tmp_path):
    import copy
    from averon_import.services.sourcing.history_identity import (
        HISTORY_IDENTITY_NORMALIZER_REVISION,
        history_name_signature_digest,
    )

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace, classification="NORMALIZED_NAME_UNIT",
        offer_overrides={"title":"Искусственный плющ"},
        offer_attributes={"characteristic":"25-40"},
        provenance_overrides={
            "normalizer_revision":HISTORY_IDENTITY_NORMALIZER_REVISION,
            "normalized_name_signature":history_name_signature_digest(source["name"]),
        },
    )
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    before = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = before["rows"][0]["candidates"][0]
    assert candidate["confirmable"] is True
    confirm = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
        }).encode(),
    )
    assert confirm.status_code == 200
    workspace_path = repository.workspace_root / workspace["tender_id"]
    run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    tampered = copy.deepcopy(run)
    tampered["rows"][0]["history_review_candidates"][0]["offer"]["history_characteristic"] = "25-41"
    from averon_import.services.manual_tenders.history_decisions import TenderHistoryDecisionStore
    restarted_store = TenderHistoryDecisionStore(repository)
    changed = restarted_store.get_snapshot(workspace_path, workspace, tampered, workspace["tender_id"], run_id)
    changed_candidate = changed["rows"][0]["candidates"][0]
    assert changed_candidate["confirmable"] is True
    assert changed_candidate["evidence_fingerprint"] != candidate["evidence_fingerprint"]
    assert changed["effective"] == {}


def test_d4a_export_rechecks_explicit_history_model_conflict(tender_api, tmp_path):
    import copy
    from averon_import.services.manual_tenders.price_export import TenderPriceResolver

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(
        main, repository, tmp_path, source_name="Насос циркуляционный", source_model="25-40",
    )
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace,
        offer_overrides={"title":source["name"]},
        offer_attributes={"characteristic":"25-40"},
    )
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    confirmed = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
        }).encode(),
    )
    assert confirmed.status_code == 200
    workspace_path = repository.workspace_root / workspace["tender_id"]
    run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    canonical = next(row for row in run["rows"] if row["source_row_id"] == source["source_row_id"])
    forged_confirmation = copy.deepcopy(confirmed.json()["effective"][source["source_row_id"]])
    forged_confirmation["candidate"]["offer"]["history_characteristic"] = "25-60"
    decision = TenderPriceResolver().resolve_human_history(
        source, canonical, run, forged_confirmation, include_historical_prices=True,
    )
    assert decision.eligible is False
    assert decision.source_unit_price is None
    assert decision.reason_code == "HISTORY_CANDIDATE_SOURCE_CONFLICT"


def test_d4a_exact_fingerprint_remains_compatible_with_legacy_candidate_contract():
    import copy
    from averon_import.services.manual_tenders.history_decisions import _candidate_fingerprint

    common = {
        "tender_id":"a" * 32,
        "workspace":{"revision":1,"source_sha256":"b" * 64},
        "run":{"run_id":"c" * 32,"history_catalog_version":"snapshot"},
        "source":{"source_row_id":"d" * 32,"name":"Насос","model":"25-40"},
        "canonical":{"physical_excel_row":2},
        "candidate":{
            "offer":{"retrieval_classification":"EXACT_NAME_UNIT","offer_id":"one_c_history:item","title":"Насос","article":"","manufacturer":"","brand":"","source_item_id":"item","price":"10","currency":"RUB","price_unit":"шт"},
            "price_provenance":{"history_item_id":"item","selected_event_id":"event","purchase_date":"2025-01-01","price_basis":"gross_including_vat","effective_unit_price_gross":"10","currency_basis":"source","unit_family":"piece"},
            "match":{"decision":"REVIEW","offer_id":"one_c_history:item","matched_attributes":[],"supporting_attributes":[],"conflicting_attributes":[],"missing_attributes":[]},
        },
    }
    old_fingerprint = _candidate_fingerprint(**common)
    enriched = copy.deepcopy(common)
    enriched["candidate"]["offer"]["history_characteristic"] = "25-40"
    assert _candidate_fingerprint(**enriched) == old_fingerprint

    normalized_legacy = copy.deepcopy(common)
    normalized_legacy["candidate"]["offer"]["retrieval_classification"] = "NORMALIZED_NAME_UNIT"
    normalized_legacy["candidate"]["price_provenance"].update({
        "normalizer_revision":"normalized-name-unit-v1",
        "normalized_name_signature":"e" * 64,
    })
    normalized_fingerprint = _candidate_fingerprint(**normalized_legacy)
    normalized_empty = copy.deepcopy(normalized_legacy)
    normalized_empty["candidate"]["offer"]["history_characteristic"] = ""
    assert _candidate_fingerprint(**normalized_empty) == normalized_fingerprint
    normalized_with_characteristic = copy.deepcopy(normalized_legacy)
    normalized_with_characteristic["candidate"]["offer"]["history_characteristic"] = "25-40"
    assert _candidate_fingerprint(**normalized_with_characteristic) != normalized_fingerprint


def test_d3_durable_human_history_confirm_restart_export_revoke(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    run_id = _persist_history_review_run(main, repository, workspace)
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    initial = _api_request(main.app, "GET", base, headers=_auth_headers())
    assert initial.status_code == 200, initial.text
    assert _api_request(main.app, "GET", base, headers=_auth_headers("other-user")).status_code == 404
    snapshot = initial.json()
    assert snapshot["decision_revision"] == 0
    assert len(snapshot["rows"][0]["candidates"]) == 1
    candidate = snapshot["rows"][0]["candidates"][0]
    assert candidate["confirmable"] is True
    assert candidate["offer"]["title"] == "Плющ искусственный"
    before_jobs = len(main.job_service.jobs)

    confirm = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE",
            "source_row_id":snapshot["rows"][0]["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"],
            "expected_revision":0,
        }, ensure_ascii=False).encode(),
    )
    assert confirm.status_code == 200, confirm.text
    assert len(main.job_service.jobs) == before_jobs
    confirmed = confirm.json()
    assert confirmed["decision_revision"] == 1
    assert len(confirmed["events"]) == 1
    assert confirmed["events"][0]["actor"] == {"username":"tender-user", "role":"user"}
    assert confirmed["effective"][snapshot["rows"][0]["source_row_id"]]["candidate_offer_id"] == candidate["candidate_offer_id"]
    durable_run = main.tender_sourcing_runs.get_public(
        repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id,
    )
    assert durable_run["rows"][0]["route"]["history_outcome"] == "REVIEW"
    assert durable_run["rows"][0]["route"]["history_safe_basis"] is None

    from averon_import.services.manual_tenders.history_decisions import TenderHistoryDecisionStore
    monkeypatch.setattr(main, "tender_history_decisions", TenderHistoryDecisionStore(repository))
    restored = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    assert restored["decision_revision"] == 1
    assert restored["effective"][snapshot["rows"][0]["source_row_id"]]["decision_id"] == confirmed["effective"][snapshot["rows"][0]["source_row_id"]]["decision_id"]

    export_endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export"
    included = _api_request(
        main.app, "POST", export_endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert included.status_code == 202, included.text
    included_job = _wait_tender_job(main, included.json()["id"])
    assert included_job["status"] == "completed", included_job
    record = included_job["result"]
    assert record["historical_count"] == 1
    assert record["automatic_historical_count"] == 0
    assert record["human_confirmed_historical_count"] == 1
    assert record["history_decision_revision"] == 1
    assert record["history_decision_digest"] == confirmed["decision_digest"]
    file_path, _record, release = main.tender_price_exports.acquire_download(
        repository.workspace_root / workspace["tender_id"], workspace["tender_id"], record["export_id"], "username:tender-user",
    )
    try:
        workbook = load_workbook(file_path, data_only=False)
        sheet = workbook[TEMPLATE_SHEET]
        assert Decimal(str(sheet["H2"].value)) == Decimal("123.450000")
        assert sheet["H2"].fill.fgColor.rgb.endswith("FFF2CC")
        assert "подтверждено пользователем" in sheet["H2"].comment.text.casefold()
        assert "tender-user" not in sheet["H2"].comment.text
        workbook.close()
    finally:
        release()

    stale = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE",
            "source_row_id":snapshot["rows"][0]["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"],
            "expected_revision":0,
        }).encode(),
    )
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "TENDER_HISTORY_DECISIONS_STALE"

    decision_id = restored["effective"][snapshot["rows"][0]["source_row_id"]]["decision_id"]
    revoked = _api_request(
        main.app, "POST", f"{base}/{decision_id}/revoke",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"decision":"REVOKE_HISTORY_CONFIRMATION","expected_revision":1}',
    )
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["decision_revision"] == 2
    assert len(revoked.json()["events"]) == 2
    assert revoked.json()["effective"] == {}
    no_history = _api_request(
        main.app, "POST", export_endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":true}',
    )
    assert no_history.status_code == 409
    assert no_history.json()["detail"]["code"] == "TENDER_EXPORT_NO_ELIGIBLE_PRICES"


def test_d3_exact_and_normalized_candidate_projection_is_bounded_and_size_safe(tmp_path):
    from averon_import.services.sourcing.history_identity import (
        HISTORY_IDENTITY_NORMALIZER_REVISION,
        history_name_signature_digest,
    )
    from averon_import.services.sourcing.models import (
        HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent,
        SourcingResult,
    )

    sources = []
    results = []
    ids = []
    for index in range(370):
        row_id = f"{index + 1:032x}"
        ids.append(row_id)
        source = {
            "source_row_id":row_id, "excel_row":index + 2, "row_type":"item",
            "name":f"Реальная позиция {index + 1}", "raw_unit":"шт", "article":"",
            "manufacturer":"", "model":"", "quantity":"1", "quantity_trusted":True,
            "unit_basis":parse_unit_basis("шт"),
        }
        sources.append(source)
        provenance = {
            "source":"one_c_history", "source_kind":"historical_purchase",
            "snapshot_version":"history-snapshot-size", "history_item_id":f"item-{index}",
            "selected_event_id":f"event-{index}", "purchase_date":"2025-04-16",
            "price_basis":"gross_including_vat", "effective_unit_price_gross":"10.00",
            "currency_basis":"company_default", "unit_family":"piece",
            "normalizer_revision":HISTORY_IDENTITY_NORMALIZER_REVISION,
            "normalized_name_signature":history_name_signature_digest(source["name"]),
        }
        offer = Offer(
            offer_id=f"one_c_history:item-{index}", provider="one_c_history", source_item_id=f"item-{index}",
            title=f"Позиция реальная {index + 1}", price=Decimal("10.00"), currency="RUB", price_unit="шт",
            attributes={"characteristic":"техническая характеристика", "unit_family":"ignored-by-projection"},
            data_provenance=provenance,
            history_retrieval_classification=HistoryRetrievalClassification.NORMALIZED_NAME_UNIT,
        )
        match = MatchResult(offer=offer, decision=MatchDecision.REVIEW, rank=1, matched_attributes=["name", "unit"])
        results.append(SourcingResult(
            intent=ProductIntent(source_row_id=row_id, source_text=source["name"], normalized_name=source["name"], unit="шт"),
            offers=[offer], match_results=[match],
            route={"source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW"},
        ).model_dump(mode="json"))
    projection = canonical_tender_projection({"results":results}, sources, ids)
    run = {
        "schema_version":1,"run_id":"a" * 32,"tender_id":"b" * 32,
        "source_sha256":"c" * 64,"workspace_revision":1,"status":"completed",
        "source_mode":"one_c_only","history_catalog_version":"history-snapshot-size",
        "selected_source_row_ids":ids,"rows":projection,
    }
    encoded = json.dumps(run, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    assert all(len(row["history_review_candidates"]) == 1 for row in projection)
    assert all(row["history_review_candidates"][0]["offer"]["history_characteristic"] == "техническая характеристика" for row in projection)
    assert all("attributes" not in row["history_review_candidates"][0]["offer"] for row in projection)
    assert len(encoded) < 1024 * 1024
    assert len(encoded) < MAX_TENDER_RUN_BYTES

    oversized_result = dict(results[0])
    oversized_result["offers"] = [{
        **results[0]["offers"][0],
        "attributes":{"characteristic":"x" * 301},
    }]
    oversized = canonical_tender_projection({"results":[oversized_result]}, sources[:1], ids[:1])
    assert oversized[0]["history_review_candidates"] == []


def test_d3_ambiguous_exact_choices_append_and_latest_confirmation_is_effective(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    run_id = _persist_history_review_run(main, repository, workspace, candidates=2)
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    candidates = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    choices = candidates["rows"][0]["candidates"]
    assert len(choices) == 2 and all(choice["confirmable"] for choice in choices)
    source_row_id = candidates["rows"][0]["source_row_id"]

    def confirm(choice, revision):
        return _api_request(
            main.app, "POST", base,
            headers={**_auth_headers(), "Content-Type":"application/json"},
            body=json.dumps({
                "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source_row_id,
                "candidate_offer_id":choice["candidate_offer_id"], "expected_revision":revision,
            }).encode(),
        )

    first = confirm(choices[0], 0)
    assert first.status_code == 200
    second = confirm(choices[1], 1)
    assert second.status_code == 200
    state = second.json()
    assert state["decision_revision"] == 2
    assert [event["decision_type"] for event in state["events"]] == ["CONFIRM_HISTORY_CANDIDATE"] * 2
    assert state["effective"][source_row_id]["candidate_offer_id"] == choices[1]["candidate_offer_id"]


def test_d3_confirmable_exact_article_and_hard_conflicts_fail_closed(tender_api, tmp_path):
    main, repository = tender_api
    article_workspace = _confirm_ivy_history_tender(main, repository, tmp_path, source_article="SKU-42")
    article_run = _persist_history_review_run(
        main, repository, article_workspace, classification="EXACT_ARTICLE",
        offer_overrides={"article":"SKU-42"},
    )
    article_base = f"/api/manual-tenders/{article_workspace['tender_id']}/runs/{article_run}/history-decisions"
    article_snapshot = _api_request(main.app, "GET", article_base, headers=_auth_headers()).json()
    assert article_snapshot["rows"][0]["candidates"][0]["confirmable"] is True
    repository.delete(article_workspace["tender_id"], "username:tender-user")

    invalid_cases = [
        ("article_mismatch", {"source_article":"SKU-SOURCE", "classification":"EXACT_NAME_UNIT", "offer_overrides":{"article":"SKU-OTHER"}}),
        ("unit_mismatch", {"offer_overrides":{"price_unit":"кг"}, "provenance_overrides":{"unit_family":"kilogram"}}),
        ("invalid_provenance", {"provenance_overrides":{"source":"untrusted"}}),
        ("missing_price", {"offer_overrides":{"price":None}}),
        ("future_date", {"provenance_overrides":{"purchase_date":"2999-01-01"}}),
        ("non_rub", {"offer_overrides":{"currency":"USD"}, "provenance_overrides":{"currency_basis":"source"}}),
        ("fuzzy", {"classification":"FUZZY"}),
    ]
    for _name, options in invalid_cases:
        source_article = options.pop("source_article", "")
        workspace = _confirm_ivy_history_tender(main, repository, tmp_path, source_article=source_article)
        run_id = _persist_history_review_run(main, repository, workspace, **options)
        base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
        state = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
        source_row_id = workspace["rows"][0]["source_row_id"]
        offer_id = "one_c_history:history-item-1"
        candidate = next((item for item in state["rows"][0]["candidates"] if item["candidate_offer_id"] == offer_id), None)
        if candidate is not None:
            assert candidate["confirmable"] is False, _name
        response = _api_request(
            main.app, "POST", base,
            headers={**_auth_headers(), "Content-Type":"application/json"},
            body=json.dumps({
                "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source_row_id,
                "candidate_offer_id":offer_id, "expected_revision":state["decision_revision"],
            }).encode(),
        )
        if _name == "fuzzy":
            assert response.status_code == 409
            assert response.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_CONFIRMABLE"
        else:
            assert response.status_code == 409, (_name, response.text)
            assert response.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_CONFIRMABLE"
        repository.delete(workspace["tender_id"], "username:tender-user")


def test_d3_stored_confirmation_is_inactive_when_candidate_evidence_changes(tender_api, tmp_path):
    import copy
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    run_id = _persist_history_review_run(main, repository, workspace)
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    confirmed = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({"decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":snapshot["rows"][0]["source_row_id"], "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0}).encode(),
    )
    assert confirmed.status_code == 200
    workspace_path = repository.workspace_root / workspace["tender_id"]
    original_run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    ledger_path = workspace_path / "history-decisions" / f"{run_id}.json"
    original_ledger = ledger_path.read_bytes()
    mutations = [
        lambda run: run["rows"][0]["history_review_candidates"][0]["offer"].update(price="999.00"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["offer"].update(price_unit="кг"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["offer"].update(offer_id="changed-offer"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["offer"].update(retrieval_classification="FUZZY"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["price_provenance"].update(history_item_id="changed-item"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["price_provenance"].update(selected_event_id="changed-event"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["price_provenance"].update(purchase_date="2025-04-17"),
        lambda run: run["rows"][0]["history_review_candidates"][0]["match"].update(offer_id="changed-match"),
        lambda run: run["rows"][0].update(physical_excel_row=99),
        lambda run: run.update(history_catalog_version="changed-snapshot"),
    ]
    for mutate in mutations:
        ledger_path.write_bytes(original_ledger)
        changed = copy.deepcopy(original_run)
        mutate(changed)
        try:
            current = main.tender_history_decisions.get_snapshot(workspace_path, workspace, changed, workspace["tender_id"], run_id)
        except TenderWorkspaceError as error:
            assert error.code in {"TENDER_RUN_CORRELATION_FAILED", "TENDER_RESULT_CORRELATION_FAILED", "TENDER_RUN_WORKSPACE_MISMATCH"}
        else:
            assert current["effective"] == {}
    ledger = json.loads(original_ledger)
    ledger["events"][0]["evidence_fingerprint"] = "f" * 64
    ledger_path.write_text(json.dumps(ledger), encoding="utf-8")
    assert main.tender_history_decisions.get_snapshot(
        workspace_path, workspace, original_run, workspace["tender_id"], run_id,
    )["effective"] == {}
    ledger_path.write_bytes(original_ledger)

    for changed_workspace in (
        {**workspace, "source_sha256":"f" * 64},
        {**workspace, "revision":int(workspace["revision"]) + 1},
    ):
        with pytest.raises(TenderWorkspaceError) as mismatch:
            main.tender_history_decisions.get_snapshot(
                workspace_path, changed_workspace, original_run, workspace["tender_id"], run_id,
            )
        assert mismatch.value.code == "TENDER_RUN_WORKSPACE_MISMATCH"
    with pytest.raises(TenderWorkspaceError):
        main.tender_history_decisions.get_snapshot(
            workspace_path, workspace, original_run, workspace["tender_id"], "d" * 32,
        )
    unknown = _api_request(
        main.app, "POST", base,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":"e" * 32,
            "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":1,
        }).encode(),
    )
    assert unknown.status_code == 404


def test_d3_human_history_ui_lifecycle_regression():
    node = shutil.which("node")
    assert node, "Node.js is required for the manual tender human-history UI regression"
    root = Path(__file__).resolve().parents[1]
    for script, expected in (
        ("manual_tender_history_decisions.cjs", "PASS: strong confirmation unchanged; fuzzy compare/assertion/two-step"),
        ("excel_tender_review_navigation.cjs", "PASS: Excel Tender review sequence"),
    ):
        result = subprocess.run(
            [node, str(root / "tests" / "js" / script), str(root / "averon_import" / "static" / "app.js")],
            capture_output=True, text=True, timeout=30, check=False,
        )
        assert result.returncode == 0, result.stdout + result.stderr
        assert expected in result.stdout


def test_excel_tender_export_run_pinning_ui_regression():
    node = shutil.which("node")
    assert node, "Node.js is required for the Excel Tender export run-pinning regression"
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [
            node,
            str(root / "tests" / "js" / "excel_tender_export_run_pinning.cjs"),
            str(root / "averon_import" / "static" / "app.js"),
        ],
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "PASS: Excel Tender run pinning" in result.stdout


def test_recent_tender_lifecycle_ui_regression():
    node = shutil.which("node")
    assert node, "Node.js is required for recent tender UI lifecycle regression"
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [
            node,
            str(root / "tests" / "js" / "recent_tender_lifecycle.cjs"),
            str(root / "averon_import" / "static" / "app.js"),
        ],
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "PASS: recent tender reload" in result.stdout


def test_d3_export_fails_if_decision_set_changes_before_runner(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    run_id = _persist_history_review_run(main, repository, workspace)
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    candidates = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = candidates["rows"][0]["candidates"][0]
    source_row_id = candidates["rows"][0]["source_row_id"]
    confirmed = main.tender_history_decisions.confirm(
        repository.workspace_root / workspace["tender_id"], workspace,
        main.tender_sourcing_runs.get_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id),
        tender_id=workspace["tender_id"], run_id=run_id, source_row_id=source_row_id,
        candidate_offer_id=candidate["candidate_offer_id"], expected_revision=0,
        actor_username="tender-user", actor_role="user",
    )
    original_submit = main.job_service.submit

    def submit_after_decision_change(*args, **kwargs):
        if kwargs.get("kind") == "manual_tender_price_export":
            run = main.tender_sourcing_runs.get_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id)
            main.tender_history_decisions.confirm(
                repository.workspace_root / workspace["tender_id"], workspace, run,
                tender_id=workspace["tender_id"], run_id=run_id, source_row_id=source_row_id,
                candidate_offer_id=candidate["candidate_offer_id"], expected_revision=confirmed["decision_revision"],
                actor_username="tender-user", actor_role="user",
            )
        return original_submit(*args, **kwargs)

    monkeypatch.setattr(main.job_service, "submit", submit_after_decision_change)
    response = _api_request(
        main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert response.status_code == 202, response.text
    job = _wait_tender_job(main, response.json()["id"])
    assert job["status"] == "failed", job
    assert job["error_code"] == "TENDER_HISTORY_DECISIONS_CHANGED"
    assert not list((repository.workspace_root / workspace["tender_id"] / "exports").glob("*.xlsx")) if (repository.workspace_root / workspace["tender_id"] / "exports").exists() else True


def test_d3_history_confirmation_is_never_projected_for_fuzzy_or_mixed_mode(tmp_path):
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent, SourcingResult, SourcingSourceMode

    source = {"source_row_id":"a" * 32, "excel_row":2, "row_type":"item", "name":"Source", "raw_unit":"шт"}
    provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase",
        "price_basis":"gross_including_vat",
    }
    offer = Offer(
        offer_id="history:fuzzy", provider="one_c_history", source_item_id="item", title="Similar",
        history_retrieval_classification=HistoryRetrievalClassification.FUZZY, data_provenance=provenance,
    )
    match = MatchResult(offer=offer, decision=MatchDecision.REVIEW, rank=1)
    intent = ProductIntent(source_row_id=source["source_row_id"], source_text="Source")
    one_c_review = SourcingResult(intent=intent, offers=[offer], match_results=[match], route={"source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW"})
    mixed_review = one_c_review.model_copy(update={"route":one_c_review.route.model_copy(update={"source_mode":SourcingSourceMode.ONE_C_THEN_PROVIDER})})
    local_history_row = canonical_tender_projection({"results":[one_c_review.model_dump(mode="json")]}, [source], [source["source_row_id"]])[0]
    assert local_history_row["history_review_candidates"] == []
    mixed_row = canonical_tender_projection({"results":[mixed_review.model_dump(mode="json")]}, [source], [source["source_row_id"]])[0]
    assert mixed_row["history_review_candidates"] == []


def test_d4b_fuzzy_confirmation_requires_explicit_two_step_and_survives_restart_export_revoke(tender_api, tmp_path, monkeypatch):
    from averon_import.services.manual_tenders.history_decisions import TenderHistoryDecisionStore
    from averon_import.services.manual_tenders.sourcing import TenderSourcingRunStore

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace, classification="FUZZY",
        offer_overrides={"title":"Плющ декоративный"},
    )
    workspace_path = repository.workspace_root / workspace["tender_id"]
    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    assert candidate["offer"]["retrieval_classification"] == "FUZZY"
    assert candidate["retrieval_rank"] == 1
    assert candidate["confirmable"] is False
    assert candidate["confirmable_for_explicit_fuzzy"] is True
    assert candidate["confirmation_basis"] == "FUZZY_MANUAL_CONFIRMATION"
    assert len(candidate["fuzzy_evidence_fingerprint"]) == 64

    normal_body = {
        "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
        "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
    }
    for extra in ({}, {"confirmation_mode":"EXPLICIT_FUZZY_IDENTITY"}, {
        "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":False,
    }):
        response = _api_request(
            main.app, "POST", base, headers={**_auth_headers(), "Content-Type":"application/json"},
            body=json.dumps({**normal_body, **extra}).encode(),
        )
        assert response.status_code == 409
        assert response.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_CONFIRMABLE"

    confirmed = _api_request(
        main.app, "POST", base, headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({**normal_body, "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True}).encode(),
    )
    assert confirmed.status_code == 200, confirmed.text
    payload = confirmed.json()
    assert payload["events"][0]["confirmation_basis"] == "FUZZY_MANUAL_CONFIRMATION"
    assert payload["events"][0]["identity_assertion"] == "SAME_PRODUCT_V1"
    effective = payload["effective"][source["source_row_id"]]
    assert effective["confirmation_basis"] == "FUZZY_MANUAL_CONFIRMATION"
    assert effective["identity_assertion"] == "SAME_PRODUCT_V1"

    # Reload both durable stores to prove there is no dependency on transient job memory.
    monkeypatch.setattr(main, "tender_history_decisions", TenderHistoryDecisionStore(repository))
    monkeypatch.setattr(main, "tender_sourcing_runs", TenderSourcingRunStore(repository))
    run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    restored = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    assert restored["effective"][source["source_row_id"]]["confirmation_basis"] == "FUZZY_MANUAL_CONFIRMATION"
    resolved = main.tender_price_resolver.resolve_run(
        workspace, run, tender_id=workspace["tender_id"], run_id=run_id,
        include_historical_prices=True, human_history_decisions=restored["effective"],
    )
    assert len(resolved) == 1 and resolved[0].eligible
    assert resolved[0].safe_summary()["authority"] == "HUMAN_CONFIRMED_HISTORY"
    assert resolved[0].audit_summary["human_confirmation_basis"] == "FUZZY_MANUAL_CONFIRMATION"

    export_response = _api_request(
        main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert export_response.status_code == 202, export_response.text
    exported = _wait_tender_job(main, export_response.json()["id"])
    assert exported["status"] == "completed", exported
    assert exported["result"]["human_confirmed_historical_count"] == 1
    export_files = list((workspace_path / "exports").glob("*.xlsx"))
    assert len(export_files) == 1
    exported_workbook = load_workbook(export_files[0])
    try:
        fuzzy_comment = exported_workbook[TEMPLATE_SHEET]["H2"].comment.text
        assert "Совпадение позиции явно подтверждено пользователем." in fuzzy_comment
        assert "tender-user" not in fuzzy_comment
    finally:
        exported_workbook.close()

    decision_id = restored["effective"][source["source_row_id"]]["decision_id"]
    revoked = _api_request(
        main.app, "POST", f"{base}/{decision_id}/revoke",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"decision":"REVOKE_HISTORY_CONFIRMATION","expected_revision":1}',
    )
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["effective"] == {}


@pytest.mark.parametrize("case", [
    "reject", "alternative", "article_mismatch", "article_missing", "manufacturer_mismatch", "model_mismatch",
    "unit_mismatch", "corrupt_provenance", "snapshot_mismatch", "event_mismatch", "price_mismatch",
    "future_date", "non_rub", "net_price",
])
def test_d4b_fuzzy_server_gate_rejects_hard_conflicts(tender_api, tmp_path, case):
    main, repository = tender_api
    source_options = {}
    offer_overrides = {"title":"Плющ декоративный", "article":"", "manufacturer":""}
    provenance_overrides = {}
    offer_attributes = {}
    if case in {"article_mismatch", "article_missing"}:
        source_options["source_article"] = "SOURCE-ARTICLE"
        offer_overrides["article"] = "SOURCE-ARTICLE"
    elif case == "manufacturer_mismatch":
        source_options["source_manufacturer"] = "Maker A"
        offer_overrides["manufacturer"] = "Maker A"
    elif case == "model_mismatch":
        source_options["source_model"] = "25-40"
        offer_attributes["characteristic"] = "25-40"

    workspace = _confirm_ivy_history_tender(main, repository, tmp_path, **source_options)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace, classification="FUZZY", offer_overrides=offer_overrides,
        provenance_overrides=provenance_overrides, offer_attributes=offer_attributes,
    )
    workspace_path = repository.workspace_root / workspace["tender_id"]
    run_path = main.tender_sourcing_runs._path(workspace_path, run_id)
    run = json.loads(run_path.read_text(encoding="utf-8"))
    candidate_projection = run["rows"][0]["history_review_candidates"][0]["fuzzy_v1"]
    if case == "reject":
        candidate_projection["match"][0] = "REJECT"
    elif case == "alternative":
        candidate_projection["match"][0] = "ALTERNATIVE"
    elif case == "event_mismatch":
        run["rows"][0]["route"]["history_selected_event_id"] = "different-event"
    elif case in {"article_mismatch", "article_missing"}:
        candidate_projection["identity"][3] = "OTHER-ARTICLE" if case == "article_mismatch" else ""
    elif case == "manufacturer_mismatch":
        candidate_projection["identity"][4] = "Maker B"
    elif case == "model_mismatch":
        candidate_projection["identity"][5] = "25-60"
    elif case == "unit_mismatch":
        candidate_projection["identity"][8] = "кг"
        candidate_projection["provenance"][9] = "kilogram"
    elif case == "corrupt_provenance":
        candidate_projection["provenance"][0] = "untrusted"
    elif case == "snapshot_mismatch":
        candidate_projection["provenance"][2] = "another-snapshot"
    elif case == "price_mismatch":
        candidate_projection["identity"][6] = "99.00"
    elif case == "future_date":
        candidate_projection["provenance"][5] = "2999-01-01"
    elif case == "non_rub":
        candidate_projection["identity"][7] = "USD"
    elif case == "net_price":
        candidate_projection["provenance"][6] = "net_excluding_vat"
    run_path.write_text(json.dumps(run, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    base = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", base, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    assert candidate["confirmable_for_explicit_fuzzy"] is False, case
    response = _api_request(
        main.app, "POST", base, headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
            "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True,
        }).encode(),
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_CONFIRMABLE"


@pytest.mark.parametrize("case", ["net_price", "wrong_source_kind"])
def test_d4b_preprojection_bad_fuzzy_provenance_is_not_confirmable_or_exportable(tender_api, tmp_path, case):
    main, repository = tender_api
    provenance_overrides = {
        "price_basis":"net_excluding_vat",
    } if case == "net_price" else {
        "source_kind":"not_a_historical_purchase",
    }
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace, classification="FUZZY",
        offer_overrides={"title":"Плющ декоративный"},
        provenance_overrides=provenance_overrides,
    )
    workspace_path = repository.workspace_root / workspace["tender_id"]
    run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    assert run["rows"][0]["history_review_candidates"] == []

    endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", endpoint, headers=_auth_headers()).json()
    assert snapshot["rows"][0]["candidates"] == []
    rejected = _api_request(
        main.app, "POST", endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
            "candidate_offer_id":"one_c_history:history-item-1", "expected_revision":0,
            "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True,
        }).encode(),
    )
    assert rejected.status_code == 404
    assert rejected.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_FOUND"

    unavailable = _api_request(
        main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert unavailable.status_code == 409
    assert unavailable.json()["detail"]["code"] == "TENDER_EXPORT_NO_ELIGIBLE_PRICES"
    assert unavailable.json()["detail"]["summary"]["historical_count"] == 0
    assert not any(job.kind == "manual_tender_price_export" for job in main.job_service.jobs.values())
    export_dir = workspace_path / "exports"
    assert not export_dir.exists() or not list(export_dir.iterdir())


@pytest.mark.parametrize("case", [
    "reject", "alternative", "match_conflict", "article_mismatch", "article_missing",
    "manufacturer_conflict", "model_conflict", "unit_conflict", "price_mismatch", "nonpositive_price",
    "non_rub", "net_price", "future_date", "snapshot_mismatch", "item_mismatch", "event_missing",
])
def test_d4b_projection_drops_non_actionable_fuzzy_candidates(tender_api, tmp_path, case):
    main, repository = tender_api
    source_options = {}
    offer_overrides = {"title":"Similar historical item"}
    provenance_overrides = {}
    offer_attributes = {}
    match_options = {}
    if case in {"reject", "alternative"}:
        match_options["decision"] = {"reject":"REJECT", "alternative":"ALTERNATIVE"}[case]
    elif case == "match_conflict":
        match_options["conflicting_attributes"] = ["article"]
    elif case in {"article_mismatch", "article_missing"}:
        source_options["source_article"] = "SOURCE-ARTICLE"
        offer_overrides["article"] = "OTHER-ARTICLE" if case == "article_mismatch" else ""
    elif case == "manufacturer_conflict":
        source_options["source_manufacturer"] = "Maker A"
        offer_overrides["manufacturer"] = "Maker B"
    elif case == "model_conflict":
        source_options["source_model"] = "25-40"
        offer_attributes["characteristic"] = "25-60"
    elif case == "unit_conflict":
        offer_overrides["price_unit"] = "кг"
        provenance_overrides["unit_family"] = "kilogram"
    elif case == "price_mismatch":
        offer_overrides["price"] = Decimal("99.00")
    elif case == "nonpositive_price":
        offer_overrides["price"] = Decimal("0")
    elif case == "non_rub":
        offer_overrides["currency"] = "USD"
    elif case == "net_price":
        provenance_overrides["price_basis"] = "net_excluding_vat"
    elif case == "future_date":
        provenance_overrides["purchase_date"] = "2999-01-01"
    elif case == "snapshot_mismatch":
        provenance_overrides["snapshot_version"] = "different-snapshot"
    elif case == "item_mismatch":
        provenance_overrides["history_item_id"] = "different-item"
    elif case == "event_missing":
        provenance_overrides["selected_event_id"] = ""

    workspace = _confirm_ivy_history_tender(main, repository, tmp_path, **source_options)
    run_id = _persist_history_review_run(
        main, repository, workspace, classification="FUZZY",
        offer_overrides=offer_overrides, provenance_overrides=provenance_overrides,
        offer_attributes=offer_attributes, candidate_specs=[match_options],
    )
    run = main.tender_sourcing_runs.get_public(
        repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id,
    )
    assert run["rows"][0]["history_review_candidates"] == [], case


@pytest.mark.parametrize("case", [
    "article_mismatch", "article_missing", "manufacturer_conflict", "model_conflict",
    "unit_conflict", "invalid_provenance", "future_date", "net_price", "non_rub", "attribute_conflict",
])
def test_d4b_likely_match_does_not_bypass_fuzzy_hard_gates(tender_api, tmp_path, case):
    main, repository = tender_api
    source_options = {}
    offer_overrides = {"title":"Similar historical item"}
    provenance_overrides = {}
    offer_attributes = {}
    match_options = {"decision":"LIKELY_MATCH", "supporting_attributes":["normalized_name"]}
    if case in {"article_mismatch", "article_missing"}:
        source_options["source_article"] = "SOURCE-ARTICLE"
        offer_overrides["article"] = "OTHER-ARTICLE" if case == "article_mismatch" else ""
    elif case == "manufacturer_conflict":
        source_options["source_manufacturer"] = "Maker A"
        offer_overrides["manufacturer"] = "Maker B"
    elif case == "model_conflict":
        source_options["source_model"] = "25-40"
        offer_attributes["characteristic"] = "25-60"
    elif case == "unit_conflict":
        offer_overrides["price_unit"] = "кг"
        provenance_overrides["unit_family"] = "kilogram"
    elif case == "invalid_provenance":
        provenance_overrides["source"] = "untrusted"
    elif case == "future_date":
        provenance_overrides["purchase_date"] = "2999-01-01"
    elif case == "net_price":
        provenance_overrides["price_basis"] = "net_excluding_vat"
    elif case == "non_rub":
        offer_overrides["currency"] = "USD"
    elif case == "attribute_conflict":
        match_options["conflicting_attributes"] = ["article"]

    workspace = _confirm_ivy_history_tender(main, repository, tmp_path, **source_options)
    run_id = _persist_history_review_run(
        main, repository, workspace, classification="FUZZY", offer_overrides=offer_overrides,
        provenance_overrides=provenance_overrides, offer_attributes=offer_attributes,
        candidate_specs=[match_options],
    )
    run = main.tender_sourcing_runs.get_public(
        repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id,
    )
    assert run["rows"][0]["history_review_candidates"] == [], case


def test_d4b_fuzzy_slot_projection_accepts_only_match_likely_match_or_review():
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent, SourcingResult

    source = {"source_row_id":"b" * 32, "excel_row":2, "row_type":"item", "name":"Source item", "raw_unit":"шт"}
    provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase", "snapshot_version":"snapshot",
        "history_item_id":"history-item", "selected_event_id":"history-event", "purchase_date":"2025-04-16",
        "price_basis":"gross_including_vat", "effective_unit_price_gross":"10.00",
        "currency_basis":"company_default", "unit_family":"piece",
    }
    decisions = [
        MatchDecision.MATCH, MatchDecision.LIKELY_MATCH, MatchDecision.REVIEW,
        MatchDecision.ALTERNATIVE, MatchDecision.REJECT,
    ]
    offers = [Offer(
        offer_id=f"one_c_history:item-{index}", provider="one_c_history", source_item_id=f"item-{index}",
        title=f"Similar source product {index}", price=Decimal("10.00"), currency="RUB", price_unit="шт",
        attributes={"characteristic":"safe specification"},
        data_provenance={**provenance, "history_item_id":f"item-{index}", "selected_event_id":f"event-{index}"},
        history_retrieval_classification=HistoryRetrievalClassification.FUZZY,
    ) for index in range(1, 6)]
    matches = [MatchResult(
        offer=offer, decision=decision, rank=index,
        matched_attributes=["article"] if decision == MatchDecision.MATCH else [],
        supporting_attributes=["model"] if decision == MatchDecision.LIKELY_MATCH else [],
        missing_attributes=["name"] if decision == MatchDecision.REVIEW else [],
    ) for index, (offer, decision) in enumerate(zip(offers, decisions, strict=True), 1)]
    result = SourcingResult(
        intent=ProductIntent(source_row_id=source["source_row_id"], source_text=source["name"], normalized_name=source["name"], unit="шт"),
        offers=offers, match_results=matches,
        route={"source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW", "history_catalog_version":"snapshot"},
    ).model_dump(mode="json")
    projected = canonical_tender_projection({"results":[result]}, [source], [source["source_row_id"]])[0]["history_review_candidates"]
    assert [candidate["fuzzy_v1"]["rank"] for candidate in projected] == [1, 2, 3]
    assert [candidate["fuzzy_v1"]["match"][0] for candidate in projected] == ["MATCH", "LIKELY_MATCH", "REVIEW"]


def test_d4b_glue_likely_match_uses_supporting_evidence_and_shared_hard_gates():
    from averon_import.services.manual_tenders.history_fuzzy_eligibility import fuzzy_confirmation_eligibility_reason
    from averon_import.services.sourcing.matching import OfferMatcher
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, Offer, ProductIntent

    source = {
        "source_row_id":"c" * 32, "excel_row":18, "name":"Клей для плитки СМ17 25 кг",
        "article":"", "manufacturer":"", "model":"СМ17", "raw_unit":"кг",
    }
    offer = Offer(
        offer_id="one_c_history:glue-item", provider="one_c_history", source_item_id="glue-item",
        title="Клей д/плитки СМ 17", price=Decimal("123.45"), currency="RUB", price_unit="кг",
        history_retrieval_classification=HistoryRetrievalClassification.FUZZY,
    )
    intent = ProductIntent(
        source_row_id=source["source_row_id"], source_text=source["name"], normalized_name=source["name"],
        model=source["model"], unit=source["raw_unit"],
    )
    actual = OfferMatcher().match(intent, [offer])[0]
    assert actual.decision.value == "LIKELY_MATCH"
    assert actual.conflicting_attributes == []
    assert actual.missing_attributes == []
    assert actual.deterministic_evidence.get("preferred_differences") == []
    assert actual.supporting_attributes and actual.deterministic_evidence.get("supporting_matches")

    provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase", "snapshot_version":"snapshot",
        "history_item_id":"glue-item", "selected_event_id":"glue-event", "purchase_date":"2025-04-16",
        "price_basis":"gross_including_vat", "effective_unit_price_gross":"123.45",
        "currency_basis":"company_default", "unit_family":"kilogram",
    }
    route = {
        "source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW",
        "history_safe_basis":None, "history_catalog_version":"snapshot",
    }
    match = {
        "decision":actual.decision.value, "offer_id":offer.offer_id,
        "conflicting_attributes":actual.conflicting_attributes,
        "missing_attributes":actual.missing_attributes,
    }
    assert fuzzy_confirmation_eligibility_reason(
        source, offer.model_dump(mode="json") | {"retrieval_classification":"FUZZY"}, provenance, match, route,
        expected_snapshot_version="snapshot", physical_excel_row=source["excel_row"], retrieval_rank=3,
    ) is None


def test_d4b_fuzzy_candidate_outside_persisted_set_cannot_be_confirmed(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    run_id = _persist_history_review_run(main, repository, workspace, classification="FUZZY")
    endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    response = _api_request(
        main.app, "POST", endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":workspace["rows"][0]["source_row_id"],
            "candidate_offer_id":"one_c_history:not-persisted", "expected_revision":0,
            "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True,
        }).encode(),
    )
    assert response.status_code == 404


def test_d4b_ambiguous_fuzzy_confirmation_binds_the_clicked_candidate(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(main, repository, tmp_path)
    run_id = _persist_history_review_run(
        main, repository, workspace, candidates=3, classification="FUZZY",
        offer_overrides={"title":"Плющ декоративный"},
    )
    endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", endpoint, headers=_auth_headers()).json()
    candidates = snapshot["rows"][0]["candidates"]
    assert [candidate["retrieval_rank"] for candidate in candidates] == [1, 2, 3]
    selected = candidates[1]
    response = _api_request(
        main.app, "POST", endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":snapshot["rows"][0]["source_row_id"],
            "candidate_offer_id":selected["candidate_offer_id"], "expected_revision":0,
            "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True,
        }).encode(),
    )
    assert response.status_code == 200, response.text
    assert response.json()["events"][0]["candidate_offer_id"] == selected["candidate_offer_id"]
    assert response.json()["events"][0]["evidence_fingerprint"] == selected["fuzzy_evidence_fingerprint"]


def test_d4b_actionable_fuzzy_slot_selection_preserves_rank_five_and_confirms(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(
        main, repository, tmp_path, source_name="Зажим троса 8 мм",
        source_article="ROPE-8", source_manufacturer="Maker A",
    )
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_history_review_run(
        main, repository, workspace, candidates=8, classification="FUZZY",
        candidate_specs=[
            {"offer_overrides":{"title":"Зажим троса 2 мм", "article":"ROPE-2"}, "decision":"ALTERNATIVE"},
            {"offer_overrides":{"title":"Зажим троса 3 мм", "article":"ROPE-8"}, "decision":"LIKELY_MATCH"},
            {"offer_overrides":{"title":"Зажим для троса 4мм", "article":"ROPE-8"}},
            {"offer_overrides":{"title":"Зажим троса подходящего типа", "article":""}},
            {"offer_overrides":{"title":"Зажим для троса 8мм", "article":"ROPE-8", "manufacturer":"Maker A"}},
            {"offer_overrides":{"title":"Зажим троса с конфликтом", "article":"ROPE-8"}, "conflicting_attributes":["article"]},
            {"offer_overrides":{"title":"Зажим троса другой марки", "article":"ROPE-8", "manufacturer":"Maker B"}},
            {"offer_overrides":{"title":"Зажим троса ещё один вариант", "article":"ROPE-8", "manufacturer":"Maker A"}},
        ],
    )
    endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", endpoint, headers=_auth_headers()).json()
    candidates = snapshot["rows"][0]["candidates"]
    assert [candidate["retrieval_rank"] for candidate in candidates] == [2, 3, 5]
    assert candidates[0]["match"]["decision"] == "LIKELY_MATCH"
    candidate = next(candidate for candidate in candidates if candidate["offer"]["title"] == "Зажим для троса 8мм")
    assert candidate["retrieval_rank"] == 5
    assert candidate["offer"]["title"] == "Зажим для троса 8мм"
    assert candidate["confirmable_for_explicit_fuzzy"] is True
    response = _api_request(
        main.app, "POST", endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
            "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
            "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True,
        }).encode(),
    )
    assert response.status_code == 200, response.text
    assert response.json()["events"][0]["candidate_offer_id"] == candidate["candidate_offer_id"]
    assert response.json()["events"][0]["evidence_fingerprint"] == candidate["fuzzy_evidence_fingerprint"]
    import copy
    tampered_effective = copy.deepcopy(response.json()["effective"][source["source_row_id"]])
    tampered_effective["candidate"]["match"]["conflicting_attributes"] = ["article"]
    run = main.tender_sourcing_runs.get_public(
        repository.workspace_root / workspace["tender_id"], workspace["tender_id"], run_id,
    )
    export_decision = main.tender_price_resolver.resolve_run(
        workspace, run, tender_id=workspace["tender_id"], run_id=run_id,
        include_historical_prices=True,
        human_history_decisions={source["source_row_id"]:tampered_effective},
    )[0]
    assert not export_decision.eligible
    assert export_decision.reason_code == "HISTORY_CANDIDATE_SOURCE_CONFLICT"


@pytest.mark.parametrize("matcher_case", ["glue_likely_match", "article_match"])
def test_d4b_fuzzy_match_and_likely_match_require_explicit_confirmation_then_export(tender_api, tmp_path, matcher_case):
    from averon_import.services.sourcing.matching import OfferMatcher
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, Offer, ProductIntent

    if matcher_case == "glue_likely_match":
        source_name = "Клей для плитки СМ17 25 кг"
        source_model = "СМ17"
        source_article = ""
        source_unit = "кг"
        candidate_title = "Клей д/плитки СМ 17"
        candidate_article = ""
        expected_match = "LIKELY_MATCH"
        unit_family = "kilogram"
    else:
        source_name = "Клапан P-17"
        source_model = ""
        source_article = "VALVE-17"
        source_unit = "шт"
        candidate_title = "Клапан P-17"
        candidate_article = "VALVE-17"
        expected_match = "MATCH"
        unit_family = "piece"

    main, repository = tender_api
    workspace = _confirm_ivy_history_tender(
        main, repository, tmp_path, source_name=source_name, source_model=source_model,
        source_article=source_article, source_unit=source_unit,
    )
    source = next(row for row in workspace["rows"] if row["row_type"] == "item")

    matcher_intent = ProductIntent(
        source_row_id=source["source_row_id"], source_text=source["name"],
        normalized_name=source["name"], unit=source["raw_unit"],
        article=source.get("article", ""), model=source.get("model", ""),
    )
    matcher_offer = Offer(
        offer_id="one_c_history:matcher-probe", provider="one_c_history", source_item_id="matcher-probe",
        title=candidate_title, article=candidate_article, price=Decimal("123.4500"), currency="RUB",
        price_unit=source_unit,
        history_retrieval_classification=HistoryRetrievalClassification.FUZZY,
    )
    matcher_result = OfferMatcher().match(matcher_intent, [matcher_offer])[0]
    assert matcher_result.decision.value == expected_match
    assert matcher_result.conflicting_attributes == []
    assert matcher_result.missing_attributes == []
    assert matcher_result.deterministic_evidence.get("preferred_differences") == []
    if expected_match == "LIKELY_MATCH":
        assert matcher_result.matched_attributes == []
        assert matcher_result.supporting_attributes
        assert matcher_result.deterministic_evidence.get("supporting_matches")

    run_id = _persist_history_review_run(
        main, repository, workspace, classification="FUZZY",
        offer_overrides={"title":candidate_title, "article":candidate_article, "price_unit":source_unit},
        provenance_overrides={"unit_family":unit_family},
        candidate_specs=[{
            "decision":matcher_result.decision.value,
            "matched_attributes":matcher_result.matched_attributes,
            "supporting_attributes":matcher_result.supporting_attributes,
            "conflicting_attributes":matcher_result.conflicting_attributes,
            "missing_attributes":matcher_result.missing_attributes,
        }],
    )
    workspace_path = repository.workspace_root / workspace["tender_id"]
    run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    assert run["rows"][0]["route"]["history_outcome"] == "REVIEW"
    assert run["rows"][0]["route"]["history_safe_basis"] is None
    compact = run["rows"][0]["history_review_candidates"][0]["fuzzy_v1"]
    assert compact["classification"] == "FUZZY"
    assert compact["match"][0] == expected_match
    assert "safe_basis" not in compact

    endpoint = f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/history-decisions"
    snapshot = _api_request(main.app, "GET", endpoint, headers=_auth_headers()).json()
    candidate = snapshot["rows"][0]["candidates"][0]
    assert candidate["offer"]["retrieval_classification"] == "FUZZY"
    assert candidate["match"]["decision"] == expected_match
    assert candidate["confirmable"] is False
    assert candidate["confirmable_for_explicit_fuzzy"] is True

    base_body = {
        "decision":"CONFIRM_HISTORY_CANDIDATE", "source_row_id":source["source_row_id"],
        "candidate_offer_id":candidate["candidate_offer_id"], "expected_revision":0,
    }
    one_step = _api_request(
        main.app, "POST", endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps(base_body).encode(),
    )
    assert one_step.status_code == 409
    assert one_step.json()["detail"]["code"] == "TENDER_HISTORY_CANDIDATE_NOT_CONFIRMABLE"

    confirmed = _api_request(
        main.app, "POST", endpoint,
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=json.dumps({
            **base_body, "confirmation_mode":"EXPLICIT_FUZZY_IDENTITY", "explicit_identity_assertion":True,
        }).encode(),
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["events"][0]["confirmation_basis"] == "FUZZY_MANUAL_CONFIRMATION"
    persisted_match = confirmed.json()["effective"][source["source_row_id"]]["candidate"]["match"]
    assert persisted_match["decision"] == expected_match

    export = _api_request(
        main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert export.status_code == 202, export.text
    job = _wait_tender_job(main, export.json()["id"])
    assert job["status"] == "completed", job
    assert job["result"]["historical_count"] == 1
    assert job["result"]["automatic_historical_count"] == 0
    assert job["result"]["human_confirmed_historical_count"] == 1


def test_d4b_fuzzy_allowlist_preserves_original_decisions_in_projection():
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent, SourcingResult

    source = {"source_row_id":"a" * 32, "excel_row":2, "row_type":"item", "name":"Source item", "raw_unit":"шт"}
    provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase", "snapshot_version":"snapshot",
        "history_item_id":"history-item", "selected_event_id":"history-event", "purchase_date":"2025-04-16",
        "price_basis":"gross_including_vat", "effective_unit_price_gross":"10.00",
        "currency_basis":"company_default", "unit_family":"piece",
    }
    offers = [Offer(
        offer_id=f"one_c_history:item-{index}", provider="one_c_history", source_item_id=f"item-{index}",
        title=f"Similar source product {index}", price=Decimal("10.00"), currency="RUB", price_unit="шт",
        attributes={"characteristic":"safe specification"},
        data_provenance={**provenance, "history_item_id":f"item-{index}", "selected_event_id":f"event-{index}"},
        history_retrieval_classification=HistoryRetrievalClassification.FUZZY,
    ) for index in range(1, 6)]
    decisions = [MatchDecision.MATCH, MatchDecision.LIKELY_MATCH, MatchDecision.REVIEW, MatchDecision.ALTERNATIVE, MatchDecision.REJECT]
    matches = [MatchResult(
        offer=offer, decision=decision, rank=index,
        matched_attributes=["article"] if decision == MatchDecision.MATCH else [],
        supporting_attributes=["model"] if decision == MatchDecision.LIKELY_MATCH else [],
        missing_attributes=["name"] if decision == MatchDecision.REVIEW else [],
    ) for index, (offer, decision) in enumerate(zip(offers, decisions, strict=True), 1)]
    result = SourcingResult(
        intent=ProductIntent(source_row_id=source["source_row_id"], source_text=source["name"], normalized_name=source["name"], unit="шт"),
        offers=offers, match_results=matches,
        route={"source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW", "history_catalog_version":"snapshot"},
    ).model_dump(mode="json")
    projected = canonical_tender_projection({"results":[result]}, [source], [source["source_row_id"]])[0]["history_review_candidates"]
    assert [candidate["fuzzy_v1"]["rank"] for candidate in projected] == [1, 2, 3]
    assert [candidate["fuzzy_v1"]["match"][0] for candidate in projected] == ["MATCH", "LIKELY_MATCH", "REVIEW"]


def test_d4b_fuzzy_projection_keeps_only_three_ranked_candidates_when_no_strong_candidate_exists():
    import copy
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent, SourcingResult

    source = {"source_row_id":"a" * 32, "excel_row":2, "row_type":"item", "name":"Source item", "raw_unit":"шт"}
    provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase", "snapshot_version":"snapshot",
        "history_item_id":"history-item", "selected_event_id":"history-event", "purchase_date":"2025-04-16",
        "price_basis":"gross_including_vat", "effective_unit_price_gross":"10.00",
        "currency_basis":"company_default", "unit_family":"piece",
    }
    fuzzy = [Offer(
        offer_id=f"one_c_history:item-{index}", provider="one_c_history", source_item_id=f"item-{index}",
        title=f"Similar source product {index}", price=Decimal("10.00"), currency="RUB", price_unit="шт",
        attributes={"characteristic":"safe bounded specification"}, data_provenance={**provenance, "history_item_id":f"item-{index}", "selected_event_id":f"event-{index}"},
        history_retrieval_classification=HistoryRetrievalClassification.FUZZY,
    ) for index in range(1, 5)]
    matches = [MatchResult(offer=offer, decision=MatchDecision.REVIEW, rank=index) for index, offer in enumerate(fuzzy, start=1)]
    result = SourcingResult(
        intent=ProductIntent(source_row_id=source["source_row_id"], source_text=source["name"], normalized_name=source["name"], unit="шт"),
        offers=fuzzy, match_results=matches,
        route={"source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW", "history_catalog_version":"snapshot"},
    ).model_dump(mode="json")
    projected = canonical_tender_projection({"results":[result]}, [source], [source["source_row_id"]])[0]["history_review_candidates"]
    assert [candidate["fuzzy_v1"]["rank"] for candidate in projected] == [1, 2, 3]
    assert all(candidate["fuzzy_v1"]["classification"] == "FUZZY" for candidate in projected)
    assert all(candidate["fuzzy_v1"]["provider"] == "1c" for candidate in projected)
    from averon_import.services.manual_tenders.history_decisions import _expand_review_candidate
    assert all(_expand_review_candidate(candidate)["offer"]["retrieval_classification"] == "FUZZY" for candidate in projected)

    malformed_facts = [
        ("source_kind", "not_a_historical_purchase"),
        ("price_basis", "net_excluding_vat"),
        ("price_basis", None),
        ("source", "untrusted"),
        ("provider", "etm_ipro"),
    ]
    for key, value in malformed_facts:
        malformed_result = copy.deepcopy(result)
        if key == "provider":
            malformed_result["offers"][0][key] = value
        elif value is None:
            malformed_result["offers"][0]["data_provenance"].pop(key, None)
        else:
            malformed_result["offers"][0]["data_provenance"][key] = value
        candidates = canonical_tender_projection(
            {"results":[malformed_result]}, [source], [source["source_row_id"]],
        )[0]["history_review_candidates"]
        assert [candidate["fuzzy_v1"]["identity"][0] for candidate in candidates] == [
            "one_c_history:item-2", "one_c_history:item-3", "one_c_history:item-4",
        ], (key, value)

    strong = fuzzy[0].model_copy(update={"history_retrieval_classification":HistoryRetrievalClassification.NORMALIZED_NAME_UNIT})
    strong_result = {**result, "offers":[strong.model_dump(mode="json"), *[offer.model_dump(mode="json") for offer in fuzzy[1:]]],
        "match_results":[matches[0].model_copy(update={"offer":strong}).model_dump(mode="json"), *[match.model_dump(mode="json") for match in matches[1:]]]}
    strong_projection = canonical_tender_projection({"results":[strong_result]}, [source], [source["source_row_id"]])[0]["history_review_candidates"]
    assert len(strong_projection) == 1
    assert strong_projection[0]["offer"]["retrieval_classification"] == "NORMALIZED_NAME_UNIT"
    assert "retrieval_rank" not in strong_projection[0]


def test_d4b_realistic_all_fuzzy_370_row_run_sizes_for_one_two_and_three_candidates():
    from averon_import.services.sourcing.models import HistoryRetrievalClassification, MatchDecision, MatchResult, Offer, ProductIntent, SourcingResult

    sources = []
    source_ids = []
    result_rows = []
    for row_index in range(370):
        row_id = f"{row_index + 1:032x}"
        source_ids.append(row_id)
        source = {
            "source_row_id":row_id, "excel_row":row_index + 2, "row_type":"item",
            "name":f"Позиция для закупки промышленного оборудования номер {row_index + 1:03d}",
            "raw_unit":"шт", "article":"", "manufacturer":"", "model":"",
            "quantity":"12", "quantity_trusted":True, "unit_basis":parse_unit_basis("шт"),
        }
        sources.append(source)
        offers = []
        matches = []
        for rank in range(1, 4):
            item_id = f"history-item-{row_index + 1:03d}-{rank}"
            event_id = f"history-event-{row_index + 1:03d}-{rank}"
            provenance = {
                "source":"one_c_history", "source_kind":"historical_purchase",
                "snapshot_version":"2026-10-shared-history-snapshot-v1", "history_item_id":item_id,
                "selected_event_id":event_id, "purchase_date":"2026-09-22",
                "price_basis":"gross_including_vat", "effective_unit_price_gross":"1234.56",
                "currency_basis":"company_default", "unit_family":"piece",
            }
            offer = Offer(
                offer_id=f"one_c_history:{item_id}", provider="one_c_history", source_item_id=item_id,
                title=f"Наименование исторической позиции {row_index + 1:03d} вариант {rank}",
                article=f"ART-{row_index + 1:03d}-{rank}", manufacturer="Производитель оборудования",
                price=Decimal("1234.56"), currency="RUB", price_unit="шт",
                attributes={"characteristic":"Характеристика оборудования, исполнение и типоразмер"},
                data_provenance=provenance,
                history_retrieval_classification=HistoryRetrievalClassification.FUZZY,
            )
            offers.append(offer)
            matches.append(MatchResult(
                offer=offer, decision=MatchDecision.REVIEW, rank=rank,
                matched_attributes=["unit"], missing_attributes=["name", "article"],
            ))
        result_rows.append(SourcingResult(
            intent=ProductIntent(
                source_row_id=row_id, source_text=source["name"], normalized_name=source["name"],
                quantity=source["quantity"], unit=source["raw_unit"],
            ),
            offers=offers, match_results=matches,
            route={
                "source_mode":"one_c_only", "final_source_kind":"history_review", "history_outcome":"REVIEW",
                "history_reason_code":"fuzzy_candidates_require_review", "history_candidate_count":20,
                "history_catalog_version":"2026-10-shared-history-snapshot-v1",
            },
        ).model_dump(mode="json"))

    sizes = {}
    for count in (1, 2, 3):
        limited_results = []
        for result in result_rows:
            limited_results.append({
                **result,
                "offers":result["offers"][:count],
                "match_results":result["match_results"][:count],
            })
        projected = canonical_tender_projection({"results":limited_results}, sources, source_ids)
        run = {
            "schema_version":1, "run_id":"a" * 32, "tender_id":"b" * 32,
            "source_sha256":"c" * 64, "workspace_revision":1, "status":"completed",
            "source_mode":"one_c_only", "history_catalog_version":"2026-10-shared-history-snapshot-v1",
            "selected_source_row_ids":source_ids, "rows":projected,
        }
        encoded = json.dumps(run, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        sizes[count] = len(encoded)
        assert all(len(row["history_review_candidates"]) == count for row in projected)
        if count == 3:
            assert set(projected[0]["route"]) == {
                "source_mode", "final_source_kind", "history_outcome",
                "history_safe_basis", "history_catalog_version", "fallback_called",
            }
    assert sizes[1] < sizes[2] < sizes[3]
    assert sizes[3] <= 1_015_909
    assert sizes[3] < MAX_TENDER_RUN_BYTES
    assert MAX_TENDER_RUN_BYTES == 1024 * 1024
    print(f"D4B 370-row fuzzy run bytes: one={sizes[1]}, two={sizes[2]}, three={sizes[3]}")


def test_tender_price_resolver_uses_only_proven_etm_pricewnds(tender_api, tmp_path):
    from averon_import.services.manual_tenders.price_export import TenderPriceResolver

    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    resolver = TenderPriceResolver()
    accepted_id = _persist_price_export_run(main, repository, workspace, source_row)
    accepted_run = main.tender_sourcing_runs.get_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"], accepted_id)
    accepted = resolver.resolve_run(workspace, accepted_run, tender_id=workspace["tender_id"], run_id=accepted_id, include_historical_prices=False)
    assert len(accepted) == 1 and accepted[0].eligible
    assert accepted[0].source_kind == "etm_ipro"
    assert accepted[0].source_unit_price == Decimal("123456.000000")

    net_id = _persist_price_export_run(
        main, repository, workspace, source_row,
        provenance={"source":"etm_ipro", "catalog_version":"etm-catalog-v4", "source_item_id":"synthetic-source-item", "price_field":"price"},
    )
    net_run = main.tender_sourcing_runs.get_public(repository.workspace_root / workspace["tender_id"], workspace["tender_id"], net_id)
    rejected = resolver.resolve_run(workspace, net_run, tender_id=workspace["tender_id"], run_id=net_id, include_historical_prices=False)
    assert len(rejected) == 1 and not rejected[0].eligible
    assert rejected[0].reason_code == "PRICE_BASIS_UNPROVEN"


def test_historical_price_export_requires_explicit_confirmation_without_ghost_job(tender_api, tmp_path):
    main, repository = tender_api
    from averon_import.services.manual_tenders.price_export import TenderPriceExportRepository

    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    provenance = {
        "source":"one_c_history", "source_kind":"historical_purchase",
        "snapshot_version":"history-snapshot-v9", "history_item_id":"history-item-1",
        "selected_event_id":"history-event-1", "purchase_date":"2025-01-24",
        "price_basis":"gross_including_vat", "effective_unit_price_gross":"1234.5600",
        "currency_basis":"RUB", "unit_family":"count",
    }
    run_id = _persist_price_export_run(main, repository, workspace, source_row, provider="one_c_history", provenance=provenance, history=True)
    path = repository.workspace_root / workspace["tender_id"]
    run = main.tender_sourcing_runs.get_public(path, workspace["tender_id"], run_id)
    historical_decision = main.tender_price_resolver.resolve_run(
        workspace, run, tender_id=workspace["tender_id"], run_id=run_id,
        include_historical_prices=True,
    )[0]
    assert historical_decision.eligible, historical_decision.safe_summary()
    response = _api_request(main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export", headers={**_auth_headers(), "Content-Type":"application/json"}, body=b'{"include_historical_prices":false,"allow_partial":false}')
    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["code"] == "TENDER_EXPORT_HISTORICAL_CONFIRMATION_REQUIRED"
    assert detail["summary"]["historical_count"] == 1
    assert not list((path / "exports").glob("*")) if (path / "exports").exists() else True
    assert not any(job.kind == "manual_tender_price_export" for job in main.job_service.jobs.values())
    assert isinstance(main.tender_price_exports, TenderPriceExportRepository)


def test_tender_price_export_job_persists_owner_bound_verified_artifact(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_price_export_run(main, repository, workspace, source_row)
    path = repository.workspace_root / workspace["tender_id"]
    response = _api_request(
        main.app, "POST",
        f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":false,"allow_partial":false}',
    )
    assert response.status_code == 202, response.text
    job_id = response.json()["id"]
    job = None
    for _ in range(200):
        snapshot = _api_request(main.app, "GET", f"/api/jobs/{job_id}", headers=_auth_headers())
        assert snapshot.status_code == 200
        job = snapshot.json()
        if job["status"] in {"completed", "failed", "expired"}:
            break
        time.sleep(0.025)
    assert job and job["status"] == "completed", job
    record = job["result"]
    assert record["run_id"] == run_id
    assert record["target_columns"] == {"unit_price":"H", "total":"I"}
    assert record["priced_count"] == 1 and record["blank_count"] == 0
    assert "C:\\" not in json.dumps(record) and str(repository.root) not in json.dumps(record)
    listed = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/exports", headers=_auth_headers())
    assert listed.status_code == 200 and listed.json()["exports"][0]["export_id"] == record["export_id"]
    wrong_owner = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/exports/{record['export_id']}", headers=_auth_headers("another-user"))
    assert wrong_owner.status_code == 404
    wrong_owner_download = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/exports/{record['export_id']}/download", headers=_auth_headers("another-user"))
    assert wrong_owner_download.status_code == 404
    wrong_owner_post = _api_request(
        main.app, "POST",
        f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers("another-user"), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":false,"allow_partial":false}',
    )
    assert wrong_owner_post.status_code == 404
    malformed_export = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/exports/not-an-id", headers=_auth_headers())
    assert malformed_export.status_code == 404
    malformed_run = _api_request(
        main.app, "POST", f"/api/manual-tenders/{workspace['tender_id']}/runs/not-a-run/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":false,"allow_partial":false}',
    )
    assert malformed_run.status_code == 404
    downloaded = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/exports/{record['export_id']}/download", headers=_auth_headers())
    assert downloaded.status_code == 200
    assert downloaded.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in downloaded.headers["content-disposition"].casefold()
    output = tmp_path / "verified-output.xlsx"
    output.write_bytes(downloaded.content)
    workbook = load_workbook(output, data_only=False, read_only=False)
    sheet = workbook[TEMPLATE_SHEET]
    assert sheet["H1"].value == "Цена за единицу"
    assert sheet["I1"].value == "Общая стоимость"
    assert Decimal(str(sheet["H2"].value)) == Decimal("1234.560000")
    assert Decimal(str(sheet["I2"].value)) == Decimal("2469.12")
    assert sheet["H2"].data_type != "f" and sheet["I2"].data_type != "f"
    assert len(sheet.tables) == 1 and next(iter(sheet.tables.values())).ref == "A1:G501"
    workbook.close()
    assert (path / "exports" / f"{record['export_id']}.json").is_file()
    assert (path / "exports" / f"{record['export_id']}.xlsx").is_file()
    from averon_import.services.manual_tenders.price_export import TenderPriceExportRepository
    restarted = TenderPriceExportRepository(repository)
    assert restarted.get_public(path, workspace["tender_id"], record["export_id"], "username:tender-user")["output_file_sha256"] == record["output_file_sha256"]
    (path / "exports" / f"{record['export_id']}.xlsx").write_bytes(b"corrupt")
    corrupt = _api_request(main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/exports/{record['export_id']}", headers=_auth_headers())
    assert corrupt.status_code == 409


def test_excel_adapter_uses_one_model_field_without_resource_code_leak():
    from averon_import.services.sourcing.product_understanding import build_fallback_intent

    source = {
        "source_row_id": "a" * 32,
        "row_type": "item",
        "name": "Synthetic pump",
        "model": "ABC-100",
        "article": "TRUE-SKU-7",
        "resource_code": "RESOURCE-CODE-9",
        "manufacturer": "Synthetic manufacturer",
        "raw_unit": "шт",
        "quantity": "2",
        "quantity_trusted": True,
    }
    adapted = TenderSourcingRowAdapter.convert(source)
    intent = build_fallback_intent(adapted)

    assert "model" not in adapted
    assert adapted["type_mark"] == "ABC-100"
    assert intent.model == "ABC-100"
    assert intent.source_text.count("ABC-100") == 1
    assert adapted["article"] == intent.article == "TRUE-SKU-7"
    assert all("RESOURCE-CODE-9" not in str(value) for value in adapted.values())
    assert "RESOURCE-CODE-9" not in intent.source_text


def test_tender_price_provenance_is_allowlisted_durable_and_self_contained(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=3)
    items = [row for row in workspace["rows"] if row["row_type"] == "item"]
    store = main.tender_sourcing_runs

    etm_result = _project_result_with_offer(
        items[0], provider="etm_ipro", source_mode="provider_only",
        provenance={
            "source": "etm_ipro", "catalog_version": "etm-mirror-2026-10",
            "source_item_id": "etm-item-42", "price_field": "pricewnds",
            "price_status": "available", "token": "private-token-value",
            "authorization": "Bearer private-auth-value", "secret": "private-secret-value",
            "raw_response": {"items": [{"secret": "nested-secret-value"}]},
            "nested": {"anything": "must not persist"},
        },
    )
    etm_run_id = _persist_project_result(
        store, repository, workspace, etm_result, source_row=items[0], source_mode="provider_only",
    )
    del etm_result

    history_result = _project_result_with_offer(
        items[1], provider="one_c_history", source_mode="one_c_only",
        provenance={
            "source": "one_c_history", "source_kind": "historical_purchase",
            "snapshot_version": "history-snapshot-v9", "history_item_id": "history-item-17",
            "selected_event_id": "history-item-17:row-51", "purchase_date": "2026-09-21",
            "price_basis": "gross_per_unit", "effective_unit_price_gross": Decimal("987.654321"),
            "currency_basis": "company_default", "unit_family": "piece",
            "counterparty": "must not persist", "document_type": "must not persist",
            "source_item_code": "must not persist",
        },
    )
    history_run_id = _persist_project_result(
        store, repository, workspace, history_result, source_row=items[1], source_mode="one_c_only",
    )
    del history_result

    lemana_result = _project_result_with_offer(
        items[2], provider="lemana_b2b", source_mode="provider_only",
        provenance={
            "source": "lemana_b2b", "product_item": "lemana-item-8",
            "mirror_revision": "lemana-mirror-v3", "region_id": "region-4",
            "price_basis": "gross_per_unit", "effective_unit_price_gross": "999.99",
            "currency_basis": "verified_gross",
        },
    )
    lemana_run_id = _persist_project_result(
        store, repository, workspace, lemana_result, source_row=items[2], source_mode="provider_only",
    )
    del lemana_result

    # A fresh store instance reads only the completed records; the in-memory
    # project result objects are gone and the short-lived Job record is unused.
    reloaded = TenderSourcingRunStore(repository)
    monkeypatch.setattr(main, "tender_sourcing_runs", reloaded)
    etm_detail = _api_request(
        main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs/{etm_run_id}",
        headers=_auth_headers(),
    )
    history_detail = _api_request(
        main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs/{history_run_id}",
        headers=_auth_headers(),
    )
    lemana_detail = _api_request(
        main.app, "GET", f"/api/manual-tenders/{workspace['tender_id']}/runs/{lemana_run_id}",
        headers=_auth_headers(),
    )
    assert etm_detail.status_code == history_detail.status_code == lemana_detail.status_code == 200
    etm_provenance = etm_detail.json()["rows"][0]["price_provenance"]
    assert etm_provenance == {
        "source": "etm_ipro", "catalog_version": "etm-mirror-2026-10",
        "source_item_id": "etm-item-42", "price_field": "pricewnds", "price_status": "available",
    }
    history_provenance = history_detail.json()["rows"][0]["price_provenance"]
    assert history_provenance == {
        "source": "one_c_history", "source_kind": "historical_purchase",
        "snapshot_version": "history-snapshot-v9", "history_item_id": "history-item-17",
        "selected_event_id": "history-item-17:row-51", "purchase_date": "2026-09-21",
        "price_basis": "gross_per_unit", "effective_unit_price_gross": "987.654321",
        "currency_basis": "company_default", "unit_family": "piece",
    }
    route = history_detail.json()["rows"][0]["route"]
    assert route["final_source_kind"] == "historical_purchase"
    assert route["history_outcome"] == "SAFE_MATCH"
    assert route["history_selected_event_id"] == "history-item-17:row-51"
    assert route["history_purchase_date"] == "2026-09-21"
    lemana_provenance = lemana_detail.json()["rows"][0]["price_provenance"]
    assert lemana_provenance == {
        "source": "lemana_b2b", "product_item": "lemana-item-8",
        "mirror_revision": "lemana-mirror-v3",
    }
    assert not any("gross" in key or "vat" in key for key in lemana_provenance)
    from averon_import.services.manual_tenders.sourcing import _safe_price_provenance
    bounded = _safe_price_provenance({
        "provider": "etm_ipro",
        "data_provenance": {
            "source": "etm_ipro", "catalog_version": "v" * 200,
            "source_item_id": {"nested": "not a scalar"},
        },
    })
    assert len(bounded["catalog_version"]) == 120
    assert "source_item_id" not in bounded

    run_dir = repository.workspace_root / workspace["tender_id"] / "runs"
    persisted_json = "\n".join(path.read_text(encoding="utf-8") for path in run_dir.glob("*.json"))
    for forbidden in (
        "private-token-value", "private-auth-value", "private-secret-value",
        "raw_response", "nested-secret-value", "must not persist", "region-4",
    ):
        assert forbidden not in persisted_json
    public_json = json.dumps([etm_detail.json(), history_detail.json(), lemana_detail.json()], ensure_ascii=False)
    for forbidden in ("token", "authorization", "raw_response", "counterparty", "document_type", "verified_gross"):
        assert forbidden not in public_json.casefold()
    assert all(
        "data_provenance" not in row
        for detail in (etm_detail, history_detail, lemana_detail)
        for row in detail.json()["rows"]
    )


def test_realistic_370_row_tender_run_fits_existing_one_mib_bound(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, count=370)
    source_rows = [row for row in workspace["rows"] if row["row_type"] == "item"]
    selected_ids = [row["source_row_id"] for row in source_rows]
    results = []
    for index, row in enumerate(source_rows, start=1):
        offer_id = f"etm_ipro:synthetic-{index}"
        results.append({
            "intent": {
                "source_row_id": row["source_row_id"],
                "normalized_name": f"Industrial centrifugal pump assembly {index:03d} for process-water service",
                "product_class": "pump equipment",
                "manufacturer": "Synthetic Industrial Equipment Works",
                "brand": "Synthetic Works",
                "model": f"SP-{index:03d}-A",
                "article": f"SYN-ARTICLE-{index:04d}",
                "quantity": str(row.get("quantity") or "2"),
                "unit": str(row.get("raw_unit") or "шт"),
            },
            "recommended_offer": {
                "offer_id": offer_id, "provider": "etm_ipro",
                "source_item_id": f"ETM-SYN-{index:06d}",
                "title": f"Synthetic industrial pump, model SP-{index:03d}-A, cast housing, standard motor",
                "article": f"SYN-ARTICLE-{index:04d}",
                "manufacturer": "Synthetic Industrial Equipment Works",
                "brand": "Synthetic Works", "price": "123456.78", "currency": "RUB",
                "price_unit": "шт", "retrieved_at": "2026-10-01T12:00:00+00:00",
                "availability": True, "availability_text": "In stock; synthetic test data",
                "data_provenance": {
                    "source": "etm_ipro", "catalog_version": "synthetic-etm-catalog-v2026-10",
                    "source_item_id": f"ETM-SYN-{index:06d}",
                    "price_field": "pricewnds", "price_status": "available",
                    "raw_response": {"excluded": "never persisted"},
                },
            },
            "match_results": [{
                "offer": {"offer_id": offer_id}, "decision": "MATCH", "rank": 1,
                "matched_attributes": ["product_class", "model", "manufacturer"],
                "supporting_attributes": ["article"],
                "conflicting_attributes": [], "missing_attributes": [],
            }],
            "route": {
                "source_mode": "provider_only", "final_source_kind": "provider",
                "fallback_status": "not_called", "history_outcome": "NO_MATCH",
                "history_catalog_version": "", "history_selected_event_id": "",
                "history_purchase_date": "", "history_age_days": None,
                "fallback_called": False, "fallback_provider_key": "",
                "fallback_catalog_version": "", "routing_policy_revision": "one-c-routing-v2",
            },
            "notices": [],
        })

    rows = canonical_tender_projection(
        {"results": results}, source_rows, selected_ids,
    )
    workspace_path = repository.workspace_root / workspace["tender_id"]
    store = main.tender_sourcing_runs
    running = store.create_running(
        workspace_path, workspace, source_mode="provider_only", provider=None,
        selected_ids=selected_ids, history_catalog_version=None,
    )
    store.complete(
        workspace_path, running["run_id"],
        summary={
            "positions_total": 370, "positions_processed": 370, "positions_matched": 370,
            "positions_review": 0, "positions_without_offers": 0,
        },
        catalog_version="synthetic-etm-catalog-v2026-10",
        history_catalog_version=None,
        rows=rows,
    )
    run_path = workspace_path / "runs" / f"{running['run_id']}.json"
    assert len(rows) == 370
    assert run_path.stat().st_size < MAX_TENDER_RUN_BYTES
    assert MAX_TENDER_RUN_BYTES == 1024 * 1024


def test_tender_ui_uses_id_only_request_default_selection_runs_and_shared_result_renderer():
    root = Path(__file__).resolve().parents[1]
    template = (root / "averon_import" / "templates" / "index.html").read_text(encoding="utf-8")
    script = (root / "averon_import" / "static" / "app.js").read_text(encoding="utf-8")
    assert 'id="tender-sourcing-button">Подобрать предложения' in template
    assert "selectedIds = new Set((workspace.rows || []).filter((row) => row.row_type === \"item\").map((row) => row.source_row_id))" in script
    assert 'body:JSON.stringify({source_row_ids:selectedIds,source_mode:mode,limit:20})' in script
    assert 'body:JSON.stringify({source_row_ids:selectedIds,source_mode:mode,limit:20})' in script
    assert "renderSourcingResult(job.result)" in script
    assert 'api(`/api/manual-tenders/${encodeURIComponent(tenderId)}/runs`)' in script
    assert "state.excelTender.pollGeneration === generation" in script
    assert "Подбор был прерван перезапуском сервера. Запустите его повторно." in script
    assert "Структура шаблона распознана." in script
    assert "Ошибок структуры нет." not in script
    assert "state.excelTender.sourcingActive = true" in script
    assert "state.excelTender.pollGeneration === generation" in script


def _direct_tender_price_decision(
    *,
    source_unit="шт",
    offer_unit="шт",
    price="2.5",
    quantity="3",
    provider="etm_ipro",
    price_field="pricewnds",
    match_decision="MATCH",
    currency="RUB",
    include_historical_prices=False,
    offer_id="synthetic-offer-1",
    offer_source_item_id="synthetic-item-1",
    provenance_patch=None,
    route_patch=None,
    run_patch=None,
    no_offer=False,
):
    from averon_import.services.manual_tenders.price_export import TenderPriceResolver

    source = {
        "source_row_id": "a" * 32,
        "excel_row": 21,
        "row_type": "item",
        "quantity": quantity,
        "quantity_trusted": True,
        "raw_unit": source_unit,
        "unit_basis": parse_unit_basis(source_unit),
    }
    historical = provider == "one_c_history"
    purchase_date = datetime.now(timezone.utc).date().isoformat()
    if historical:
        offer_id = offer_id or "one_c_history:history-item-1"
        offer_source_item_id = offer_source_item_id or "history-item-1"
        provenance = {
            "source": "one_c_history",
            "source_kind": "historical_purchase",
            "snapshot_version": "history-snapshot-1",
            "history_item_id": "history-item-1",
            "selected_event_id": "history-event-1",
            "purchase_date": purchase_date,
            "price_basis": "gross_including_vat",
            "effective_unit_price_gross": str(price),
            "currency_basis": "company_default",
            "unit_family": "count",
        }
        route = {
            "final_source_kind": "historical_purchase",
            "history_outcome": "SAFE_MATCH",
            "history_safe_basis": "EXACT_ARTICLE",
            "history_catalog_version": "history-snapshot-1",
            "history_selected_event_id": "history-event-1",
            "history_purchase_date": purchase_date,
        }
        run = {"source_mode": "one_c_only", "history_catalog_version": "history-snapshot-1"}
    else:
        provenance = {
            "source": provider,
            "catalog_version": "etm-catalog-1",
            "source_item_id": "synthetic-item-1",
            "price_field": price_field,
        }
        route = {"final_source_kind": "provider"}
        run = {"source_mode": "provider_only"}
    if provenance_patch:
        provenance.update(provenance_patch)
    if route_patch:
        route.update(route_patch)
    if run_patch:
        run.update(run_patch)

    offer = None if no_offer else {
        "offer_id": offer_id,
        "provider": provider,
        "source_item_id": offer_source_item_id,
        "price": price,
        "currency": currency,
        "price_unit": offer_unit,
    }
    canonical = {
        "recommended_offer": offer,
        "recommended_match": None if no_offer else {
            "offer_id": offer_id,
            "decision": match_decision,
        },
        "provider_source": {"provider": provider},
        "route": route,
        "price_provenance": provenance,
    }
    return TenderPriceResolver().resolve(
        source, canonical, run,
        include_historical_prices=include_historical_prices,
    )


def test_price_export_provider_and_match_authority_is_fail_closed():
    for decision in ("MATCH", "LIKELY_MATCH"):
        accepted = _direct_tender_price_decision(match_decision=decision)
        assert accepted.eligible and accepted.source_kind == "etm_ipro"

    for decision in ("ALTERNATIVE", "REVIEW", "REJECT"):
        rejected = _direct_tender_price_decision(match_decision=decision)
        assert not rejected.eligible and rejected.reason_code == "MATCH_NOT_EXPORTABLE"

    cases = (
        (_direct_tender_price_decision(price_field="price"), "PRICE_BASIS_UNPROVEN"),
        (_direct_tender_price_decision(currency="USD"), "CURRENCY_UNSUPPORTED"),
        (_direct_tender_price_decision(provider="lemana_b2b", price_field="pricewnds"), "PRICE_BASIS_UNPROVEN"),
        (_direct_tender_price_decision(provider="future_provider", price_field="pricewnds"), "PRICE_BASIS_UNPROVEN"),
        (_direct_tender_price_decision(no_offer=True), "NO_RECOMMENDED_OFFER"),
        (_direct_tender_price_decision(offer_id="offer-a", provenance_patch={"source_item_id":"other-item"}), "PROVENANCE_MISMATCH"),
        (_direct_tender_price_decision(offer_id="offer-a", provenance_patch={"source":"other-provider"}), "PROVENANCE_MISMATCH"),
    )
    for rejected, reason in cases:
        assert not rejected.eligible and rejected.reason_code == reason


def test_price_export_historical_1c_requires_complete_safe_provenance_and_consent():
    accepted = _direct_tender_price_decision(
        provider="one_c_history", offer_id="one_c_history:history-item-1",
        offer_source_item_id="history-item-1", include_historical_prices=True,
    )
    assert accepted.eligible and accepted.historical

    no_consent = _direct_tender_price_decision(
        provider="one_c_history", offer_id="one_c_history:history-item-1",
        offer_source_item_id="history-item-1", include_historical_prices=False,
    )
    assert no_consent.reason_code == "HISTORICAL_PRICE_NOT_INCLUDED"

    rejected_cases = (
        (_direct_tender_price_decision(
            provider="one_c_history", offer_id="one_c_history:history-item-1",
            offer_source_item_id="history-item-1", include_historical_prices=True,
            route_patch={"history_outcome":"REVIEW"},
        ), "HISTORY_NOT_SAFE"),
        (_direct_tender_price_decision(
            provider="one_c_history", offer_id="one_c_history:history-item-1",
            offer_source_item_id="history-item-1", include_historical_prices=True,
            provenance_patch={"history_item_id":"different-item"},
        ), "PROVENANCE_MISMATCH"),
        (_direct_tender_price_decision(
            provider="one_c_history", offer_id="one_c_history:history-item-1",
            offer_source_item_id="history-item-1", include_historical_prices=True,
            route_patch={"history_selected_event_id":"different-event"},
        ), "PROVENANCE_MISMATCH"),
        (_direct_tender_price_decision(
            provider="one_c_history", offer_id="one_c_history:history-item-1",
            offer_source_item_id="history-item-1", include_historical_prices=True,
            run_patch={"history_catalog_version":"different-snapshot"},
        ), "PROVENANCE_MISMATCH"),
        (_direct_tender_price_decision(
            provider="one_c_history", offer_id="one_c_history:history-item-1",
            offer_source_item_id="history-item-1", include_historical_prices=True,
            provenance_patch={"price_basis":"gross_per_unit"},
        ), "PRICE_BASIS_UNPROVEN"),
        (_direct_tender_price_decision(
            provider="one_c_history", offer_id="one_c_history:history-item-1",
            offer_source_item_id="history-item-1", include_historical_prices=True,
            provenance_patch={"purchase_date":"2999-01-01"},
        ), "PROVENANCE_MISMATCH"),
    )
    for rejected, reason in rejected_cases:
        assert not rejected.eligible and rejected.reason_code == reason


@pytest.mark.parametrize(
    ("source_unit", "offer_unit", "expected_price"),
    [
        ("шт", "шт", "2.500000"),
        ("100 шт", "шт", "250.000000"),
        ("10 шт", "шт", "25.000000"),
        ("1000 шт", "шт", "2500.000000"),
        ("м", "м", "2.500000"),
        ("10 м", "м", "25.000000"),
        ("1000 м", "м", "2500.000000"),
        ("т", "кг", "2500.000000"),
        ("кг", "т", "0.002500"),
    ],
)
def test_price_export_unit_scales_use_decimal(source_unit, offer_unit, expected_price):
    decision = _direct_tender_price_decision(source_unit=source_unit, offer_unit=offer_unit)
    assert decision.eligible
    assert decision.source_unit_price == Decimal(expected_price)
    assert isinstance(decision.source_unit_price, Decimal)
    assert isinstance(decision.total_price, Decimal)


def test_price_export_unknown_package_incompatible_and_rounding_policies():
    for source_unit, offer_unit, reason in (
        ("затрата", "шт", "SOURCE_UNIT_UNTRUSTED"),
        ("компл", "компл", "SOURCE_UNIT_UNTRUSTED"),
        ("уп", "уп", "SOURCE_UNIT_UNTRUSTED"),
        ("шт", "unknown", "OFFER_UNIT_UNTRUSTED"),
        ("шт", "м", "UNIT_INCOMPATIBLE"),
    ):
        decision = _direct_tender_price_decision(source_unit=source_unit, offer_unit=offer_unit)
        assert not decision.eligible and decision.reason_code == reason

    half_up = _direct_tender_price_decision(price="1.2345675", quantity="3")
    assert half_up.source_unit_price == Decimal("1.234568")
    assert half_up.total_price == Decimal("3.70")

    rounded_first = _direct_tender_price_decision(price="1.0000004", quantity="12500")
    assert rounded_first.source_unit_price == Decimal("1.000000")
    assert rounded_first.total_price == Decimal("12500.00")


def test_export_repository_retention_download_pin_restart_and_quota(tender_api, tmp_path, monkeypatch):
    from averon_import.services.manual_tenders.price_export import (
        MAX_COMPLETED_EXPORTS_PER_WORKSPACE,
        MAX_TENDER_STORAGE_BYTES as EXPORT_STORAGE_LIMIT,
        TenderPriceExportRepository,
    )

    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_price_export_run(main, repository, workspace, source_row)
    workspace_path = repository.workspace_root / workspace["tender_id"]
    workspace = repository.verify_sourcing_snapshot(
        workspace_path, workspace["tender_id"], "username:tender-user",
        source_sha256=workspace["source_sha256"], revision=workspace["revision"],
    )
    run = main.tender_sourcing_runs.get_public(workspace_path, workspace["tender_id"], run_id)
    decisions = main.tender_price_resolver.resolve_run(
        workspace, run, tender_id=workspace["tender_id"], run_id=run_id,
        include_historical_prices=False,
    )

    records = [
        main.tender_price_exports.create_completed(
            workspace_path, workspace, run, decisions,
            owner_id="username:tender-user", allow_partial=False,
            include_historical_prices=False, builder=main.tender_xlsx_price_exporter,
        )
        for _ in range(MAX_COMPLETED_EXPORTS_PER_WORKSPACE)
    ]
    first_id, second_id, third_id = [item["export_id"] for item in records]
    _download_path, _metadata, release = main.tender_price_exports.acquire_download(
        workspace_path, workspace["tender_id"], first_id, "username:tender-user",
    )
    fourth = main.tender_price_exports.create_completed(
        workspace_path, workspace, run, decisions,
        owner_id="username:tender-user", allow_partial=False,
        include_historical_prices=False, builder=main.tender_xlsx_price_exporter,
    )
    retained = {item["export_id"] for item in main.tender_price_exports.list_public(
        workspace_path, workspace["tender_id"], "username:tender-user",
    )["exports"]}
    assert retained == {first_id, third_id, fourth["export_id"]}
    assert second_id not in retained

    # A completed file and its checksum-backed metadata remain usable after
    # reconstructing the repository service, like they would after app restart.
    restarted = TenderPriceExportRepository(repository)
    assert restarted.get_public(
        workspace_path, workspace["tender_id"], first_id, "username:tender-user",
    )["output_file_sha256"] == records[0]["output_file_sha256"]

    export_dir = workspace_path / "exports"
    orphan = export_dir / f".{uuid.uuid4().hex}.{uuid.uuid4().hex}.tmp.xlsx"
    orphan.write_bytes(b"incomplete")
    old = datetime.now(timezone.utc).timestamp() - 7200
    os.utime(orphan, (old, old))
    TenderPriceExportRepository(repository)
    assert not orphan.exists()

    # The output XLSX is already part of workspace bytes; the new metadata JSON
    # must be included too, even when no older export can be reclaimed.
    release()
    for existing in export_dir.iterdir():
        if existing.is_file():
            existing.unlink()
    monkeypatch.setattr(repository, "_workspace_storage_bytes", lambda: EXPORT_STORAGE_LIMIT - 1)
    with pytest.raises(TenderWorkspaceError) as quota:
        main.tender_price_exports.create_completed(
            workspace_path, workspace, run, decisions,
            owner_id="username:tender-user", allow_partial=False,
            include_historical_prices=False, builder=main.tender_xlsx_price_exporter,
        )
    assert quota.value.code == "TENDER_DISK_QUOTA"
    assert not list(export_dir.glob("[a-f0-9]" * 32 + ".json"))


def test_export_semantic_verifier_tracks_workbook_controls(tmp_path):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from averon_import.services.manual_tenders.price_export import _verify_semantic_roundtrip, _workbook_snapshot

    source = tmp_path / "controls-source.xlsx"
    output = tmp_path / "controls-output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Synthetic"
    sheet.auto_filter.ref = "A1:A2"
    validation = DataValidation(type="whole", operator="greaterThan", formula1="0")
    validation.add("A2")
    sheet.add_data_validation(validation)
    sheet.conditional_formatting.add(
        "A2", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill(fill_type="solid", fgColor="FFFF0000")),
    )
    sheet.protection.sheet = True
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    before = _workbook_snapshot(source)
    after = _workbook_snapshot(output)
    assert before["sheets"][0]["data_validations"] == after["sheets"][0]["data_validations"]
    assert before["sheets"][0]["conditional_formatting"] == after["sheets"][0]["conditional_formatting"]
    _verify_semantic_roundtrip(before, after, target_sheet="Sheet", allowed_cells=set(), target_columns=set())


def test_historical_price_export_with_consent_is_visibly_marked(tender_api, tmp_path):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    provenance = {
        "source": "one_c_history", "source_kind": "historical_purchase",
        "snapshot_version": "history-snapshot-v9", "history_item_id": "history-item-1",
        "selected_event_id": "history-event-1", "purchase_date": "2025-01-24",
        "price_basis": "gross_including_vat", "effective_unit_price_gross": "1234.5600",
        "currency_basis": "RUB", "unit_family": "count",
    }
    run_id = _persist_price_export_run(
        main, repository, workspace, source_row, provider="one_c_history",
        provenance=provenance, history=True,
    )
    response = _api_request(
        main.app, "POST",
        f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
        headers={**_auth_headers(), "Content-Type":"application/json"},
        body=b'{"include_historical_prices":true,"allow_partial":false}',
    )
    assert response.status_code == 202, response.text
    job = _wait_tender_job(main, response.json()["id"])
    assert job["status"] == "completed", job.get("error") or job
    assert job["result"]["historical_count"] == 1
    workspace_path = repository.workspace_root / workspace["tender_id"]
    file_path, _record, release = main.tender_price_exports.acquire_download(
        workspace_path, workspace["tender_id"], job["result"]["export_id"], "username:tender-user",
    )
    try:
        workbook = load_workbook(file_path, data_only=False)
        sheet = workbook[TEMPLATE_SHEET]
        assert isinstance(sheet["H2"].value, (int, float, Decimal))
        assert isinstance(sheet["I2"].value, (int, float, Decimal))
        assert sheet["H2"].comment and "2025-01-24" in sheet["H2"].comment.text
        assert len(sheet["H2"].comment.text) > 75
        assert sheet["I2"].comment and "2025-01-24" in sheet["I2"].comment.text
        assert len(sheet["I2"].comment.text) > 75
        assert sheet["H2"].fill.fill_type == "solid"
        assert sheet["H1"].comment and len(sheet["H1"].comment.text) > 75
        workbook.close()
    finally:
        release()


def test_price_export_job_holds_lease_through_deduped_completion(tender_api, tmp_path, monkeypatch):
    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_price_export_run(main, repository, workspace, source_row)
    entered = threading.Event()
    release_writer = threading.Event()
    original_write = main.tender_xlsx_price_exporter.write

    def delayed_write(*args, **kwargs):
        entered.set()
        assert release_writer.wait(5)
        return original_write(*args, **kwargs)

    monkeypatch.setattr(main.tender_xlsx_price_exporter, "write", delayed_write)

    def post(username="tender-user"):
        return _api_request(
            main.app, "POST",
            f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
            headers={**_auth_headers(username), "Content-Type":"application/json"},
            body=b'{"include_historical_prices":false,"allow_partial":false}',
        )

    first = post()
    assert first.status_code == 202 and entered.wait(2)
    duplicate = post()
    assert duplicate.status_code == 202 and duplicate.json()["id"] == first.json()["id"]
    assert post("another-user").status_code == 404
    assert repository.activity.active(workspace["tender_id"])
    blocked_delete = _api_request(
        main.app, "DELETE", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers(),
    )
    assert blocked_delete.status_code == 409

    release_writer.set()
    job = _wait_tender_job(main, first.json()["id"])
    assert job["status"] == "completed", job
    assert not repository.activity.active(workspace["tender_id"])
    assert len(list((repository.workspace_root / workspace["tender_id"] / "exports").glob("*.json"))) == 1
    assert _api_request(
        main.app, "DELETE", f"/api/manual-tenders/{workspace['tender_id']}", headers=_auth_headers(),
    ).status_code == 200


def test_price_export_failure_admission_and_queue_expiry_release_leases(tender_api, tmp_path, monkeypatch):
    from averon_import.services.jobs import DOCUMENT_PROCESSING, JobService, SOURCING

    main, repository = tender_api
    workspace = _confirm_synthetic_tender(main, repository, tmp_path, official=True)
    source_row = next(row for row in workspace["rows"] if row["row_type"] == "item")
    run_id = _persist_price_export_run(main, repository, workspace, source_row)

    def post():
        return _api_request(
            main.app, "POST",
            f"/api/manual-tenders/{workspace['tender_id']}/runs/{run_id}/export",
            headers={**_auth_headers(), "Content-Type":"application/json"},
            body=b'{"include_historical_prices":false,"allow_partial":false}',
        )

    export_dir = repository.workspace_root / workspace["tender_id"] / "exports"
    monkeypatch.setattr(
        main.tender_xlsx_price_exporter, "write",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            TenderWorkspaceError("Книга не прошла проверку сохранения.", 409, "TENDER_EXPORT_PRESERVATION_FAILED")
        ),
    )
    failed = post()
    assert failed.status_code == 202
    failed_job = _wait_tender_job(main, failed.json()["id"])
    assert failed_job["status"] == "failed"
    assert failed_job["error_code"] == "TENDER_EXPORT_PRESERVATION_FAILED"
    assert not repository.activity.active(workspace["tender_id"])
    assert not list(export_dir.glob("*.json")) and not list(export_dir.glob("*.xlsx"))

    admission_jobs = JobService(capacities={DOCUMENT_PROCESSING:(1,0), SOURCING:(1,0)})
    monkeypatch.setattr(main, "job_service", admission_jobs)
    blocker_started = threading.Event()
    release_blocker = threading.Event()
    admission_jobs.submit(
        lambda _progress: (blocker_started.set(), release_blocker.wait(5))[1],
        lane=DOCUMENT_PROCESSING, owner_id="synthetic-admission-blocker",
    )
    assert blocker_started.wait(2)
    rejected = post()
    assert rejected.status_code == 409 and rejected.json()["detail"]["code"] == "JOB_LANE_BUSY"
    assert not repository.activity.active(workspace["tender_id"])
    assert not list(export_dir.glob("*.json")) and not list(export_dir.glob("*.xlsx"))
    release_blocker.set()
    admission_jobs.executor.shutdown(wait=True)

    expiry_jobs = JobService(queue_wait_seconds=0.05)
    monkeypatch.setattr(main, "job_service", expiry_jobs)
    blocker_started = threading.Event()
    release_blocker = threading.Event()
    expiry_jobs.submit(
        lambda _progress: (blocker_started.set(), release_blocker.wait(5))[1],
        lane=DOCUMENT_PROCESSING, owner_id="synthetic-expiry-blocker",
    )
    assert blocker_started.wait(2)
    queued = post()
    assert queued.status_code == 202
    time.sleep(0.08)
    expired = _wait_tender_job(main, queued.json()["id"])
    assert expired["status"] == "expired"
    assert not repository.activity.active(workspace["tender_id"])
    assert not list(export_dir.glob("*.json")) and not list(export_dir.glob("*.xlsx"))
    release_blocker.set()
    expiry_jobs.executor.shutdown(wait=True)


def test_external_xlsx_noop_roundtrip_accepts_equivalent_ooxml_encodings(tmp_path):
    from xml.etree import ElementTree as ET
    from averon_import.services.manual_tenders.price_export import (
        _verify_package_roundtrip,
        _workbook_snapshot,
        _verify_semantic_roundtrip,
    )

    source = tmp_path / "external-source.xlsx"
    output = tmp_path / "external-output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "External source"
    sheet.append(["Description", "Quantity"])
    sheet.append(["A source row", 4])
    from openpyxl.worksheet.table import Table
    sheet.add_table(Table(displayName="ExternalTable", ref="A1:B2"))
    workbook.save(source)
    workbook.close()

    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    ET.register_namespace("", main_ns)
    ET.register_namespace("r", office_rel_ns)
    ET.register_namespace("pkg", rel_ns)

    with zipfile.ZipFile(source, "r") as archive:
        parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
    # Model an external producer: shared strings rather than inline strings,
    # no generated document properties, alternate relationship IDs and a
    # prefixed table namespace. All are valid equivalent OOXML encodings.
    sheet_root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    shared_values = []
    shared_index = {}
    for cell in sheet_root.iter(f"{{{main_ns}}}c"):
        inline = cell.find(f"{{{main_ns}}}is")
        if inline is None:
            continue
        text_node = inline.find(f"{{{main_ns}}}t")
        value = text_node.text if text_node is not None and text_node.text is not None else ""
        if value not in shared_index:
            shared_index[value] = len(shared_values)
            shared_values.append(value)
        cell.attrib["t"] = "s"
        cell.remove(inline)
        ET.SubElement(cell, f"{{{main_ns}}}v").text = str(shared_index[value])
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
    sst = ET.Element(f"{{{main_ns}}}sst", {"count":str(len(shared_values)), "uniqueCount":str(len(shared_values))})
    for value in shared_values:
        item = ET.SubElement(sst, f"{{{main_ns}}}si")
        ET.SubElement(item, f"{{{main_ns}}}t").text = value
    parts["xl/sharedStrings.xml"] = ET.tostring(sst, encoding="utf-8", xml_declaration=True)

    workbook_rels_name = "xl/_rels/workbook.xml.rels"
    workbook_rels = ET.fromstring(parts[workbook_rels_name])
    original_to_new = {}
    for index, relation in enumerate(workbook_rels.findall(f"{{{rel_ns}}}Relationship"), start=1):
        old = relation.attrib["Id"]
        new = f"externalRel{index}"
        original_to_new[old] = new
        relation.attrib["Id"] = new
    ET.SubElement(workbook_rels, f"{{{rel_ns}}}Relationship", {
        "Id":"externalSharedStrings",
        "Type":f"{office_rel_ns}/sharedStrings",
        "Target":"sharedStrings.xml",
    })
    parts[workbook_rels_name] = ET.tostring(workbook_rels, encoding="utf-8", xml_declaration=True)
    workbook_root = ET.fromstring(parts["xl/workbook.xml"])
    for element in workbook_root.iter():
        relation_id = element.attrib.get(f"{{{office_rel_ns}}}id")
        if relation_id in original_to_new:
            element.attrib[f"{{{office_rel_ns}}}id"] = original_to_new[relation_id]
    parts["xl/workbook.xml"] = ET.tostring(workbook_root, encoding="utf-8", xml_declaration=True)

    table_name = next(name for name in parts if name.casefold().startswith("xl/tables/table") and name.endswith(".xml"))
    table_root = ET.fromstring(parts[table_name])
    table_namespace = table_root.tag.split("}", 1)[0][1:]
    ET.register_namespace("tblx", table_namespace)
    parts[table_name] = ET.tostring(table_root, encoding="utf-8", xml_declaration=True)

    package_rels_name = "_rels/.rels"
    package_rels = ET.fromstring(parts[package_rels_name])
    parts[package_rels_name] = ET.tostring(package_rels, encoding="utf-8", xml_declaration=True)
    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for child in list(content_types):
        if child.tag.endswith("Override") and child.attrib.get("PartName", "").casefold() in {"/docprops/app.xml", "/docprops/core.xml"}:
            content_types.remove(child)
    ET.SubElement(content_types, f"{{{content_ns}}}Override", {
        "PartName":"/xl/sharedStrings.xml",
        "ContentType":"application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    })
    parts["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8", xml_declaration=True)
    for name in ("docProps/app.xml", "docProps/core.xml"):
        parts.pop(name, None)
    package_rels = ET.fromstring(parts[package_rels_name])
    for child in list(package_rels):
        if child.attrib.get("Type", "").endswith(("/extended-properties", "/metadata/core-properties")):
            package_rels.remove(child)
    parts[package_rels_name] = ET.tostring(package_rels, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)

    assert "xl/sharedStrings.xml" in zipfile.ZipFile(source).namelist()
    assert not any(name.casefold().startswith("docprops/") for name in zipfile.ZipFile(source).namelist())
    with zipfile.ZipFile(source) as archive:
        source_ids = {
            item.attrib["Id"]
            for item in ET.fromstring(archive.read(workbook_rels_name)).findall(f"{{{rel_ns}}}Relationship")
        }
    opened = load_workbook(source)
    opened.save(output)
    opened.close()
    with zipfile.ZipFile(output) as archive:
        output_ids = {
            item.attrib["Id"]
            for item in ET.fromstring(archive.read(workbook_rels_name)).findall(f"{{{rel_ns}}}Relationship")
        }
    assert source_ids != output_ids
    before = _workbook_snapshot(source)
    after = _workbook_snapshot(output)
    _verify_semantic_roundtrip(before, after, target_sheet="External source", allowed_cells=set(), target_columns=set())
    _verify_package_roundtrip(source, output, semantic_before=before, semantic_after=after)


@pytest.mark.parametrize(
    "mutation",
    ["cell_value", "formula", "style", "row_style", "column_style", "merge", "defined_name_target", "defined_name_scope"],
)
def test_export_semantic_verifier_rejects_source_changes(tmp_path, mutation):
    from openpyxl.styles import PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from averon_import.services.manual_tenders.price_export import _verify_semantic_roundtrip, _workbook_snapshot

    source = tmp_path / "semantic-source.xlsx"
    output = tmp_path / "semantic-output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Source value"
    workbook.defined_names.add(DefinedName("SourceTarget", attr_text="'Sheet'!$A$1"))
    workbook.save(source)
    workbook.close()
    before = _workbook_snapshot(source)

    workbook = load_workbook(source)
    sheet = workbook["Sheet"]
    if mutation == "cell_value":
        sheet["A1"] = "Changed source value"
    elif mutation == "formula":
        sheet["A1"] = "=1+1"
    elif mutation == "style":
        sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    elif mutation == "row_style":
        sheet.row_dimensions[1].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    elif mutation == "column_style":
        sheet.column_dimensions["A"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    elif mutation == "merge":
        sheet.merge_cells("C3:D3")
    elif mutation == "defined_name_target":
        workbook.defined_names["SourceTarget"].attr_text = "'Sheet'!$B$1"
    elif mutation == "defined_name_scope":
        del workbook.defined_names["SourceTarget"]
        sheet.defined_names.add(DefinedName("SourceTarget", attr_text="'Sheet'!$A$1"))
    workbook.save(output)
    workbook.close()
    with pytest.raises(TenderWorkspaceError) as error:
        _verify_semantic_roundtrip(
            before, _workbook_snapshot(output), target_sheet="Sheet", allowed_cells=set(), target_columns=set(),
        )
    assert error.value.code == "TENDER_EXPORT_PRESERVATION_FAILED"


@pytest.mark.parametrize(
    "mutation",
    ["table_ref", "table_column_name_order", "relationship_target", "relationship_type", "effective_content_type", "opaque_loss"],
)
def test_export_package_verifier_rejects_ooxml_mutations(tmp_path, mutation):
    from xml.etree import ElementTree as ET
    from averon_import.services.manual_tenders.price_export import _verify_package_roundtrip

    source = tmp_path / "package-source.xlsx"
    output = tmp_path / "package-output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Description", "Quantity"])
    sheet.append(["A source row", 4])
    from openpyxl.worksheet.table import Table
    sheet.add_table(Table(displayName="PackageTable", ref="A1:B2"))
    workbook.save(source)
    workbook.close()
    opened = load_workbook(source)
    opened.save(output)
    opened.close()

    def entries(path):
        with zipfile.ZipFile(path) as archive:
            return {item.filename: archive.read(item.filename) for item in archive.infolist()}

    parts = entries(output)
    if mutation in {"table_ref", "table_column_name_order"}:
        name = next(part for part in parts if re.fullmatch(r"xl/tables/table\d+\.xml", part, re.IGNORECASE))
        root = ET.fromstring(parts[name])
        if mutation == "table_ref":
            root.attrib["ref"] = "A1:B1"
        else:
            columns = next(node for node in root.iter() if node.tag.endswith("tableColumns"))
            first, second = list(columns)
            first.attrib["name"], second.attrib["name"] = second.attrib["name"], first.attrib["name"]
        parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    elif mutation in {"relationship_target", "relationship_type"}:
        name = "xl/_rels/workbook.xml.rels"
        root = ET.fromstring(parts[name])
        relation = next(node for node in root.iter() if node.tag.endswith("Relationship") and node.attrib.get("Type", "").endswith("/worksheet"))
        if mutation == "relationship_target":
            relation.attrib["Target"] = "worksheets/missing-sheet.xml"
        else:
            relation.attrib["Type"] = "urn:unexpected:worksheet"
        parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    elif mutation == "effective_content_type":
        name = "[Content_Types].xml"
        root = ET.fromstring(parts[name])
        override = next(node for node in root if node.tag.endswith("Override") and node.attrib.get("PartName", "").casefold() == "/xl/workbook.xml")
        override.attrib["ContentType"] = "application/x-unexpected-workbook"
        parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    elif mutation == "opaque_loss":
        parts["custom/opaque.bin"] = b"source-owned opaque package data"
        types = ET.fromstring(parts["[Content_Types].xml"])
        ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", {
            "Extension":"bin", "ContentType":"application/vnd.example.opaque",
        })
        parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
        # The source includes an opaque part the ordinary workbook writer drops.
        source_parts = entries(source)
        source_parts.update({key: value for key, value in parts.items() if key in {"custom/opaque.bin", "[Content_Types].xml"}})
        with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, payload in source_parts.items():
                archive.writestr(name, payload)
        parts.pop("custom/opaque.bin")
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)

    with pytest.raises(TenderWorkspaceError) as error:
        _verify_package_roundtrip(source, output)
    assert error.value.code in {"TENDER_EXPORT_PRESERVATION_FAILED", "TENDER_EXPORT_PRESERVATION_UNSUPPORTED"}


def _rewrite_xlsx_parts(path, transform):
    with zipfile.ZipFile(path) as archive:
        parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
    transform(parts)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)


def _add_known_xlsx_extensions(parts):
    from xml.etree import ElementTree as ET

    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    styles = ET.fromstring(parts["xl/styles.xml"])
    styles_ext = ET.SubElement(styles, f"{{{spreadsheet_ns}}}extLst")
    slicer = ET.SubElement(styles_ext, f"{{{spreadsheet_ns}}}ext", {"uri": "{EB79DEF2-80B8-43e5-95BD-54CBDDF9020C}"})
    ET.SubElement(slicer, "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}slicerStyles", {"defaultSlicerStyle": "DefaultSlicerStyle1"})
    timeline = ET.SubElement(styles_ext, f"{{{spreadsheet_ns}}}ext", {"uri": "{9260A510-F301-46a8-8635-F512D64BE5F5}"})
    ET.SubElement(timeline, "{http://schemas.microsoft.com/office/spreadsheetml/2010/11/main}timelineStyles", {"defaultTimelineStyle": "DefaultTimelineStyle1"})
    parts["xl/styles.xml"] = ET.tostring(styles, encoding="utf-8", xml_declaration=True)

    theme_part = next(name for name in parts if name.casefold() == "xl/theme/theme1.xml")
    theme = ET.fromstring(parts[theme_part])
    theme_ext = ET.SubElement(theme, f"{{{drawing_ns}}}extLst")
    family = ET.SubElement(theme_ext, f"{{{drawing_ns}}}ext", {"uri": "{05A4C25C-085E-4340-85A3-A5531E510DB2}"})
    ET.SubElement(family, "{http://schemas.microsoft.com/office/thememl/2012/main}themeFamily", {
        "id": "{00000000-0000-0000-0000-000000000001}", "name": "SyntheticTheme", "vid": "{00000000-0000-0000-0000-000000000002}",
    })
    parts[theme_part] = ET.tostring(theme, encoding="utf-8", xml_declaration=True)


def test_known_inert_extensions_are_accepted_and_preserved_with_source_docprops(tmp_path):
    from xml.etree import ElementTree as ET
    from averon_import.services.manual_tenders.parser import _package_preservation_inventory
    from averon_import.services.manual_tenders.price_export import (
        _restore_source_owned_package_metadata,
        _verify_package_roundtrip,
    )

    source = tmp_path / "known-extensions.xlsx"
    output = tmp_path / "known-extensions-output.xlsx"
    workbook = Workbook()
    workbook.active.append(["Наименование", "Количество"])
    workbook.active.append(["Тестовая позиция", 1])
    workbook.save(source)
    workbook.close()
    _rewrite_xlsx_parts(source, _add_known_xlsx_extensions)

    # Give the source-owned properties distinct values so their exact bytes
    # must survive openpyxl's save.
    with zipfile.ZipFile(source) as archive:
        app_name = next(name for name in archive.namelist() if name.casefold() == "docprops/app.xml")
        core_name = next(name for name in archive.namelist() if name.casefold() == "docprops/core.xml")
        app_bytes = archive.read(app_name)
        core_bytes = archive.read(core_name)
    app_root = ET.fromstring(app_bytes)
    core_root = ET.fromstring(core_bytes)
    for node in app_root.iter():
        if node.tag.endswith("}Application"):
            node.text = "SourceApplicationMetadata"
        elif node.tag.endswith("}AppVersion"):
            node.text = "17.42"
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    vt_ns = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    headings = ET.SubElement(app_root, f"{{{app_ns}}}HeadingPairs")
    heading_vector = ET.SubElement(headings, f"{{{vt_ns}}}vector", {"size": "2", "baseType": "variant"})
    first_heading = ET.SubElement(heading_vector, f"{{{vt_ns}}}variant")
    ET.SubElement(first_heading, f"{{{vt_ns}}}lpstr").text = "Synthetic heading"
    second_heading = ET.SubElement(heading_vector, f"{{{vt_ns}}}variant")
    ET.SubElement(second_heading, f"{{{vt_ns}}}i4").text = "1"
    titles = ET.SubElement(app_root, f"{{{app_ns}}}TitlesOfParts")
    title_vector = ET.SubElement(titles, f"{{{vt_ns}}}vector", {"size": "1", "baseType": "lpstr"})
    ET.SubElement(title_vector, f"{{{vt_ns}}}lpstr").text = "Synthetic title"
    for node in core_root.iter():
        if node.tag.endswith("}creator"):
            node.text = "SourceCreatorMetadata"
        elif node.tag.endswith("}lastModifiedBy"):
            node.text = "SourceEditorMetadata"
        elif node.tag.endswith("}created"):
            node.text = "2001-02-03T04:05:06Z"
    ET.SubElement(core_root, "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastPrinted").text = "2002-03-04T05:06:07Z"
    app_bytes = ET.tostring(app_root, encoding="utf-8", xml_declaration=True)
    core_bytes = ET.tostring(core_root, encoding="utf-8", xml_declaration=True)
    _rewrite_xlsx_parts(source, lambda parts: parts.update({app_name: app_bytes, core_name: core_bytes}))

    parser = TenderWorkbookParser().parse(source)
    assert parser["manifest"]["unsupported_preservation_sensitive_objects"] == []
    source_extensions, unsupported = _package_preservation_inventory(source)
    assert not unsupported
    assert set(source_extensions) == {"xl/styles.xml", "xl/theme/theme1.xml"}

    workbook = load_workbook(source)
    workbook.save(output)
    workbook.close()
    _restore_source_owned_package_metadata(source, output)
    with zipfile.ZipFile(output) as archive:
        assert archive.read(app_name) == app_bytes
        assert archive.read(core_name) == core_bytes
    output_extensions, output_unsupported = _package_preservation_inventory(output)
    assert not output_unsupported
    assert set(output_extensions) == set(source_extensions)
    _verify_package_roundtrip(source, output)


def test_export_package_verifier_rejects_dropped_known_extension_lists(tmp_path):
    from averon_import.services.manual_tenders.price_export import _verify_package_roundtrip

    source = tmp_path / "known-extension-source.xlsx"
    output = tmp_path / "known-extension-dropped.xlsx"
    workbook = Workbook()
    workbook.save(source)
    workbook.close()
    _rewrite_xlsx_parts(source, _add_known_xlsx_extensions)
    workbook = load_workbook(source)
    workbook.save(output)
    workbook.close()
    with pytest.raises(TenderWorkspaceError) as error:
        _verify_package_roundtrip(source, output)
    assert error.value.code == "TENDER_EXPORT_PRESERVATION_FAILED"


def test_newly_generated_docprops_allow_standard_values_but_reject_unexpected_fields(tmp_path):
    from xml.etree import ElementTree as ET
    from averon_import.services.manual_tenders.price_export import (
        _generated_docprops_are_safe,
        _workbook_snapshot,
    )

    source = tmp_path / "no-docprops-source.xlsx"
    output = tmp_path / "generated-docprops.xlsx"
    workbook = Workbook()
    workbook.active.append(["Наименование", "Количество"])
    workbook.active.append(["Синтетическая строка", 1])
    workbook.save(source)
    workbook.close()
    workbook = load_workbook(source)
    workbook.save(output)
    workbook.close()
    source_snapshot = _workbook_snapshot(source)
    assert _generated_docprops_are_safe(output, {"docprops/app.xml", "docprops/core.xml"}, source_snapshot)

    with zipfile.ZipFile(output) as archive:
        app_name = next(name for name in archive.namelist() if name.casefold() == "docprops/app.xml")
        app_bytes = archive.read(app_name)
    app_root = ET.fromstring(app_bytes)
    ET.SubElement(app_root, "{urn:unexpected}EmbeddedSourceData").text = "unexpected"
    malicious = ET.tostring(app_root, encoding="utf-8", xml_declaration=True)
    _rewrite_xlsx_parts(output, lambda parts: parts.update({app_name: malicious}))
    assert not _generated_docprops_are_safe(output, {"docprops/app.xml"}, source_snapshot)


@pytest.mark.parametrize("extension_mutation", ["unknown_uri", "unexpected_child"])
def test_unknown_or_malformed_extension_lists_fail_closed(tmp_path, extension_mutation):
    from xml.etree import ElementTree as ET
    from averon_import.services.manual_tenders.parser import _package_preservation_inventory

    path = tmp_path / f"bad-extension-{extension_mutation}.xlsx"
    workbook = Workbook()
    workbook.active.append(["Наименование", "Количество"])
    workbook.active.append(["Тестовая позиция", 1])
    workbook.save(path)
    workbook.close()

    def mutate(parts):
        root = ET.fromstring(parts["xl/styles.xml"])
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        ext_list = ET.SubElement(root, f"{{{ns}}}extLst")
        if extension_mutation == "unknown_uri":
            extension = ET.SubElement(ext_list, f"{{{ns}}}ext", {"uri": "{00000000-0000-0000-0000-000000000000}"})
            ET.SubElement(extension, "{urn:unknown-extension}extra")
        else:
            extension = ET.SubElement(ext_list, f"{{{ns}}}ext", {"uri": "{EB79DEF2-80B8-43e5-95BD-54CBDDF9020C}"})
            ET.SubElement(extension, "{urn:unexpected}customChild")
        parts["xl/styles.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_xlsx_parts(path, mutate)
    extensions, unsupported = _package_preservation_inventory(path)
    assert extensions == {}
    assert "xl/styles.xml#unknown-extension-list" in unsupported
    parsed = TenderWorkbookParser().parse(path)
    assert "xl/styles.xml#unknown-extension-list" in parsed["manifest"]["unsupported_preservation_sensitive_objects"]


@pytest.mark.parametrize(
    "active_part,relationship_type",
    [
        ("xl/slicers/slicer1.xml", None),
        ("xl/timelines/timeline1.xml", None),
        ("xl/pivotTables/pivotTable1.xml", None),
        (None, "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main/slicer"),
    ],
)
def test_known_style_extensions_reject_active_slicer_timeline_or_pivot_objects(tmp_path, active_part, relationship_type):
    from xml.etree import ElementTree as ET
    from averon_import.services.manual_tenders.parser import _package_preservation_inventory

    path = tmp_path / "active-object.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()

    def mutate(parts):
        _add_known_xlsx_extensions(parts)
        if active_part:
            parts[active_part] = b"<?xml version='1.0' encoding='UTF-8'?><activeObject/>"
        if relationship_type:
            rel_name = "xl/worksheets/_rels/sheet1.xml.rels"
            rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            root = ET.Element(f"{{{rel_ns}}}Relationships")
            ET.SubElement(root, f"{{{rel_ns}}}Relationship", {
                "Id": "rIdActive", "Type": relationship_type, "Target": "../worksheets/sheet1.xml",
            })
            parts[rel_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_xlsx_parts(path, mutate)
    _extensions, unsupported = _package_preservation_inventory(path)
    assert any("active" in item.casefold() or any(token in item.casefold() for token in ("slicer", "timeline", "pivot")) for item in unsupported)


def test_added_target_comments_follow_sheet_relationships_and_preserve_other_sheet_comment(tmp_path):
    from averon_import.services.manual_tenders.price_export import (
        _restore_source_owned_package_metadata,
        _verify_package_roundtrip,
        _workbook_snapshot,
    )

    source = tmp_path / "comments-source.xlsx"
    output = tmp_path / "comments-output.xlsx"
    workbook = Workbook()
    other = workbook.active
    other.title = "Other"
    other["B2"] = "Unchanged"
    other["B2"].comment = Comment("Existing note", "Source author")
    target = workbook.create_sheet("Target")
    target["A1"] = "Header"
    target["F2"] = None
    workbook.save(source)
    workbook.close()
    before = _workbook_snapshot(source)

    workbook = load_workbook(source)
    workbook["Target"]["F2"].comment = Comment("New export note", "Averon Import")
    workbook.save(output)
    workbook.close()
    _restore_source_owned_package_metadata(source, output)
    after = _workbook_snapshot(output)
    _verify_package_roundtrip(
        source,
        output,
        target_sheet="Target",
        allowed_cells={"F2"},
        target_comment_cells={"F2"},
        semantic_before=before,
        semantic_after=after,
    )

    with zipfile.ZipFile(output) as archive:
        from averon_import.services.manual_tenders.price_export import _package_snapshot, _worksheet_relationship_part_map
        package = _package_snapshot(output)
        rel_parts = _worksheet_relationship_part_map(output)
        target_comments = {
            target for _rid, rel_type, target, mode in package["relationships"][rel_parts["Target"]]
            if rel_type.casefold().endswith("/comments") and not mode
        }
        other_comments = {
            target for _rid, rel_type, target, mode in package["relationships"][rel_parts["Other"]]
            if rel_type.casefold().endswith("/comments") and not mode
        }
        assert len(target_comments) == len(other_comments) == 1
        assert target_comments.isdisjoint(other_comments)
        assert any(name.casefold().endswith("comment1.xml") for name in archive.namelist())


def test_export_failure_panel_is_visible_safe_and_retryable():
    root = Path(__file__).resolve().parents[1]
    template = (root / "averon_import" / "templates" / "index.html").read_text(encoding="utf-8")
    script = (root / "averon_import" / "static" / "app.js").read_text(encoding="utf-8")
    assert 'id="tender-export-failure" role="alert" aria-live="assertive" hidden' in template
    assert "Не удалось создать Excel" in template and "Повторить экспорт" in template
    assert 'id="tender-export-retry"' in template
    assert "function safeTenderExportErrorMessage(error)" in script
    assert "function showTenderExportFailure(error)" in script
    assert '$("#tender-export-retry").addEventListener("click", () => { void startExcelTenderPriceExport(); })' in script
    assert "localStorage" not in script[script.index("function safeTenderExportErrorMessage(error)"):script.index("function chooseTenderHistoricalPricePolicy")]
