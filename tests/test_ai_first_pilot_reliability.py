"""Synthetic API/storage regressions; no cloud transport or customer fixtures."""
from io import BytesIO
import json

import fitz
import pytest
from openpyxl import load_workbook

from test_ai_first_user_workflow import ai_first_client, _make_document, _make_sized_document, _start_fake, _wait_for_job, _ApiClient
from averon_import.services.ai_first.local_workflow import AIWorkspaceStore
from averon_import.services.ai_first.orchestration import FitZPageRenderer, OrchestrationConfig, TokenPricing
from test_account_auth import session_api, _session_headers


@pytest.fixture()
def private_session_api(monkeypatch, tmp_path, request):
    # session_api imports the app; select test storage before that import even
    # when this test is the first/only collected test in a fresh process.
    monkeypatch.setenv("AVERON_DATA_DIR", str(tmp_path / "session-app-data"))
    return request.getfixturevalue("session_api")


def state(client, workspace):
    response = client.get(f"/api/documents/{workspace.document_id}/ai-first")
    assert response.status_code == 200, response.text
    return response.json()


def post(client, workspace, action, payload):
    return client.post(f"/api/documents/{workspace.document_id}/ai-first/{action}", json=payload)


def edits(**values):
    return {"row_values": {}, "row_roles": {}, "issue_notes": {}, "add_row_page": None,
            "add_row_values": {}, **values}


def test_save_restores_all_editor_buffers_without_confirming_or_changing_ocr(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main)
    _start_fake(client, workspace.document_id)
    before = state(client, workspace)
    page = before["draft"]["pages"][0]
    row_a, row_b = page["tables"][0]["rows"]
    issue = page["issues"][0]
    pending = edits(row_values={row_a["id"]: {"cell:1": "Synthetic edited name", "cell:4": "0"}, row_b["id"]: {"cell:5": ""}},
                    row_roles={row_a["id"]: {"role": "standalone_item", "parent_record_id": None},
                               row_b["id"]: {"role": "component", "parent_record_id": ""}},
                    issue_notes={issue["id"]: "Synthetic note not yet resolved"},
                    add_row_page=1, add_row_values={"synthetic name#1": "Unfinished manual row"})
    response = post(client, workspace, "review/save", {"review_token": before["review_token"], "edits": pending})
    assert response.status_code == 200, response.text
    reopened = state(_ApiClient(main.app), workspace)
    assert reopened["editor_draft"]["edits"] == pending
    assert reopened["editor_draft"]["revision"] == 1
    assert reopened["editor_draft"]["stale"] is False
    assert reopened["draft"] == before["draft"]
    assert reopened["review"] == before["review"]
    assert len(AIWorkspaceStore(workspace.root).read("editor-draft.json")["history"]) == 1
    assert reopened["review_token"] != before["review_token"]
    conflict = post(client, workspace, "review/save", {"review_token": before["review_token"], "edits": edits()})
    assert conflict.status_code == 409
    assert state(client, workspace)["editor_draft"] == reopened["editor_draft"]
    invalid = post(client, workspace, "review/save", {"review_token": reopened["review_token"], "edits": edits(row_values={"foreign-row": {"cell:0": "x"}})})
    assert invalid.status_code == 409
    assert invalid.json()["detail"]["record_id"] == "foreign-row"


def _checked_fragment(main, client, page_count=20, page_number=10):
    workspace = _make_document(main, page_count)
    _start_fake(client, workspace.document_id, pages=[page_number])
    data = state(client, workspace)
    page = next(p for p in data["draft"]["pages"] if p["page_number"] == page_number)
    first, service = page["tables"][0]["rows"]
    for row, role in [(first, "standalone_item"), (service, "header")]:
        assert post(client, workspace, "review/row-role", {"record_id": row["id"], "role": role}).status_code == 200
    assert post(client, workspace, "review/row", {"record_id": first["id"], "disposition": "corrected", "confirmed": True,
                                               "values": {"cell:1": "Synthetic checked name", "cell:4": "0", "cell:5": ""}}).status_code == 200
    assert post(client, workspace, "review/complete-page", {"page_number": page_number}).status_code == 200
    return workspace, first


def test_page10_of20_exports_without_confirming_other_pages_and_buffered_edits_fail_closed(ai_first_client):
    main, client = ai_first_client
    workspace, row = _checked_fragment(main, client)
    preview = post(client, workspace, "pilot-export/preview", {"page_numbers": [10]})
    assert preview.status_code == 200, preview.text
    assert preview.json()["verified_position_count"] == 1
    assert preview.json()["ocr_row_count"] == 2
    book = load_workbook(BytesIO(post(client, workspace, "pilot-export", {"page_numbers": [10]}).content))
    assert book.sheetnames[0] == "Спецификация"
    assert book["Спецификация"].cell(5, 2).value == "Synthetic checked name"
    assert book["Спецификация"].cell(5, 5).value == "0"
    assert book["Спецификация"].cell(5, 6).value in (None, "")
    assert "Частичный" in book["Спецификация"]["A1"].value
    assert "10" in book["Спецификация"]["A2"].value
    assert all(cell.data_type != "f" for sheet in book for cells in sheet for cell in cells)
    before = state(client, workspace)
    pending = edits(row_values={row["id"]: {"cell:1": "=SYNTHETIC_UNCONFIRMED()"}})
    assert post(client, workspace, "review/save", {"review_token": before["review_token"], "edits": pending}).status_code == 200
    assert post(client, workspace, "pilot-export/preview", {"page_numbers": [10]}).json()["verified_position_count"] == 0
    book = load_workbook(BytesIO(post(client, workspace, "pilot-export", {"page_numbers": [10]}).content))
    assert book["Спецификация"].cell(5, 1).value.startswith("Нет позиций")
    assert "Черновик редактора" in book.sheetnames
    assert all(cell.data_type != "f" for sheet in book for cells in sheet for cell in cells)
    assert state(client, workspace)["review"] == before["review"]
    # Explicit confirmation consumes the buffer, while the private editor history remains.
    current = state(client, workspace)
    assert post(client, workspace, "review/row", {"review_token": current["review_token"], "record_id": row["id"],
        "disposition": "corrected", "values": pending["row_values"][row["id"]], "confirmed": True}).status_code == 200
    after = state(client, workspace)
    assert after["editor_draft"]["edits"]["row_values"] == {}
    assert len(after["editor_draft"]["history"]) == 1
    assert post(client, workspace, "pilot-export/preview", {"page_numbers": [10]}).json()["verified_position_count"] == 1


@pytest.mark.parametrize("change", ["context", "dpi", "model", "output"])
def test_changed_identity_requires_explicit_new_run_preserves_old_checkpoint_and_other_document(ai_first_client, monkeypatch, change):
    main, client = ai_first_client
    workspace = _make_document(main)
    _start_fake(client, workspace.document_id, pages=[1])
    checked_row = state(client, workspace)["draft"]["pages"][0]["tables"][0]["rows"][0]["id"]
    assert post(client, workspace, "review/row-role", {"record_id": checked_row, "role": "standalone_item"}).status_code == 200
    assert post(client, workspace, "review/row", {"record_id": checked_row, "disposition": "confirmed"}).status_code == 200
    assert post(client, workspace, "review/complete-page", {"page_number": 1}).status_code == 200
    current = state(client, workspace)
    assert post(client, workspace, "review/save", {"review_token": current["review_token"],
        "edits": edits(row_values={checked_row: {"cell:1": "Unconfirmed old-run draft"}})}).status_code == 200
    store = AIWorkspaceStore(workspace.root)
    manifest = (store.checkpoint_directory / "manifest.json").read_bytes()
    original = main._build_ai_first_runtime
    if change != "context":
        def changed(request, count):
            adapter, config = original(request, count)
            if change == "dpi":
                config = config.model_copy(update={"dpi": 115})
            else:
                adapter._descriptor = adapter.descriptor.model_copy(update={"parameters": {change: "synthetic-new-setting"}})
            return adapter, config
        monkeypatch.setattr(main, "_build_ai_first_runtime", changed)
    payload = {"mode": "fake", "pages": [1], "user_context": "Synthetic different hint" if change == "context" else None}
    mismatch = post(client, workspace, "recognize", payload)
    assert mismatch.status_code == 409, mismatch.text
    assert mismatch.json()["detail"]["code"] == "checkpoint_identity_mismatch"
    before = state(client, workspace)
    fresh = _make_document(main)
    assert fresh.document_id != workspace.document_id
    response = post(client, fresh, "recognize", payload)
    assert response.status_code == 200, response.text
    assert _wait_for_job(client, response.json()["id"])["status"] == "completed"
    response = post(client, workspace, "recognize", {**payload, "start_new_run": True})
    assert response.status_code == 200, response.text
    assert _wait_for_job(client, response.json()["id"])["status"] == "completed"
    assert (store.checkpoint_directory / "manifest.json").read_bytes() == manifest
    active = store.read("active-run.json")["run_id"]
    assert (store.root / "runs" / active / "checkpoint" / "manifest.json").exists()
    archived = AIWorkspaceStore(store.root / "runs" / active)
    assert archived.read("transcription.json") == before["draft"]
    assert archived.read("review.json") == before["review"]
    after = state(client, workspace)
    assert after["review"]["decisions"][checked_row]["active"] is False
    assert after["review"]["row_roles"][checked_row]["active"] is False
    assert after["review"]["confirmed_complete_pages"] == []
    assert after["editor_draft"]["edits"] == edits()
    assert len(after["editor_draft"]["history"]) == 1
    assert archived.read("editor-draft.json")["edits"]["row_values"][checked_row]["cell:1"] == "Unconfirmed old-run draft"


def test_real_crop_renders_small_region_before_allocation_and_marks_scope(ai_first_client, monkeypatch):
    main, client = ai_first_client
    workspace = _make_sized_document(main, [(10000, 10000)])
    sizes = []
    original = main.SyntheticPageAdapter.extract
    async def capture(adapter, request):
        sizes.append((request.image.width, request.image.height))
        return await original(adapter, request)
    monkeypatch.setattr(main.SyntheticPageAdapter, "extract", capture)
    crop = {"page_number": 1, "x": 0.01, "y": 0.01, "width": 0.02, "height": 0.01}
    response = post(client, workspace, "recognize", {"mode": "fake", "pages": [1], "crop": crop})
    assert response.status_code == 200, response.text
    job = _wait_for_job(client, response.json()["id"])
    assert job["status"] == "completed", job
    assert sizes and sizes[0][0] < 420 and sizes[0][1] < 212
    data = state(client, workspace)
    envelope = data["draft"]["source_envelope"]
    assert envelope["page_regions"] == {"1": [0.01, 0.01, 0.02, 0.01]}
    assert envelope["page_region_pdf_points"]["1"] == [100.0, 100.0, 300.0, 200.0]
    assert data["review_ready_for_guard"]["code"] == "partial_page_scope"
    book = load_workbook(BytesIO(post(client, workspace, "pilot-export", {"page_numbers": [1]}).content))
    assert "Выбранная область" in book["Спецификация"]["A1"].value
    assert "Области страниц" in book.sheetnames
    for bad in [{**crop, "page_number": 2}, {**crop, "width": 2}, {**crop, "x": float("inf")}]:
        assert post(client, workspace, "recognize", {"mode": "fake", "pages": [1], "crop": bad}).status_code in {400, 422}
    assert len(sizes) == 1


def test_session_users_cannot_read_modify_or_export_another_ai_first_workspace(private_session_api, monkeypatch, tmp_path):
    from averon_import import main
    from averon_import.services.workspace import WorkspaceService
    from averon_import.services.ocr.qwen_vision_client import _HttpsTransport
    client, accounts = private_session_api
    monkeypatch.setattr(main, "workspace_service", WorkspaceService(tmp_path / "private-documents"))
    monkeypatch.setattr(main, "AI_FIRST_UI_ENABLED", True)
    monkeypatch.setattr(main, "AI_FIRST_FAKE_OCR_ENABLED", True)
    monkeypatch.setattr(main, "AI_FIRST_LIVE_ENABLED", False)
    monkeypatch.setattr(_HttpsTransport, "post", lambda *a, **k: pytest.fail("No cloud transport is allowed"))
    owner = accounts.create_user("synthetic-owner", "synthetic-test-password-1")
    other = accounts.create_user("synthetic-other", "synthetic-test-password-2")
    owner_session = accounts.create_session(owner["user_id"])
    other_session = accounts.create_session(other["user_id"])
    class SessionClient:
        def __init__(self, session):
            self.headers = _session_headers(session["token"], session["csrf_token"])
        def get(self, path):
            return client.get(path, headers=self.headers)
        def post(self, path, *, json=None):
            return client.post(path, json=json, headers=self.headers)
    a, b = SessionClient(owner_session), SessionClient(other_session)
    legacy = _make_document(main)
    assert b.get(f"/api/documents/{legacy.document_id}").status_code == 200
    uploaded = client.post_multipart("/api/documents", headers=a.headers,
        files={"file": ("synthetic-upload.pdf", legacy.pdf_path.read_bytes(), "application/pdf")})
    assert uploaded.status_code == 200
    assert "uploaded_by_user_id" not in uploaded.json()
    workspace = main.workspace_service.get(uploaded.json()["document_id"])
    assert post(b, workspace, "recognize", {"mode": "fake", "pages": [1]}).status_code == 404
    _start_fake(a, workspace.document_id, pages=[1])
    current = state(a, workspace)
    base = f"/api/documents/{workspace.document_id}"
    for path in [base, base + "/ai-first", base + "/ai-first/progress", base + "/page/1"]:
        assert b.get(path).status_code == 404, path
    assert workspace.document_id not in json.dumps(b.get("/api/documents").json())
    for action, payload in [("review/save", {"review_token": current["review_token"], "edits": edits()}),
                            ("pilot-export/preview", {"page_numbers": [1]}),
                            ("pilot-export", {"page_numbers": [1]}),
                            ("recognize", {"mode": "fake", "pages": [1]})]:
        assert post(b, workspace, action, payload).status_code == 404
    assert state(a, workspace)["draft"] == current["draft"]


def test_consuming_saved_role_preserves_saved_values_and_stale_token_cannot_modify(ai_first_client):
    main, client = ai_first_client
    workspace = _make_document(main)
    _start_fake(client, workspace.document_id, pages=[1])
    before = state(client, workspace)
    row = before["draft"]["pages"][0]["tables"][0]["rows"][0]
    pending = edits(row_values={row["id"]: {"cell:1": "Still pending values"}},
                    row_roles={row["id"]: {"role": "standalone_item", "parent_record_id": None}})
    assert post(client, workspace, "review/save", {"review_token": before["review_token"], "edits": pending}).status_code == 200
    saved = state(client, workspace)
    assert post(client, workspace, "review/row-role", {"review_token": before["review_token"], "record_id": row["id"], "role": "section"}).status_code == 409
    assert post(client, workspace, "review/row-role", {"review_token": saved["review_token"], "record_id": row["id"], "role": "standalone_item"}).status_code == 200
    updated = state(client, workspace)
    assert updated["editor_draft"]["edits"]["row_values"] == pending["row_values"]
    assert updated["editor_draft"]["edits"]["row_roles"] == {}
    assert len(updated["editor_draft"]["history"]) == 1


def test_unfinished_saved_manual_row_blocks_completeness_only_in_export_view(ai_first_client):
    main, client = ai_first_client
    workspace, row = _checked_fragment(main, client)
    before = state(client, workspace)
    pending = edits(add_row_page=10, add_row_values={"name#1": "Unfinished synthetic position"})
    assert post(client, workspace, "review/save", {"review_token": before["review_token"], "edits": pending}).status_code == 200
    assert post(client, workspace, "pilot-export/preview", {"page_numbers": [10]}).json()["verified_position_count"] == 0
    after = state(client, workspace)
    assert after["review"] == before["review"]
    assert after["review"]["confirmed_complete_pages"] == [10]


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_crop_matches_the_same_region_of_displayed_rotated_pdf(tmp_path, rotation):
    from decimal import Decimal
    from PIL import Image, ImageChops
    path = tmp_path / "synthetic-rotated.pdf"
    pdf = fitz.open()
    page = pdf.new_page(width=200, height=100)
    page.draw_rect(fitz.Rect(0, 0, 100, 50), color=None, fill=(1, 0, 0))
    page.draw_rect(fitz.Rect(100, 0, 200, 50), color=None, fill=(0, 1, 0))
    page.draw_rect(fitz.Rect(0, 50, 100, 100), color=None, fill=(0, 0, 1))
    page.set_rotation(rotation)
    pdf.save(path)
    pdf.close()
    config = OrchestrationConfig(max_budget_rub=Decimal("1"), pricing=TokenPricing(
        input_rub_per_1k=Decimal("1"), output_rub_per_1k=Decimal("1")))
    full = FitZPageRenderer().render(path, 1, 72, config)
    crop = FitZPageRenderer().render(path, 1, 72, config.model_copy(update={"page_regions": {1: (0.25, 0.2, 0.5, 0.5)}}))
    image = Image.open(BytesIO(full.data))
    expected = image.crop((int(full.width * 0.25), int(full.height * 0.2), int(full.width * 0.75), int(full.height * 0.7)))
    actual = Image.open(BytesIO(crop.data))
    assert actual.size == expected.size
    assert ImageChops.difference(actual, expected).getbbox() is None


@pytest.mark.parametrize("invalid", ["stale_revision", "unknown_schema"])
def test_incompatible_editor_file_cannot_silently_admit_old_confirmations(ai_first_client, invalid):
    main, client = ai_first_client
    workspace, row = _checked_fragment(main, client)
    current = state(client, workspace)
    assert post(client, workspace, "review/save", {"review_token": current["review_token"],
        "edits": edits(row_values={row["id"]: {"cell:1": "Unconfirmed synthetic revision"}})}).status_code == 200
    store = AIWorkspaceStore(workspace.root)
    saved = store.read("editor-draft.json")
    saved["draft_revision_id" if invalid == "stale_revision" else "version"] = "old-revision" if invalid == "stale_revision" else 2
    store.write("editor-draft.json", saved)
    response = client.get(f"/api/documents/{workspace.document_id}/ai-first")
    if invalid == "stale_revision":
        assert response.status_code == 200
        assert response.json()["editor_draft"]["stale"] is True
        assert response.json()["review_ready_for_guard"]["ready"] is False
    else:
        assert response.status_code == 409
    for action in ["pilot-export/preview", "pilot-export"]:
        assert post(client, workspace, action, {"page_numbers": [10]}).status_code == 409
    assert store.read("editor-draft.json") == saved


def test_fragment_workbook_does_not_include_buffered_text_from_other_pages(ai_first_client):
    from zipfile import ZipFile
    main, client = ai_first_client
    workspace = _make_document(main, 20)
    _start_fake(client, workspace.document_id, pages=[1, 10])
    before = state(client, workspace)
    first = before["draft"]["pages"][0]["tables"][0]["rows"][0]
    tenth = before["draft"]["pages"][9]["tables"][0]["rows"][0]
    other_text = "SYNTHETIC-OTHER-PAGE-DO-NOT-EXPORT"
    selected_text = "SYNTHETIC-SELECTED-PAGE-DRAFT"
    pending = edits(row_values={first["id"]: {"cell:1": other_text}, tenth["id"]: {"cell:1": selected_text}},
                    add_row_page=1, add_row_values={"name#1": other_text})
    assert post(client, workspace, "review/save", {"review_token": before["review_token"], "edits": pending}).status_code == 200
    response = post(client, workspace, "pilot-export", {"page_numbers": [10]})
    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive:
        xml = "".join(archive.read(name).decode("utf8") for name in archive.namelist() if name.endswith(".xml"))
    assert selected_text in xml
    assert other_text not in xml
    assert state(client, workspace)["editor_draft"]["edits"] == pending
