from __future__ import annotations

from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from averon_import.services.ai_first.contracts import (
    AICell, AIColumn, AIDocumentDraft, AIPageExtraction, AIRow, AITable,
)
from averon_import.services.ai_first.pilot_export import (
    PilotExportError, build_ai_first_diagnostic_xlsx, build_ai_first_pilot_xlsx,
)
from averon_import.services.ai_first.review import (
    confirm_page_complete, new_review_state, record_row_decision,
    record_row_role_classification, resolve_review_issue,
)


HEADERS = ["Поз.", "Наименование", "Марка", "Производитель", "Артикул", "Ед. изм.", "Количество", "Примечание"]
PAGE16_HEADERS = [
    "Позиция", "Наименование и техническая характеристика",
    "Тип, марка, обозначение\nдокумента, опросного\nлиста",
    "Код\nоборудования,\nматериала\nизделия", "Завод-\nизготовитель",
    "Единица\nизмерения", "Кол-во", "Масса\nединицы,\nкг.", "Примечание",
]


def _row(page_number: int, table_id: str, row_index: int, values: list[str], *, kind: str = "item") -> AIRow:
    record_id = f"p{page_number}-r{row_index}"
    cells = [AICell(
        id=f"{record_id}-c{index}", source_document_id="synthetic-doc",
        page_number=page_number, table_id=table_id, row_id=record_id,
        column_index=index, raw_value=value,
    ) for index, value in enumerate(values)]
    return AIRow(
        id=record_id, source_document_id="synthetic-doc", page_number=page_number,
        table_id=table_id, index=row_index, model_kind=kind, raw_cells=cells,
    )


def _fixture(*, issue: bool = False):
    table_id = "p1-t1"
    columns = [AIColumn(
        id=f"{table_id}-c{index}", source_document_id="synthetic-doc",
        page_number=1, table_id=table_id, index=index, raw_header=header,
    ) for index, header in enumerate(HEADERS)]
    parent = _row(1, table_id, 0, ["1", "Синтетический насос", "М-1", "Примерный завод", "SYN-1", "шт.", "2", "Исходная заметка"])
    section = _row(1, table_id, 1, ["", "Раздел В2", "", "", "", "", "", ""], kind="section")
    component = _row(1, table_id, 2, ["1.1", "Компонент комплекта", "М-2", "Завод 2", "SYN-2", "компл.", "1", "В составе комплекта"], kind="component")
    table = AITable(
        id=table_id, source_document_id="synthetic-doc", page_number=1,
        index=0, columns=columns, rows=[parent, section, component],
    )
    page = AIPageExtraction(
        source_document_id="synthetic-doc", page_number=1,
        extraction_fingerprint="page-fingerprint-1", status="extracted", tables=[table],
    )
    draft = AIDocumentDraft(
        source_document_id="synthetic-doc", revision_id="revision-1",
        pages=[page], expected_page_numbers=[1, 2],
    )
    review = new_review_state(draft)
    review = record_row_role_classification(draft, review, parent.id, role="standalone_item", reviewer_id="synthetic-reviewer")
    review = record_row_role_classification(draft, review, section.id, role="section", reviewer_id="synthetic-reviewer")
    review = record_row_role_classification(
        draft, review, component.id, role="component", parent_record_id=parent.id,
        reviewer_id="synthetic-reviewer",
    )
    review = record_row_decision(
        draft, review, parent.id, disposition="corrected", reviewer_id="synthetic-reviewer",
        values={"cell:6": "3"},
    )
    review = record_row_decision(
        draft, review, section.id, disposition="confirmed", reviewer_id="synthetic-reviewer",
    )
    review = record_row_decision(
        draft, review, component.id, disposition="confirmed", reviewer_id="synthetic-reviewer",
    )
    if issue:
        from averon_import.services.ai_first.contracts import AIReviewIssue
        draft = draft.model_copy(update={"issues": [AIReviewIssue(
            id="issue-section", source_document_id="synthetic-doc", code="synthetic_review",
            severity="review", page_number=1, record_ids=[parent.id], detail="synthetic fixture",
        )]})
    review = confirm_page_complete(draft, review, 1, reviewer_id="synthetic-reviewer")
    report = {
        "adapter": {"provider_id": "offline_synthetic", "model_id": "fake_vision"},
        "pages": [
            {"page_number": 1, "state": "completed"},
            {"page_number": 2, "state": "pending"},
        ],
    }
    return draft, review, report, parent, section, component


def _schema_fixture(headers: list[str], values: list[str]):
    draft, review, report, parent, *_ = _fixture()
    table_id = "p1-t1"
    row_id = parent.id
    columns = [AIColumn(
        id=f"{table_id}-c{index}", source_document_id="synthetic-doc",
        page_number=1, table_id=table_id, index=index, raw_header=header,
    ) for index, header in enumerate(headers)]
    row = AIRow(
        id=row_id, source_document_id="synthetic-doc", page_number=1,
        table_id=table_id, index=0, model_kind="item",
        raw_cells=[AICell(
            id=f"{row_id}-c{index}", source_document_id="synthetic-doc",
            page_number=1, table_id=table_id, row_id=row_id,
            column_index=index, raw_value=value,
        ) for index, value in enumerate(values)],
    )
    table = AITable(
        id=table_id, source_document_id="synthetic-doc", page_number=1,
        index=0, columns=columns, rows=[row],
    )
    page = draft.pages[0].model_copy(update={"tables": [table], "issues": []})
    draft = draft.model_copy(update={"pages": [page], "issues": []})
    return draft, review, report, row


def _sheet(workbook_bytes: bytes, name: str):
    book = load_workbook(BytesIO(workbook_bytes), data_only=False)
    return book[name]


def test_pilot_xlsx_separates_verified_standalone_rows_and_preserves_all_source_data():
    draft, review, report, parent, section, component = _fixture()
    content = build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )
    book = load_workbook(BytesIO(content), data_only=False)
    positions = book["Спецификация"]
    diagnostic = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )), data_only=False)
    all_rows = diagnostic["Все распознанные строки"]

    assert book.sheetnames[0] == "Спецификация"
    assert positions.max_row == 5  # compact metadata, source headers, one admitted row
    assert positions.max_column == len(HEADERS)
    assert [positions.cell(4, column).value for column in range(1, positions.max_column + 1)] == HEADERS
    assert [positions.cell(5, column).value for column in range(1, positions.max_column + 1)] == [
        "1", "Синтетический насос", "М-1", "Примерный завод", "SYN-1", "шт.", "3", "Исходная заметка",
    ]
    assert positions["A1"].value == "ТЕСТ · Частичный охват"
    assert positions.freeze_panes == "A4"
    assert positions.auto_filter.ref == "A4:H5"
    assert "UUID" not in " ".join(str(positions.cell(4, column).value or "") for column in range(1, positions.max_column + 1))
    assert "ID строки" not in " ".join(str(positions.cell(4, column).value or "") for column in range(1, positions.max_column + 1))
    assert book.sheetnames == ["Спецификация"]
    assert "коммерческая" not in " ".join(str(cell.value or "") for row in positions.iter_rows() for cell in row).casefold()
    source = diagnostic["P001-T001"]
    assert source.max_column == 6 + len(HEADERS) * 2
    assert [source.cell(5, 7 + index * 2).value or "" for index in range(len(HEADERS))] == HEADERS
    assert source.cell(6, 7 + 6 * 2).value == "2"
    assert source.cell(6, 8 + 6 * 2).value == "3"
    assert source.cell(6, 7 + 1 * 2).value == "Синтетический насос"
    assert source.cell(6, 8 + 1 * 2).value == "Синтетический насос"

    all_headers = {all_rows.cell(4, column).value: column for column in range(1, all_rows.max_column + 1)}
    records = {
        all_rows.cell(row, all_headers["ID строки"]).value: row
        for row in range(5, all_rows.max_row + 1)
    }
    assert set(records) == {parent.id, section.id, component.id}
    component_row = records[component.id]
    assert all_rows.cell(component_row, all_headers["Роль Human Review"]).value == "Компонент комплекта"
    section_row = records[section.id]
    assert all_rows.cell(section_row, all_headers["Роль Human Review"]).value == "Раздел"
    assert all_rows.cell(records[parent.id], all_headers["Лист исходной таблицы"]).value == "P001-T001"
    assert all_rows.cell(records[parent.id], all_headers["Строка на листе исходной таблицы"]).value == "6"
    assert all_rows.cell(records[parent.id], all_headers["Лист исходной таблицы"]).hyperlink.target == "#'P001-T001'!A6"
    source_detail = json.loads(all_rows.cell(records[parent.id], all_headers["Исходная строка и таблица (JSON)"]).value)
    assert len(source_detail["row"]["raw_cell_ids"]) == len(HEADERS)
    assert source_detail["full_raw_payload_sheet"] == "Исходная структура"
    assert "Исходная структура" in diagnostic.sheetnames
    source_sheet = diagnostic["Исходная структура"]
    chunks = [source_sheet.cell(row, 4).value for row in range(5, source_sheet.max_row + 1)]
    assert json.loads("".join(chunks))["tables"][0]["rows"][0]["id"] == parent.id


def test_unclassified_rows_and_unresolved_issues_never_enter_verified_positions():
    draft, review, report, parent, section, component = _fixture(issue=True)
    # An unresolved row-level issue keeps even a manually confirmed row out.
    with pytest.raises(PilotExportError, match="no_verified_positions"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    book = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")), data_only=False)
    all_rows = book["Все распознанные строки"]
    headers = {all_rows.cell(4, column).value: column for column in range(1, all_rows.max_column + 1)}
    parent_row = next(row for row in range(5, all_rows.max_row + 1) if all_rows.cell(row, headers["ID строки"]).value == parent.id)
    assert all_rows.cell(parent_row, headers["Статус проверки"]).value == "Есть неразрешённое замечание"


def test_document_note_is_kept_in_separate_block_and_never_as_position():
    draft, review, report, parent, *_ = _fixture()
    table = draft.pages[0].tables[0]
    note = AIRow(
        id="synthetic-note", source_document_id=draft.source_document_id,
        page_number=1, table_id=table.id, index=3, model_kind="project_note",
        raw_cells=[AICell(
            id=f"synthetic-note-c{index}", source_document_id=draft.source_document_id,
            page_number=1, table_id=table.id, row_id="synthetic-note",
            column_index=index, raw_value=("Согласование замены с заказчиком" if index == 1 else ""),
        ) for index in range(len(HEADERS))],
    )
    table.rows.append(note)
    workbook = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )), data_only=False)
    sheet = workbook["Спецификация"]
    assert workbook.sheetnames == ["Спецификация"]
    assert sheet.cell(5, 2).value == "Синтетический насос"
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "Примечания из документа · страница 1 · таблица 1" in values
    assert "Согласование замены с заказчиком" in values
    note_index = values.index("Согласование замены с заказчиком")
    assert "Требует проверки" in values[max(0, note_index - 4):note_index]


def test_resolved_row_level_visual_issue_can_be_human_verified():
    from averon_import.services.ai_first.contracts import AIReviewIssue

    draft, review, report, parent, *_ = _fixture()
    page = draft.pages[0]
    table = page.tables[0]
    row = table.rows[0]
    row.raw_cells[6] = row.raw_cells[6].model_copy(update={"struck_through":True})
    issue = AIReviewIssue(
        id="struck-quantity", source_document_id="synthetic-doc", code="row_has_uncertain_visual_evidence",
        severity="review", page_number=1, record_ids=[parent.id], detail="synthetic fixture",
    )
    page.issues.append(issue)
    draft = draft.model_copy(update={"issues":[issue]})
    review = resolve_review_issue(draft, review, issue.id, reviewer_id="synthetic-reviewer", detail="Сверено с PDF")
    content = build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook["Спецификация"].max_row == 5


def test_role_without_human_decision_is_exported_as_unverified_only():
    draft, review, report, parent, section, component = _fixture()
    review = review.model_copy(deep=True)
    review.decisions[parent.id] = review.decisions[parent.id].model_copy(update={"active": False})
    with pytest.raises(PilotExportError, match="no_verified_positions"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    book = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")), data_only=False)
    all_rows = book["Все распознанные строки"]
    headers = {all_rows.cell(4, column).value: column for column in range(1, all_rows.max_column + 1)}
    parent_row = next(row for row in range(5, all_rows.max_row + 1) if all_rows.cell(row, headers["ID строки"]).value == parent.id)
    assert all_rows.cell(parent_row, headers["Статус проверки"]).value == "Значения требуют проверки"


def test_invalidated_role_change_keeps_saved_values_visible_but_not_verified():
    draft, review, report, parent, *_ = _fixture()
    review = review.model_copy(deep=True)
    review.decisions[parent.id] = review.decisions[parent.id].model_copy(update={
        "active": False, "invalidation_reason": "role_change",
    })
    with pytest.raises(PilotExportError, match="no_verified_positions"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    book = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")), data_only=False)
    all_rows = book["Все распознанные строки"]
    headers = {all_rows.cell(4, column).value: column for column in range(1, all_rows.max_column + 1)}
    parent_row = next(row for row in range(5, all_rows.max_row + 1) if all_rows.cell(row, headers["ID строки"]).value == parent.id)
    saved = json.loads(all_rows.cell(parent_row, headers["Последние сохранённые значения Human Review (JSON)"]).value)
    source = book["P001-T001"]
    assert source.cell(6, 7 + 6 * 2).value == "2"
    assert source.cell(6, 8 + 6 * 2).value == "3"
    assert saved["cell:6"] == "3"
    assert json.loads(all_rows.cell(parent_row, headers["Последнее решение Human Review (JSON)"]).value)["invalidation_reason"] == "role_change"


def test_export_is_scoped_to_at_most_three_real_document_pages():
    draft, review, report, *_ = _fixture()
    with pytest.raises(PilotExportError, match="invalid_page_scope"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[], document_name="synthetic.pdf")
    with pytest.raises(PilotExportError, match="invalid_page_scope"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1, 1], document_name="synthetic.pdf")
    with pytest.raises(PilotExportError, match="page_outside_document"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[3], document_name="synthetic.pdf")


def test_selected_page_export_does_not_require_all_28_document_pages():
    draft, review, _report, *_ = _fixture()
    draft = draft.model_copy(update={"expected_page_numbers": list(range(1, 29))})
    report = {"pages": [
        {"page_number": 1, "state": "completed"},
        *[{"page_number": page, "state": "pending"} for page in range(2, 29)],
    ]}

    content = build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic-28-pages.pdf",
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    positions = workbook["Спецификация"]
    assert positions.max_row == 5
    assert positions["A1"].value == "ТЕСТ · Частичный охват"
    assert workbook.sheetnames == ["Спецификация"]


def test_ambiguous_duplicate_column_headers_do_not_enter_verified_positions():
    draft, review, report, parent, *_ = _fixture()
    table = draft.pages[0].tables[0]
    table.columns[7] = table.columns[7].model_copy(update={"raw_header":"Количество"})
    with pytest.raises(PilotExportError, match="no_verified_positions"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    book = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")), data_only=False)
    all_rows = book["Все распознанные строки"]
    headers = {all_rows.cell(4, column).value: column for column in range(1, all_rows.max_column + 1)}
    parent_row = next(row for row in range(5, all_rows.max_row + 1) if all_rows.cell(row, headers["ID строки"]).value == parent.id)
    assert all_rows.cell(parent_row, headers["Статус проверки"]).value == "Неоднозначное сопоставление колонок"


def test_ocr_values_are_written_as_text_and_not_excel_formulas():
    draft, review, report, parent, *_ = _fixture()
    row = draft.pages[0].tables[0].rows[0]
    row.raw_cells[1] = row.raw_cells[1].model_copy(update={"raw_value": "=HYPERLINK(\"https://example.invalid\",\"click\")"})
    content = build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    book = load_workbook(BytesIO(content), data_only=False)
    source = book["Спецификация"]
    source_cell = source.cell(5, 2)
    assert source_cell.value.startswith("=HYPERLINK")
    assert source_cell.data_type == "s"

    assert source_cell.value.startswith("=HYPERLINK")
    assert source_cell.data_type == "s"


def test_32767_character_value_remains_in_ordinary_workbook_cell():
    draft, review, report, parent, *_ = _fixture()
    value = "я" * 32767
    parent.raw_cells[1] = parent.raw_cells[1].model_copy(update={"raw_value": value})
    content = build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert "Длинные значения" not in workbook.sheetnames
    source_sheet = workbook["Спецификация"]
    assert any(cell.value == value for row in source_sheet.iter_rows() for cell in row)
    assert all(cell.data_type == "s" for row in source_sheet.iter_rows() for cell in row if cell.value == value)


def test_32768_character_value_is_losslessly_chunked_as_text():
    draft, review, report, parent, *_ = _fixture()
    value = "=" + "x" * 32767
    assert len(value) == 32768
    parent.raw_cells[1] = parent.raw_cells[1].model_copy(update={"raw_value": value})
    content = build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert "Длинные значения" in workbook.sheetnames

    source_sheet = workbook["Спецификация"]
    marker = next(
        cell for row in source_sheet.iter_rows() for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("Длинный текст · запись ")
    )
    assert marker.hyperlink is not None
    assert marker.data_type == "s"

    long_sheet = workbook["Длинные значения"]
    parts = [row for row in long_sheet.iter_rows(min_row=2, values_only=True) if row[1] == marker.coordinate]
    parts.sort(key=lambda row: row[2])
    assert parts
    assert [int(row[2]) for row in parts] == list(range(1, len(parts) + 1))
    assert all(int(row[3]) == len(parts) for row in parts)
    assert "".join(row[4] for row in parts) == value
    assert all(len(row[4]) <= 32767 for row in parts)
    assert all(long_sheet.cell(row_number, 5).data_type == "s" for row_number in range(2, long_sheet.max_row + 1))
    assert not any(cell.data_type == "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)


def test_export_source_label_comes_from_saved_adapter_or_falls_back_neutrally():
    draft, review, report, *_ = _fixture()
    synthetic = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )), data_only=False)
    assert "Синтетический тест" in synthetic["Информация"]["A1"].value

    report["adapter"] = {"provider_id": "yandex_ai_studio", "model_id": "qwen_vision"}
    qwen = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )), data_only=False)
    assert "Qwen Vision" in qwen["Информация"]["A1"].value

    report["adapter"] = {"provider_id": "unrecognized", "model_id": "unknown"}
    unknown = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )), data_only=False)
    assert "Источник распознавания не определён" in unknown["Информация"]["A1"].value


def test_admin_diagnostic_export_reports_usage_without_credential_fields():
    draft, review, report, *_ = _fixture()
    report.update({
        "status": "completed_with_errors", "elapsed_s": 4.5,
        "actual_input_tokens": 1250, "actual_output_tokens": 990,
        "actual_reasoning_tokens": None, "actual_cost_rub": "0.547",
        "source_integrity_verified": True, "expected_page_count": 2,
        "config": {"dpi": 150, "concurrency": 1, "max_budget_rub": "9", "api_key": "SYNTHETIC-SECRET"},
        "pages": [{"page_number": 1, "state": "completed", "attempts": [{
            "state": "completed", "http_status": 200, "finish_reason": "stop",
            "elapsed_s": 2.0, "usage": {"input_tokens": 1250, "output_tokens": 990, "reasoning_tokens": None},
            "safe_error_code": None,
        }]}],
    })
    workbook = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic.pdf",
    )), data_only=False)
    diag_text = " ".join(str(cell.value or "") for sheet in workbook for row in sheet.iter_rows() for cell in row)
    assert "1250" in diag_text and "990" in diag_text and "4.5" in diag_text
    assert "SYNTHETIC-SECRET" not in diag_text
    assert "api_key" not in diag_text.casefold()
    assert "Диагностика запуска" in workbook.sheetnames
    assert "Попытки страниц" in workbook.sheetnames


def test_large_model_payload_is_chunked_without_breaking_row_review_export():
    draft, review, report, *_ = _fixture()
    table = draft.pages[0].tables[0]
    table.raw_model_payload = {"synthetic_payload": "x" * 100_000}

    content = build_ai_first_diagnostic_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic.pdf")
    workbook = load_workbook(BytesIO(content), data_only=False)
    all_rows = workbook["Все распознанные строки"]
    headers = {all_rows.cell(4, column).value: column for column in range(1, all_rows.max_column + 1)}
    row_evidence = json.loads(all_rows.cell(5, headers["Исходная строка и таблица (JSON)"]).value)
    assert row_evidence["full_raw_payload_sheet"] == "Исходная структура"

    source = workbook["Исходная структура"]
    chunks = [source.cell(row, 4).value for row in range(5, source.max_row + 1)]
    assert all(len(chunk) <= 30000 for chunk in chunks)
    assert "synthetic_payload" in "".join(chunks)


def test_page16_nine_column_shape_remains_visible_without_template_mapping():
    values = [f"SYNTHETIC-VALUE-{index + 1}" for index in range(9)]
    draft, review, report, _row = _schema_fixture(PAGE16_HEADERS, values)
    content = build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="synthetic-page16.pdf")
    book = load_workbook(BytesIO(content), data_only=False)
    positions = book["Спецификация"]
    assert positions.max_column == 9
    assert [positions.cell(4, column).value for column in range(1, 10)] == PAGE16_HEADERS
    assert positions.cell(5, 2).value == "SYNTHETIC-VALUE-2"
    assert positions.cell(5, 3).value == "SYNTHETIC-VALUE-3"
    expected = list(values)
    expected[6] = "3"  # saved and explicitly confirmed Human Review correction
    assert [positions.cell(5, column).value for column in range(1, 10)] == expected


def test_page16_nine_column_export_has_twenty_three_verified_rows_and_preserves_audit():
    table_id = "synthetic-page16-table"
    values_by_row = [
        [str(index + 1), f"Synthetic item {index + 1:02d}", f"Type / mark {index + 1:02d}",
         f"Synthetic maker {index + 1:02d}", f"SYN-ART-{index + 1:03d}", "pcs", str(index + 2),
         f"{index + 1}.25", f"Synthetic note {index + 1:02d}"]
        for index in range(23)
    ]
    rows = [_row(1, table_id, index, values) for index, values in enumerate(values_by_row)]
    rows.extend([
        _row(1, table_id, 23, ["", "Section A", "", "", "", "", "", "", ""], kind="section"),
        _row(1, table_id, 24, ["23.1", "Synthetic component", "Type C", "Maker C", "SYN-C", "pcs", "1", "0.5", "Part of set"], kind="component"),
    ])
    columns = [AIColumn(
        id=f"{table_id}-column-{index}", source_document_id="synthetic-doc",
        page_number=1, table_id=table_id, index=index, raw_header=header,
    ) for index, header in enumerate(PAGE16_HEADERS)]
    table = AITable(id=table_id, source_document_id="synthetic-doc", page_number=1, index=0, columns=columns, rows=rows)
    page = AIPageExtraction(
        source_document_id="synthetic-doc", page_number=1,
        extraction_fingerprint="synthetic-page16-fingerprint", status="extracted", tables=[table],
    )
    draft = AIDocumentDraft(
        source_document_id="synthetic-doc", revision_id="synthetic-page16-revision",
        pages=[page], expected_page_numbers=list(range(1, 29)),
    )
    review = new_review_state(draft)
    for row in rows:
        role = "standalone_item" if row.index < 23 else ("section" if row.index == 23 else "component")
        review = record_row_role_classification(
            draft, review, row.id, role=role,
            parent_record_id=rows[0].id if role == "component" else None,
            reviewer_id="synthetic-reviewer",
        )
        review = record_row_decision(
            draft, review, row.id, disposition="confirmed", reviewer_id="synthetic-reviewer",
        )
    review = confirm_page_complete(draft, review, 1, reviewer_id="synthetic-reviewer")
    report = {
        "adapter": {"provider_id": "offline_synthetic", "model_id": "fake_vision"},
        "pages": [
            {"page_number": 1, "state": "completed"},
            *[{"page_number": number, "state": "pending"} for number in range(2, 29)],
        ],
    }

    workbook = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic-page16.pdf",
    )), data_only=False)
    specification = workbook["Спецификация"]
    assert workbook.sheetnames[0] == "Спецификация"
    assert specification.max_column == 9
    assert specification.max_row == 27
    assert [specification.cell(4, column).value for column in range(1, 10)] == PAGE16_HEADERS
    assert [specification.cell(5, column).value for column in range(1, 10)] == values_by_row[0]
    assert [specification.cell(27, column).value for column in range(1, 10)] == values_by_row[22]
    assert all(specification.cell(row, 2).value for row in range(5, 28))
    assert specification["A1"].value == "ТЕСТ · Частичный охват"
    assert specification.freeze_panes == "A4"
    assert specification.auto_filter.ref == "A4:I27"
    assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    diagnostic = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="synthetic-page16.pdf",
    )), data_only=False)
    source = diagnostic["P001-T001"]
    assert source.max_column == 6 + 9 * 2
    assert source.max_row == 30
    assert source.cell(29, 4).value == "Раздел"
    assert source.cell(30, 4).value == "Компонент комплекта"
    assert source.cell(30, 7 + 8 * 2).value == "Part of set"
    complete_source_values = values_by_row + [
        ["", "Section A", "", "", "", "", "", "", ""],
        ["23.1", "Synthetic component", "Type C", "Maker C", "SYN-C", "pcs", "1", "0.5", "Part of set"],
    ]
    assert [
        [(source.cell(6 + row_index, 7 + column_index * 2).value or "") for column_index in range(9)]
        for row_index in range(25)
    ] == complete_source_values
    assert workbook.sheetnames == ["Спецификация"]
    assert "Исходная структура" in diagnostic.sheetnames


def test_unknown_headers_are_exported_as_ordinary_cells():
    headers = ["Диаметр наружный (условн.)", "Код позиции заказчика", "Группа / узел"]
    values = ["SYN-D-14", "SYN-C-02", "SYN-G-1"]
    draft, review, report, _row = _schema_fixture(headers, values)
    book = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="unknown-headers.pdf",
    )), data_only=False)
    specification = book["Спецификация"]
    assert [specification.cell(4, index).value for index in range(1, 4)] == headers
    assert [specification.cell(5, index).value for index in range(1, 4)] == values


def test_cell_without_column_descriptor_is_still_visible_by_source_index():
    draft, review, report, parent, *_ = _fixture()
    parent.raw_cells.append(AICell(
        id=f"{parent.id}-undeclared-c8", source_document_id="synthetic-doc",
        page_number=1, table_id=parent.table_id, row_id=parent.id,
        column_index=8, raw_value="SYN-UNDECLARED-CELL",
    ))
    book = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="missing-column-metadata.pdf",
    )), data_only=False)
    sheet = book["Спецификация"]
    assert sheet.cell(4, 9).value in (None, "")
    assert sheet.cell(5, 9).value == "SYN-UNDECLARED-CELL"


def test_blank_and_repeated_headers_keep_source_indexes_and_exact_values():
    headers = ["", "Повтор", "Повтор", "   "]
    values = ["SYN-EMPTY", "SYN-A", "SYN-B", "SYN-SPACE"]
    draft, review, report, _row = _schema_fixture(headers, values)
    book = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="duplicate-headers.pdf",
    )), data_only=False)
    specification = book["Спецификация"]
    assert [specification.cell(4, index).value or "" for index in range(1, 5)] == headers
    assert [specification.cell(5, index).value for index in range(1, 5)] == values


def test_different_table_shapes_on_separate_pages_use_separate_blocks_on_one_sheet():
    draft, review, _report, *_ = _fixture()
    headers = ["Неизвестный столбец", "Другой столбец"]
    table_id, row_id = "p2-t1", "p2-r0"
    columns = [AIColumn(
        id=f"{table_id}-c{index}", source_document_id="synthetic-doc",
        page_number=2, table_id=table_id, index=index, raw_header=header,
    ) for index, header in enumerate(headers)]
    row = AIRow(
        id=row_id, source_document_id="synthetic-doc", page_number=2,
        table_id=table_id, index=0, model_kind="item",
        raw_cells=[AICell(
            id=f"{row_id}-c{index}", source_document_id="synthetic-doc",
            page_number=2, table_id=table_id, row_id=row_id,
            column_index=index, raw_value=value,
        ) for index, value in enumerate(["SYN-P2-A", "SYN-P2-B"])],
    )
    table = AITable(id=table_id, source_document_id="synthetic-doc", page_number=2, index=0, columns=columns, rows=[row])
    page2 = AIPageExtraction(
        source_document_id="synthetic-doc", page_number=2,
        extraction_fingerprint="page-fingerprint-2", status="extracted", tables=[table],
    )
    draft = draft.model_copy(update={"pages": [draft.pages[0], page2], "expected_page_numbers": [1, 2, 3]})
    report = {"pages": [{"page_number": 1, "state": "completed"}, {"page_number": 2, "state": "completed"}]}
    book = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1, 2], document_name="different-tables.pdf",
    )), data_only=False)
    assert book.sheetnames == ["Спецификация"]
    assert "SYN-P2-A" not in [cell.value for row_values in book["Спецификация"].iter_rows() for cell in row_values]

    review = record_row_role_classification(
        draft, review, row.id, role="standalone_item", reviewer_id="synthetic-reviewer",
    )
    review = record_row_decision(draft, review, row.id, disposition="confirmed", reviewer_id="synthetic-reviewer")
    review = confirm_page_complete(draft, review, 2, reviewer_id="synthetic-reviewer")
    workbook = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1, 2], document_name="different-tables.pdf",
    )), data_only=False)
    assert workbook.sheetnames == ["Спецификация"]
    cells = [cell.value for row_values in workbook["Спецификация"].iter_rows() for cell in row_values]
    assert cells.count("Неизвестный столбец") == 1
    assert cells.count("Другой столбец") == 1
    assert "SYN-P2-A" in cells and "SYN-P2-B" in cells
    assert "Страница 1 · Таблица 1" in cells and "Страница 2 · Таблица 1" in cells


def test_unresolved_table_level_continuation_issue_blocks_confirmed_positions():
    from averon_import.services.ai_first.contracts import AIReviewIssue

    draft, review, report, parent, *_ = _fixture()
    issue = AIReviewIssue(
        id="table-continuation", source_document_id="synthetic-doc", code="continuation_unverified",
        severity="review", page_number=1, record_ids=[parent.table_id], detail="synthetic continuation review",
    )
    # Page-level issue storage is supported independently of the draft index.
    draft.pages[0].issues.append(issue)
    with pytest.raises(PilotExportError, match="no_verified_positions"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="table-issue.pdf")
    book = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(draft, review, report, page_numbers=[1], document_name="table-issue.pdf")), data_only=False)
    summary = book["Все распознанные строки"]
    headers = {summary.cell(4, column).value: column for column in range(1, summary.max_column + 1)}
    parent_row = next(row for row in range(5, summary.max_row + 1) if summary.cell(row, headers["ID строки"]).value == parent.id)
    assert summary.cell(parent_row, headers["Статус проверки"]).value == "Есть неразрешённое замечание"

    review = resolve_review_issue(draft, review, issue.id, reviewer_id="synthetic-reviewer", detail="Проверено вручную")
    content = build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="table-issue.pdf")
    assert load_workbook(BytesIO(content), data_only=False)["Спецификация"].max_row == 5


def test_saved_correction_is_visible_but_not_in_confirmed_positions_until_confirmed():
    draft, review, report, parent, *_ = _fixture()
    saved_values = {"cell:6": "3", "cell:2": "Saved synthetic correction"}
    review = record_row_decision(
        draft, review, parent.id, disposition="corrected", reviewer_id="synthetic-reviewer",
        values=saved_values, human_confirmed=False,
    )
    with pytest.raises(PilotExportError, match="no_verified_positions"):
        build_ai_first_pilot_xlsx(draft, review, report, page_numbers=[1], document_name="saved-correction.pdf")
    workbook = load_workbook(BytesIO(build_ai_first_diagnostic_xlsx(
        draft, review, report, page_numbers=[1], document_name="saved-correction.pdf",
    )), data_only=False)
    assert "saved synthetic correction" not in str(workbook["Спецификация"].values)
    summary = workbook["Все распознанные строки"]
    headers = {summary.cell(4, column).value: column for column in range(1, summary.max_column + 1)}
    exported_parent = next(row for row in range(5, summary.max_row + 1) if summary.cell(row, headers["ID строки"]).value == parent.id)
    assert summary.cell(exported_parent, headers["Статус проверки"]).value == "Исправления сохранены, но не подтверждены"

    review = record_row_decision(draft, review, parent.id, disposition="confirmed", reviewer_id="synthetic-reviewer")
    assert review.decisions[parent.id].values == saved_values
    workbook = load_workbook(BytesIO(build_ai_first_pilot_xlsx(
        draft, review, report, page_numbers=[1], document_name="saved-correction.pdf",
    )), data_only=False)
    assert workbook["Спецификация"].max_row == 5
    assert workbook["Спецификация"].cell(5, 7).value == "3"
    assert workbook["Спецификация"].cell(5, 3).value == "Saved synthetic correction"
